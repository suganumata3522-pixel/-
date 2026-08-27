# -*- coding: utf-8 -*-
"""第一・第二構造設計G 若手合同勉強会 カリキュラム表を作成する。
出力: docs/joint_study/合同勉強会カリキュラム.xlsx
"""
import os
from datetime import date, timedelta
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = ("/root/.claude/uploads/e311f4a7-bf71-5958-8017-03a2b7df06f5/"
       "4fcf4979-260521_____G___________Ver.1.1_.xlsx")
OUT = "docs/joint_study"; os.makedirs(OUT, exist_ok=True)

# ---- 元資料から大項目を抽出 ----
src = load_workbook(SRC, data_only=True)


def items(sheet):
    ws = src[sheet]; out = []
    for r in range(1, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if a and isinstance(a, str) and a.strip().startswith("NO.") and b:
            out.append((a.strip(), str(b).strip()))
    return out


ITEMS = {s: items(s) for s in
         ["③-1_RC基本", "③-2_S基本", "④-1_RC二次部材", "④-2_S二次部材",
          "⑤-1_RC上部構造", "⑤-2_S上部構造", "⑥_地盤基本・基礎構造",
          "⑨-1_RC監理", "⑨-2_S監理"]}

G2, G1, MX = "第二構造設計G", "第一構造設計G", "第一・第二 合同"

# (区分, シート, 表示名, 対象index(0始まり), 担当G, 備考)
PLAN = [
    ("③-1", "③-1_RC基本", "RC造基本", range(0, 9), G2, ""),
    ("③-2", "③-2_S基本", "S造基本", range(0, 9), G1, ""),
    ("④-1", "④-1_RC二次部材", "RC造二次部材 ①", range(0, 3), G2, "手計算の3本柱"),
    ("④-2", "④-2_S二次部材", "S造二次部材", range(0, 10), G1, "項目数が多く要事前配布"),
    ("④-1", "④-1_RC二次部材", "RC造二次部材 ②", range(3, 7), G2, ""),
    ("⑤-2", "⑤-2_S上部構造", "S造上部構造 ①", range(0, 5), G1, "柱脚〜ルート"),
    ("⑤-1", "⑤-1_RC上部構造", "RC造上部構造 ①", range(0, 5), G2, "モデル化〜偏心率"),
    ("⑤-2", "⑤-2_S上部構造", "S造上部構造 ②", range(5, 10), G1, "保有耐力〜横補剛"),
    ("⑤-1", "⑤-1_RC上部構造", "RC造上部構造 ②", range(5, 10), G2, "保有水平耐力〜配筋調整"),
    ("⑥", "⑥_地盤基本・基礎構造", "地盤基本・基礎構造 ①", range(0, 5), MX, "両G混合。地盤調査〜支持力"),
    ("⑥", "⑥_地盤基本・基礎構造", "地盤基本・基礎構造 ②", range(5, 10), MX, "両G混合。杭〜基礎設計"),
    ("⑨-2", "⑨-2_S監理", "S造監理", range(0, 6), G1, "検査G 天野氏より製品検査のチェックポイント解説"),
    ("⑨-1", "⑨-1_RC監理", "RC造監理", range(0, 7), G2, "監理G 宮嵜氏より配筋検査のチェックポイント解説"),
]

# ---- 日程（2026/9/17 木を初回、以降隔週木曜、2027/3末まで）----
WD = "月火水木金土日"
d = date(2026, 9, 17); DATES = []
while d <= date(2027, 3, 31):
    DATES.append(d); d += timedelta(days=14)
assert len(DATES) == 1 + len(PLAN), (len(DATES), len(PLAN))

# ---- 書式 ----
wb = Workbook()
C_T, C_H = "1F4E79", "2E75B6"
COL = {G2: "DEEBF7", G1: "E2EFDA", MX: "FFF2CC"}
thin = Side(style="thin", color="B4C6E7")
bd = Border(left=thin, right=thin, top=thin, bottom=thin)
f_t = Font(name="MS PGothic", size=15, bold=True, color="FFFFFF")
f_h = Font(name="MS PGothic", size=10, bold=True, color="FFFFFF")
f_b = Font(name="MS PGothic", size=10)
f_s = Font(name="MS PGothic", size=9, color="555555")
f_bd = Font(name="MS PGothic", size=10, bold=True)
wrap = Alignment(wrap_text=True, vertical="center")
ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)
lft = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)


def setup(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False


def title(ws, row, text, span):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_t
    c.fill = PatternFill("solid", fgColor=C_T); c.alignment = lft
    ws.row_dimensions[row].height = 30


def note(ws, row, text, span, h=None, font=None):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = font or f_b
    c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    if h:
        ws.row_dimensions[row].height = h


# ===================== Sheet1 年間スケジュール =====================
ws = wb.active; ws.title = "年間スケジュール"
setup(ws, [6, 13, 6, 8, 30, 12, 16, 14, 14, 34])
title(ws, 1, "第一・第二構造設計G 若手合同勉強会　年間カリキュラム（2026年度下期）", 10)
note(ws, 2, "会場・時間：別途調整　／　隔週木曜　／　"
            "担当は第二構造設計G・第一構造設計Gが交互（⑥は両G混合）", 10, h=20, font=f_s)

HDR = ["回", "実施日", "曜日", "区分", "解説項目", "対象NO.", "担当G", "担当者", "指導者", "備考"]
r = 4
for j, h in enumerate(HDR):
    c = ws.cell(r, 1 + j, h); c.font = f_h
    c.fill = PatternFill("solid", fgColor=C_H); c.alignment = ctr; c.border = bd
ws.row_dimensions[r].height = 22

# ガイダンス行
r += 1
gd = DATES[0]
vals = ["―", gd.strftime("%Y/%m/%d"), WD[gd.weekday()], "―", "ガイダンス（趣旨説明・進め方・担当割当）",
        "―", "第一・第二 合同", "", "", "スキルマップ確認、年間スケジュール共有、発表要領の説明"]
for j, v in enumerate(vals):
    c = ws.cell(r, 1 + j, v); c.font = f_bd if j == 4 else f_b
    c.alignment = lft if j in (4, 9) else ctr
    c.border = bd; c.fill = PatternFill("solid", fgColor="FCE4D6")
ws.row_dimensions[r].height = 30

for i, (kind, sheet, disp, idx, grp, memo) in enumerate(PLAN):
    r += 1
    dt = DATES[i + 1]
    nos = [ITEMS[sheet][k][0] for k in idx]
    nolabel = f"{nos[0]}〜{nos[-1]}" if len(nos) > 1 else nos[0]
    vals = [f"第{i+1}回", dt.strftime("%Y/%m/%d"), WD[dt.weekday()], kind, disp,
            nolabel, grp, "", "", memo]
    for j, v in enumerate(vals):
        c = ws.cell(r, 1 + j, v)
        c.font = f_bd if j == 4 else f_b
        c.alignment = lft if j in (4, 9) else ctr
        c.border = bd
        if j in (6,):
            c.fill = PatternFill("solid", fgColor=COL[grp])
        elif j in (7, 8):
            c.fill = PatternFill("solid", fgColor="FFFFFF")
        else:
            c.fill = PatternFill("solid", fgColor="F7FAFD" if i % 2 == 0 else "FFFFFF")
    ws.row_dimensions[r].height = 30

r += 2
note(ws, r, "【担当の考え方】\n"
            "・RC造の内容（③-1／④-1／⑤-1／⑨-1）は第二構造設計Gの若手が第一構造設計Gへ解説\n"
            "・S造の内容（③-2／④-2／⑤-2／⑨-2）は第一構造設計Gの若手が第二構造設計Gへ解説\n"
            "・⑥地盤基本・基礎構造は両Gが共通で使う知識のため、両Gの若手が混合で担当\n"
            "・初回を第二構造設計Gとし、以降は第二→第一→第二…と交互（⑥の2回は交互の対象外）",
     10, h=90)
r += 1
note(ws, r, "【備考】\n"
            "・⑨監理の2回は、検査G・監理Gのご協力をいただき、実務のチェックポイントを併せて解説\n"
            "・各回の後半に「担当物件の紹介／設計時・施工時の失敗事例・不具合事例」の時間を設ける\n"
            "・12/24 と 1/7 は年末年始に近接するため、業務状況により振替を検討",
     10, h=70, font=f_s)

# ===================== Sheet2 各回の詳細項目 =====================
ws = wb.create_sheet("各回の詳細項目")
setup(ws, [6, 13, 10, 26, 10, 62])
title(ws, 1, "各回で解説する項目（人材育成ガイドライン スキルマップ 大項目）", 6)
note(ws, 2, "※ 各大項目の下にある小項目（■＝習得済／□＝未習得）は元資料の各シートを参照", 6,
     h=18, font=f_s)
r = 4
for j, h in enumerate(["回", "実施日", "区分", "解説項目", "NO.", "大項目"]):
    c = ws.cell(r, 1 + j, h); c.font = f_h
    c.fill = PatternFill("solid", fgColor=C_H); c.alignment = ctr; c.border = bd
ws.row_dimensions[r].height = 22

for i, (kind, sheet, disp, idx, grp, memo) in enumerate(PLAN):
    dt = DATES[i + 1]
    idx = list(idx)
    start = r + 1
    for k, ii in enumerate(idx):
        r += 1
        no, txt = ITEMS[sheet][ii]
        for j, v in enumerate([f"第{i+1}回", dt.strftime("%Y/%m/%d"), kind, disp, no, txt]):
            c = ws.cell(r, 1 + j, v); c.font = f_b
            c.alignment = lft if j == 5 else ctr; c.border = bd
            c.fill = PatternFill("solid", fgColor=COL[grp] if j <= 4 else "FFFFFF")
        ws.row_dimensions[r].height = 20
    for col in range(1, 5):
        ws.merge_cells(start_row=start, start_column=col, end_row=r, end_column=col)

# ===================== Sheet3 運営要領 =====================
ws = wb.create_sheet("運営要領")
setup(ws, [4, 22, 76])
title(ws, 1, "運営要領（案）", 3)
rows = [
    ("目的", "・グループ間の人材異動を円滑にするため、互いの構造形式への理解を深める\n"
             "・勉強会を通じて両グループの若手が意見交換できる関係性を築く\n"
             "・若手が他者に説明する機会をつくり、自らの理解を確かなものにする"),
    ("対象", "第一構造設計G・第二構造設計G の若手職員"),
    ("期間・頻度", "2026年9月17日（木）ガイダンス、以降 隔週木曜、2027年3月18日（木）まで 全14回"),
    ("1回の構成（案）", "① 解説（60分）　担当Gの若手がスキルマップの項目を解説\n"
                  "② 質疑・意見交換（15分）　聞き手側のGから質問\n"
                  "③ 物件紹介・事例共有（15分）　担当物件の概要、苦労した点、工夫した点、\n"
                  "　　審査機関・施主の指摘対応、設計時・施工時の失敗事例や不具合事例"),
    ("担当者", "各回の解説担当は、担当Gの若手職員が持ち回りで務める"),
    ("指導者", "担当者と同一Gの中堅職員が、資料作成と当日の補足を担当する"),
    ("特別解説", "⑨監理の回は、検査G（天野氏）に製品検査、監理G（宮嵜氏）に配筋検査の\n"
              "チェックポイントを解説いただく"),
    ("資料", "解説資料は前週までに共有し、聞き手側は事前に目を通したうえで参加する"),
    ("記録", "各回の資料・質疑内容を蓄積し、次年度以降の教育資料として活用する"),
]
r = 3
for k, v in rows:
    r += 1
    c = ws.cell(r, 2, k); c.font = f_bd; c.alignment = ctr
    c.fill = PatternFill("solid", fgColor="DEEBF7"); c.border = bd
    c2 = ws.cell(r, 3, v); c2.font = f_b
    c2.alignment = Alignment(wrap_text=True, vertical="center", indent=1); c2.border = bd
    ws.row_dimensions[r].height = 20 + 14 * v.count("\n")

XLSX = os.path.join(OUT, "合同勉強会カリキュラム.xlsx")
wb.save(XLSX)
print("saved:", XLSX)
print(f"ガイダンス {DATES[0]} ＋ 講義 {len(PLAN)} 回（最終 {DATES[-1]}）")
grps = [p[4] for p in PLAN]
print("担当順:", " → ".join("二" if g == G2 else "一" if g == G1 else "混" for g in grps))
