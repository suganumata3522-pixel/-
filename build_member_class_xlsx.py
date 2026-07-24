# -*- coding: utf-8 -*-
"""部材種別 問題集（図つき）Excel 生成スクリプト。
出力: docs/member_class/部材種別問題集.xlsx
NO.8 2次設計における部材種別を理解している
1 柱部材種別 / 2 梁部材種別 / 3 耐震壁部材種別 / 4 部材群種別の設定
RC造マンションの設計担当を想定。数値は build 時に検算済み。
"""
import os
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib import font_manager as fm

FONT_PATH = "/usr/share/fonts/truetype/fonts-japanese-gothic.ttf"
jp = fm.FontProperties(fname=FONT_PATH)
fm.fontManager.addfont(FONT_PATH)
plt.rcParams["font.family"] = jp.get_name()
plt.rcParams["axes.unicode_minus"] = False

OUT = "docs/member_class"
FIG = os.path.join(OUT, "figures")
os.makedirs(FIG, exist_ok=True)
C_BLUE = "#2a78d6"; C_PINK = "#d55181"


def save(fig, name):
    p = os.path.join(FIG, name)
    fig.savefig(p, dpi=130, bbox_inches="tight")
    plt.close(fig)
    return p


# 図1: 部材種別の全体像（FA〜FD → Ds）
def fig_overview():
    fig, ax = plt.subplots(figsize=(12, 5.0))
    classes = [("FA", "靭性 大", "#a8dadc", 0.30),
               ("FB", "靭性 やや大", "#cfe0c0", 0.35),
               ("FC", "靭性 中", "#fde2c4", 0.40),
               ("FD", "脆性（靭性小）", "#f8c0c0", 0.45)]
    x = 0.5
    for name, desc, color, ds in classes:
        ax.add_patch(mpatches.FancyBboxPatch((x, 2.5), 2.3, 1.6,
                     boxstyle="round,pad=0.05", fc=color, ec="k", lw=1.2))
        ax.text(x + 1.15, 3.7, name, fontproperties=jp, ha="center",
                fontsize=14, fontweight="bold")
        ax.text(x + 1.15, 3.1, desc, fontproperties=jp, ha="center",
                fontsize=9)
        ax.text(x + 1.15, 2.0, f"Ds≒{ds}", fontproperties=jp, ha="center",
                fontsize=9, color="#c00000")
        x += 2.6
    ax.annotate("", xy=(10.5, 1.6), xytext=(0.5, 1.6),
                arrowprops=dict(arrowstyle="-|>", color="#c00000", lw=2))
    ax.text(5.5, 1.2, "靭性が小さいほど（FA→FD）Ds が大きい＝必要保有水平耐力が増える",
            fontproperties=jp, ha="center", fontsize=9.5, color="#c00000")
    ax.text(5.5, 4.6,
            "部材種別＝各部材の『粘り強さ（靭性）』のランク。"
            "軸力比・せん断余裕・帯筋等で判定",
            fontproperties=jp, ha="center", fontsize=10, color="#1f4e79")
    ax.set_xlim(0, 11); ax.set_ylim(0.8, 5.0)
    ax.axis("off")
    ax.set_title("図 1  部材種別 FA〜FD と Ds",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig1_overview.png")


# 図2: 柱の部材種別パラメータ
def fig_column():
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 5.0))
    # (a) 断面と軸力
    ax = axes[0]
    ax.add_patch(mpatches.Rectangle((0, 0), 3, 3, fc="#d9d9d9", ec="k",
                 lw=1.2))
    for px in [0.5, 1.5, 2.5]:
        for py in [0.5, 2.5]:
            ax.add_patch(plt.Circle((px, py), 0.16, fc="k"))
    for py in [1.5]:
        for px in [0.5, 2.5]:
            ax.add_patch(plt.Circle((px, py), 0.16, fc="k"))
    ax.annotate("", xy=(1.5, 3.3), xytext=(1.5, 4.0),
                arrowprops=dict(arrowstyle="-|>", color="#c00000", lw=2.5))
    ax.text(1.5, 4.2, "軸力 N", fontproperties=jp, ha="center", fontsize=9,
            color="#c00000")
    ax.text(1.5, -0.7, "柱600角、帯筋（せん断補強）",
            fontproperties=jp, ha="center", fontsize=8.5, color="#444")
    ax.set_xlim(-0.5, 3.5); ax.set_ylim(-1.2, 4.4)
    ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("(a) 柱断面", fontproperties=jp, fontsize=10,
                 fontweight="bold")
    # (b) パラメータ
    ax = axes[1]
    ax.text(0.5, 0.93, "柱の部材種別 判定パラメータ", transform=ax.transAxes,
            ha="center", fontproperties=jp, fontsize=11, fontweight="bold",
            color="#1f4e79")
    txt = (
        "① 軸力比 η = N/(b·D·Fc)\n"
        "  小さいほど良い（FA<=0.35, FB<=0.45, FC<=0.55）\n"
        "  軸力大→圧縮破壊しやすく脆性\n\n"
        "② せん断余裕（Qsu/Qmu）\n"
        "  せん断耐力に余裕→曲げ降伏先行→靭性◎\n\n"
        "③ 破壊形式（曲げ/せん断）\n"
        "  せん断スパン比 h0/D 大→曲げ支配→良い\n"
        "  短柱(h0/D 小)→せん断支配→FD に落ちる\n\n"
        "④ 帯筋比 pw（せん断補強・拘束）\n"
        "  多いほど靭性・じん性が上がる")
    ax.text(0.03, 0.83, txt, transform=ax.transAxes, fontproperties=jp,
            fontsize=9, color="#333", va="top")
    ax.axis("off")
    ax.set_title("(b) 判定パラメータと理由", fontproperties=jp, fontsize=10,
                 fontweight="bold")
    fig.suptitle("図 2  柱部材種別のパラメータ（軸力比・せん断余裕・帯筋）",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig2_column.png")


# 図3: 梁の部材種別パラメータ
def fig_beam():
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 5.0))
    ax = axes[0]
    ax.add_patch(mpatches.Rectangle((0, 0), 2, 3.5, fc="#bcd2ea", ec="k",
                 lw=1.2))
    for px in [0.5, 1.0, 1.5]:
        ax.add_patch(plt.Circle((px, 3.1), 0.13, fc="k"))
        ax.add_patch(plt.Circle((px, 0.4), 0.13, fc="#1f7a1f"))
    ax.add_patch(mpatches.Rectangle((0.2, 0.2), 1.6, 3.1, fc="none",
                 ec="#c00000", lw=1))
    ax.text(1, -0.7, "梁：上端筋・下端筋・あばら筋",
            fontproperties=jp, ha="center", fontsize=8.5, color="#444")
    ax.set_xlim(-0.6, 2.6); ax.set_ylim(-1.2, 4.0)
    ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("(a) 梁断面", fontproperties=jp, fontsize=10,
                 fontweight="bold")
    ax = axes[1]
    ax.text(0.5, 0.93, "梁の部材種別 判定パラメータ", transform=ax.transAxes,
            ha="center", fontproperties=jp, fontsize=11, fontweight="bold",
            color="#1f4e79")
    txt = (
        "① 引張鉄筋比 pt\n"
        "  過大（pb に近い）→ 脆性化 → FD 寄り\n"
        "  適量（釣合鉄筋比の半分以下）→ 靭性◎\n\n"
        "② せん断余裕（Qsu/Qmu）\n"
        "  せん断耐力に余裕→曲げ降伏先行→FA\n"
        "  せん断破壊が先→FD（脆性）\n\n"
        "③ あばら筋比 pw（せん断補強）\n"
        "  多いほどせん断破壊を防ぎ靭性↑\n\n"
        "④ 破壊形式\n"
        "  曲げ降伏型が望ましい（梁は柱より\n"
        "  先に降伏させるのが強柱弱梁）")
    ax.text(0.03, 0.83, txt, transform=ax.transAxes, fontproperties=jp,
            fontsize=9, color="#333", va="top")
    ax.axis("off")
    ax.set_title("(b) 判定パラメータと理由", fontproperties=jp, fontsize=10,
                 fontweight="bold")
    fig.suptitle("図 3  梁部材種別のパラメータ（引張鉄筋比・せん断余裕・あばら筋）",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig3_beam.png")


# 図4: 壁・部材群種別
def fig_wall_group():
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 5.0))
    # (a) 壁の種別（WA〜WD）
    ax = axes[0]
    ax.text(0.5, 0.93, "耐震壁の部材種別（WA〜WD）", transform=ax.transAxes,
            ha="center", fontproperties=jp, fontsize=11, fontweight="bold",
            color="#1f4e79")
    txt = (
        "① 破壊形式（曲げ/せん断）\n"
        "  曲げ降伏型→WA（靭性大）\n"
        "  せん断破壊型→WC・WD（脆性）\n\n"
        "② せん断余裕（Qsu/Qmu）\n"
        "  せん断耐力に余裕→曲げ降伏先行\n\n"
        "③ 壁筋比 ps・開口\n"
        "  壁筋が十分・開口が小さいほど良い\n\n"
        "④ 軸力・基礎の状態\n"
        "  脚部で曲げ降伏できる条件（基礎の\n"
        "  引抜き耐力）も影響")
    ax.text(0.03, 0.83, txt, transform=ax.transAxes, fontproperties=jp,
            fontsize=9, color="#333", va="top")
    ax.axis("off")
    ax.set_title("(a) 耐震壁の判定パラメータ", fontproperties=jp,
                 fontsize=10, fontweight="bold")
    # (b) 部材群種別（各階でまとめる）
    ax = axes[1]
    ax.text(0.5, 0.93, "部材群種別の設定（階ごと）", transform=ax.transAxes,
            ha="center", fontproperties=jp, fontsize=11, fontweight="bold",
            color="#1f4e79")
    # 柱群・梁群を集計してその階の Ds を決めるイメージ
    ax.text(0.5, 0.72,
            "各部材の種別（FA〜FD）を、その階の\n"
            "柱群・梁群・壁群ごとに集計して\n"
            "『部材群としての種別』を決める",
            transform=ax.transAxes, ha="center", fontproperties=jp,
            fontsize=9.5, color="#333")
    ax.text(0.5, 0.42,
            "・悪い種別（FD 等）の部材が多いと\n"
            "  その群・その階の Ds が大きくなる\n"
            "・水平力の負担割合（β）で重み付けして\n"
            "  各階の Ds を決定する",
            transform=ax.transAxes, ha="center", fontproperties=jp,
            fontsize=9.5, color="#333")
    ax.text(0.5, 0.12,
            "その階の Ds ＝ 部材群種別と\n負担割合から決まる（各階ごと）",
            transform=ax.transAxes, ha="center", fontproperties=jp,
            fontsize=9.5, color="#c00000", fontweight="bold")
    ax.axis("off")
    ax.set_title("(b) 部材群種別 → 階の Ds", fontproperties=jp,
                 fontsize=10, fontweight="bold")
    fig.suptitle("図 4  耐震壁部材種別 と 部材群種別の設定",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig4_wall_group.png")


figs = {"overview": fig_overview(), "column": fig_column(),
        "beam": fig_beam(), "wg": fig_wall_group()}
print("figures:", list(figs.keys()))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

wb = Workbook()
C_TITLE = "1F4E79"; C_HEAD = "2E75B6"; C_ANS = "E2EFDA"
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
f_title = Font(name="MS PGothic", size=15, bold=True, color="FFFFFF")
f_head = Font(name="MS PGothic", size=11, bold=True, color="FFFFFF")
f_body = Font(name="MS PGothic", size=10)
f_ans = Font(name="MS PGothic", size=10, color="375623")
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)


def setup(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False


def title_row(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text)
    c.font = f_title; c.fill = PatternFill("solid", fgColor=C_TITLE)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[row].height = 30


def head(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text)
    c.font = f_head; c.fill = PatternFill("solid", fgColor=C_HEAD)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[row].height = 22


def body(ws, row, text, span=8, ans=False, h=None):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text)
    c.font = f_ans if ans else f_body
    c.alignment = wrap
    if ans:
        c.fill = PatternFill("solid", fgColor=C_ANS)
    if h:
        ws.row_dimensions[row].height = h


def table(ws, start_row, headers, rows, col1=1):
    r = start_row
    for j, htxt in enumerate(headers):
        c = ws.cell(r, col1 + j, htxt)
        c.font = Font(name="MS PGothic", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C_HEAD)
        c.alignment = center; c.border = border
    for data in rows:
        r += 1
        for j, v in enumerate(data):
            c = ws.cell(r, col1 + j, v)
            c.font = f_body
            c.alignment = center if j > 0 else wrap
            c.border = border
            if r % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F2F7FC")
    return r


def put_img(ws, path, anchor, w=None):
    img = XLImage(path)
    if w:
        ratio = w / img.width
        img.width = w; img.height = int(img.height * ratio)
    ws.add_image(img, anchor)


ws = wb.active
ws.title = "目次"
setup(ws, [4, 22, 60, 14])
title_row(ws, 1, "部材種別 問題集（RC マンション設計担当・新人向け）", span=4)
body(ws, 2,
     "目標：二次設計における部材種別を理解すること。各部材の靭性（粘り強さ）を"
     "FA〜FD にランク分けし、これが Ds（構造特性係数）を決める。柱・梁・壁の"
     "判定パラメータと理由、部材群種別の設定を通す。崩壊形・保有水平耐力教材の続き。",
     span=4, h=44)
r = table(ws, 4, ["No.", "シート", "到達目標", "図"],
          [["1", "1 部材種別の全体像", "FA〜FD と Ds の関係を理解する", "FA〜FD"],
           ["2", "2 柱部材種別", "柱のパラメータと設定理由を理解する",
            "軸力比・せん断"],
           ["3", "3 梁部材種別", "梁のパラメータと設定理由を理解する",
            "引張鉄筋比"],
           ["4", "4 壁・部材群種別",
            "壁のパラメータ・部材群種別の設定を理解する", "壁・群"]])
body(ws, r + 2,
     "部材種別が良い（FA・FB）ほど靭性が高く、その階の Ds を小さくできる"
     "（＝経済的で安全）。No.7 崩壊形・保有水平耐力・保証設計と一体で学ぶこと。",
     span=4, h=32)

# 1
ws = wb.create_sheet("1 全体像")
setup(ws, [8, 18, 18, 16, 12, 12])
title_row(ws, 1, "1  部材種別の全体像（FA〜FD）")
head(ws, 3, "■ 図 1  FA〜FD と Ds")
put_img(ws, figs["overview"], "A4", w=800)
head(ws, 27, "■ 問題 1  部材種別とは")
body(ws, 28, "部材種別は各部材の【 ① 】（粘り強さ）のランクで、FA（靭性【 ② 】）から"
             "FD（靭性【 ③ 】＝脆性）まである。良い種別（FA・FB）ほど、その階の"
             "Ds が【 ④ 】くなる。", h=40)
head(ws, 30, "■ 問題 2  種別と Ds")
r = table(ws, 31, ["種別", "靭性", "Ds の目安（記入）"],
          [["FA", "大", ""],
           ["FB", "やや大", ""],
           ["FC", "中", ""],
           ["FD", "小（脆性）", ""]])
head(ws, r + 2, "■ 問題 3  なぜ種別を判定するか")
body(ws, r + 3, "部材種別を判定して Ds を決める意義を述べよ（粘る部材で構成された建物は"
                "Ds が小さく必要保有水平耐力が小さくて済む＝靭性を評価して"
                "合理的に設計する）。", h=40)

# 2
ws = wb.create_sheet("2 柱部材種別")
setup(ws, [8, 18, 18, 16, 12, 12])
title_row(ws, 1, "2  柱部材種別のパラメータ")
head(ws, 3, "■ 図 2  柱のパラメータ")
put_img(ws, figs["column"], "A4", w=820)
head(ws, 30, "■ 問題 1  軸力比")
body(ws, 31, "柱600角、N=2,000kN、Fc=24。軸力比 η=N/(b·D·Fc) を求めよ。"
             "FA（≤0.35）・FB（≤0.45）のどれに該当するか。"
             "軸力比が大きいと脆性化する理由も述べよ。", h=44)
head(ws, 33, "■ 問題 2  破壊形式（せん断スパン比）")
body(ws, 34, "せん断スパン比（h0/D 等）が小さい短柱は、せん断破壊しやすく"
             "部材種別が FD に落ちる。腰壁・垂れ壁による短柱化を防ぐ対策"
             "（No.耐震スリット教材参照）を述べよ。", h=40)
head(ws, 36, "■ 問題 3  パラメータと設定理由")
r = table(ws, 37, ["パラメータ", "良い方向（記入）", "理由（記入）"],
          [["軸力比 η", "", ""],
           ["せん断余裕 Qsu/Qmu", "", ""],
           ["せん断スパン比 h0/D", "", ""],
           ["帯筋比 pw", "", ""]])

# 3
ws = wb.create_sheet("3 梁部材種別")
setup(ws, [8, 18, 18, 16, 12, 12])
title_row(ws, 1, "3  梁部材種別のパラメータ")
head(ws, 3, "■ 図 3  梁のパラメータ")
put_img(ws, figs["beam"], "A4", w=820)
head(ws, 30, "■ 問題 1  引張鉄筋比")
body(ws, 31, "梁の引張鉄筋比 pt が過大（釣合鉄筋比 pb に近い）だと部材種別が"
             "悪くなる（脆性化）理由を述べよ（No.鉄筋比・釣合鉄筋比教材参照）。"
             "適量（pb の半分以下程度）が靭性に有利。", h=44)
head(ws, 33, "■ 問題 2  せん断余裕")
body(ws, 34, "梁のせん断余裕（Qsu/Qmu＝せん断耐力/曲げ終局時せん断力）を確保すると"
             "曲げ降伏が先行し靭性が高くなる。せん断破壊が先だとどうなるか"
             "（脆性・FD）を述べよ。", h=40)
head(ws, 36, "■ 問題 3  パラメータと設定理由")
r = table(ws, 37, ["パラメータ", "良い方向（記入）", "理由（記入）"],
          [["引張鉄筋比 pt", "", ""],
           ["せん断余裕 Qsu/Qmu", "", ""],
           ["あばら筋比 pw", "", ""],
           ["破壊形式", "", ""]])
head(ws, r + 2, "■ 問題 4  柱との関係")
body(ws, r + 3, "梁は柱より先に降伏させる（強柱弱梁）のが望ましい。"
                "梁を良い部材種別（FA・FB＝曲げ降伏型）にすることが"
                "全体崩壊形につながることを述べよ。", h=40)

# 4
ws = wb.create_sheet("4 壁・部材群種別")
setup(ws, [8, 18, 18, 16, 12, 12])
title_row(ws, 1, "4  耐震壁部材種別 と 部材群種別の設定")
head(ws, 3, "■ 図 4  壁のパラメータ・部材群種別")
put_img(ws, figs["wg"], "A4", w=820)
head(ws, 30, "■ 問題 1  耐震壁のパラメータ")
r = table(ws, 31, ["パラメータ", "良い方向（記入）", "理由（記入）"],
          [["破壊形式（曲げ/せん断）", "", ""],
           ["せん断余裕 Qsu/Qmu", "", ""],
           ["壁筋比 ps・開口", "", ""]])
body(ws, r + 2, "耐震壁も曲げ降伏型（WA）が良く、せん断破壊型（WD）は脆性で悪い。"
                "曲げ降伏に誘導する設計を 1 行で。", h=32)
head(ws, r + 4, "■ 問題 2  部材群種別とは")
body(ws, r + 5, "個々の部材種別（FA〜FD）を、その階の柱群・梁群・壁群ごとに"
                "集計したものが『部材群種別』。悪い種別の部材が多いと群の種別が"
                "悪くなり、その階の Ds が大きくなる仕組みを述べよ。", h=44)
head(ws, r + 7, "■ 問題 3  階ごとの Ds")
body(ws, r + 8, "Ds は建物全体で 1 つでなく『各階ごと』に決まる。"
                "各階の部材群種別と、水平力の負担割合（β）で重み付けして"
                "その階の Ds を決定することを述べよ。", h=40)
head(ws, r + 10, "■ 問題 4  設計へのフィードバック")
body(ws, r + 11, "ある階の Ds が大きい（部材種別が悪い）と分かったら、"
                 "設計者は何をするか（悪い部材＝FD の柱の軸力を下げる・"
                 "せん断補強を増やす・短柱化を解消する等で種別を改善し Ds を下げる）。"
                 "設計は部材種別の改善を繰り返して収束させることを述べよ。", h=44)

ws = wb.create_sheet("解答")
setup(ws, [4, 22, 60, 14])
title_row(ws, 1, "解答・解説", span=4)
row = 2


def ah(t):
    global row
    head(ws, row, t, span=4); row += 1


def an(t, h=None):
    global row
    body(ws, row, t, span=4, ans=True, h=h); row += 1


ah("1  全体像")
an("問1：①靭性 ②大 ③小 ④小さ。FA（靭性大）ほどその階の Ds が小さい。", h=32)
an("問2：FA≒0.30、FB≒0.35、FC≒0.40、FD≒0.45〜0.5（構造種別・規準により差）。",
   h=32)
an("問3：部材種別を判定して Ds を決めることで、粘り強い（靭性の高い）部材で"
   "構成された建物は Ds を小さくでき、必要保有水平耐力が小さくて済む。"
   "靭性を適切に評価して合理的・経済的に設計するため。", h=44)

ah("2  柱部材種別")
an("問1：η=N/(b·D·Fc)=2,000×10³/(600×600×24)=2,000,000/8,640,000=0.231。"
   "0.231≤0.35 なので FA 相当。軸力比が大きいと圧縮側が先に潰れ（圧壊）"
   "変形能力が小さく脆性化するため、軸力比は小さいほど良い。", h=44)
an("問2：短柱（h0/D 小）はせん断力が集中しせん断破壊（脆性）しやすく FD に落ちる。"
   "腰壁・垂れ壁による短柱化を耐震スリットで縁切りするか、せん断補強を密にして"
   "防ぐ（No.耐震スリット教材）。", h=44)
an("問3：軸力比 η→小さいほど良い（圧壊・脆性を防ぐ）。"
   "せん断余裕→大きいほど良い（曲げ降伏先行）。"
   "せん断スパン比→大きいほど良い（曲げ支配、短柱回避）。"
   "帯筋比 pw→大きいほど良い（せん断補強・コア拘束で靭性↑）。", h=44)

ah("3  梁部材種別")
an("問1：pt が過大（pb に近い）だと、鉄筋降伏前にコンクリートが圧壊する"
   "過大鉄筋・脆性破壊になり靭性が低下（FD 寄り）。適量（pb の半分以下）なら"
   "鉄筋が十分降伏してから壊れ靭性が高い（No.鉄筋比教材）。", h=44)
an("問2：せん断余裕を確保すると曲げ降伏がせん断破壊に先行し、粘り強く壊れる（FA）。"
   "せん断破壊が先だと予兆なく急激に耐力を失う脆性破壊で FD になる。", h=40)
an("問3：引張鉄筋比 pt→適量（過大にしない）。せん断余裕→大きく。"
   "あばら筋比 pw→大きく（せん断補強）。破壊形式→曲げ降伏型が良い。", h=40)
an("問4：梁を良い種別（曲げ降伏型 FA・FB）にすると、梁が柱より先に降伏（強柱弱梁）"
   "してヒンジが梁に分散し全体崩壊形になる。梁の種別が全体崩壊形の要。", h=40)

ah("4  壁・部材群種別")
an("問1：破壊形式→曲げ降伏型が良い（WA）。せん断余裕→大きく（曲げ降伏先行）。"
   "壁筋比 ps→十分に、開口→小さく。壁のせん断耐力を曲げ降伏時せん断力より"
   "大きくして曲げ降伏に誘導する（保証設計）。", h=44)
an("問2：個々の部材種別を階の柱群・梁群・壁群ごとに集計したのが部材群種別。"
   "FD 等の悪い部材が多いと群種別が悪くなり、その階の Ds が大きくなる"
   "（悪い部材に足を引っ張られる）。", h=44)
an("問3：Ds は各階ごとに決まる。各階の部材群種別と、その部材が負担する"
   "水平力の割合（β：柱・壁の負担割合）で重み付けして、その階の Ds を決定する。",
   h=40)
an("問4：Ds が大きい階は、悪い部材（FD）を改善する：柱の軸力を下げる（断面増）、"
   "せん断補強を増やす、短柱化を解消（スリット）等で種別を FA・FB に上げ Ds を下げる。"
   "部材種別の改善→再計算を繰り返して設計を収束させる。", h=44)

XLSX = os.path.join(OUT, "部材種別問題集.xlsx")
wb.save(XLSX)
print("saved:", XLSX)
