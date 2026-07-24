# -*- coding: utf-8 -*-
"""配筋調整 問題集（図つき）Excel。出力: docs/rebar_adjust/配筋調整問題集.xlsx
1 配筋調整の目的 / 2 配筋変更の影響 / 3 効率的な調整方法 / 4 納まりを意識した調整
"""
import os
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib import font_manager as fm
FONT = "/usr/share/fonts/truetype/fonts-japanese-gothic.ttf"
jp = fm.FontProperties(fname=FONT); fm.fontManager.addfont(FONT)
plt.rcParams["font.family"] = jp.get_name(); plt.rcParams["axes.unicode_minus"] = False
OUT = "docs/rebar_adjust"; FIG = os.path.join(OUT, "figures"); os.makedirs(FIG, exist_ok=True)


def save(fig, n):
    p = os.path.join(FIG, n); fig.savefig(p, dpi=130, bbox_inches="tight"); plt.close(fig); return p


def fig_purpose():
    fig, ax = plt.subplots(figsize=(12, 5.0))
    ax.text(0.5, 0.93, "配筋調整の目的（どの検討のために変えるか）", transform=ax.transAxes,
            ha="center", fontproperties=jp, fontsize=13, fontweight="bold", color="#1f4e79")
    boxes = [(0.04, "一次設計（許容応力度）", "#cfe0f0",
              "常時・中地震で応力度を\n許容内に。曲げ・せん断・\nたわみを満たす配筋"),
             (0.36, "二次設計（保有水平耐力）", "#cfead4",
              "靭性確保・部材種別 FA/FB。\nせん断余裕・帯筋・pt 調整で\nDs を下げる"),
             (0.68, "接合部・納まり", "#fde2c4",
              "定着・重ね継手・あき・\nかぶり。主筋段数・干渉を\n解消して施工可能に")]
    for x, t, c, d in boxes:
        ax.add_patch(mpatches.FancyBboxPatch((x, 0.3), 0.28, 0.44,
                     boxstyle="round,pad=0.02", transform=ax.transAxes, fc=c, ec="k", lw=1.2))
        ax.text(x + 0.14, 0.66, t, transform=ax.transAxes, ha="center",
                fontproperties=jp, fontsize=10, fontweight="bold")
        ax.text(x + 0.14, 0.45, d, transform=ax.transAxes, ha="center", va="center",
                fontproperties=jp, fontsize=8.5, color="#333")
    ax.text(0.5, 0.12, "配筋調整は『何のために変えるか（目的）』を明確にしてから行う\n"
                       "目的が違えば変える方向も違う（応力満足／靭性確保／納まり）",
            transform=ax.transAxes, ha="center", fontproperties=jp, fontsize=9.5, color="#c00000")
    ax.axis("off"); ax.set_title("図 1  配筋調整の目的", fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig1_purpose.png")


def fig_impact():
    fig, ax = plt.subplots(figsize=(12, 5.2))
    steps = [("梁の主筋を増やす", 0.82), ("梁の曲げ耐力↑", 0.66),
             ("柱梁耐力比↓（強柱弱梁が崩れる）", 0.50),
             ("崩壊形が悪化（柱降伏の恐れ）", 0.34), ("Ds が上がり必要保有耐力↑", 0.18)]
    for i, (t, y) in enumerate(steps):
        col = "#f8d0d0" if i >= 2 else "#cfe0f0"
        ax.add_patch(mpatches.FancyBboxPatch((0.2, y), 0.6, 0.11,
                     boxstyle="round,pad=0.01", transform=ax.transAxes, fc=col, ec="k", lw=1))
        ax.text(0.5, y + 0.055, t, transform=ax.transAxes, ha="center", va="center",
                fontproperties=jp, fontsize=9.5)
        if i < len(steps) - 1:
            ax.annotate("", xy=(0.5, y - 0.02), xytext=(0.5, y),
                        xycoords="axes fraction", arrowprops=dict(arrowstyle="-|>", color="#c00000", lw=1.8))
    ax.text(0.5, 0.03, "1 か所の配筋変更が耐力比・崩壊形・Ds まで波及する（影響を追う）",
            transform=ax.transAxes, ha="center", fontproperties=jp, fontsize=9.5, color="#c00000")
    ax.axis("off"); ax.set_title("図 2  配筋変更の影響が波及する事象",
                                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig2_impact.png")


def fig_efficient():
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 4.8))
    ax = axes[0]
    members = ["梁A", "梁B", "梁C", "柱D", "柱E"]
    ratio = [0.98, 0.72, 0.55, 0.91, 0.60]
    y = np.arange(len(members))
    ax.barh(y, ratio, color=["#f8c0c0" if r > 0.9 else "#a8dadc" for r in ratio], ec="k", lw=0.6)
    ax.axvline(1.0, color="red", ls="--", lw=1.5)
    ax.text(1.01, 4.2, "許容 1.0", fontproperties=jp, fontsize=8, color="red")
    ax.set_yticks(y); ax.set_yticklabels(members, fontproperties=jp)
    ax.invert_yaxis(); ax.set_xlabel("応力比（余裕度）", fontproperties=jp); ax.set_xlim(0, 1.2)
    ax.set_title("(a) 応力比の高い部材から調整", fontproperties=jp, fontsize=10, fontweight="bold")
    ax.text(0.5, 4.7, "応力比 0.9 超（梁A・柱D）を優先", fontproperties=jp, fontsize=8,
            color="#c00000", ha="center")
    ax = axes[1]
    ax.text(0.5, 0.92, "効率的な配筋調整のコツ", transform=ax.transAxes, ha="center",
            fontproperties=jp, fontsize=11, fontweight="bold", color="#1f4e79")
    ax.text(0.05, 0.78, "① 応力比の高い（1.0 に近い）部材から着手\n\n"
            "② 配筋を標準化・グループ化\n  （同じ符号で径・本数をそろえる）\n\n"
            "③ 径を上げる vs 本数を増やす を使い分け\n  （あき・段数・定着への影響を見る）\n\n"
            "④ 変更の影響（耐力比・種別・Ds）を確認\n  してから次へ（手戻りを減らす）\n\n"
            "⑤ 断面変更（せい・幅）も選択肢に入れる",
            transform=ax.transAxes, fontproperties=jp, fontsize=9.5, color="#333", va="top")
    ax.axis("off"); ax.set_title("(b) 調整のコツ", fontproperties=jp, fontsize=10, fontweight="bold")
    fig.suptitle("図 3  効率的な配筋調整方法", fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig3_efficient.png")


def fig_detail():
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 5.0))
    ax = axes[0]
    ax.add_patch(mpatches.Rectangle((0, 0), 4, 6, fc="#d9d9d9", ec="k", lw=1.2))
    for px in np.linspace(0.6, 3.4, 5):
        ax.add_patch(plt.Circle((px, 5.4), 0.22, fc="k"))
    for px in np.linspace(0.9, 3.1, 4):
        ax.add_patch(plt.Circle((px, 4.7), 0.22, fc="#c00000"))
    ax.annotate("あき（鉄筋間隔）\n径×1.5 以上・粗骨材+α", xy=(1.2, 5.4),
                xytext=(4.4, 5.2), fontproperties=jp, fontsize=8, color="#c00000",
                arrowprops=dict(arrowstyle="->", color="#c00000"))
    ax.annotate("2 段配筋（1段で入らない時）", xy=(2, 4.7), xytext=(4.4, 4.0),
                fontproperties=jp, fontsize=8, color="#7a3b3b",
                arrowprops=dict(arrowstyle="->", color="#7a3b3b"))
    ax.annotate("かぶり厚", xy=(0.15, 3), xytext=(-1.6, 3), fontproperties=jp, fontsize=8,
                color="#1f4e79", va="center", arrowprops=dict(arrowstyle="->", color="#1f4e79"))
    ax.set_xlim(-2, 7.5); ax.set_ylim(-0.5, 6.5); ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("(a) 梁断面の納まり（あき・段数・かぶり）", fontproperties=jp,
                 fontsize=10, fontweight="bold")
    ax = axes[1]
    ax.add_patch(mpatches.Rectangle((0, 0), 2.6, 6, fc="#d9d9d9", ec="k", lw=1.2))
    ax.add_patch(mpatches.Rectangle((2.6, 2.4), 4, 1.4, fc="#bcd2ea", ec="k", lw=1))
    ax.plot([6.5, 0.7], [3.5, 3.5], color="#c00000", lw=2.5)
    ax.plot([0.7, 0.7], [3.5, 1.3], color="#c00000", lw=2.5)
    ax.plot([6.5, 1.1], [2.7, 2.7], color="#1f7a1f", lw=2.5)
    ax.plot([1.1, 1.1], [2.7, 4.6], color="#1f7a1f", lw=2.5)
    ax.text(3, -0.7, "接合部：上下梁主筋の定着・干渉を\nイメージしながら本数・段数を調整",
            fontproperties=jp, ha="center", fontsize=8.5, color="#444")
    ax.set_xlim(-0.5, 7); ax.set_ylim(-1.3, 6.5); ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("(b) 接合部の納まり（定着・干渉）", fontproperties=jp,
                 fontsize=10, fontweight="bold")
    fig.suptitle("図 4  納まりをイメージした配筋調整", fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig4_detail.png")


figs = {"purpose": fig_purpose(), "impact": fig_impact(),
        "efficient": fig_efficient(), "detail": fig_detail()}
print("figs:", list(figs.keys()))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
wb = Workbook()
C_TITLE = "1F4E79"; C_HEAD = "2E75B6"; C_ANS = "E2EFDA"
thin = Side(style="thin", color="BFBFBF"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
f_title = Font(name="MS PGothic", size=15, bold=True, color="FFFFFF")
f_head = Font(name="MS PGothic", size=11, bold=True, color="FFFFFF")
f_body = Font(name="MS PGothic", size=10); f_ans = Font(name="MS PGothic", size=10, color="375623")
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)


def setup(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False


def title_row(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_title; c.fill = PatternFill("solid", fgColor=C_TITLE)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1); ws.row_dimensions[row].height = 30


def head(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_head; c.fill = PatternFill("solid", fgColor=C_HEAD)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1); ws.row_dimensions[row].height = 22


def body(ws, row, text, span=8, ans=False, h=None):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_ans if ans else f_body; c.alignment = wrap
    if ans:
        c.fill = PatternFill("solid", fgColor=C_ANS)
    if h:
        ws.row_dimensions[row].height = h


def table(ws, start_row, headers, rows, col1=1):
    r = start_row
    for j, htxt in enumerate(headers):
        c = ws.cell(r, col1 + j, htxt); c.font = Font(name="MS PGothic", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C_HEAD); c.alignment = center; c.border = border
    for data in rows:
        r += 1
        for j, v in enumerate(data):
            c = ws.cell(r, col1 + j, v); c.font = f_body
            c.alignment = center if j > 0 else wrap; c.border = border
            if r % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F2F7FC")
    return r


def put_img(ws, path, anchor, w=None):
    img = XLImage(path)
    if w:
        ratio = w / img.width; img.width = w; img.height = int(img.height * ratio)
    ws.add_image(img, anchor)


ws = wb.active; ws.title = "目次"; setup(ws, [4, 22, 60, 14])
title_row(ws, 1, "配筋調整 問題集（RC マンション設計担当・新人向け）", span=4)
body(ws, 2, "目標：目的・影響を理解した上で配筋調整ができること。一次設計・二次設計・"
            "接合部で目的が違う。配筋変更の波及、効率的な進め方、納まりを通す。", span=4, h=32)
r = table(ws, 4, ["No.", "シート", "到達目標", "図"],
          [["1", "1 配筋調整の目的", "調整の目的を明確にできる", "一次/二次/接合部"],
           ["2", "2 配筋変更の影響", "変更が及ぼす事象を理解する", "波及"],
           ["3", "3 効率的な調整方法", "効率的な調整方法を理解する", "応力比"],
           ["4", "4 納まりを意識した調整", "納まりをイメージし調整できる", "あき・接合部"]])
body(ws, r + 2, "配筋調整は No.鉄筋比・定着・継手・部材種別・保証設計の各教材の知識を"
                "総動員する実務スキル。", span=4, h=32)

ws = wb.create_sheet("1 調整の目的"); setup(ws, [8, 18, 18, 16, 12, 12])
title_row(ws, 1, "1  配筋調整の目的（一次設計・二次設計・接合部）")
head(ws, 3, "■ 図 1  配筋調整の目的"); put_img(ws, figs["purpose"], "A4", w=800)
head(ws, 27, "■ 問題 1  目的の対比")
r = table(ws, 28, ["検討", "目的（記入）", "調整する内容（記入）"],
          [["一次設計", "", ""], ["二次設計", "", ""], ["接合部・納まり", "", ""]])
body(ws, r + 2, "選択肢：応力度を許容内に／靭性確保（部材種別・Ds）／定着・あき・干渉の解消。", h=32)
head(ws, r + 4, "■ 問題 2  目的を明確にする意義")
body(ws, r + 5, "配筋調整を『目的を決めずに』行うと何が起こるか（あちこち直して手戻り・"
                "別の検討で NG）。まず目的（何の検討のために何を満たすか）を"
                "明確にする重要性を述べよ。", h=44)
head(ws, r + 7, "■ 問題 3  目的別の方向")
body(ws, r + 8, "同じ『梁の配筋を変える』でも、(1) 曲げ応力が NG（一次）なら主筋を増やす、"
                "(2) 靭性不足（二次・pt 過大）なら主筋を減らす／断面を大きくする、"
                "と逆方向になることがある。目的で方向が変わる例を述べよ。", h=44)

ws = wb.create_sheet("2 変更の影響"); setup(ws, [8, 18, 18, 16, 12, 12])
title_row(ws, 1, "2  配筋変更により影響が及ぶ事象")
head(ws, 3, "■ 図 2  影響の波及"); put_img(ws, figs["impact"], "A4", w=760)
head(ws, 28, "■ 問題 1  波及の理解")
body(ws, 29, "梁の主筋を増やす（曲げ耐力↑）と、柱梁耐力比・崩壊形・Ds へどう波及するか"
             "図の流れで説明せよ。1 か所の変更が全体に影響することを理解せよ。", h=44)
head(ws, 31, "■ 問題 2  影響する事象の列挙")
r = table(ws, 32, ["変更", "影響する事象（記入）"],
          [["梁主筋を増やす", ""],
           ["柱主筋を増やす", ""],
           ["あばら筋・帯筋を増やす", ""],
           ["主筋径を上げる（本数減）", ""]])
head(ws, r + 2, "■ 問題 3  接合部への影響")
body(ws, r + 3, "梁主筋を増やす・太くすると、接合部（定着・あき・干渉）にどう影響するか。"
                "断面内に納まらない・接合部で定着長が取れない等の問題を述べよ。", h=40)
head(ws, r + 5, "■ 問題 4  影響を追う習慣")
body(ws, r + 6, "配筋変更後に確認すべき事項を 4 つ挙げよ"
                "（応力度／柱梁耐力比・崩壊形／部材種別・Ds／納まり・定着）。", h=40)

ws = wb.create_sheet("3 効率的な調整"); setup(ws, [8, 18, 18, 16, 12, 12])
title_row(ws, 1, "3  効率的な配筋調整方法")
head(ws, 3, "■ 図 3  効率的な調整"); put_img(ws, figs["efficient"], "A4", w=820)
head(ws, 30, "■ 問題 1  優先順位")
body(ws, 31, "多数の部材を調整するとき、どの部材から着手すべきか"
             "（応力比が 1.0 に近い＝余裕のない部材から）。理由も述べよ。", h=32)
head(ws, 33, "■ 問題 2  径 vs 本数")
body(ws, 34, "鉄筋量を増やすとき『径を上げる』と『本数を増やす』の使い分けを述べよ"
             "（本数増→あき・段数に影響、径増→あき確保しやすいが定着長増）。", h=40)
head(ws, 36, "■ 問題 3  標準化")
body(ws, 37, "配筋を標準化・グループ化（同符号で径・本数をそろえる）する利点を述べよ"
             "（図面・積算・施工の効率、ミス減）。No.符号割教材とも関連。", h=40)
head(ws, 39, "■ 問題 4  手戻りを減らす")
body(ws, 40, "配筋調整の手戻りを減らす進め方を述べよ（変更の影響を確認してから次へ、"
             "断面変更も選択肢、一貫計算の再計算サイクルを回す）。", h=40)

ws = wb.create_sheet("4 納まりを意識"); setup(ws, [8, 18, 18, 16, 12, 12])
title_row(ws, 1, "4  納まりをイメージした配筋調整")
head(ws, 3, "■ 図 4  納まり（あき・段数・接合部）"); put_img(ws, figs["detail"], "A4", w=820)
head(ws, 30, "■ 問題 1  あき・かぶり・段数")
body(ws, 31, "鉄筋のあき（間隔）の最小値の考え方（径×1.5・粗骨材径+α 等）を述べよ。"
             "1 段で納まらない場合の 2 段配筋、かぶり厚の確保にも触れよ。", h=40)
head(ws, 33, "■ 問題 2  接合部の干渉")
body(ws, 34, "柱梁接合部で、直交する梁主筋・柱主筋が干渉する。"
             "本数・段数・径をどう調整して納めるか（No.柱梁接合部・定着教材参照）。", h=40)
head(ws, 36, "■ 問題 3  納まりを先読み")
body(ws, 37, "配筋を決める段階で『施工できるか（納まるか）』をイメージする重要性を述べよ。"
             "計算上 OK でも納まらない配筋は成立しない（過密配筋・定着不足）。", h=40)
head(ws, 39, "■ 問題 4  実務まとめ")
body(ws, 40, "配筋調整で確認する項目を 4 つ挙げよ（①目的（一次/二次/接合部）②応力・"
             "耐力・種別への影響 ③あき・かぶり・段数の納まり ④定着・継手・干渉）。", h=40)

ws = wb.create_sheet("解答"); setup(ws, [4, 22, 60, 14])
title_row(ws, 1, "解答・解説", span=4); row = 2


def ah(t):
    global row; head(ws, row, t, span=4); row += 1


def an(t, h=None):
    global row; body(ws, row, t, span=4, ans=True, h=h); row += 1


ah("1  調整の目的")
an("問1：一次設計＝応力度を許容内に（曲げ・せん断・たわみ）／主筋・あばら筋の量。"
   "二次設計＝靭性確保・部材種別 FA/FB・Ds 低減／pt・せん断余裕・帯筋。"
   "接合部・納まり＝定着・あき・干渉の解消／本数・段数・径。", h=44)
an("問2：目的を決めずに直すと、ある検討で OK にしても別の検討（靭性・納まり）で NG になり"
   "手戻りが増える。まず『何の検討のために何を満たすか』を明確にしてから調整する。", h=44)
an("問3：(1)曲げ応力 NG（一次）→主筋を増やす。(2)pt 過大で靭性不足（二次）→"
   "主筋を減らす or 断面を大きくする。目的が違えば変える方向も逆になり得る。", h=44)

ah("2  変更の影響")
an("問1：梁主筋増→梁曲げ耐力↑→柱梁耐力比↓（強柱弱梁が崩れる）→崩壊形悪化"
   "（柱降伏の恐れ）→Ds↑→必要保有耐力↑。1 か所の変更が全体に波及する。", h=44)
an("問2：梁主筋増→曲げ耐力・柱梁耐力比・接合部。柱主筋増→柱耐力・軸力比・接合部。"
   "帯筋増→せん断耐力・部材種別（靭性）。径増（本数減）→あき・定着長・接合部。", h=44)
an("問3：梁主筋を増やす・太くすると、断面内のあきが不足して 2 段配筋が必要になったり、"
   "接合部で定着長・のみ込みが取れなくなる。計算と納まりは常にセットで確認する。", h=44)
an("問4：①応力度 ②柱梁耐力比・崩壊形 ③部材種別・Ds ④納まり・定着・継手。", h=32)

ah("3  効率的な調整")
an("問1：応力比が 1.0 に近い（余裕のない）部材から着手する。余裕のある部材を先に"
   "触っても NG は解消せず、ボトルネックの部材から直すのが効率的。", h=40)
an("問2：本数増＝あき・段数に影響（過密化）。径増＝あきは確保しやすいが定着長・重ね継手長が"
   "伸びる。納まりと定着のバランスで使い分ける。", h=40)
an("問3：標準化・グループ化すると図面・積算・施工が効率化しミスも減る"
   "（同符号は同配筋、数量集計が容易）。No.符号割の集約と同じ考え方。", h=40)
an("問4：変更のたびに影響（応力・耐力・種別・納まり）を確認してから次へ進む。"
   "配筋で無理なら断面変更（せい・幅）も選択肢に。一貫計算の再計算を回して収束させる。",
   h=44)

ah("4  納まりを意識")
an("問1：あきは鉄筋径×1.5・粗骨材最大径+α・25mm 等の最大値以上を確保。"
   "1 段で入らなければ 2 段配筋（段間隔も確保）。土に接する面はかぶりを大きく。", h=44)
an("問2：直交梁主筋・柱主筋が同じ高さで交差し干渉する。主筋を細径多本数にする、"
   "段をずらす、本数を減らして径を上げる等で納める（No.柱梁接合部・定着教材）。", h=44)
an("問3：計算上 OK でも、あき不足・過密で施工できない配筋は成立しない。"
   "コンクリートが充填できず豆板、定着不足で耐力不足になる。納まりを先読みして決める。", h=44)
an("問4：①目的（一次/二次/接合部）②応力・耐力・種別への影響 ③あき・かぶり・段数 "
   "④定着・継手・干渉。計算と納まりを両輪で確認するのが配筋調整の要。", h=44)

XLSX = os.path.join(OUT, "配筋調整問題集.xlsx"); wb.save(XLSX); print("saved:", XLSX)
