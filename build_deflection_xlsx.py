# -*- coding: utf-8 -*-
"""たわみ・層間変形角 問題集（図つき）Excel 生成スクリプト。
出力: docs/deflection/たわみ層間変形角問題集.xlsx
1 たわみ・層間変形角のイメージ / 2 RC造でのたわみの影響
3 変形増大率K(クリープ・乾燥収縮) / 4 たわみ抑制の対応 / 5 制限値
RC造マンションの設計担当を想定。数値は build 時に検算済み。
"""
import os
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib import font_manager as fm

FONT_PATH = "/usr/share/fonts/truetype/fonts-japanese-gothic.ttf"
jp = fm.FontProperties(fname=FONT_PATH)
fm.fontManager.addfont(FONT_PATH)
plt.rcParams["font.family"] = jp.get_name()
plt.rcParams["axes.unicode_minus"] = False

OUT = "docs/deflection"
FIG = os.path.join(OUT, "figures")
os.makedirs(FIG, exist_ok=True)

C_BLUE = "#2a78d6"
C_PINK = "#d55181"


def save(fig, name):
    p = os.path.join(FIG, name)
    fig.savefig(p, dpi=130, bbox_inches="tight")
    plt.close(fig)
    return p


# ===========================================================================
# 図 1: たわみ と 層間変形角
# ===========================================================================
def fig_image():
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 5.0))
    # (a) たわみ（梁の鉛直下がり）
    ax = axes[0]
    L = 6
    # 支持
    ax.add_patch(mpatches.Rectangle((-0.3, 1.5), 0.3, 0.6, fc="#d9d9d9",
                                     ec="k"))
    ax.add_patch(mpatches.Rectangle((L, 1.5), 0.3, 0.6, fc="#d9d9d9", ec="k"))
    # 変形前（点線）
    ax.plot([0, L], [2.0, 2.0], color="#999", lw=1.2, ls="--")
    # 変形後（たわみ曲線）
    x = np.linspace(0, L, 50)
    d = 0.7 * np.sin(np.pi * x / L)
    ax.plot(x, 2.0 - d, color="#c00000", lw=2.5)
    # 鉛直荷重
    for xx in np.linspace(0.5, L - 0.5, 6):
        ax.annotate("", xy=(xx, 2.0 - 0.7 * np.sin(np.pi * xx / L)),
                    xytext=(xx, 2.7),
                    arrowprops=dict(arrowstyle="-|>", color="#1f7a1f", lw=1))
    ax.annotate("たわみ δ（鉛直の下がり）", xy=(L / 2, 2.0 - 0.7),
                xytext=(L / 2 + 0.3, 0.7), fontproperties=jp, fontsize=9.5,
                color="#c00000",
                arrowprops=dict(arrowstyle="->", color="#c00000"))
    ax.annotate("", xy=(L / 2, 2.0), xytext=(L / 2, 2.0 - 0.7),
                arrowprops=dict(arrowstyle="<|-|>", color="#c00000", lw=1.2))
    ax.text(L / 2 - 0.15, 1.6, "δ", fontproperties=jp, fontsize=11,
            color="#c00000", ha="right")
    ax.text(L / 2, 3.0, "鉛直荷重（自重・積載）", fontproperties=jp,
            ha="center", fontsize=8.5, color="#1f7a1f")
    ax.text(L / 2, -0.1,
            "たわみ＝梁・スラブが鉛直方向に下がる量\n"
            "制限：δ/スパン <= 1/250（長期）",
            fontproperties=jp, ha="center", fontsize=9, color="#444")
    ax.set_xlim(-0.8, L + 0.9); ax.set_ylim(-0.7, 3.4)
    ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("(a) たわみ（鉛直方向）", fontproperties=jp, fontsize=11,
                 fontweight="bold")

    # (b) 層間変形角（水平のずれ）
    ax = axes[1]
    h = 3.0
    w = 3.0
    # 変形前（点線の四角）
    ax.plot([0, w, w, 0, 0], [0, 0, h, h, 0], color="#999", lw=1.2, ls="--")
    # 変形後（平行四辺形＝せん断変形）
    sh = 0.9
    ax.plot([0, w, w + sh, sh, 0], [0, 0, h, h, 0], color="#c00000", lw=2.5)
    # 柱
    ax.plot([0, sh], [0, h], color="#1f4e79", lw=2)
    ax.plot([w, w + sh], [0, h], color="#1f4e79", lw=2)
    # 地震力
    ax.annotate("", xy=(w + sh + 0.5, h), xytext=(w + sh - 0.5, h),
                arrowprops=dict(arrowstyle="-|>", color="#c00000", lw=2.5))
    ax.text(w / 2 + sh, h + 0.35, "地震力（水平）", fontproperties=jp,
            ha="center", fontsize=8.5, color="#c00000")
    # 層間変位 δ
    ax.annotate("", xy=(sh, h + 0.05), xytext=(0, h + 0.05),
                arrowprops=dict(arrowstyle="<|-|>", color="#7a3b3b", lw=1.2))
    ax.text(sh / 2, h + 0.25, "層間変位 δ", fontproperties=jp, ha="center",
            fontsize=8.5, color="#7a3b3b")
    ax.annotate("", xy=(-0.4, 0), xytext=(-0.4, h),
                arrowprops=dict(arrowstyle="<|-|>", color="#1f4e79", lw=1))
    ax.text(-0.7, h / 2, "階高 h", fontproperties=jp, rotation=90,
            va="center", fontsize=8.5, color="#1f4e79")
    ax.text(w / 2 + sh / 2, -0.7,
            "層間変形角 R = δ / h（水平のずれ／階高）\n"
            "制限：R <= 1/200",
            fontproperties=jp, ha="center", fontsize=9, color="#444")
    ax.set_xlim(-1.0, w + sh + 1.0); ax.set_ylim(-1.3, h + 0.9)
    ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("(b) 層間変形角（水平方向）", fontproperties=jp, fontsize=11,
                 fontweight="bold")
    fig.suptitle("図 1  たわみ（鉛直）と 層間変形角（水平）のイメージ",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig1_image.png")


# ===========================================================================
# 図 2: RC造でのたわみの影響
# ===========================================================================
def fig_effect():
    fig, ax = plt.subplots(figsize=(12, 5.2))
    items = [
        ("建具・サッシの不具合", "扉が閉まらない・開かない、\nサッシのガラス割れ", "#f8d0d0"),
        ("仕上げのひび割れ", "天井・壁のクロス・\nタイルのひび、剥落", "#fde2c4"),
        ("床の水勾配・排水", "バルコニー・水回りで\n水が溜まる・逆勾配", "#cfe0f0"),
        ("振動・使用感", "歩行振動・上階の\n足音が響く", "#e8e0f0"),
        ("間仕切り壁の損傷", "たわみで乾式間仕切りが\n押されてひび", "#cfead4"),
    ]
    x = 0.3
    for label, desc, color in items:
        ax.add_patch(mpatches.FancyBboxPatch((x, 2.2), 2.1, 2.2,
                     boxstyle="round,pad=0.05", fc=color, ec="k", lw=1))
        ax.text(x + 1.05, 3.9, label, fontproperties=jp, ha="center",
                fontsize=9.5, fontweight="bold")
        ax.text(x + 1.05, 2.9, desc, fontproperties=jp, ha="center",
                va="center", fontsize=8.2, color="#333")
        x += 2.35
    ax.text(6, 1.2,
            "RC は自重が大きく・クリープ乾燥収縮でたわみが進行するため、\n"
            "『構造的には安全でも使用上・意匠上の不具合』が問題になりやすい\n"
            "（たわみは強度より使用性・居住性の問題）",
            fontproperties=jp, ha="center", fontsize=10, color="#1f4e79")
    ax.set_xlim(0, 12); ax.set_ylim(0.5, 4.8)
    ax.axis("off")
    ax.set_title("図 2  RC 造におけるたわみの影響",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig2_effect.png")


# ===========================================================================
# 図 3: 変形増大率 K（時間経過）
# ===========================================================================
def fig_K():
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 4.8))
    # (a) 時間-たわみ曲線
    ax = axes[0]
    t = np.linspace(0, 100, 200)
    de = 1.0
    # クリープ・乾燥収縮でたわみが増大（漸近）
    d = de * (1 + 15 * (1 - np.exp(-t / 25)))
    ax.plot(t, d, color="#c00000", lw=2.5)
    ax.axhline(de, color=C_BLUE, lw=1.5, ls="--")
    ax.text(60, de + 0.6, "弾性たわみ δe（載荷直後）", fontproperties=jp,
            fontsize=9, color=C_BLUE)
    ax.axhline(16 * de, color="#7a3b3b", lw=1, ls=":")
    ax.text(50, 16 * de - 1.3, "長期たわみ δL = K·δe（K≒16）",
            fontproperties=jp, fontsize=9, color="#7a3b3b")
    ax.annotate("クリープ＋乾燥収縮で\n時間とともに増大",
                xy=(30, d[60]), xytext=(45, 6), fontproperties=jp,
                fontsize=9, color="#c00000",
                arrowprops=dict(arrowstyle="->", color="#c00000"))
    ax.set_xlabel("時間（年）", fontproperties=jp)
    ax.set_ylabel("たわみ（弾性たわみの倍率）", fontproperties=jp)
    ax.set_xlim(0, 100); ax.set_ylim(0, 18)
    ax.grid(alpha=0.25)
    ax.set_title("(a) 時間経過とたわみの増大", fontproperties=jp, fontsize=10,
                 fontweight="bold")

    # (b) 式まとめ
    ax = axes[1]
    ax.text(0.5, 0.9, "変形増大率 K", transform=ax.transAxes, ha="center",
            fontproperties=jp, fontsize=13, fontweight="bold", color="#1f4e79")
    ax.text(0.5, 0.72, "長期たわみ δL = K × 弾性たわみ δe",
            transform=ax.transAxes, ha="center", fontproperties=jp,
            fontsize=13, color="#c00000", fontweight="bold")
    ax.text(0.5, 0.5,
            "K：変形増大率（RC 規準）\n"
            "・クリープ（持続荷重で徐々に変形）\n"
            "・乾燥収縮（コンクリートが乾いて縮む）\n"
            "を考慮した割増し係数",
            transform=ax.transAxes, ha="center", fontproperties=jp,
            fontsize=10, color="#333")
    ax.text(0.5, 0.2,
            "一般に K = 8 〜 16（長期・持続荷重）\n"
            "圧縮鉄筋（複筋）を増やすと K は小さくできる\n"
            "（クリープ・収縮を圧縮筋が拘束）",
            transform=ax.transAxes, ha="center", fontproperties=jp,
            fontsize=9.5, color="#1f7a1f")
    ax.axis("off")
    ax.set_title("(b) 変形増大率 K の式と特性", fontproperties=jp,
                 fontsize=10, fontweight="bold")
    fig.suptitle("図 3  たわみ要因 ── 変形増大率 K（クリープ・乾燥収縮）",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig3_K.png")


# ===========================================================================
# 図 4: たわみ抑制の対応
# ===========================================================================
def fig_measures():
    fig, axes = plt.subplots(1, 3, figsize=(14.5, 4.2))
    # (a) 梁せいを大きく（D↑ → I↑）
    ax = axes[0]
    ax.add_patch(mpatches.Rectangle((0, 0), 1.0, 1.4, fc="#bcd2ea", ec="k",
                 lw=1.2))
    ax.add_patch(mpatches.Rectangle((2.0, 0), 1.0, 2.2, fc="#9ec6e8", ec="k",
                 lw=1.5))
    ax.annotate("", xy=(2.0, 2.4), xytext=(3.0, 2.4),
                arrowprops=dict(arrowstyle="<|-|>", color="#c00000"))
    ax.text(2.5, 2.6, "せい D↑", fontproperties=jp, ha="center", fontsize=8,
            color="#c00000")
    ax.text(1.5, -0.7, "① 梁せいを大きく\nI∝D³ でたわみ激減\n（δ∝1/D³）",
            fontproperties=jp, ha="center", fontsize=8.5, color="#444")
    ax.set_xlim(-0.4, 3.4); ax.set_ylim(-1.5, 2.9)
    ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("(a) 断面を大きく", fontproperties=jp, fontsize=10,
                 fontweight="bold")

    # (b) むくり（キャンバー）
    ax = axes[1]
    x = np.linspace(0, 3, 50)
    ax.plot(x, 0.4 * np.sin(np.pi * x / 3), color="#1f7a1f", lw=2.5)
    ax.plot([0, 3], [0, 0], color="#999", lw=1, ls="--")
    ax.text(1.5, 0.55, "施工時に上げておく（むくり）", fontproperties=jp,
            ha="center", fontsize=8, color="#1f7a1f")
    ax.annotate("たわむと水平に", xy=(1.5, 0.4), xytext=(1.5, -0.8),
                fontproperties=jp, fontsize=8, ha="center", color="#c00000",
                arrowprops=dict(arrowstyle="->", color="#c00000"))
    ax.text(1.5, -1.4, "② むくり（キャンバー）\n予めたわみ分\n上げて打設",
            fontproperties=jp, ha="center", fontsize=8.5, color="#444")
    ax.set_xlim(-0.3, 3.3); ax.set_ylim(-2.1, 1.0)
    ax.axis("off")
    ax.set_title("(b) むくりをつける", fontproperties=jp, fontsize=10,
                 fontweight="bold")

    # (c) その他対策リスト
    ax = axes[2]
    ax.text(0.05, 0.9, "③ その他の対応", transform=ax.transAxes,
            fontproperties=jp, fontsize=10, fontweight="bold", color="#1f4e79")
    txt = ("・スパンを短くする（δ∝L^4）\n"
           "・圧縮鉄筋（複筋）を増やし K を下げる\n"
           "・コンクリート強度 Fc を上げ Ec↑\n"
           "・プレストレス（PC）でたわみ相殺\n"
           "・小梁を入れてスラブスパンを分割\n"
           "・早期の過大載荷を避ける（養生）")
    ax.text(0.05, 0.72, txt, transform=ax.transAxes, fontproperties=jp,
            fontsize=9.5, color="#333", va="top")
    ax.axis("off")
    ax.set_title("(c) その他の対応", fontproperties=jp, fontsize=10,
                 fontweight="bold")
    fig.suptitle("図 4  たわみを抑制するための対応",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig4_measures.png")


# ===========================================================================
# 図 5: 制限値
# ===========================================================================
def fig_limits():
    fig, ax = plt.subplots(figsize=(11.5, 5.0))
    ax.text(0.5, 0.93, "たわみ・層間変形角の制限値", transform=ax.transAxes,
            ha="center", fontproperties=jp, fontsize=13, fontweight="bold",
            color="#1f4e79")
    tbl = [
        ("項目", "制限値", "備考"),
        ("梁・スラブのたわみ（長期）", "δ / スパン <= 1/250",
         "変形増大率 K を考慮した長期たわみで"),
        ("片持ち梁・スラブ", "δ / 出寸法 <= 1/250",
         "先端たわみ／出寸法。振動も確認"),
        ("層間変形角（地震時・一次設計）", "R = δ/h <= 1/200",
         "全層で確認"),
        ("層間変形角（緩和）", "1/120 まで可",
         "帳壁・非構造部材に著しい損傷なきこと"),
    ]
    yy = 0.78
    for i, (a, b, c) in enumerate(tbl):
        col = "#dbe8f5" if i == 0 else ("#f7f7f7" if i % 2 else "white")
        for xx, w, txt, al in [(0.03, 0.36, a, "left"),
                               (0.39, 0.26, b, "center"),
                               (0.65, 0.33, c, "left")]:
            ax.add_patch(mpatches.Rectangle((xx, yy), w, 0.13,
                         transform=ax.transAxes, fc=col, ec="#bbb", lw=0.6))
            ax.text(xx + (0.02 if al == "left" else w / 2), yy + 0.065, txt,
                    transform=ax.transAxes,
                    ha=al, va="center", fontproperties=jp,
                    fontsize=8.8, fontweight="bold" if i == 0 else "normal",
                    color="#c00000" if (i in (3, 4) and xx == 0.39) else "#333")
        yy -= 0.145
    ax.text(0.5, 0.08,
            "たわみ＝使用性（δ/L<=1/250）、層間変形角＝地震時の変形制限（R<=1/200）。\n"
            "※具体値は RC 規準・建築基準法施行令で確認すること",
            transform=ax.transAxes, ha="center", fontproperties=jp,
            fontsize=9, color="#7a3b3b")
    ax.axis("off")
    return save(fig, "fig5_limits.png")


figs = {
    "image": fig_image(),
    "effect": fig_effect(),
    "K": fig_K(),
    "measures": fig_measures(),
    "limits": fig_limits(),
}
print("figures:", list(figs.keys()))

# ===========================================================================
# Excel 構築
# ===========================================================================
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

wb = Workbook()
C_TITLE = "1F4E79"; C_HEAD = "2E75B6"; C_ANS = "E2EFDA"
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
f_title = Font(name="MS PGothic", size=15, bold=True, color="FFFFFF")
f_head = Font(name="MS PGothic", size=11, bold=True, color="FFFFFF")
f_body = Font(name="MS PGothic", size=10)
f_ans = Font(name="MS PGothic", size=10, color="375623")
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)


def setup(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False


def title_row(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text)
    c.font = f_title; c.fill = PatternFill("solid", fgColor=C_TITLE)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[row].height = 30


def head(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text)
    c.font = f_head; c.fill = PatternFill("solid", fgColor=C_HEAD)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[row].height = 22


def body(ws, row, text, span=8, ans=False, h=None):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text)
    c.font = f_ans if ans else f_body
    c.alignment = wrap
    if ans:
        c.fill = PatternFill("solid", fgColor=C_ANS)
    if h:
        ws.row_dimensions[row].height = h


def table(ws, start_row, headers, rows, col1=1):
    r = start_row
    for j, htxt in enumerate(headers):
        c = ws.cell(r, col1 + j, htxt)
        c.font = Font(name="MS PGothic", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C_HEAD)
        c.alignment = center; c.border = border
    for data in rows:
        r += 1
        for j, v in enumerate(data):
            c = ws.cell(r, col1 + j, v)
            c.font = f_body
            c.alignment = center if j > 0 else wrap
            c.border = border
            if r % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F2F7FC")
    return r


def put_img(ws, path, anchor, w=None):
    img = XLImage(path)
    if w:
        ratio = w / img.width
        img.width = w; img.height = int(img.height * ratio)
    ws.add_image(img, anchor)


# ---- 目次 ----
ws = wb.active
ws.title = "目次"
setup(ws, [4, 22, 60, 14])
title_row(ws, 1, "たわみ・層間変形角 問題集（RC マンション設計担当・新人向け）",
          span=4)
body(ws, 2,
     "目標：たわみ・層間変形角を理解しイメージできること。たわみ（鉛直）と"
     "層間変形角（水平）の違い、RC 造でのたわみの影響、変形増大率 K、抑制対応、"
     "制限値を通す。『たわみは強度でなく使用性の問題』が勘所。",
     span=4, h=44)
r = table(ws, 4,
          ["No.", "シート", "到達目標", "図"],
          [["1", "1 たわみ・層間変形角のイメージ",
            "たわみ・層間変形角を理解しイメージできる", "鉛直・水平"],
           ["2", "2 RC造でのたわみの影響",
            "RC 造におけるたわみの影響を理解している", "影響 5 例"],
           ["3", "3 変形増大率 K",
            "たわみ要因（変形増大率 K）を理解している", "時間-たわみ"],
           ["4", "4 たわみ抑制の対応",
            "たわみを抑制するための対応を理解している", "せい増・むくり"],
           ["5", "5 制限値",
            "たわみ・層間変形角の制限値を理解している", "制限値表"]])
body(ws, r + 2,
     "共通例：RC 大梁 400×700、L=6.0m、長期 w=30 kN/m。弾性たわみ δe≒2.2mm、"
     "変形増大率 K=16 で長期たわみ δL≒35mm。階高 h=3.5m。"
     "制限：たわみ δ/L≤1/250、層間変形角 R≤1/200。数値は本教材作成時に検算済み。"
     "実務は RC 規準・建築基準法施行令で確認すること。",
     span=4, h=58)

# ---- 1 イメージ ----
ws = wb.create_sheet("1 イメージ")
setup(ws, [8, 18, 18, 16, 12, 12])
title_row(ws, 1, "1  たわみ・層間変形角のイメージ")
head(ws, 3, "■ 図 1  たわみ（鉛直）と 層間変形角（水平）")
put_img(ws, figs["image"], "A4", w=820)
head(ws, 30, "■ 問題 1  定義の対比")
r = table(ws, 31,
          ["項目", "たわみ（記入）", "層間変形角（記入）"],
          [["方向", "", ""],
           ["原因となる荷重", "", ""],
           ["定義（式）", "", ""],
           ["主な制限値", "", ""]])
head(ws, r + 2, "■ 問題 2  層間変形角の計算")
body(ws, r + 3, "階高 h=3,500mm、地震時の層間変位 δ=10mm のとき、"
                "層間変形角 R=δ/h を分数（1/○）で求めよ。制限 1/200 と比較せよ。",
     h=32)
head(ws, r + 5, "■ 問題 3  なぜ両方を確認するか")
body(ws, r + 6, "(1) たわみ（鉛直）は主に何の荷重で生じ、いつ問題になるか"
                "（常時・長期／使用性）。"
                "(2) 層間変形角（水平）は主に何で生じ、いつ問題になるか"
                "（地震時／非構造部材の損傷・脱落）。", h=44)
head(ws, r + 8, "■ 問題 4  イメージの言語化")
body(ws, r + 9, "『たわみ』と『層間変形角』を、それぞれ一言でイメージ説明せよ"
                "（たわみ＝梁が下がる、層間変形角＝階が横にずれる）。"
                "マンションで実際に体感する場面（床の振動／地震時の建具の"
                "きしみ）も挙げよ。", h=44)

# ---- 2 RC造でのたわみの影響 ----
ws = wb.create_sheet("2 たわみの影響")
setup(ws, [8, 18, 18, 16, 12, 12])
title_row(ws, 1, "2  RC 造におけるたわみの影響")
head(ws, 3, "■ 図 2  たわみの影響（5 例）")
put_img(ws, figs["effect"], "A4", w=800)
head(ws, 28, "■ 問題 1  影響の分類")
body(ws, 29, "過大なたわみが引き起こす不具合を 5 つ挙げよ"
             "（建具・サッシ／仕上げのひび／床の水勾配／振動／間仕切り損傷）。"
             "これらが『強度の問題ではなく使用性・意匠の問題』である点を述べよ。",
     h=44)
head(ws, 31, "■ 問題 2  なぜ RC はたわみやすいか")
body(ws, 32, "RC 造がたわみで問題になりやすい理由を 2 つ挙げよ"
             "（自重が大きい／クリープ・乾燥収縮で長期に進行する）。"
             "鉄骨造・木造との比較にも触れよ。", h=40)
head(ws, 34, "■ 問題 3  マンション特有の影響")
body(ws, 35, "RC マンションでたわみが特に問題になる部位・場面を挙げよ"
             "（大スパンリビングの梁／バルコニーの片持ち／水回りの床勾配／"
             "上下階の音・振動）。", h=40)
head(ws, 37, "■ 問題 4  たわみと安全性の関係")
body(ws, 38, "『たわみが大きい＝すぐ壊れる』ではない。たわみは主に使用性の問題だが、"
             "放置するとどんな二次的問題につながるか（ひび割れからの中性化・"
             "鉄筋腐食、仕上げ剥落による危険）を述べよ。", h=44)

# ---- 3 変形増大率K ----
ws = wb.create_sheet("3 変形増大率K")
setup(ws, [8, 18, 16, 14, 14, 12, 12])
title_row(ws, 1, "3  たわみ要因 ── 変形増大率 K")
head(ws, 3, "■ 図 3  時間経過とたわみ・変形増大率 K")
put_img(ws, figs["K"], "A4", w=820)
head(ws, 28, "■ 問題 1  変形増大率とは（穴埋め）")
body(ws, 29, "コンクリートは持続荷重で徐々に変形する【 ① 】と、乾いて縮む"
             "【 ② 】により、載荷直後の弾性たわみ δe よりも長期のたわみが大きくなる。"
             "この割増しを変形増大率 K といい、長期たわみ δL =【 ③ 】で表す。"
             "一般に K =【 ④ 】程度。", h=44)
head(ws, 31, "■ 問題 2  弾性たわみの計算")
body(ws, 32, "大梁 400×700（I=1.14×10^10）、L=6,000、長期等分布 w=30 kN/m。"
             "単純梁の弾性たわみ δe=5wL⁴/(384EcI)（Ec=2.05×10⁴）を求めよ。", h=32)
r = table(ws, 35,
          ["項目", "式", "値（記入）"],
          [["弾性たわみ δe", "5×30×6000⁴/(384×2.05×10⁴×1.14×10^10)", ""],
           ["δe/L", "—", ""]])
head(ws, r + 2, "■ 問題 3  長期たわみ")
body(ws, r + 3, "変形増大率 K=16 のときの長期たわみ δL=K·δe を求めよ。"
                "δL/L を分数で表し、制限 1/250 と比較して OK/NG を判定せよ。", h=32)
head(ws, r + 5, "■ 問題 4  K を小さくするには")
body(ws, r + 6, "変形増大率 K を小さくする方法を述べよ"
                "（圧縮鉄筋（複筋）を増やす → クリープ・収縮を圧縮筋が拘束）。"
                "複筋梁が長期たわみに有利な理由を 1 行で。", h=40)

# ---- 4 たわみ抑制の対応 ----
ws = wb.create_sheet("4 たわみ抑制の対応")
setup(ws, [8, 18, 18, 16, 12, 12])
title_row(ws, 1, "4  たわみを抑制するための対応")
head(ws, 3, "■ 図 4  たわみ抑制の対応")
put_img(ws, figs["measures"], "A4", w=820)
head(ws, 28, "■ 問題 1  最も効く対策")
body(ws, 29, "たわみは δ∝L⁴/(Ec·I)。次の対策の効果を大きい順に整理せよ。", h=20)
r = table(ws, 31,
          ["対策", "効くパラメータ", "効果（記入）"],
          [["梁せい D を大きく", "I ∝ D³", ""],
           ["スパン L を短く", "δ ∝ L⁴", ""],
           ["コンクリート強度 Fc↑", "Ec", ""],
           ["圧縮鉄筋（複筋）", "K（変形増大率）", ""]])
body(ws, r + 2, "問題 3-2 の梁（δL=35mm・NG）で、梁せいを 700→800 にすると"
                "たわみは何倍になるか（δ∝1/D³）。OK になるか確認せよ。", h=32)
head(ws, r + 4, "■ 問題 2  むくり（キャンバー）")
body(ws, r + 5, "むくり（キャンバー）とは何か。予めどの向きに上げておくか。"
                "たわみの『発生を防ぐ』のではなく『見かけ上打ち消す』対策である"
                "点を述べよ。", h=40)
head(ws, r + 7, "■ 問題 3  設計段階 vs 施工段階")
body(ws, r + 8, "たわみ対策を『設計段階』（断面・スパン・複筋・PC）と"
                "『施工段階』（むくり・支保工・養生・過大載荷を避ける）に"
                "分類せよ。", h=40)
head(ws, r + 10, "■ 問題 4  スパン割りの工夫")
body(ws, r + 11, "大スパンの部屋で、小梁を入れてスラブを分割するとたわみが"
                 "どうなるか（スラブの実質スパンが短くなり δ 激減）。"
                 "意匠（梁型が出る）とのトレードオフにも触れよ。", h=40)

# ---- 5 制限値 ----
ws = wb.create_sheet("5 制限値")
setup(ws, [8, 18, 18, 16, 12, 12])
title_row(ws, 1, "5  たわみ・層間変形角の制限値")
head(ws, 3, "■ 図 5  制限値の一覧")
put_img(ws, figs["limits"], "A4", w=800)
head(ws, 28, "■ 問題 1  制限値の暗記")
r = table(ws, 29,
          ["項目", "制限値（記入）"],
          [["梁・スラブのたわみ（長期）", ""],
           ["片持ち梁・スラブのたわみ", ""],
           ["層間変形角（地震時・一次設計）", ""],
           ["層間変形角（緩和できる場合）", ""]])
head(ws, r + 2, "■ 問題 2  たわみ制限の判定")
body(ws, r + 3, "問題 3 の梁：長期たわみ δL=35mm、L=6,000mm。"
                "δL/L を求め、制限 1/250 と比較して判定せよ。"
                "NG の場合、問題 4 の対策でどう改善するか。", h=40)
head(ws, r + 5, "■ 問題 3  層間変形角の判定")
body(ws, r + 6, "階高 h=3,500mm。層間変位 δ=10mm と δ=20mm の 2 ケースで"
                "層間変形角 R を求め、制限 1/200（=17.5mm）と比較せよ。", h=32)
r = table(ws, r + 9,
          ["ケース", "δ", "R=δ/h（記入）", "判定（記入）"],
          [["ケース1", "10mm", "", ""],
           ["ケース2", "20mm", "", ""]])
head(ws, r + 2, "■ 問題 4  層間変形角の緩和")
body(ws, r + 3, "層間変形角は原則 1/200 だが、条件を満たせば 1/120 まで緩和できる。"
                "その条件（帳壁・非構造部材が変形に追従し著しい損傷が生じないこと）"
                "を述べよ。なぜ非構造部材が絡むのか（変形で建具・外装が壊れる）。",
     h=44)
head(ws, r + 5, "■ 問題 5  実務まとめ")
body(ws, r + 6, "たわみ・層間変形角のチェックで確認する項目を 4 つ挙げよ"
                "（長期たわみに変形増大率 K を見込んだか／δ/L≤1/250／"
                "層間変形角 R≤1/200／片持ち・大スパン部の重点確認）。", h=44)

# ---- 解答 ----
ws = wb.create_sheet("解答")
setup(ws, [4, 22, 60, 14])
title_row(ws, 1, "解答・解説", span=4)
row = 2


def ah(t):
    global row
    head(ws, row, t, span=4); row += 1


def an(t, h=None):
    global row
    body(ws, row, t, span=4, ans=True, h=h); row += 1


ah("1  イメージ")
an("問1：たわみ＝鉛直方向／鉛直荷重（自重・積載）／δ（下がり量）、"
   "制限 δ/スパン≤1/250。層間変形角＝水平方向／地震力（水平）／R=δ/h、"
   "制限 R≤1/200。", h=44)
an("問2：R=δ/h=10/3500=1/350。制限 1/200（=17.5mm）より小さく OK。", h=32)
an("問3：(1)たわみは鉛直荷重（常時・長期）で生じ、クリープ・収縮で進行、"
   "使用性（建具・仕上げ・振動）で問題になる。"
   "(2)層間変形角は地震力（水平）で生じ、地震時に非構造部材（帳壁・建具・外装）の"
   "損傷・脱落として問題になる。", h=58)
an("問4：たわみ＝『梁・スラブが下がる』、層間変形角＝『階が横にずれる（平行四辺形に）』。"
   "体感例：大スパンリビングの床の歩行振動（たわみ系）、地震時の建具のきしみ・"
   "隙間（層間変形系）。", h=44)

ah("2  たわみの影響")
an("問1：①建具・サッシが閉まらない/ガラス割れ ②天井・壁のクロス・タイルのひび "
   "③バルコニー・水回りの逆勾配・水溜まり ④歩行振動・音 ⑤乾式間仕切りの損傷。"
   "いずれも強度（安全性）でなく使用性・意匠・居住性の問題。", h=58)
an("問2：①RC は自重が大きく曲げ・たわみが出やすい ②クリープ（持続荷重の変形）・"
   "乾燥収縮でたわみが長期に進行する。鉄骨造は軽く弾性的、木造も軽量で、"
   "RC のような長期増大は相対的に小さい。", h=44)
an("問3：大スパンリビングの梁（たわみ大）、バルコニー片持ち（先端たわみ・振動）、"
   "水回りの床（逆勾配で排水不良）、上下階の音・振動（薄い大スパン床）。", h=44)
an("問4：たわみ自体は使用性の問題だが、放置するとひび割れ→中性化・塩害で"
   "鉄筋腐食、仕上げのひび・剥落による落下（第三者被害）につながり得る。"
   "早期の是正・適切な制限が重要。", h=44)

ah("3  変形増大率K")
an("問1：①クリープ ②乾燥収縮 ③K×δe ④8〜16。", h=32)
an("問2：δe=5×30×6000⁴/(384×2.05×10⁴×1.14×10^10)。"
   "分子=5×30×1.296×10^15=1.944×10^17。分母=384×2.05×10⁴×1.14×10^10=8.975×10^16。"
   "δe=2.17mm。δe/L=2.17/6000=1/2,770。（弾性だけなら十分小さい）", h=44)
an("問3：δL=16×2.17=34.6mm。δL/L=34.6/6000=1/173。"
   "制限 1/250 に対し 1/173 は大きい（変形が制限超）→ NG。"
   "弾性たわみは小さくても、変形増大率で長期に NG になり得る＝K の確認が重要。",
   h=44)
an("問4：圧縮鉄筋（複筋）を増やすと、圧縮側コンクリートのクリープ・乾燥収縮ひずみを"
   "鉄筋が拘束し、長期たわみの増大を抑える → K が小さくなる。"
   "複筋梁は長期たわみに有利。", h=44)

ah("4  たわみ抑制の対応")
an("問1：梁せい D↑＝I∝D³ で最も効く。スパン L↓＝δ∝L⁴ で非常に効く"
   "（ただしスパンは計画で決まる）。Fc↑＝Ec 増だが効果は限定的（Ec∝√Fc 程度）。"
   "複筋＝K を下げ長期たわみに効く。"
   "梁せい 700→800：δ×(700/800)³=0.67。δL=34.6×0.67=23.2mm→δL/L=1/259<1/250 OK。",
   h=58)
an("問2：むくり＝施工時に梁・スラブを予め上向きに上げておくこと。"
   "たわむと水平になるよう、たわみ予測分だけ上げる。"
   "たわみの発生を防ぐのでなく『見かけ上打ち消す』対策（応力は変わらない）。",
   h=44)
an("問3：設計段階＝断面（せい）増、スパン短縮、複筋、PC 導入、Fc アップ。"
   "施工段階＝むくり、支保工の適切な存置、養生（早期乾燥防止）、"
   "強度発現前の過大載荷を避ける。", h=44)
an("問4：小梁を入れるとスラブの実質スパンが短くなり、δ∝L⁴ で"
   "たわみが激減する（例：スパン半分→たわみ 1/16）。"
   "ただし小梁の梁型が天井に出る意匠上のデメリットがあり、"
   "梁レス（フラット）志向とトレードオフ。", h=44)

ah("5  制限値")
an("問1：梁・スラブ長期＝δ/スパン≤1/250。片持ち＝δ/出寸法≤1/250。"
   "層間変形角（地震時）＝R=δ/h≤1/200。緩和＝非構造部材に著しい損傷なければ"
   "1/120 まで。", h=44)
an("問2：δL/L=35/6000≒1/171＜1/250 は満たさない → NG。"
   "対策：梁せい 700→800 で δL≒23mm→1/259 で OK。または複筋で K を下げる、"
   "スパン短縮、PC 等。", h=44)
an("問3：ケース1：R=10/3500=1/350＜1/200 OK。"
   "ケース2：R=20/3500=1/175＞1/200 NG（17.5mm 超）。"
   "剛性を上げる（壁・ブレース追加、断面増）必要。", h=44)
an("問4：緩和条件＝地震時の層間変形に対し、帳壁・建具・外装などの非構造部材が"
   "追従でき、著しい損傷（脱落・ガラス割れ等）が生じないこと。"
   "層間変形は非構造部材を直接変形させるため、部材の追従性能とセットで"
   "制限が決まる。", h=44)
an("問5：①長期たわみに変形増大率 K を見込んだか ②δ/L≤1/250 を満たすか "
   "③層間変形角 R≤1/200（緩和時 1/120）④片持ち・大スパン部を重点確認。", h=44)

XLSX = os.path.join(OUT, "たわみ層間変形角問題集.xlsx")
wb.save(XLSX)
print("saved:", XLSX)
