# -*- coding: utf-8 -*-
"""鉄筋の継手 問題集（図つき）Excel 生成スクリプト。
出力: docs/rebar_splice/鉄筋継手問題集.xlsx
No.7-1 継手の応力伝達機構(重ね・溶接・圧接・機械式) / 7-2 L1・L1h の説明・算出
7-3 継手長一覧表の暗記 / 7-4 計算仮定条件 / 7-5 各部材の継手可能範囲
RC造マンションの設計担当を想定。
"""
import os
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib import font_manager as fm

FONT_PATH = "/usr/share/fonts/truetype/fonts-japanese-gothic.ttf"
jp = fm.FontProperties(fname=FONT_PATH)
fm.fontManager.addfont(FONT_PATH)
plt.rcParams["font.family"] = jp.get_name()
plt.rcParams["axes.unicode_minus"] = False

OUT = "docs/rebar_splice"
FIG = os.path.join(OUT, "figures")
os.makedirs(FIG, exist_ok=True)


def save(fig, name):
    p = os.path.join(FIG, name)
    fig.savefig(p, dpi=130, bbox_inches="tight")
    plt.close(fig)
    return p


def draw_bar(ax, x1, x2, y, lw=5, color="#888", ribs=True):
    """異形鉄筋を描く（水平）。"""
    ax.plot([x1, x2], [y, y], color=color, lw=lw, solid_capstyle="butt",
            zorder=3)
    if ribs:
        for x in np.arange(min(x1, x2) + 0.15, max(x1, x2) - 0.05, 0.3):
            ax.plot([x, x], [y - 0.09, y + 0.09], color="#555", lw=1.2,
                    zorder=4)


# ===========================================================================
# 図 7-1: 4種類の継手の応力伝達機構
# ===========================================================================
def fig_7_1():
    fig, axes = plt.subplots(2, 2, figsize=(13.5, 7.6))

    # --- (a) 重ね継手 ---
    ax = axes[0, 0]
    ax.add_patch(mpatches.Rectangle((0, 0), 10, 2.6, fc="#e8e8e8",
                                     ec="k", lw=0.8))
    draw_bar(ax, 0.3, 6.5, 1.7)
    draw_bar(ax, 3.5, 9.7, 0.9)
    # 引張力
    ax.annotate("", xy=(0.3, 1.7), xytext=(-1.3, 1.7),
                arrowprops=dict(arrowstyle="<|-", color="#c00000", lw=2.5))
    ax.text(-1.4, 1.7, "T", fontproperties=jp, ha="right", va="center",
            fontsize=11, color="#c00000")
    ax.annotate("", xy=(9.7, 0.9), xytext=(11.3, 0.9),
                arrowprops=dict(arrowstyle="<|-", color="#c00000", lw=2.5))
    ax.text(11.4, 0.9, "T", fontproperties=jp, ha="left", va="center",
            fontsize=11, color="#c00000")
    # 重ね長さ
    ax.annotate("", xy=(3.5, 2.15), xytext=(6.5, 2.15),
                arrowprops=dict(arrowstyle="<|-|>", color="#1f4e79", lw=1.3))
    ax.text(5.0, 2.32, "重ね継手長さ L1", fontproperties=jp, ha="center",
            fontsize=9, color="#1f4e79", fontweight="bold")
    # 伝達経路
    for x in np.arange(3.9, 6.3, 0.6):
        ax.annotate("", xy=(x + 0.25, 1.05), xytext=(x, 1.55),
                    arrowprops=dict(arrowstyle="-|>", color="#1f7a1f",
                                    lw=1.2))
    ax.text(5.0, -0.55, "応力は 鉄筋 → 付着 → コンクリート → 付着 → 鉄筋 と間接伝達",
            fontproperties=jp, ha="center", fontsize=9, color="#1f7a1f")
    ax.set_xlim(-2.2, 12.2)
    ax.set_ylim(-1.1, 3.0)
    ax.set_aspect("equal")
    ax.axis("off")
    ax.set_title("(a) 重ね継手（間接伝達）", fontproperties=jp,
                 fontsize=11, fontweight="bold")

    # --- (b) ガス圧接継手 ---
    ax = axes[0, 1]
    draw_bar(ax, 0.3, 4.7, 1.3)
    draw_bar(ax, 5.3, 9.7, 1.3)
    # ふくらみ
    ax.add_patch(mpatches.Ellipse((5.0, 1.3), 1.4, 0.75, fc="#aaa",
                                   ec="k", lw=1.0, zorder=5))
    ax.annotate("", xy=(0.3, 1.3), xytext=(-1.3, 1.3),
                arrowprops=dict(arrowstyle="<|-", color="#c00000", lw=2.5))
    ax.annotate("", xy=(9.7, 1.3), xytext=(11.3, 1.3),
                arrowprops=dict(arrowstyle="<|-", color="#c00000", lw=2.5))
    ax.annotate("ふくらみ（こぶ）\n径 1.4db 以上\n長さ 1.1db 以上",
                xy=(5.0, 1.68), xytext=(6.4, 2.6),
                fontproperties=jp, fontsize=8.5, color="#1f4e79",
                arrowprops=dict(arrowstyle="->", color="#1f4e79"))
    ax.text(5.0, -0.1, "端面同士を加熱・加圧して一体化 ── 応力は断面で直接伝達",
            fontproperties=jp, ha="center", fontsize=9, color="#1f7a1f")
    ax.set_xlim(-2.2, 12.2)
    ax.set_ylim(-0.7, 3.6)
    ax.set_aspect("equal")
    ax.axis("off")
    ax.set_title("(b) ガス圧接継手（直接伝達）", fontproperties=jp,
                 fontsize=11, fontweight="bold")

    # --- (c) 溶接継手 ---
    ax = axes[1, 0]
    draw_bar(ax, 0.3, 4.85, 1.3)
    draw_bar(ax, 5.15, 9.7, 1.3)
    # 開先溶接マーク（V形）
    ax.fill([4.85, 5.15, 5.3, 4.7], [1.05, 1.05, 1.62, 1.62],
            color="#e0a800", zorder=6)
    ax.annotate("", xy=(0.3, 1.3), xytext=(-1.3, 1.3),
                arrowprops=dict(arrowstyle="<|-", color="#c00000", lw=2.5))
    ax.annotate("", xy=(9.7, 1.3), xytext=(11.3, 1.3),
                arrowprops=dict(arrowstyle="<|-", color="#c00000", lw=2.5))
    ax.annotate("突合せ溶接\n（開先＋溶着金属）", xy=(5.0, 1.62),
                xytext=(6.4, 2.5), fontproperties=jp, fontsize=8.5,
                color="#1f4e79",
                arrowprops=dict(arrowstyle="->", color="#1f4e79"))
    ax.text(5.0, -0.1, "溶着金属を介して断面で直接伝達（品質は溶接施工管理に依存）",
            fontproperties=jp, ha="center", fontsize=9, color="#1f7a1f")
    ax.set_xlim(-2.2, 12.2)
    ax.set_ylim(-0.7, 3.4)
    ax.set_aspect("equal")
    ax.axis("off")
    ax.set_title("(c) 溶接継手（直接伝達）", fontproperties=jp,
                 fontsize=11, fontweight="bold")

    # --- (d) 機械式継手 ---
    ax = axes[1, 1]
    draw_bar(ax, 0.3, 4.6, 1.3, ribs=True)
    draw_bar(ax, 5.4, 9.7, 1.3, ribs=True)
    # カプラー（スリーブ）
    ax.add_patch(mpatches.Rectangle((3.4, 0.95), 3.2, 0.7, fc="#c9a227",
                                     ec="k", lw=1.2, zorder=5))
    # グラウト点
    for x in np.arange(3.6, 6.5, 0.35):
        ax.plot(x, 1.3, ".", color="#7a5a00", ms=3, zorder=6)
    ax.annotate("", xy=(0.3, 1.3), xytext=(-1.3, 1.3),
                arrowprops=dict(arrowstyle="<|-", color="#c00000", lw=2.5))
    ax.annotate("", xy=(9.7, 1.3), xytext=(11.3, 1.3),
                arrowprops=dict(arrowstyle="<|-", color="#c00000", lw=2.5))
    ax.annotate("カプラー（スリーブ）\nねじ節・モルタル充填 等",
                xy=(5.0, 1.65), xytext=(6.2, 2.55),
                fontproperties=jp, fontsize=8.5, color="#1f4e79",
                arrowprops=dict(arrowstyle="->", color="#1f4e79"))
    ax.text(5.0, -0.1, "鉄筋 → カプラー → 鉄筋 と機械的に伝達（PCa 接合部で多用）",
            fontproperties=jp, ha="center", fontsize=9, color="#1f7a1f")
    ax.set_xlim(-2.2, 12.2)
    ax.set_ylim(-0.7, 3.4)
    ax.set_aspect("equal")
    ax.axis("off")
    ax.set_title("(d) 機械式継手（機械的伝達）", fontproperties=jp,
                 fontsize=11, fontweight="bold")

    fig.suptitle("図 7-1  鉄筋継手 4 種類の応力伝達機構",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig7-1_types.png")


# ===========================================================================
# 図 7-2: L1・L1h の定義
# ===========================================================================
def fig_7_2():
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 4.6))

    # --- (a) L1: 直線重ね継手 ---
    ax = axes[0]
    ax.add_patch(mpatches.Rectangle((0, 0), 10, 2.6, fc="#e8e8e8",
                                     ec="k", lw=0.8))
    draw_bar(ax, 0.3, 6.5, 1.7)
    draw_bar(ax, 3.5, 9.7, 0.9)
    ax.annotate("", xy=(3.5, 2.15), xytext=(6.5, 2.15),
                arrowprops=dict(arrowstyle="<|-|>", color="#1f4e79", lw=1.5))
    ax.text(5.0, 2.35, "L1（直線重ね継手長さ）", fontproperties=jp,
            ha="center", fontsize=10, color="#1f4e79", fontweight="bold")
    ax.text(5.0, -0.55, "フックなしで所定長さを重ねる。太径（D35 以上）には使えない",
            fontproperties=jp, ha="center", fontsize=9, color="#444")
    ax.set_xlim(-0.5, 10.5)
    ax.set_ylim(-1.1, 3.1)
    ax.set_aspect("equal")
    ax.axis("off")
    ax.set_title("(a) L1：直線の重ね継手", fontproperties=jp,
                 fontsize=11, fontweight="bold")

    # --- (b) L1h: フック付き重ね継手 ---
    ax = axes[1]
    ax.add_patch(mpatches.Rectangle((0, 0), 10, 2.6, fc="#e8e8e8",
                                     ec="k", lw=0.8))
    # 左の鉄筋（右端にフック＝下向き）
    draw_bar(ax, 0.3, 6.0, 1.7, ribs=False)
    ax.plot([6.0, 6.0], [1.7, 1.0], color="#888", lw=5,
            solid_capstyle="butt", zorder=3)
    # 右の鉄筋（左端にフック＝上向き）
    draw_bar(ax, 4.0, 9.7, 0.9, ribs=False)
    ax.plot([4.0, 4.0], [0.9, 1.6], color="#888", lw=5,
            solid_capstyle="butt", zorder=3)
    ax.annotate("", xy=(4.0, 2.15), xytext=(6.0, 2.15),
                arrowprops=dict(arrowstyle="<|-|>", color="#1f4e79", lw=1.5))
    ax.text(5.0, 2.35, "L1h（フック付き重ね継手長さ）", fontproperties=jp,
            ha="center", fontsize=10, color="#1f4e79", fontweight="bold")
    ax.text(1.0, 0.45, "※ L1h はフックの折り曲げ起点間の長さで測る（余長は含まない）",
            fontproperties=jp, fontsize=8, color="#c00000")
    ax.text(5.0, -0.55, "フックの機械的抵抗が加わるため L1h < L1（同条件で 10db 程度短い）",
            fontproperties=jp, ha="center", fontsize=9, color="#444")
    ax.set_xlim(-0.5, 10.5)
    ax.set_ylim(-1.1, 3.1)
    ax.set_aspect("equal")
    ax.axis("off")
    ax.set_title("(b) L1h：フック付き重ね継手", fontproperties=jp,
                 fontsize=11, fontweight="bold")

    fig.suptitle("図 7-2  重ね継手長さ L1・L1h の定義",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig7-2_L1.png")


# ===========================================================================
# 図 7-5: 部材別の継手可能範囲（梁・柱）
# ===========================================================================
def fig_7_5():
    fig, axes = plt.subplots(1, 2, figsize=(14, 5.8),
                              gridspec_kw={"width_ratios": [1.5, 1.0]})

    # --- (a) 大梁：モーメント図と継手範囲 ---
    ax = axes[0]
    L = 10.0
    yb = 3.0   # 梁軸
    # 柱
    for xc in [0, L]:
        ax.add_patch(mpatches.Rectangle((xc - 0.45, yb - 2.2), 0.9, 4.4,
                                         fc="#d9d9d9", ec="k", lw=1.0))
    # 梁
    ax.add_patch(mpatches.Rectangle((0.45, yb - 0.45), L - 0.9, 0.9,
                                     fc="#bcd2ea", ec="k", lw=1.0))
    # M 図（引張側に描く：端部負曲げ→上、中央正曲げ→下）
    xs = np.linspace(0.45, L - 0.45, 100)
    M = -1.0 + 6.0 * (xs / L) * (1 - xs / L)   # 端 -1, 中央 +0.5
    y_m = yb + 0.45 - M * 1.1 * (-1)  # 負(引張上)→上に描く
    y_m = yb - M * 1.1
    ax.plot(xs, yb - M * 1.1, color="#777", lw=1.5, ls="--")
    ax.axhline(yb, color="#999", lw=0.5)
    ax.text(1.0, yb + 1.35, "端部：上端引張\n（負曲げ大）", fontproperties=jp,
            fontsize=8, color="#777")
    ax.text(L / 2, yb - 1.1, "中央：下端引張（正曲げ）", fontproperties=jp,
            fontsize=8, color="#777", ha="center")
    # 上端筋の継手範囲（中央 L0/2）
    ax.add_patch(mpatches.Rectangle((L / 4, yb + 0.5), L / 2, 0.28,
                                     fc="#7fbf7f", ec="k", lw=0.6, zorder=5))
    ax.text(L / 2, yb + 0.95, "上端筋の継手範囲：中央 L0/2（端部の負曲げを避ける）",
            fontproperties=jp, ha="center", fontsize=8.5, color="#1f7a1f")
    # 下端筋の継手範囲（柱面から L0/4 付近、両側）
    for x0 in [0.45 + L / 10, L - 0.45 - L / 10 - L / 5]:
        ax.add_patch(mpatches.Rectangle((x0, yb - 0.78), L / 5, 0.28,
                                         fc="#7fbf7f", ec="k", lw=0.6,
                                         zorder=5))
    ax.text(L / 2, yb - 1.65,
            "下端筋の継手範囲：柱面から L0/4 付近（中央の正曲げと柱面直近のヒンジ域を避ける）",
            fontproperties=jp, ha="center", fontsize=8.5, color="#1f7a1f")
    ax.text(L / 2, yb - 2.35, "L0＝梁の内法スパン。隣り合う継手は 0.5·L1 以上ずらす（千鳥）",
            fontproperties=jp, ha="center", fontsize=8.5, color="#c00000")
    ax.set_xlim(-0.8, L + 0.8)
    ax.set_ylim(yb - 2.8, yb + 2.6)
    ax.set_aspect("equal")
    ax.axis("off")
    ax.set_title("(a) 大梁の継手可能範囲（曲げモーメントと対応）",
                 fontproperties=jp, fontsize=11, fontweight="bold")

    # --- (b) 柱：継手範囲 ---
    ax = axes[1]
    h0 = 6.0    # 内法高さ
    yb0 = 0.0   # 下階梁天端
    # 梁（上下）
    for y in [yb0 - 0.8, yb0 + h0]:
        ax.add_patch(mpatches.Rectangle((-1.6, y), 4.4, 0.8,
                                         fc="#bcd2ea", ec="k", lw=0.8))
    # 柱
    ax.add_patch(mpatches.Rectangle((0, yb0), 1.2, h0, fc="#d9d9d9",
                                     ec="k", lw=1.2))
    # 柱の M 図（柱頭・柱脚で大、中央で small）
    ys = np.linspace(yb0, yb0 + h0, 50)
    Mc = (2 * (ys - yb0) / h0 - 1)  # -1(脚)〜+1(頭)
    ax.plot(1.2 + np.abs(Mc) * 1.0 + 0.15, ys, color="#777", lw=1.5, ls="--")
    ax.text(3.0, yb0 + h0 - 0.4, "柱頭：曲げ大", fontproperties=jp,
            fontsize=8, color="#777")
    ax.text(3.0, yb0 + 0.3, "柱脚：曲げ大", fontproperties=jp,
            fontsize=8, color="#777")
    # 継手範囲（梁天端+500 〜 h0·3/4）
    y1 = yb0 + 0.9   # 500mm 相当
    y2 = yb0 + h0 * 0.75
    ax.add_patch(mpatches.Rectangle((-0.45, y1), 0.28, y2 - y1,
                                     fc="#7fbf7f", ec="k", lw=0.6, zorder=5))
    ax.annotate("継手可能範囲\n梁天端から 500mm 以上\nかつ h0×3/4 以下",
                xy=(-0.31, (y1 + y2) / 2), xytext=(-3.4, (y1 + y2) / 2),
                fontproperties=jp, fontsize=9, color="#1f7a1f",
                va="center",
                arrowprops=dict(arrowstyle="->", color="#1f7a1f"))
    ax.plot([-0.45, 1.4], [y1, y1], color="#c00000", lw=0.8, ls=":")
    ax.plot([-0.45, 1.4], [y2, y2], color="#1f4e79", lw=0.8, ls=":")
    ax.text(0.6, -1.6, "柱主筋（圧接等）は\n梁天端から 500mm 以上・内法高さの 3/4 以下",
            fontproperties=jp, ha="center", fontsize=8.5, color="#444")
    ax.set_xlim(-3.6, 4.6)
    ax.set_ylim(-2.2, h0 + 1.2)
    ax.set_aspect("equal")
    ax.axis("off")
    ax.set_title("(b) 柱の継手可能範囲", fontproperties=jp,
                 fontsize=11, fontweight="bold")

    fig.suptitle("図 7-5  部材別の継手可能範囲 ── 応力の小さい位置で継ぐ",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig7-5_zones.png")


figs = {
    "f1": fig_7_1(),
    "f2": fig_7_2(),
    "f5": fig_7_5(),
}
print("figures:", list(figs.keys()))

# ===========================================================================
# Excel 構築
# ===========================================================================
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

wb = Workbook()
C_TITLE = "1F4E79"; C_HEAD = "2E75B6"; C_ANS = "E2EFDA"
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
f_title = Font(name="MS PGothic", size=15, bold=True, color="FFFFFF")
f_head = Font(name="MS PGothic", size=11, bold=True, color="FFFFFF")
f_body = Font(name="MS PGothic", size=10)
f_ans = Font(name="MS PGothic", size=10, color="375623")
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)


def setup(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False


def title_row(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text)
    c.font = f_title; c.fill = PatternFill("solid", fgColor=C_TITLE)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[row].height = 30


def head(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text)
    c.font = f_head; c.fill = PatternFill("solid", fgColor=C_HEAD)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[row].height = 22


def body(ws, row, text, span=8, ans=False, h=None):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text)
    c.font = f_ans if ans else f_body
    c.alignment = wrap
    if ans:
        c.fill = PatternFill("solid", fgColor=C_ANS)
    if h:
        ws.row_dimensions[row].height = h


def table(ws, start_row, headers, rows, col1=1):
    r = start_row
    for j, htxt in enumerate(headers):
        c = ws.cell(r, col1 + j, htxt)
        c.font = Font(name="MS PGothic", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C_HEAD)
        c.alignment = center; c.border = border
    for data in rows:
        r += 1
        for j, v in enumerate(data):
            c = ws.cell(r, col1 + j, v)
            c.font = f_body
            c.alignment = center if j > 0 else wrap
            c.border = border
            if r % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F2F7FC")
    return r


def put_img(ws, path, anchor, w=None):
    img = XLImage(path)
    if w:
        ratio = w / img.width
        img.width = w; img.height = int(img.height * ratio)
    ws.add_image(img, anchor)


# ---- 目次 ----
ws = wb.active
ws.title = "目次"
setup(ws, [4, 22, 60, 14])
title_row(ws, 1, "鉄筋の継手 問題集（RC マンション設計担当・新人向け）", span=4)
body(ws, 2,
     "テーマ No.7：鉄筋継手。4 種類の継手（重ね・ガス圧接・溶接・機械式）の応力伝達機構、"
     "重ね継手長さ L1・L1h の説明・算出・暗記、一覧表の計算仮定、部材別の継手可能範囲までを"
     "通す。定着（No.6 教材）とセットで学ぶこと。",
     span=4, h=44)
r = table(ws, 4,
          ["No.", "シート", "到達目標", "図"],
          [["7-1", "7-1 応力伝達機構",
            "重ね・溶接・圧接・機械式継手の伝達機構を説明できる", "4 種比較"],
           ["7-2", "7-2 L1・L1h",
            "L1・L1h の説明・算出ができる", "重ね継手の定義"],
           ["7-3", "7-3 L1・L1h 暗記",
            "継手長一覧表の L1・L1h を暗記している", "一覧表"],
           ["7-4", "7-4 計算仮定条件",
            "一覧表の計算仮定条件を理解している", "—"],
           ["7-5", "7-5 継手可能範囲",
            "各部材の継手可能範囲を説明できる", "梁・柱の範囲"]])
body(ws, r + 2,
     "凡例：L1=直線重ね継手長さ、L1h=フック付き重ね継手長さ、db=鉄筋径。"
     "数値は『RC 規準』『公共建築工事標準仕様書』等に基づく目安（Fc・鉄筋種別で変わる）。"
     "最新版規準・標準図・自社基準を必ず併用すること。",
     span=4, h=58)

# ---- 7-1 ----
ws = wb.create_sheet("7-1 応力伝達機構")
setup(ws, [8, 16, 18, 16, 14, 12, 12, 12])
title_row(ws, 1, "7-1  継手の応力伝達機構（重ね・圧接・溶接・機械式）")
head(ws, 3, "■ 図 7-1  4 種類の継手")
put_img(ws, figs["f1"], "A4", w=800)
head(ws, 32, "■ 問題 1  伝達機構の対応")
r = table(ws, 33,
          ["継手の種類", "応力の伝達経路（記入）", "直接/間接（記入）"],
          [["重ね継手", "", ""],
           ["ガス圧接継手", "", ""],
           ["溶接継手", "", ""],
           ["機械式継手", "", ""]])
head(ws, r + 2, "■ 問題 2  重ね継手の限界")
body(ws, r + 3, "(1) 重ね継手が『太径鉄筋（D35 以上）』に使えない理由を、"
                "付着で伝達できる力の限界と割裂ひび割れの観点から述べよ。", h=40)
body(ws, r + 4, "(2) 実務では D29 以上はガス圧接・機械式が主流となる。"
                "重ね継手に対する利点を 2 つ挙げよ（継手長さ不要・断面の過密回避）。",
     h=40)
head(ws, r + 6, "■ 問題 3  ガス圧接の検査基準")
body(ws, r + 7, "ガス圧接部の外観検査項目を埋めよ：ふくらみの径は【 ① 】db 以上、"
                "ふくらみの長さは【 ② 】db 以上、鉄筋中心軸の偏心量は【 ③ 】db 以下、"
                "圧接面のずれは【 ④ 】db 以下。", h=44)
head(ws, r + 9, "■ 問題 4  継手の選定")
body(ws, r + 10, "RC マンションの次の部位で、どの継手を選ぶのが一般的か。"
                 "理由とともに答えよ。", h=22)
r = table(ws, r + 12,
          ["No.", "部位", "継手（記入）", "理由（記入）"],
          [["(1)", "スラブ筋 D13", "", ""],
           ["(2)", "大梁主筋 D25", "", ""],
           ["(3)", "柱主筋 D29", "", ""],
           ["(4)", "PCa 部材の接合部 D22", "", ""]])

# ---- 7-2 ----
ws = wb.create_sheet("7-2 L1・L1h")
setup(ws, [8, 16, 18, 16, 14, 12, 12, 12])
title_row(ws, 1, "7-2  重ね継手長さ L1・L1h の説明と算出")
head(ws, 3, "■ 図 7-2  L1・L1h の定義")
put_img(ws, figs["f2"], "A4", w=800)
head(ws, 26, "■ 問題 1  定義の穴埋め")
body(ws, 27, "L1 は【 ① 】の重ね継手長さ、L1h は【 ② 】付きの重ね継手長さである。"
             "重ね継手は 2 本の鉄筋がそれぞれコンクリートへの【 ③ 】で応力を"
             "受け渡すため、同条件の定着長 L2 より【 ④ 】く設定される"
             "（一般に L1 ≒ L2 + 5db 程度）。", h=44)
head(ws, 29, "■ 問題 2  L1 の算出")
body(ws, 30, "Fc=24、SD345、D22（db=22mm）のとき、一覧表（7-3 シート）より "
             "L1（db の倍数と mm）を求めよ。フック付き L1h なら何 mm か。", h=32)
head(ws, 32, "■ 問題 3  L2 との比較")
body(ws, 33, "同条件（Fc24・SD345）の定着長 L2=35db と比較して、L1 が長い理由を"
             "応力の伝達経路（1 本 vs 2 本直列）から説明せよ。", h=40)
head(ws, 35, "■ 問題 4  L1h の測り方")
body(ws, 36, "L1h はどこからどこまでの長さか（フックの折り曲げ起点間）。"
             "フックの余長を L1h に含めてよいか。", h=32)
head(ws, 38, "■ 問題 5  径が異なる鉄筋の重ね継手")
body(ws, 39, "D22 と D25 を重ね継手する場合、L1 はどちらの径を基準に算定するか。"
             "理由とともに答えよ（一般に細い方の径 × 倍数）。", h=40)

# ---- 7-3 ----
ws = wb.create_sheet("7-3 L1・L1h 暗記")
setup(ws, [10, 14, 12, 12, 12, 12, 14])
title_row(ws, 1, "7-3  継手長一覧表（L1・L1h）の暗記", span=7)
head(ws, 3, "■ 継手長一覧表（Fc 別・鉄筋種別、db の倍数）", span=7)
body(ws, 4, "下表は『直線重ね継手 L1 / フック付き L1h』の代表値（db の倍数）。"
            "定着長（L2・L2h）より 5db 長い、と覚えると整理しやすい。", span=7, h=32)
r = table(ws, 6,
          ["Fc (N/mm²)", "SD295 L1", "SD295 L1h",
           "SD345 L1", "SD345 L1h", "SD390 L1", "SD390 L1h"],
          [["18", "45db", "35db", "45db", "35db", "—", "—"],
           ["21", "40db", "30db", "45db", "35db", "50db", "40db"],
           ["24〜27", "35db", "25db", "40db", "30db", "45db", "35db"],
           ["30〜36", "35db", "25db", "35db", "25db", "40db", "30db"],
           ["39〜45", "30db", "20db", "35db", "25db", "40db", "30db"]])
body(ws, r + 2,
     "※ 上表は『一般的な目安値』。実際の設計では使用する規準・標準仕様書・自社標準図の"
     "値を用いること。数値は版・条件で異なるため、必ず最新の表で確認する。"
     "（No.6 教材の定着長一覧表と見比べると、同条件で L1 = L2 + 5db になっている）",
     span=7, h=44)
head(ws, r + 4, "■ 問題 1  表の穴埋め暗記", span=7)
body(ws, r + 5, "次の条件の L1・L1h を上表から答えよ（db の倍数）。", span=7, h=20)
r2 = table(ws, r + 7,
           ["No.", "Fc", "鉄筋", "L1（記入）", "L1h（記入）"],
           [["(1)", "21", "SD345", "", ""],
            ["(2)", "24", "SD295", "", ""],
            ["(3)", "30", "SD345", "", ""],
            ["(4)", "21", "SD390", "", ""],
            ["(5)", "36", "SD390", "", ""]])
head(ws, r2 + 2, "■ 問題 2  実寸法の算出", span=7)
body(ws, r2 + 3, "Fc=24、SD345 の場合：(1) D19 の L1 は何 mm か。"
                 "(2) D22 の L1h は何 mm か。(3) スラブ筋 D13（SD295）の L1 は何 mm か。",
     span=7, h=32)
head(ws, r2 + 5, "■ 問題 3  定着長との関係", span=7)
body(ws, r2 + 6, "同じ Fc・鉄筋種別で L1 と L2 の関係はどうなっているか。"
                 "その理由（重ね継手は 2 本の付着の直列＝乗り換えロス）とともに述べよ。",
     span=7, h=40)

# ---- 7-4 ----
ws = wb.create_sheet("7-4 計算仮定条件")
setup(ws, [8, 18, 20, 16, 14, 12, 12])
title_row(ws, 1, "7-4  継手長一覧表における計算仮定条件")
head(ws, 3, "■ 問題 1  一覧表の前提（理解の確認）")
body(ws, 4, "継手長一覧表の数値は、定着長と同じく基本定着長 Lb をベースに、"
            "重ね継手特有の仮定を加えて丸めたもの。主な仮定条件を確認する。", h=32)
r = table(ws, 6,
          ["項目", "一般的な仮定", "理解（記入）"],
          [["鉄筋の応力度", "鉄筋の許容引張応力度（全強）を伝達できること", ""],
           ["継手位置", "応力の小さい位置に設けることが前提", ""],
           ["隣接継手のずらし", "隣り合う継手は 0.5·L1 以上ずらす（千鳥配置）", ""],
           ["鉄筋位置", "上端筋は付着低下を考慮（定着と同様の割増し）", ""],
           ["かぶり・あき", "標準的なかぶり厚・鉄筋あきを前提", ""],
           ["適用径", "太径（D35 以上）は重ね継手不可", ""]])
head(ws, r + 2, "■ 問題 2  千鳥配置（ずらし）の理由")
body(ws, r + 3, "(1) 同一断面に継手を集中させてはいけない理由を、断面の弱点集中と"
                "割裂ひび割れの観点から述べよ。", h=40)
body(ws, r + 4, "(2) 隣接継手のずらし量の目安（0.5·L1 以上）を、D22・L1=880mm の場合の"
                "mm 数で答えよ。", h=32)
head(ws, r + 6, "■ 問題 3  仮定から外れる場合")
r = table(ws, r + 8,
          ["No.", "状況", "扱い（記入）"],
          [["(1)", "応力の大きい位置にやむを得ず継手を設ける", ""],
           ["(2)", "軽量コンクリートを使用", ""],
           ["(3)", "上端筋（打設面から下に多くのコンクリート）", ""],
           ["(4)", "D38 の柱主筋", ""]])
head(ws, r + 2, "■ 問題 4  なぜ仮定を理解する必要があるか")
body(ws, r + 3, "『一覧表の数字だけ』を使って継手を設計すると危険な理由を、"
                "継手位置・千鳥・太径不可の 3 つの仮定をふまえて述べよ。", h=44)

# ---- 7-5 ----
ws = wb.create_sheet("7-5 継手可能範囲")
setup(ws, [8, 16, 18, 16, 14, 12, 12, 12])
title_row(ws, 1, "7-5  各部材の継手可能範囲")
head(ws, 3, "■ 図 7-5  梁・柱の継手可能範囲（応力の小さい位置で継ぐ）")
put_img(ws, figs["f5"], "A4", w=820)
head(ws, 30, "■ 問題 1  大原則の穴埋め")
body(ws, 31, "継手は部材の【 ① 】が小さい位置に設けるのが大原則。"
             "梁の上端筋は【 ② 】部（端部の負曲げを避ける）、下端筋は柱面から"
             "【 ③ 】付近（中央の正曲げと柱面直近のヒンジ域を避ける）に設ける。"
             "柱主筋は梁天端から【 ④ 】mm 以上、かつ内法高さの【 ⑤ 】以下の範囲で継ぐ。",
     h=58)
head(ws, 33, "■ 問題 2  部材別の継手範囲")
r = table(ws, 34,
          ["部材・鉄筋", "継手可能範囲（記入）", "避ける位置（記入）"],
          [["大梁 上端筋", "", ""],
           ["大梁 下端筋", "", ""],
           ["柱 主筋", "", ""],
           ["スラブ 上端筋", "", ""],
           ["スラブ 下端筋", "", ""],
           ["壁 縦筋", "", ""]])
head(ws, r + 2, "■ 問題 3  理由の理解")
body(ws, r + 3, "(1) 梁の上端筋と下端筋で継手範囲が逆になる理由を、"
                "曲げモーメント図（どこでどちら側が引張になるか）から説明せよ。", h=40)
body(ws, r + 4, "(2) 柱頭・柱脚部で継手を避ける理由を、地震時の応力と"
                "塑性ヒンジの観点から述べよ。", h=40)
body(ws, r + 5, "(3) 柱の継手位置が『梁天端から 500mm 以上』とされる理由を、"
                "施工（圧接作業スペース・コンクリート打継ぎ）の観点から 1 行で。", h=32)
head(ws, r + 7, "■ 問題 4  マンション実務総合")
body(ws, r + 8, "RC マンション基準階の配筋図をチェックする際、継手について確認する"
                "項目を 4 つ挙げよ（①位置が可能範囲内か ②長さ L1 が一覧表どおりか "
                "③千鳥（0.5L1 ずらし）か ④太径に重ね継手を使っていないか）。", h=44)
body(ws, r + 9, "また、圧接継手の場合に配筋検査で確認する項目（ふくらみ径・長さ・偏心・"
                "有資格者施工）も挙げよ。", h=32)

# ---- 解答 ----
ws = wb.create_sheet("解答")
setup(ws, [4, 22, 60, 14])
title_row(ws, 1, "解答・解説", span=4)
row = 2


def ah(t):
    global row
    head(ws, row, t, span=4); row += 1


def an(t, h=None):
    global row
    body(ws, row, t, span=4, ans=True, h=h); row += 1


ah("7-1  応力伝達機構")
an("問1：重ね継手＝鉄筋→付着→コンクリート→付着→鉄筋（間接伝達）。"
   "ガス圧接＝加熱・加圧で一体化した断面で直接伝達。"
   "溶接＝溶着金属を介して断面で直接伝達。"
   "機械式＝鉄筋→カプラー（ねじ・モルタル）→鉄筋の機械的伝達。", h=72)
an("問2：(1)太径ほど伝達すべき力（As·σy）は径の 2 乗で増えるが、付着面積は径の"
   "1 乗でしか増えないため、必要な重ね長さが非現実的に長くなる。また太径の付着力は"
   "かぶりコンクリートを割裂させやすい。よって D35 以上は重ね継手不可。"
   "(2)圧接・機械式の利点：①継手長さが不要で鉄筋量・断面の過密を減らせる "
   "②応力伝達が直接的で信頼性が高い（重ね部の割裂リスクがない）。", h=100)
an("問3：①1.4 ②1.1 ③1/5（0.2）④1/4（0.25）。"
   "外観検査の不合格は切断・再圧接。超音波探傷検査も併用される。", h=44)
an("問4：(1)スラブ筋 D13＝重ね継手（細径で L1 が短く経済的・施工簡単）。"
   "(2)大梁主筋 D25＝ガス圧接 or 重ね継手（現場方針による。D25 は両方あり得る）。"
   "(3)柱主筋 D29＝ガス圧接（太径で重ね継手は長大・過密になるため）。"
   "(4)PCa 接合部 D22＝機械式継手（スリーブ継手。PCa は現場で圧接できないため）。",
   h=86)

ah("7-2  L1・L1h")
an("問1：①直線 ②フック ③付着 ④長（L1 ≒ L2 + 5db 程度）。", h=22)
an("問2：Fc24・SD345 → L1=40db。D22 なら 40×22=880mm。"
   "L1h=30db → 30×22=660mm。", h=32)
an("問3：定着は鉄筋 1 本の力をコンクリートへ伝えるだけだが、重ね継手は"
   "『鉄筋 A →コンクリート→鉄筋 B』と付着の乗り換えが直列に 2 回起こり、"
   "重ね区間で 2 本分の付着応力が集中する。安全率を確保するため L1 > L2 とする。",
   h=58)
an("問4：L1h はフックの折り曲げ起点から相手鉄筋のフック折り曲げ起点までの長さで測る。"
   "フックの余長（先端の直線部）は L1h に含めない（形状保持のための最小寸法）。",
   h=44)
an("問5：細い方（D22）の径を基準に算定する。継手で伝達すべき力は細い方の鉄筋の"
   "耐力で決まる（細い方が先に降伏する）ため、L1 = 倍数 × 22mm でよい。", h=44)

ah("7-3  L1・L1h 暗記")
an("問1：(1)Fc21/SD345：L1=45db、L1h=35db。(2)Fc24/SD295：L1=35db、L1h=25db。"
   "(3)Fc30/SD345：L1=35db、L1h=25db。(4)Fc21/SD390：L1=50db、L1h=40db。"
   "(5)Fc36/SD390：L1=40db、L1h=30db。", h=58)
an("問2：(1)Fc24/SD345/D19：L1=40db=40×19=760mm。"
   "(2)D22 の L1h=30db=660mm。"
   "(3)Fc24/SD295/D13：L1=35db=35×13=455mm。", h=44)
an("問3：同条件で L1 = L2 + 5db（例：Fc24/SD345 → L2=35db、L1=40db）。"
   "重ね継手は 2 本の鉄筋間で付着を乗り換えるため、単独の定着より余裕を持たせている。",
   h=44)

ah("7-4  計算仮定条件")
an("問1：全強伝達＝継手部でも母材と同等の力を伝える前提。"
   "応力の小さい位置＝一覧表の長さは『適切な位置に設けること』とセット。"
   "千鳥＝同一断面への集中を避ける前提。上端筋割増し・標準かぶり・太径不可も"
   "定着一覧表と同様の前提。", h=72)
an("問2：(1)同一断面に継手が集中すると、その断面の伝達能力が同時に低下し"
   "弱点断面になる。また重ね部の付着応力が重なって割裂ひび割れを誘発する。"
   "(2)0.5×880=440mm 以上ずらす。", h=58)
an("問3：(1)原則避ける。やむを得ない場合は継手長さの割増しや機械式への変更を検討。"
   "(2)軽量コンクリートは付着が小さく割増し必要。"
   "(3)上端筋はブリーディングで付着低下→割増し必要（定着と同じ）。"
   "(4)D38 は重ね継手不可→ガス圧接・機械式・溶接を選定。", h=72)
an("問4：一覧表の長さは『応力の小さい位置・千鳥・標準条件・適用径内』が前提。"
   "数字だけ覚えて位置やずらしを無視すると、長さが合っていても弱点断面をつくる。"
   "『長さ』と『位置・配置』はセットで初めて成立する、と理解する。", h=58)

ah("7-5  継手可能範囲")
an("問1：①応力（曲げモーメント） ②中央（スパン中央 L0/2） ③L0/4 "
   "④500 ⑤3/4。", h=32)
an("問2：大梁上端筋＝中央 L0/2（避ける：端部の負曲げ域）。"
   "大梁下端筋＝柱面から L0/4 付近（避ける：中央の正曲げ域と柱面直近のヒンジ域）。"
   "柱主筋＝梁天端+500mm 以上〜内法高さ 3/4 以下（避ける：柱頭・柱脚）。"
   "スラブ上端筋＝中央部（避ける：端部）。スラブ下端筋＝端部（避ける：中央）。"
   "壁縦筋＝各階の中間部（床上から立ち上げて応力の小さい位置で継ぐ）。", h=100)
an("問3：(1)梁の曲げは端部で上端引張（負曲げ）、中央で下端引張（正曲げ）。"
   "引張応力の大きい位置を避けると、上端筋は中央・下端筋は端部寄りが継手適所になる。"
   "(2)柱頭・柱脚は地震時曲げが最大で塑性ヒンジが想定される領域。ヒンジ域に継手が"
   "あると変形能力・耐力の低下に直結するため避ける。"
   "(3)梁天端（打継ぎ面）直上は圧接器具が使えない・コンクリートの品質が不安定なため、"
   "500mm 以上離して作業性と品質を確保する。", h=100)
an("問4：チェック項目：①継手位置が可能範囲内か ②L1・L1h の長さが一覧表どおりか "
   "③隣接継手が 0.5L1 以上ずれているか（千鳥） ④D35 以上に重ね継手を使っていないか。"
   "圧接検査：ふくらみ径 1.4db 以上・長さ 1.1db 以上・偏心 1/5db 以下・"
   "圧接技量資格者の施工・（必要に応じ）超音波探傷。", h=86)

XLSX = os.path.join(OUT, "鉄筋継手問題集.xlsx")
wb.save(XLSX)
print("saved:", XLSX)
