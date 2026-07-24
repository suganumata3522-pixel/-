# -*- coding: utf-8 -*-
"""地盤の支持力（Terzaghi）問題集（図つき）Excel。出力: docs/bearing_terzaghi/地盤の支持力Terzaghi問題集.xlsx
1 支持力式と3つの項 / 2 支持力係数と地盤種別の目安
3 形状係数・寸法効果・荷重傾斜補正 / 4 砂質土・粘性土の算定例
※ 支持力係数・形状係数・目安値は規準（告示1113号/建築基礎指針）で確認要
"""
import os
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib import font_manager as fm
FONT = "/usr/share/fonts/truetype/fonts-japanese-gothic.ttf"
jp = fm.FontProperties(fname=FONT); fm.fontManager.addfont(FONT)
plt.rcParams["font.family"] = jp.get_name(); plt.rcParams["axes.unicode_minus"] = False
OUT = "docs/bearing_terzaghi"; FIG = os.path.join(OUT, "figures"); os.makedirs(FIG, exist_ok=True)

NF = {0: (5.7, 1.0, 0.0), 5: (7.3, 1.6, 0.5), 10: (9.6, 2.7, 1.2),
      15: (12.9, 4.4, 2.5), 20: (17.7, 7.4, 5.0), 25: (25.1, 12.7, 9.7),
      30: (37.2, 22.5, 19.7), 35: (57.8, 41.4, 42.4), 40: (95.7, 81.3, 100.4)}


def save(fig, n):
    p = os.path.join(FIG, n); fig.savefig(p, dpi=130, bbox_inches="tight"); plt.close(fig); return p


def fig_formula():
    fig, ax = plt.subplots(figsize=(12.5, 5.6))
    # 基礎とすべり面
    ax.add_patch(mpatches.Rectangle((4.0, 3.0), 2.0, 0.7, fc="#d9d9d9", ec="k", lw=1.2))
    ax.text(5.0, 3.35, "基礎 B", ha="center", fontproperties=jp, fontsize=9)
    ax.axhline(3.0, color="#8B5A2B", lw=1, ls=":")
    ax.annotate("", xy=(3.7, 3.0), xytext=(3.7, 4.0), arrowprops=dict(arrowstyle="<->", color="#333"))
    ax.text(3.4, 3.5, "Df", ha="right", fontproperties=jp, fontsize=9)
    # すべり面（三角くさび＋対数らせん風）
    th = np.linspace(0, 1, 30)
    for sgn in [-1, 1]:
        x0 = 5.0 + sgn * 1.0
        xx = x0 + sgn * 2.2 * th
        yy = 3.0 - 1.6 * np.sin(th * np.pi / 1.5)
        ax.plot(xx, yy, color="#c00000", lw=1.6)
    ax.plot([4.0, 5.0, 6.0], [3.0, 1.6, 3.0], color="#c00000", lw=1.6)
    ax.text(5.0, 1.3, "せん断すべり面", ha="center", fontproperties=jp, fontsize=8, color="#c00000")
    # 荷重
    ax.annotate("", xy=(5.0, 3.0), xytext=(5.0, 4.6), arrowprops=dict(arrowstyle="-|>", color="#1f4e79", lw=3))
    ax.text(5.2, 4.3, "極限支持力 qu", fontproperties=jp, fontsize=10, color="#1f4e79")
    # 式
    ax.text(0.5, 0.9, "qu = α・c・Nc  +  β・γ1・B・Nγ  +  γ2・Df・Nq",
            transform=ax.transAxes, ha="center", fontproperties=jp, fontsize=13, fontweight="bold")
    ax.add_patch(mpatches.FancyBboxPatch((0.02, 0.02), 0.30, 0.32, boxstyle="round,pad=0.01",
                 transform=ax.transAxes, fc="#e8f0fa", ec="#2e75b6"))
    ax.text(0.17, 0.28, "第1項  α・c・Nc", transform=ax.transAxes, ha="center",
            fontproperties=jp, fontsize=9.5, fontweight="bold", color="#1f4e79")
    ax.text(0.17, 0.12, "粘着力による抵抗\n（粘性土で主役）", transform=ax.transAxes, ha="center",
            va="center", fontproperties=jp, fontsize=8.5)
    ax.add_patch(mpatches.FancyBboxPatch((0.35, 0.02), 0.30, 0.32, boxstyle="round,pad=0.01",
                 transform=ax.transAxes, fc="#eaf7ea", ec="#548235"))
    ax.text(0.50, 0.28, "第2項  β・γ1・B・Nγ", transform=ax.transAxes, ha="center",
            fontproperties=jp, fontsize=9.5, fontweight="bold", color="#548235")
    ax.text(0.50, 0.12, "基礎幅・土の自重\n（砂質土で効く）", transform=ax.transAxes, ha="center",
            va="center", fontproperties=jp, fontsize=8.5)
    ax.add_patch(mpatches.FancyBboxPatch((0.68, 0.02), 0.30, 0.32, boxstyle="round,pad=0.01",
                 transform=ax.transAxes, fc="#fdefe0", ec="#c55a11"))
    ax.text(0.83, 0.28, "第3項  γ2・Df・Nq", transform=ax.transAxes, ha="center",
            fontproperties=jp, fontsize=9.5, fontweight="bold", color="#c55a11")
    ax.text(0.83, 0.12, "根入れ（上載圧）\n深いほど大", transform=ax.transAxes, ha="center",
            va="center", fontproperties=jp, fontsize=8.5)
    ax.set_xlim(0, 10); ax.set_ylim(0.5, 5.2); ax.axis("off")
    ax.set_title("図 1  Terzaghi の支持力式と3つの項", fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig1_formula.png")


def fig_factors():
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 4.8))
    ax = axes[0]
    phi = list(NF.keys())
    Nc = [NF[p][0] for p in phi]; Nq = [NF[p][1] for p in phi]; Ng = [NF[p][2] for p in phi]
    ax.semilogy(phi, Nc, "o-", color="#c00000", lw=1.8, label="Nc（粘着）")
    ax.semilogy(phi, Nq, "s-", color="#c55a11", lw=1.8, label="Nq（根入れ）")
    ax.semilogy(phi, [max(g, 0.1) for g in Ng], "^-", color="#548235", lw=1.8, label="Nγ（幅）")
    ax.set_xlabel("内部摩擦角 φ (度)", fontproperties=jp, fontsize=9)
    ax.set_ylabel("支持力係数（対数）", fontproperties=jp, fontsize=9)
    ax.set_title("(a) 支持力係数 Nc・Nq・Nγ と φ", fontproperties=jp, fontsize=9.5, fontweight="bold")
    ax.legend(prop=jp, fontsize=8.5); ax.grid(alpha=0.3, which="both")
    ax.text(2, 0.15, "※Terzaghi値は目安。規準で確認要", fontproperties=jp, fontsize=7.5, color="#833c00")
    ax = axes[1]
    ax.text(0.5, 0.96, "地盤種別と許容支持力度の目安", transform=ax.transAxes, ha="center",
            fontproperties=jp, fontsize=10.5, fontweight="bold", color="#1f4e79")
    ax.text(0.04, 0.85,
            "  硬岩            : 1000 kN/m2 以上\n"
            "  軟岩・土丹      : 500〜1000\n"
            "  密な礫層        : 500〜600\n"
            "  密な砂質地盤    : 200〜300\n"
            "  中位の砂質地盤  : 100〜200\n"
            "  硬い粘性土      : 100〜200\n"
            "  中位の粘性土    : 50〜100\n"
            "  軟らかい粘性土  : 20〜50",
            transform=ax.transAxes, fontproperties=jp, fontsize=9.2, va="top", family="monospace")
    ax.text(0.04, 0.1, "※ 目安値。実際は支持力式・平板載荷試験・規準で確認要。",
            transform=ax.transAxes, fontproperties=jp, fontsize=7.8, color="#833c00")
    ax.axis("off")
    ax.set_title("(b) 地盤種別の目安（暗記）", fontproperties=jp, fontsize=9.5, fontweight="bold")
    fig.suptitle("図 2  支持力係数と地盤種別の目安", fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig2_factors.png")


def fig_correct():
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 4.8))
    # (a) 形状係数
    ax = axes[0]
    shapes = [("連続基礎", 1.0, 0.5), ("正方形", 1.3, 0.4), ("円形", 1.3, 0.3),
              ("長方形", "1.0+0.3B/L", "0.5-0.1B/L")]
    ax.text(0.5, 0.95, "形状係数 α・β", transform=ax.transAxes, ha="center",
            fontproperties=jp, fontsize=10.5, fontweight="bold", color="#1f4e79")
    ax.text(0.06, 0.8, "基礎形状        α        β", transform=ax.transAxes,
            fontproperties=jp, fontsize=9.5, fontweight="bold", family="monospace")
    for i, (nm, a, b) in enumerate(shapes):
        ax.text(0.06, 0.68 - i * 0.13, f"{nm:<8}    {a:<12} {b}", transform=ax.transAxes,
                fontproperties=jp, fontsize=9, family="monospace")
    ax.text(0.06, 0.1, "第1項に α、第2項に β を乗じる。\n正方形/円形は α 大（c項有利）、β 小。",
            transform=ax.transAxes, fontproperties=jp, fontsize=8.5, va="top", color="#333")
    ax.axis("off")
    ax.set_title("(a) 基礎形状による形状係数", fontproperties=jp, fontsize=9.5, fontweight="bold")
    # (b) 荷重傾斜・寸法効果
    ax = axes[1]
    ax.text(0.5, 0.95, "荷重傾斜補正・寸法効果", transform=ax.transAxes, ha="center",
            fontproperties=jp, fontsize=10.5, fontweight="bold", color="#1f4e79")
    # 傾斜荷重の図
    ax.annotate("", xy=(0.25, 0.55), xytext=(0.15, 0.8), xycoords="axes fraction",
                arrowprops=dict(arrowstyle="-|>", color="#c00000", lw=2.5))
    ax.text(0.05, 0.82, "傾斜荷重", transform=ax.transAxes, fontproperties=jp, fontsize=8.5, color="#c00000")
    ax.add_patch(mpatches.Rectangle((0.18, 0.48), 0.16, 0.06, transform=ax.transAxes, fc="#d9d9d9", ec="k"))
    ax.text(0.04, 0.34,
            "・荷重傾斜補正 ic・iγ・iq（<=1）\n"
            "  水平力があると支持力は低下\n"
            "  → 各項に i を乗じて低減\n\n"
            "・寸法効果\n"
            "  第2項 β・γ・B・Nγ は幅Bで増加\n"
            "  だが実地盤では頭打ち（低減）",
            transform=ax.transAxes, fontproperties=jp, fontsize=8.8, va="top")
    ax.axis("off")
    ax.set_title("(b) 荷重傾斜・寸法効果", fontproperties=jp, fontsize=9.5, fontweight="bold")
    fig.suptitle("図 3  形状係数・寸法効果・荷重傾斜補正", fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig3_correct.png")


def fig_compare():
    fig, ax = plt.subplots(figsize=(11.5, 5.2))
    labels = ["第1項\nα・c・Nc", "第2項\nβ・γ・B・Nγ", "第3項\nγ・Df・Nq"]
    sand = [0, 284, 608]
    clay = [370, 0, 26]
    x = np.arange(len(labels)); w = 0.36
    ax.bar(x - w / 2, sand, w, color="#e7c86a", ec="k", label="砂質土 (φ=30, c=0)")
    ax.bar(x + w / 2, clay, w, color="#8fb08f", ec="k", label="粘性土 (φ=0, c=50)")
    for xi, (s, c) in enumerate(zip(sand, clay)):
        ax.text(xi - w / 2, s + 12, f"{s}", ha="center", fontproperties=jp, fontsize=9)
        ax.text(xi + w / 2, c + 12, f"{c}", ha="center", fontproperties=jp, fontsize=9)
    ax.set_xticks(x); ax.set_xticklabels(labels, fontproperties=jp, fontsize=9)
    ax.set_ylabel("各項の支持力 (kN/m2)", fontproperties=jp, fontsize=9)
    ax.legend(prop=jp, fontsize=9); ax.grid(alpha=0.3, axis="y")
    ax.text(0.5, 0.84, "砂質土：qu=891（第2・3項が主）\n粘性土：qu=396（第1項が主）",
            transform=ax.transAxes, ha="center", va="top", fontproperties=jp, fontsize=9.5, color="#c00000")
    ax.set_title("図 4  砂質土と粘性土の各項寄与（B=2.0m, Df=1.5m, 正方形）",
                 fontproperties=jp, fontsize=12, fontweight="bold")
    return save(fig, "fig4_compare.png")


figs = {"formula": fig_formula(), "factors": fig_factors(),
        "correct": fig_correct(), "compare": fig_compare()}
print("figs:", list(figs.keys()))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
wb = Workbook()
C_TITLE = "1F4E79"; C_HEAD = "2E75B6"; C_ANS = "E2EFDA"; C_WARN = "FCE4D6"
thin = Side(style="thin", color="BFBFBF"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
f_title = Font(name="MS PGothic", size=15, bold=True, color="FFFFFF")
f_head = Font(name="MS PGothic", size=11, bold=True, color="FFFFFF")
f_body = Font(name="MS PGothic", size=10); f_ans = Font(name="MS PGothic", size=10, color="375623")
f_warn = Font(name="MS PGothic", size=9, color="833C00", italic=True)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)


def setup(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False


def title_row(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_title; c.fill = PatternFill("solid", fgColor=C_TITLE)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1); ws.row_dimensions[row].height = 30


def head(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_head; c.fill = PatternFill("solid", fgColor=C_HEAD)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1); ws.row_dimensions[row].height = 22


def body(ws, row, text, span=8, ans=False, h=None):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_ans if ans else f_body; c.alignment = wrap
    if ans:
        c.fill = PatternFill("solid", fgColor=C_ANS)
    if h:
        ws.row_dimensions[row].height = h


def warn(ws, row, text, span=8, h=None):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_warn; c.alignment = wrap
    c.fill = PatternFill("solid", fgColor=C_WARN)
    if h:
        ws.row_dimensions[row].height = h


def tbl(ws, start_row, headers, rows, col1=1):
    r = start_row
    for j, htxt in enumerate(headers):
        c = ws.cell(r, col1 + j, htxt); c.font = Font(name="MS PGothic", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C_HEAD); c.alignment = center; c.border = border
    for data in rows:
        r += 1
        for j, v in enumerate(data):
            c = ws.cell(r, col1 + j, v); c.font = f_body
            c.alignment = wrap if j == 0 else center; c.border = border
            if r % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F2F7FC")
    return r


def put_img(ws, path, anchor, w=None):
    img = XLImage(path)
    if w:
        ratio = w / img.width; img.width = w; img.height = int(img.height * ratio)
    ws.add_image(img, anchor)


# ===== 目次 =====
ws = wb.active; ws.title = "目次"; setup(ws, [4, 24, 54, 16])
title_row(ws, 1, "地盤の支持力（Terzaghi）問題集（RC マンション設計担当・新人向け）", span=4)
body(ws, 2, "目標：Terzaghi の支持力式の3項の意味を理解し、支持力係数・形状係数・補正を用いて"
            "砂質土・粘性土の極限支持力・許容支持力度を手計算できること。", span=4, h=32)
r = tbl(ws, 4, ["No.", "シート", "到達目標", "図"],
        [["1", "1 支持力式と3つの項", "第1〜3項の意味を説明できる", "式と項"],
         ["2", "2 支持力係数と地盤種別", "係数・地盤種別の目安を暗記", "係数/目安"],
         ["3", "3 形状・寸法・傾斜補正", "形状係数・寸法効果・傾斜補正", "補正"],
         ["4", "4 砂質土・粘性土の算定", "手計算で qu・qa を求める", "各項寄与"]])
warn(ws, r + 2, "※ 支持力係数（Nc・Nq・Nγ）、形状係数、地盤種別の目安値は文献・規準により異なる。"
                "設計値は告示1113号・建築基礎構造設計指針の最新版で確認すること。", span=4, h=40)

# ===== 1 支持力式 =====
ws = wb.create_sheet("1 支持力式と3項"); setup(ws, [8, 16, 18, 16, 12, 12])
title_row(ws, 1, "1  Terzaghi の支持力式と3つの項")
head(ws, 3, "■ 図 1  支持力式の構造"); put_img(ws, figs["formula"], "A4", w=840)
head(ws, 30, "■ 問題 1  3つの項の意味")
r = tbl(ws, 31, ["項", "式", "意味（記入）"],
        [["第1項", "α・c・Nc", "粘着力による抵抗（記入）"],
         ["第2項", "β・γ1・B・Nγ", "基礎幅・土の自重（記入）"],
         ["第3項", "γ2・Df・Nq", "根入れ（上載圧）（記入）"]])
head(ws, r + 2, "■ 問題 2  記号の意味")
body(ws, r + 3, "式中の記号（α・β：形状係数、c：粘着力、γ1：底面下の単位体積重量、"
                "γ2：根入れ部の単位体積重量、B：基礎幅、Df：根入れ深さ、Nc・Nγ・Nq：支持力係数）"
                "の意味を答えよ。", h=40)
head(ws, r + 5, "■ 問題 3  砂と粘土での主役")
body(ws, r + 6, "砂質土（c≒0）と粘性土（φ≒0）で、支持力式のどの項が主役になるか説明せよ"
                "（砂＝第2・3項、粘土＝第1項）。理由（c・φの大小、Nγ・Nqの値）も述べよ。", h=40)

# ===== 2 支持力係数と地盤種別 =====
ws = wb.create_sheet("2 係数と地盤種別"); setup(ws, [8, 16, 18, 16, 12, 12])
title_row(ws, 1, "2  支持力係数と地盤種別の目安")
head(ws, 3, "■ 図 2  支持力係数・地盤種別目安"); put_img(ws, figs["factors"], "A4", w=860)
head(ws, 27, "■ 問題 1  支持力係数とφ")
body(ws, 28, "支持力係数 Nc・Nq・Nγ が内部摩擦角φの関数であること、φが大きいほど急増することを述べよ。"
             "φ=0（粘土）のとき Nq=1・Nγ=0・Nc=5.7 となる意味を説明せよ。", h=40)
head(ws, 30, "■ 問題 2  地盤種別の目安")
r = tbl(ws, 31, ["地盤種別", "許容支持力度の目安 (kN/m2)"],
        [["硬岩", "1000 以上"], ["密な礫層", "500〜600"],
         ["密な砂質地盤", "200〜300"], ["中位の砂質地盤", "100〜200"],
         ["硬い粘性土", "100〜200"], ["軟らかい粘性土", "20〜50"]])
head(ws, r + 2, "■ 問題 3  目安の使い方")
body(ws, r + 3, "地盤種別の許容支持力度の『目安』を暗記しておく意義を述べよ"
                "（オーダー確認・基礎形式の一次判断）。ただし設計値は支持力式・載荷試験で"
                "確認することにも触れよ。", h=40)
warn(ws, r + 5, "※ 係数・目安値は規準で確認要。ここでの数値は概略把握のための例示。", h=24)

# ===== 3 形状・寸法・傾斜 =====
ws = wb.create_sheet("3 形状寸法傾斜"); setup(ws, [8, 16, 18, 16, 12, 12])
title_row(ws, 1, "3  形状係数・寸法効果・荷重傾斜補正")
head(ws, 3, "■ 図 3  形状係数・補正"); put_img(ws, figs["correct"], "A4", w=860)
head(ws, 27, "■ 問題 1  形状係数")
body(ws, 28, "連続基礎・正方形・円形・長方形の形状係数α・βの違いを述べよ。"
             "第1項にα、第2項にβを乗じる意味を説明せよ。", h=36)
head(ws, 30, "■ 問題 2  寸法効果")
body(ws, 31, "第2項 β・γ・B・Nγ が基礎幅Bに比例して増える一方、実地盤では支持力度が頭打ち"
             "（幅が大きいほど平均支持力度が下がる）になる『寸法効果』を説明せよ。", h=40)
head(ws, 33, "■ 問題 3  荷重傾斜補正")
body(ws, 34, "水平力を伴う傾斜荷重では、支持力が低下するため補正係数 ic・iγ・iq（<=1）を各項に乗じる。"
             "なぜ傾斜すると支持力が下がるか（すべり面の変化）を説明せよ。", h=40)
head(ws, 36, "■ 問題 4  補正の実務")
body(ws, 37, "実務で支持力を評価するとき、形状・根入れ・荷重傾斜・地下水位（γの水中重量）を"
             "どう考慮するか整理せよ。安全側の設定に触れよ。", h=40)

# ===== 4 算定例 =====
ws = wb.create_sheet("4 砂と粘土の算定"); setup(ws, [8, 16, 18, 16, 12, 12])
title_row(ws, 1, "4  砂質土・粘性土の支持力算定")
head(ws, 3, "■ 図 4  各項の寄与（砂 vs 粘土）"); put_img(ws, figs["compare"], "A4", w=760)
head(ws, 28, "■ 問題 1  砂質土の算定")
body(ws, 29, "正方形フーチング B=2.0m、Df=1.5m、φ=30°（Nc=37.2・Nq=22.5・Nγ=19.7）、c=0、"
             "γ1=γ2=18kN/m3、形状係数α=1.3・β=0.4。極限支持力 qu と許容支持力度 qa=qu/3 を求めよ。", h=48)
head(ws, 31, "■ 問題 2  粘性土の算定")
body(ws, 32, "同じ基礎で φ=0°（Nc=5.7・Nq=1.0・Nγ=0）、c=50kN/m2、γ=17kN/m3 のとき、"
             "qu と qa=qu/3 を求めよ。どの項が支配的か述べよ。", h=40)
head(ws, 34, "■ 問題 3  砂と粘土の比較")
body(ws, 35, "問1・問2の結果を比較し、砂質土と粘性土で支配項が異なる理由を述べよ。"
             "基礎幅Bを大きくしたとき、砂・粘土でquがどう変わるか（第2項の効き方）も述べよ。", h=40)
head(ws, 37, "■ 問題 4  許容支持力度と設計")
body(ws, 38, "求めた qa を接地圧 q（前教材の上下分離モデル参照）と比較し、q<=qa を確認する流れを述べよ。"
             "安全率Fs（長期3・短期1.5程度）の考え方に触れよ。", h=40)
warn(ws, 40, "※ 支持力係数・安全率・形状係数は規準（告示1113号/建築基礎指針）で確認要。", h=24)

# ===== 解答 =====
ws = wb.create_sheet("解答"); setup(ws, [4, 22, 60, 14])
title_row(ws, 1, "解答・解説", span=4); row = 2


def ah(t):
    global row; head(ws, row, t, span=4); row += 1


def an(t, h=None):
    global row; body(ws, row, t, span=4, ans=True, h=h); row += 1


def aw(t, h=None):
    global row; warn(ws, row, t, span=4, h=h); row += 1


ah("1  支持力式と3項")
an("問1：第1項 α・c・Nc＝粘着力による支持（粘性土で主役）。第2項 β・γ1・B・Nγ＝基礎幅と"
   "底面下の土の自重による支持（砂質土で効く）。第3項 γ2・Df・Nq＝根入れ部の上載圧による支持"
   "（根入れが深いほど大）。", h=44)
an("問2：α・β＝基礎形状の係数、c＝粘着力、γ1＝基礎底面下の単位体積重量、"
   "γ2＝根入れ部の単位体積重量、B＝基礎幅、Df＝根入れ深さ、Nc・Nγ・Nq＝φで決まる支持力係数。", h=40)
an("問3：砂質土はc≒0なので第1項が消え、φが大きくNγ・Nqが大きいため第2・3項が主役。"
   "粘性土（非排水）はφ≒0でNγ=0・Nq=1のため第2項が消え、Ncが効く第1項が主役。", h=40)

ah("2  係数と地盤種別")
an("問1：Nc・Nq・Nγはφの関数でφとともに急増（指数的）。φ=0ではせん断抵抗が粘着のみで、"
   "Nq=1（上載圧そのまま）・Nγ=0（幅の効果なし）・Nc=5.7。砂ではφが大きく各係数が大きい。", h=44)
an("問2：（目安）硬岩1000以上、密な礫層500〜600、密な砂200〜300、中位の砂100〜200、"
   "硬い粘土100〜200、軟らかい粘土20〜50 kN/m2。オーダー感を持つことが重要。", h=40)
an("問3：目安を覚えておくと、地盤種別から支持力度のオーダーを即座に把握でき、"
   "基礎形式（直接/杭）の一次判断や検算に使える。ただし設計値は支持力式・平板載荷試験で確認する。", h=40)
aw("※ 係数・目安値は規準で確認要。", h=22)

ah("3  形状・寸法・傾斜")
an("問1：連続α=1.0・β=0.5、正方形α=1.3・β=0.4、円形α=1.3・β=0.3、"
   "長方形α=1.0+0.3B/L・β=0.5-0.1B/L。第1項（粘着）にα、第2項（幅）にβを乗じ、"
   "基礎形状による3次元的な支持効果を補正する。", h=44)
an("問2：第2項はBに比例して増えるが、実地盤ではBが大きいほど平均支持力度が頭打ち・低減する"
   "（応力の及ぶ深さが増え軟弱層の影響、進行性破壊等）。大きな基礎ほど支持力度を安全側に見る。", h=44)
an("問3：傾斜荷重（水平力併存）ではすべり面が浅く非対称になり抵抗が減るため、ic・iγ・iq（<=1）"
   "を乗じて低減する。水平力が大きいほど補正係数は小さくなる。", h=40)
an("問4：形状（α・β）、根入れ（第3項・Df）、荷重傾斜（i係数）、地下水位（水中はγ'=γsat-γwを使用）"
   "を考慮する。不確実性が大きいので安全率を確保し安全側に評価する。", h=40)

ah("4  砂と粘土の算定")
an("問1（砂）：qu=α・c・Nc+β・γ1・B・Nγ+γ2・Df・Nq"
   "=1.3×0×37.2 + 0.4×18×2.0×19.7 + 18×1.5×22.5 = 0+284+608 = 891kN/m2。qa=891/3≒297kN/m2。", h=44)
an("問2（粘土）：qu=1.3×50×5.7 + 0.4×17×2.0×0 + 17×1.5×1.0 = 370+0+26 = 396kN/m2。"
   "qa=396/3≒132kN/m2。第1項（粘着 c・Nc）が支配的。", h=40)
an("問3：砂は第2・3項（幅・根入れ）、粘土は第1項（粘着）が支配。"
   "Bを大きくすると砂は第2項が増えquが増加するが、粘土は第2項=0なのでほぼ変わらない"
   "（粘土の支持力は基礎幅にほぼ依存しない）。", h=44)
an("問4：接地圧 q（＝上部反力＋自重を底面積で割った分布の最大値）が qa 以下（q<=qa）であることを確認。"
   "満たさなければ基礎拡大・杭基礎へ。安全率Fsは長期3・短期1.5程度（規準で確認）。", h=44)
aw("※ 支持力係数・安全率・形状係数は規準（告示1113号/建築基礎指針）で確認要。", h=22)

XLSX = os.path.join(OUT, "地盤の支持力Terzaghi問題集.xlsx"); wb.save(XLSX); print("saved:", XLSX)
