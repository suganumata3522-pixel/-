# -*- coding: utf-8 -*-
"""各地盤定数 問題集（図つき）Excel。出力: docs/soil_params/各地盤定数問題集.xlsx
1 地盤定数の一覧（定義・調査・用途）/ 2 せん断強度（φ・C・qu）
3 変形・物理・圧密（Eo・FC・ρ・pc）/ 4 N値換算式（暗記）
※ 換算式は目安。設計値は規準（建築基礎指針/道示）で確認要
"""
import os
import math
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib import font_manager as fm
FONT = "/usr/share/fonts/truetype/fonts-japanese-gothic.ttf"
jp = fm.FontProperties(fname=FONT); fm.fontManager.addfont(FONT)
plt.rcParams["font.family"] = jp.get_name(); plt.rcParams["axes.unicode_minus"] = False
OUT = "docs/soil_params"; FIG = os.path.join(OUT, "figures"); os.makedirs(FIG, exist_ok=True)


def save(fig, n):
    p = os.path.join(FIG, n); fig.savefig(p, dpi=130, bbox_inches="tight"); plt.close(fig); return p


def fig_map():
    fig, ax = plt.subplots(figsize=(12.5, 5.6))
    # 中央：N値・各種試験
    ax.add_patch(mpatches.FancyBboxPatch((0.38, 0.42), 0.24, 0.16, boxstyle="round,pad=0.01",
                 transform=ax.transAxes, fc="#2e75b6", ec="k", lw=1.2))
    ax.text(0.5, 0.5, "地盤調査\n（原位置試験・室内試験）", transform=ax.transAxes, ha="center",
            va="center", fontproperties=jp, fontsize=9.5, fontweight="bold", color="white")
    params = [
        (0.10, 0.85, "N値", "標準貫入試験", "支持力・液状化・換算"),
        (0.42, 0.88, "φ（内部摩擦角）", "三軸/N値換算", "砂のせん断強度・支持力・土圧"),
        (0.74, 0.85, "C（粘着力）", "一軸/三軸", "粘土のせん断強度・支持力"),
        (0.06, 0.12, "qu（一軸圧縮強さ）", "一軸圧縮試験", "C=qu/2・許容支持力"),
        (0.34, 0.06, "Eo（変形係数）", "孔内水平載荷/N換算", "沈下・杭の水平ばね"),
        (0.62, 0.06, "FC（細粒分含有率）", "粒度試験", "液状化判定・土質分類"),
        (0.86, 0.12, "ρ・γ（密度）", "物理試験", "土被り圧・土圧・液状化"),
    ]
    for x, y, name, method, use in params:
        ax.add_patch(mpatches.FancyBboxPatch((x - 0.005, y - 0.02), 0.2, 0.13,
                     boxstyle="round,pad=0.008", transform=ax.transAxes, fc="#e8f0fa", ec="#2e75b6", lw=1))
        ax.text(x + 0.095, y + 0.075, name, transform=ax.transAxes, ha="center",
                fontproperties=jp, fontsize=8.5, fontweight="bold", color="#1f4e79")
        ax.text(x + 0.095, y + 0.02, f"調査:{method}\n用途:{use}", transform=ax.transAxes, ha="center",
                va="center", fontproperties=jp, fontsize=6.8, color="#444")
        ax.annotate("", xy=(x + 0.095, y), xytext=(0.5, 0.5), xycoords="axes fraction",
                    arrowprops=dict(arrowstyle="-", color="#9dc3e6", lw=0.8), zorder=0)
    # pc は別枠
    ax.add_patch(mpatches.FancyBboxPatch((0.40, 0.7), 0.2, 0.11,
                 boxstyle="round,pad=0.008", transform=ax.transAxes, fc="#fde2c4", ec="#c55a11", lw=1))
    ax.text(0.5, 0.79, "pc（圧密降伏応力）", transform=ax.transAxes, ha="center",
            fontproperties=jp, fontsize=8.5, fontweight="bold", color="#833c00")
    ax.text(0.5, 0.735, "調査:圧密試験  用途:圧密沈下", transform=ax.transAxes, ha="center",
            fontproperties=jp, fontsize=6.8, color="#833c00")
    ax.axis("off")
    ax.set_title("図 1  主要な地盤定数と調査方法・用途の関係",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig1_map.png")


def fig_shear():
    fig, axes = plt.subplots(1, 2, figsize=(13.0, 4.8))
    ax = axes[0]
    s = np.linspace(0, 200, 100)
    ax.plot(s, 0 + s * math.tan(math.radians(30)), color="#c00000", lw=2, label="砂質土 (c≒0, φ=30°)")
    ax.plot(s, 40 + s * 0, color="#2a78d6", lw=2, label="粘性土 (φ≒0, c=一定)")
    ax.plot(s, 20 + s * math.tan(math.radians(20)), color="#548235", lw=2, ls="--",
            label="c-φ 材 (c,φ)")
    ax.text(150, 150 * math.tan(math.radians(30)) + 5, "φ", fontproperties=jp, fontsize=11, color="#c00000")
    ax.annotate("c（粘着力）", xy=(0, 40), xytext=(35, 70), fontproperties=jp, fontsize=9,
                color="#2a78d6", arrowprops=dict(arrowstyle="->", color="#2a78d6"))
    ax.set_xlabel("垂直応力 σ (kN/m2)", fontproperties=jp, fontsize=9)
    ax.set_ylabel("せん断強さ τ (kN/m2)", fontproperties=jp, fontsize=9)
    ax.set_title("(a) モール・クーロンの破壊基準 τ=c+σtanφ", fontproperties=jp, fontsize=9.5, fontweight="bold")
    ax.legend(prop=jp, fontsize=8, loc="upper left"); ax.grid(alpha=0.3)
    ax.set_xlim(0, 200); ax.set_ylim(0, 140)
    ax = axes[1]
    # 一軸圧縮：応力-ひずみでqu、モール円でC=qu/2
    ax.text(0.5, 0.93, "一軸圧縮試験（粘性土）", transform=ax.transAxes, ha="center",
            fontproperties=jp, fontsize=10, fontweight="bold", color="#1f4e79")
    eps = np.linspace(0, 8, 100)
    sig = 120 * (1 - np.exp(-eps / 1.5)) * np.exp(-eps / 30)
    ax.plot(eps, sig, color="#c00000", lw=2, transform=ax.transData)
    ax.axhline(max(sig), color="#888", ls=":", lw=1)
    ax.text(5.5, max(sig) + 3, "qu（一軸圧縮強さ）", fontproperties=jp, fontsize=9, color="#c00000")
    ax.text(4.5, 30, "C = qu / 2\n（非排水せん断強さ）", fontproperties=jp, fontsize=9.5, color="#1f4e79")
    ax.set_xlabel("ひずみ ε (%)", fontproperties=jp, fontsize=9)
    ax.set_ylabel("圧縮応力 σ (kN/m2)", fontproperties=jp, fontsize=9)
    ax.set_xlim(0, 8); ax.set_ylim(0, 140); ax.grid(alpha=0.3)
    ax.set_title("(b) qu と C の関係", fontproperties=jp, fontsize=9.5, fontweight="bold")
    fig.suptitle("図 2  せん断強度定数（φ・C・qu）", fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig2_shear.png")


def fig_other():
    fig, axes = plt.subplots(1, 3, figsize=(14.5, 4.4))
    # (a) 圧密 e-logP
    ax = axes[0]
    P = np.array([10, 30, 60, 100, 150, 200, 400, 800])
    e = np.array([1.20, 1.19, 1.17, 1.14, 1.06, 0.98, 0.86, 0.74])
    ax.semilogx(P, e, color="#c00000", lw=2, marker="o", ms=3)
    ax.axvline(120, color="#548235", ls="--", lw=1.4)
    ax.text(125, 1.15, "pc（圧密降伏応力）", fontproperties=jp, fontsize=8, color="#548235", rotation=90, va="top")
    ax.text(15, 1.05, "過圧密\n（緩やか）", fontproperties=jp, fontsize=8, color="#555")
    ax.text(300, 0.95, "正規圧密\n（急）", fontproperties=jp, fontsize=8, color="#555")
    ax.set_xlabel("圧密圧力 P (kN/m2)", fontproperties=jp, fontsize=8.5)
    ax.set_ylabel("間隙比 e", fontproperties=jp, fontsize=8.5)
    ax.set_title("(a) 圧密試験 e-logP（pc）", fontproperties=jp, fontsize=9, fontweight="bold")
    ax.grid(alpha=0.3, which="both")
    # (b) 応力-ひずみ Eo
    ax = axes[1]
    eps = np.linspace(0, 1.0, 100); sig = 8000 * eps * (1 - 0.3 * eps)
    ax.plot(eps, sig, color="#2a78d6", lw=2)
    ax.plot([0, 0.5], [0, 8000 * 0.5 * (1 - 0.3 * 0.5)], color="#c00000", lw=1.4, ls="--")
    ax.text(0.52, 2600, "Eo=Δσ/Δε\n（割線変形係数）", fontproperties=jp, fontsize=8, color="#c00000")
    ax.set_xlabel("ひずみ ε", fontproperties=jp, fontsize=8.5)
    ax.set_ylabel("応力 σ (kN/m2)", fontproperties=jp, fontsize=8.5)
    ax.set_title("(b) 変形係数 Eo", fontproperties=jp, fontsize=9, fontweight="bold")
    ax.grid(alpha=0.3); ax.set_xlim(0, 1.05)
    # (c) 粒度加積曲線 FC
    ax = axes[2]
    d = np.array([0.001, 0.005, 0.02, 0.075, 0.25, 0.85, 2, 4.75, 19])
    passing = np.array([5, 12, 22, 35, 55, 75, 88, 96, 100])
    ax.semilogx(d, passing, color="#548235", lw=2, marker="o", ms=3)
    ax.axvline(0.075, color="#c00000", ls="--", lw=1.4)
    ax.axhline(35, color="#c00000", ls=":", lw=1)
    ax.text(0.002, 40, "FC=35%\n（0.075mm通過率）", fontproperties=jp, fontsize=8, color="#c00000")
    ax.set_xlabel("粒径 (mm)", fontproperties=jp, fontsize=8.5)
    ax.set_ylabel("通過質量百分率 (%)", fontproperties=jp, fontsize=8.5)
    ax.set_title("(c) 粒度加積曲線（FC）", fontproperties=jp, fontsize=9, fontweight="bold")
    ax.grid(alpha=0.3, which="both"); ax.set_ylim(0, 100)
    fig.suptitle("図 3  変形・物理・圧密（Eo・FC・ρ・pc）", fontproperties=jp, fontsize=13, fontweight="bold")
    fig.subplots_adjust(wspace=0.35)
    return save(fig, "fig3_other.png")


def fig_conv():
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 5.0), gridspec_kw={"width_ratios": [1.1, 1.0]})
    ax = axes[0]
    ax.text(0.5, 0.96, "N値換算式（暗記）", transform=ax.transAxes, ha="center",
            fontproperties=jp, fontsize=11, fontweight="bold", color="#1f4e79")
    ax.text(0.03, 0.85,
            "【砂質土 内部摩擦角 φ】\n"
            "  大崎式：φ = √(20N) + 15  (度)\n"
            "  Dunham：φ = √(12N) + 15〜25\n\n"
            "【粘性土 一軸圧縮強さ qu・粘着力 C】\n"
            "  qu ≒ 12.5 N  (kN/m2)\n"
            "  C = qu/2 ≒ 6.25 N  (kN/m2)\n\n"
            "【変形係数 Eo】\n"
            "  Eo ≒ 700 N  (kN/m2)  ※式により2800N等\n\n"
            "【せん断波速度 Vs（道示）】\n"
            "  砂 Vs=80・N^(1/3)、粘性土 Vs=100・N^(1/3) (m/s)",
            transform=ax.transAxes, fontproperties=jp, fontsize=9, va="top")
    ax.text(0.03, 0.06, "※ いずれも目安。設計値は規準・室内試験で確認要。",
            transform=ax.transAxes, fontproperties=jp, fontsize=8, color="#833c00")
    ax.axis("off")
    ax.set_title("(a) 換算式一覧", fontproperties=jp, fontsize=10, fontweight="bold")
    ax = axes[1]
    N = np.arange(2, 41)
    phi = np.sqrt(20 * N) + 15
    qu = 12.5 * N
    l1, = ax.plot(N, phi, color="#c00000", lw=2, label="φ=√(20N)+15 (°)")
    ax.set_xlabel("N値", fontproperties=jp, fontsize=9)
    ax.set_ylabel("φ (度)", fontproperties=jp, fontsize=9, color="#c00000")
    ax.tick_params(axis="y", labelcolor="#c00000")
    ax.set_ylim(15, 55)
    ax2 = ax.twinx()
    l2, = ax2.plot(N, qu, color="#2a78d6", lw=2, ls="--", label="qu=12.5N (kN/m2)")
    ax2.set_ylabel("qu (kN/m2)", fontproperties=jp, fontsize=9, color="#2a78d6")
    ax2.tick_params(axis="y", labelcolor="#2a78d6")
    ax.legend(handles=[l1, l2], prop=jp, fontsize=8, loc="upper left")
    ax.grid(alpha=0.3)
    ax.set_title("(b) N→φ・qu の換算", fontproperties=jp, fontsize=10, fontweight="bold")
    fig.suptitle("図 4  N値換算式（砂→φ、粘土→qu・C、Eo）",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig4_conv.png")


figs = {"map": fig_map(), "shear": fig_shear(), "other": fig_other(), "conv": fig_conv()}
print("figs:", list(figs.keys()))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
wb = Workbook()
C_TITLE = "1F4E79"; C_HEAD = "2E75B6"; C_ANS = "E2EFDA"; C_WARN = "FCE4D6"
thin = Side(style="thin", color="BFBFBF"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
f_title = Font(name="MS PGothic", size=15, bold=True, color="FFFFFF")
f_head = Font(name="MS PGothic", size=11, bold=True, color="FFFFFF")
f_body = Font(name="MS PGothic", size=10); f_ans = Font(name="MS PGothic", size=10, color="375623")
f_warn = Font(name="MS PGothic", size=9, color="833C00", italic=True)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)


def setup(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False


def title_row(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_title; c.fill = PatternFill("solid", fgColor=C_TITLE)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1); ws.row_dimensions[row].height = 30


def head(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_head; c.fill = PatternFill("solid", fgColor=C_HEAD)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1); ws.row_dimensions[row].height = 22


def body(ws, row, text, span=8, ans=False, h=None):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_ans if ans else f_body; c.alignment = wrap
    if ans:
        c.fill = PatternFill("solid", fgColor=C_ANS)
    if h:
        ws.row_dimensions[row].height = h


def warn(ws, row, text, span=8, h=None):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_warn; c.alignment = wrap
    c.fill = PatternFill("solid", fgColor=C_WARN)
    if h:
        ws.row_dimensions[row].height = h


def tbl(ws, start_row, headers, rows, col1=1):
    r = start_row
    for j, htxt in enumerate(headers):
        c = ws.cell(r, col1 + j, htxt); c.font = Font(name="MS PGothic", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C_HEAD); c.alignment = center; c.border = border
    for data in rows:
        r += 1
        for j, v in enumerate(data):
            c = ws.cell(r, col1 + j, v); c.font = f_body
            c.alignment = wrap if j == 0 else center; c.border = border
            if r % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F2F7FC")
    return r


def put_img(ws, path, anchor, w=None):
    img = XLImage(path)
    if w:
        ratio = w / img.width; img.width = w; img.height = int(img.height * ratio)
    ws.add_image(img, anchor)


# ===== 目次 =====
ws = wb.active; ws.title = "目次"; setup(ws, [4, 22, 40, 14, 14])
title_row(ws, 1, "各地盤定数 問題集（RC マンション設計担当・新人向け）", span=5)
body(ws, 2, "目標：N値・φ・C・qu・Eo・FC・ρ・pc の定義・調査方法・用途を理解し、"
            "N値換算式を暗記して概略値を求められること。地盤定数は設計の入力値そのもの。", span=5, h=32)
r = tbl(ws, 4, ["No.", "シート", "到達目標", "対象定数"],
        [["1", "1 地盤定数の一覧", "定義・調査方法・用途を対応づけ", "全定数"],
         ["2", "2 せん断強度", "φ・C・qu の意味を理解", "φ・C・qu"],
         ["3", "3 変形・物理・圧密", "Eo・FC・ρ・pc を理解", "Eo・FC・ρ・pc"],
         ["4", "4 N値換算式", "換算式を暗記し概算できる", "N→φ,qu,Eo"]])
warn(ws, r + 2, "※ 換算式（大崎・Dunham・qu=12.5N・Eo=700N 等）は概略の目安。"
                "設計に用いる定数は室内試験・規準（建築基礎構造設計指針/道示）で確認すること。", span=5, h=40)

# ===== 1 一覧 =====
ws = wb.create_sheet("1 地盤定数の一覧"); setup(ws, [8, 14, 18, 18, 20, 12])
title_row(ws, 1, "1  主要な地盤定数の定義・調査方法・用途")
head(ws, 3, "■ 図 1  地盤定数マップ"); put_img(ws, figs["map"], "A4", w=820)
head(ws, 30, "■ 問題 1  定義・調査・用途の対応")
r = tbl(ws, 31, ["定数", "定義", "主な調査方法", "主な用途"],
        [["N値", "標準貫入試験の打撃回数", "標準貫入試験(SPT)", "硬さ・支持力・液状化・換算"],
         ["φ", "内部摩擦角（記入）", "（記入）", "（記入）"],
         ["C", "粘着力（記入）", "（記入）", "（記入）"],
         ["qu", "一軸圧縮強さ（記入）", "（記入）", "（記入）"],
         ["Eo", "変形係数（記入）", "（記入）", "（記入）"],
         ["FC", "細粒分含有率（記入）", "（記入）", "（記入）"],
         ["ρ・γ", "密度・単位体積重量（記入）", "（記入）", "（記入）"],
         ["pc", "圧密降伏応力（記入）", "（記入）", "（記入）"]])
body(ws, r + 2, "各定数が『どの試験で得られ、何の設計に使うか』を対応づけて覚える。"
                "地盤定数は支持力・沈下・土圧・液状化すべての入力値になる。", h=32)

# ===== 2 せん断強度 =====
ws = wb.create_sheet("2 せん断強度"); setup(ws, [8, 16, 18, 16, 12, 12])
title_row(ws, 1, "2  せん断強度定数（φ・C・qu）")
head(ws, 3, "■ 図 2  せん断強度"); put_img(ws, figs["shear"], "A4", w=840)
head(ws, 27, "■ 問題 1  モール・クーロン")
body(ws, 28, "破壊基準 τ=c+σtanφ の意味を説明せよ。砂質土（c≒0）と粘性土（φ≒0、非排水）"
             "でせん断強度がどう表されるか述べよ。", h=40)
head(ws, 30, "■ 問題 2  qu と C")
body(ws, 31, "一軸圧縮試験で得られる qu と粘着力 C の関係（C=qu/2）を説明せよ。"
             "なぜ 1/2 になるか（φ≒0 のモール円の半径＝せん断強さ）に触れよ。", h=40)
head(ws, 33, "■ 問題 3  調査方法")
body(ws, 34, "φ・C を求める試験（三軸圧縮・一軸圧縮・直接せん断）を挙げ、"
             "砂質土・粘性土それぞれで用いる試験を述べよ。", h=32)
head(ws, 36, "■ 問題 4  用途")
body(ws, 37, "せん断強度定数（φ・C）が使われる設計を挙げよ（支持力、土圧、斜面安定、"
             "杭周面摩擦）。強度が支持力・土圧にどう効くか概説せよ。", h=32)

# ===== 3 変形・物理・圧密 =====
ws = wb.create_sheet("3 変形物理圧密"); setup(ws, [8, 16, 18, 16, 12, 12])
title_row(ws, 1, "3  変形・物理・圧密（Eo・FC・ρ・pc）")
head(ws, 3, "■ 図 3  Eo・FC・pc"); put_img(ws, figs["other"], "A4", w=880)
head(ws, 26, "■ 問題 1  変形係数 Eo")
body(ws, 27, "変形係数 Eo の定義（応力-ひずみの割線勾配）と用途（沈下量・杭の水平地盤反力）を"
             "述べよ。調査方法（孔内水平載荷試験・N値換算 Eo≒700N）にも触れよ。", h=40)
head(ws, 29, "■ 問題 2  細粒分含有率 FC")
body(ws, 30, "FC（粒径0.075mm未満の質量百分率）の定義と用途（液状化判定・土質分類）を述べよ。"
             "粒度加積曲線からFCを読む方法を説明せよ。", h=40)
head(ws, 32, "■ 問題 3  密度 ρ・単位体積重量 γ")
body(ws, 33, "湿潤密度・飽和密度の違い、有効単位体積重量 γ'=γsat-γw（水中重量）を説明せよ。"
             "土被り圧・有効応力の計算にどう使うか述べよ。", h=40)
head(ws, 35, "■ 問題 4  圧密降伏応力 pc")
body(ws, 36, "圧密試験の e-logP 曲線から pc を求める意味を説明せよ。"
             "過圧密（現在の有効応力<pc）と正規圧密の違い、圧密沈下との関係を述べよ。", h=40)

# ===== 4 N値換算式 =====
ws = wb.create_sheet("4 N値換算式"); setup(ws, [8, 16, 18, 16, 12, 12])
title_row(ws, 1, "4  N値換算式（暗記）")
head(ws, 3, "■ 図 4  N値換算式と換算グラフ"); put_img(ws, figs["conv"], "A4", w=840)
head(ws, 28, "■ 問題 1  φ の換算")
body(ws, 29, "砂質土 N=15 のとき、大崎式 φ=√(20N)+15 で内部摩擦角を求めよ。", h=28)
head(ws, 31, "■ 問題 2  qu・C の換算")
body(ws, 32, "粘性土 N=5 のとき、qu≒12.5N（kN/m2）と C=qu/2 を求めよ。", h=28)
head(ws, 34, "■ 問題 3  Eo の換算")
body(ws, 35, "N=10 のとき、変形係数 Eo≒700N（kN/m2）を求めよ。杭の水平ばねに使うことに触れよ。", h=28)
head(ws, 37, "■ 問題 4  換算値の注意")
body(ws, 38, "N値換算はあくまで概略の目安であり、実設計では室内試験・原位置試験で確認する理由を述べよ"
             "（式の適用範囲、土質・粒度・応力状態による差、換算式が複数あること）。", h=40)
warn(ws, 40, "※ 換算式・係数は規準（建築基礎構造設計指針/道路橋示方書）の最新版で確認。", h=24)

# ===== 解答 =====
ws = wb.create_sheet("解答"); setup(ws, [4, 22, 60, 14])
title_row(ws, 1, "解答・解説", span=4); row = 2


def ah(t):
    global row; head(ws, row, t, span=4); row += 1


def an(t, h=None):
    global row; body(ws, row, t, span=4, ans=True, h=h); row += 1


def aw(t, h=None):
    global row; warn(ws, row, t, span=4, h=h); row += 1


ah("1  地盤定数の一覧")
an("問1：φ＝内部摩擦角（砂のせん断強度／三軸・N値換算／支持力・土圧）。"
   "C＝粘着力（粘土の強度／一軸・三軸／支持力・土圧）。qu＝一軸圧縮強さ（一軸／C=qu/2）。"
   "Eo＝変形係数（孔内水平載荷・N換算／沈下・杭水平ばね）。FC＝細粒分含有率（粒度／液状化・分類）。"
   "ρ・γ＝密度・単位体積重量（物理試験／土被り圧・土圧）。pc＝圧密降伏応力（圧密試験／圧密沈下）。", h=60)

ah("2  せん断強度")
an("問1：τ=c+σtanφ。せん断強さは粘着力cと、垂直応力σに比例する摩擦成分σtanφの和。"
   "砂質土はc≒0でτ=σtanφ（摩擦材）、粘性土（非排水）はφ≒0でτ=c（粘着材）。", h=44)
an("問2：一軸圧縮試験（側圧0）でqu。非排水（φ≒0）ではモール円の半径＝せん断強さ＝qu/2。"
   "よってC=qu/2。粘性土の簡便な強度指標として使われる。", h=40)
an("問3：三軸圧縮試験（c,φ・応力状態を再現、砂・粘土とも）、一軸圧縮（粘性土のqu）、"
   "直接せん断/一面せん断。砂は原位置N値・三軸、粘土は一軸・三軸を用いることが多い。", h=44)
an("問4：支持力（Terzaghi式のc,φ）、土圧（主働・受働のKa,Kp）、斜面安定、杭の周面摩擦。"
   "cやφが大きいほど支持力・受働土圧は大きく、主働土圧は小さくなる。", h=40)

ah("3  変形・物理・圧密")
an("問1：Eo＝応力-ひずみ曲線の割線勾配（Δσ/Δε）。沈下量の算定や杭の水平地盤反力係数の"
   "算定に使う。孔内水平載荷試験で直接測定、またはN値換算 Eo≒700N（kN/m2）。", h=44)
an("問2：FC＝粒径0.075mm未満（細粒分）の質量百分率。粒度加積曲線で0.075mmの通過率を読む。"
   "液状化判定（FCが大きいと液状化しにくい）・土質分類に用いる。", h=40)
an("問3：湿潤密度は自然状態、飽和密度は間隙が水で満たされた状態。水中では浮力を受けるため"
   "有効単位体積重量 γ'=γsat-γw（γw≒9.8kN/m3）を用いる。地下水位以下の有効応力計算に使う。", h=44)
an("問4：pcは過去に受けた最大有効応力（圧密降伏応力）。現在の有効応力<pcなら過圧密（沈下小）、"
   "≒pcで正規圧密（載荷で大きく沈下）。pcと増加応力の関係で圧密沈下量を評価する。", h=44)

ah("4  N値換算式")
an("問1：φ=√(20×15)+15=√300+15=17.3+15≒32.3°。（大崎式・砂質土）", h=28)
an("問2：qu≒12.5×5=62.5kN/m2。C=qu/2=31.25kN/m2。（Terzaghi-Peck・粘性土の目安）", h=28)
an("問3：Eo≒700×10=7000kN/m2（=7.0MN/m2）。杭の水平地盤反力係数kh（→水平ばね）の算定に使う。"
   "（式により2800N等もあり、確認要）", h=32)
an("問4：換算式は土質・粒度・応力状態・拘束圧で変わり、式も複数ある（大崎・Dunham等）。"
   "概略把握には有効だが、設計値は室内試験・原位置試験と規準で確認する。過信は禁物。", h=44)
aw("※ 換算式・係数は規準（建築基礎構造設計指針/道示）の最新版で確認。", h=24)

XLSX = os.path.join(OUT, "各地盤定数問題集.xlsx"); wb.save(XLSX); print("saved:", XLSX)
