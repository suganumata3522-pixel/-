# -*- coding: utf-8 -*-
"""
年最高外気温度の極値統計解析ブック（再現期間100年）を生成するスクリプト。

出力: 年最高外気温_ワイブル極値解析.xlsx

すべての計算は Excel の数式として書き込む（Python 側で値を焼き込まない）。
気象庁の年最高気温データを「データ入力」シートに貼り替えれば自動で再計算される。
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.marker import Marker
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.comments import Comment
from openpyxl.worksheet.properties import PageSetupProperties

# ---------------------------------------------------------------- 定数
FONT = "Yu Gothic"                 # 日本語が確実に表示される汎用ゴシック
NDATA = 200                        # 入力できる年数（行）
DROW0 = 12                         # データ入力シートの先頭データ行
DROW1 = DROW0 + NDATA - 1          # 211
RROW0 = 7                          # 極値統計計算シートの順位表 先頭行
RROW1 = RROW0 + NDATA - 1          # 206
NGRID = 200                        # ω 探索の候補数（各段）
GROW0 = 7
GROW1 = GROW0 + NGRID - 1          # 206
UROW0 = 7                          # 適合曲線用 u グリッド 先頭行
NUGRID = 41
UROW1 = UROW0 + NUGRID - 1         # 47

RNG_T = "データ入力!$B$%d:$B$%d" % (DROW0, DROW1)
RNG_Y = "データ入力!$A$%d:$A$%d" % (DROW0, DROW1)

# 動作確認用のダミー標本（気象庁の実測値ではない）。
# グンベル分布に緩やかな昇温トレンドを与えて生成した 75 年分。
SAMPLE_START = 1950
SAMPLE = [36.0, 33.3, 33.0, 34.7, 34.9, 34.4, 35.6, 32.7, 33.2, 33.0,
          35.0, 35.5, 33.8, 32.9, 35.3, 33.5, 33.1, 33.6, 33.5, 34.2,
          35.7, 35.1, 36.9, 34.9, 34.7, 33.4, 35.0, 34.9, 38.5, 32.4,
          34.2, 35.3, 33.9, 39.4, 33.8, 34.6, 34.8, 34.2, 37.6, 34.7,
          35.5, 35.2, 36.8, 36.8, 36.3, 35.5, 33.1, 36.4, 36.1, 36.1,
          35.1, 35.3, 34.2, 36.5, 34.3, 37.2, 35.8, 35.0, 36.9, 39.1,
          34.7, 34.4, 34.0, 36.2, 36.0, 34.6, 33.6, 34.6, 34.9, 35.9,
          35.6, 36.9, 35.3, 37.5, 38.3]

RETURN_PERIODS = [2, 5, 10, 20, 25, 30, 50, 75, 100, 150, 200]
TBL_ROW0 = 14                                  # 結果シート 再現期間表の先頭データ行
ROW_T100 = TBL_ROW0 + RETURN_PERIODS.index(100)  # 結果シートで T=100 の行

# ---------------------------------------------------------------- 書式
C_TITLE = "1F3864"
C_HEAD = "2E5C8A"
C_BAND = "DCE6F1"
C_INPUT = "FFF2CC"
C_RESULT = "FCE4D6"
C_NOTE = "F2F2F2"

f_title = Font(name=FONT, size=16, bold=True, color="FFFFFF")
f_h1 = Font(name=FONT, size=11, bold=True, color="FFFFFF")
f_h2 = Font(name=FONT, size=11, bold=True, color=C_TITLE)
f_base = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_small = Font(name=FONT, size=9, color="595959")
f_warn = Font(name=FONT, size=10, bold=True, color="C00000")
f_big = Font(name=FONT, size=24, bold=True, color="C00000")
f_input = Font(name=FONT, size=10, color="0000FF")

fill_title = PatternFill("solid", fgColor=C_TITLE)
fill_head = PatternFill("solid", fgColor=C_HEAD)
fill_band = PatternFill("solid", fgColor=C_BAND)
fill_input = PatternFill("solid", fgColor=C_INPUT)
fill_result = PatternFill("solid", fgColor=C_RESULT)
fill_note = PatternFill("solid", fgColor=C_NOTE)

thin = Side(style="thin", color="A6A6A6")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
ctr = Alignment(horizontal="center", vertical="center")


def put(ws, ref, value, font=f_base, fill=None, fmt=None,
        align=None, border=None, comment=None):
    c = ws[ref]
    c.value = value
    c.font = font
    if fill is not None:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if align is not None:
        c.alignment = align
    if border is not None:
        c.border = border
    if comment:
        c.comment = Comment(comment, "極値統計解析ブック", width=320, height=110)
    return c


def banner(ws, text, last_col):
    ws.merge_cells("A1:%s1" % last_col)
    put(ws, "A1", text, f_title, fill_title,
        align=Alignment(horizontal="left", vertical="center"))
    ws.row_dimensions[1].height = 26


def headers(ws, row, labels, col0=1):
    for i, lab in enumerate(labels):
        ref = "%s%d" % (get_column_letter(col0 + i), row)
        put(ws, ref, lab, f_h1, fill_head, align=ctr, border=box)
    ws.row_dimensions[row].height = 32


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def printing(ws, area, title_rows=None, landscape=True):
    """A4・幅1ページに収める印刷設定。"""
    ws.print_area = area
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    if title_rows:
        ws.print_title_rows = title_rows


wb = Workbook()

# ================================================================ 1) 使い方
ws = wb.active
ws.title = "使い方"
widths(ws, {"A": 4, "B": 26, "C": 96})
banner(ws, "年最高外気温度の極値統計解析（再現期間100年）　― 使い方 ―", "C")

L = [
    ("目的", "気象庁が公開している「年最高気温」の観測記録に極値統計（ワイブル分布）を"
             "あてはめ、再現期間 100 年に相当する年最高外気温度（100年に1回程度の確率で"
             "起こる猛暑の気温）を求めます。空調負荷計算・設備機器の耐熱設計・屋外構造物の"
             "温度応力検討などの設計用外気温度の設定に使います。"),
    ("使う順序",
     "① 「データ入力」シートに 年 と 年最高気温(℃) を貼り付ける\n"
     "② 「結果」シートの上部に 再現期間100年の年最高外気温度 が表示される\n"
     "③ 「極値統計計算」シートで分布パラメータ・相関係数（あてはまりの良さ）を確認する\n"
     "※ 数式はすべて生きています。データを入れ替えれば自動的に再計算されます。"),
    ("気象庁データの取り方",
     "【方法A：過去の気象データ検索（年ごとの値）】\n"
     "  https://www.data.jma.go.jp/obd/stats/etrn/index.php\n"
     "  → 地点を選ぶ → 「年ごとの値を表示」 → 表の『最高気温(℃)／最高』の列が年最高気温です。\n"
     "【方法B：過去の気象データ・ダウンロード（CSV）】\n"
     "  https://www.data.jma.go.jp/risk/obsdl/index.php\n"
     "  → 地点を選ぶ → 項目を選ぶ →『年別値』→『最高気温の最高』→ CSVをダウンロード\n"
     "  → CSV は Shift_JIS。Excelで開き、年 と 気温 の2列を「データ入力」に貼り付け。\n"
     "※ CSVには品質情報（均質番号・現象なし情報）の列が付きます。値が不完全な年は除いてください。"),
    ("計算方法（本ブックの中身）",
     "① 標本を降順に並べ、順位 m を付ける（m=1 が最大値）\n"
     "② 非超過確率 F をワイブルのプロッティング・ポジション F = 1 − m/(n+1) で与える\n"
     "   （再現期間は T = 1/(1−F) = (n+1)/m）\n"
     "③ ワイブル分布（第III種極値分布・上限型）を最小二乗であてはめる\n"
     "   F(x) = exp[ −((ω−x)/a)^k ]　　x ≦ ω\n"
     "   ln(−ln F) = k·ln(ω−x) − k·ln a　… ω を仮定すると直線になる\n"
     "   → ω を2段階のグリッド探索で振り、相関係数 r が最大になる ω を採用\n"
     "   → 傾き k（形状）と切片から a（尺度）を決定\n"
     "④ 再現期間 T の値：　x(T) = ω − a·[ −ln(1 − 1/T) ]^(1/k)\n"
     "⑤ 検証用に グンベル分布（第I種極値分布）も併記\n"
     "   F(x)=exp[−exp(−(x−β)/α)]　x(T) = β + α·[ −ln(−ln(1−1/T)) ]"),
    ("なぜワイブル（第III種）か",
     "気温には物理的な上限が存在すると考えられるため、上に有界な第III種極値分布"
     "（逆ワイブル分布）は年最高気温に適した形をしています。\n"
     "形状母数 k が大きく、ω が標本最大値からかけ離れて大きくなる場合、分布は"
     "グンベル分布（k→∞ の極限）に漸近しています。その場合はグンベルの値を採用してください。"
     "「極値統計計算」シートに自動判定を表示します。"),
    ("結果を使うときの注意",
     "・標本数 n：n≧30、できれば n≧50 を推奨します。n が小さいまま 100 年を外挿すると"
     "誤差が大きくなります（「結果」シートに 95%信頼区間の目安を表示）。\n"
     "・温暖化トレンド：極値統計は「母集団が変わらない（定常）」ことを前提にします。"
     "昇温トレンドがある地点では過去の長期データをそのまま使うと過小評価になります。"
     "直近30〜50年に期間を絞る、あるいはトレンド補正した値を入力してください。"
     "（「データ入力」シートに経年変化グラフと回帰トレンドを表示します）\n"
     "・観測所の移転・周辺環境の変化（都市化・ヒートアイランド）があると系列が不連続に"
     "なります。気象庁の観測所メタデータで移転履歴を確認してください。\n"
     "・ここで求まるのは「観測露場の気温」です。建物周辺の実効的な外気温は日射・輻射・"
     "排熱の影響で数℃高くなることがあります。設計値には別途割増しを検討してください。\n"
     "・空調設計では TAC 温度（超過確率を許容する設計外気温）など別の考え方もあります。"
     "用途に応じて使い分けてください。"),
    ("シート構成",
     "使い方　　　　… このシート\n"
     "データ入力　　… 年 と 年最高気温を入れる（黄色いセルが入力欄）\n"
     "極値統計計算　… 順位統計・確率プロット・ω探索・分布母数の算出（通常は触らない）\n"
     "結果　　　　　… 再現期間表・確率プロット・再現期間曲線\n"
     "サンプルデータ… 動作確認用のダミー標本の控え（実測値ではありません）"),
    ("記号", "n：標本数（年数）　m：降順順位　x(m)：m番目に大きい年最高気温\n"
             "F：非超過確率　T：再現期間(年)　ω：上限母数(℃)　a：尺度母数(℃)　k：形状母数\n"
             "α：グンベル尺度母数(℃)　β：グンベル位置母数(℃)　r：確率プロット上の相関係数"),
]

r = 3
for head, body in L:
    put(ws, "B%d" % r, head, f_h2, fill_band, align=wrap, border=box)
    put(ws, "C%d" % r, body, f_base, align=wrap, border=box)
    ws.row_dimensions[r].height = 15.5 * (body.count("\n") + 1) + 8
    r += 1

put(ws, "B%d" % (r + 1), "重要", f_warn, fill_input, align=wrap, border=box)
put(ws, "C%d" % (r + 1),
    "出荷時の「データ入力」シートには動作確認用のダミー値が入っています。"
    "気象庁の実測値に必ず置き換えてください。ダミーのままの結果は設計に使えません。",
    f_warn, fill_input, align=wrap, border=box)
ws.row_dimensions[r + 1].height = 34
printing(ws, "A1:C%d" % (r + 1))
ws.sheet_view.showGridLines = False

# ================================================================ 2) データ入力
ws = wb.create_sheet("データ入力")
widths(ws, {"A": 10, "B": 17, "C": 30, "D": 20, "E": 13, "F": 3,
            "G": 10, "H": 10})
banner(ws, "データ入力　― 気象庁の年最高気温を貼り付けてください ―", "H")

ws.merge_cells("A2:P2")
put(ws, "A2",
    "⚠ 現在入っている値は動作確認用のダミーです。気象庁の実測値に置き換えてください"
    "（黄色いセルが入力欄）。行が足りない場合は %d 行目までのどこかに追記してください。" % DROW1,
    f_warn, fill_input, align=Alignment(vertical="center"))
ws.row_dimensions[2].height = 20

meta = [
    ("観測地点名", "（ダミー）東京 相当", True),
    ("観測所番号", "（ダミー）47662 相当", True),
    ("データ出典", "※ダミーデータ／実データは気象庁 過去の気象データ検索に置換", True),
]
r = 4
for lab, val, is_in in meta:
    put(ws, "A%d" % r, lab, f_bold, fill_band, border=box)
    ws.merge_cells("B%d:C%d" % (r, r))
    put(ws, "B%d" % r, val, f_input if is_in else f_base,
        fill_input if is_in else None, border=box)
    ws["C%d" % r].border = box
    r += 1

put(ws, "A7", "統計期間", f_bold, fill_band, border=box)
ws.merge_cells("B7:C7")
put(ws, "B7", '=IF(B8=0,"データを入力してください",TEXT(MIN(%s),"0")&" 〜 "'
              '&TEXT(MAX(%s),"0")&"　（%s年間）")' % (RNG_Y, RNG_Y, '"&B8&"'),
    f_base, border=box)
ws["C7"].border = box

put(ws, "A8", "標本数 n", f_bold, fill_band, border=box)
put(ws, "B8", "=COUNT(%s)" % RNG_T, f_bold, border=box, fmt="0")
put(ws, "C8", "← 年最高気温が入力されている年数（自動）", f_small)

put(ws, "A9", "標本数の判定", f_bold, fill_band, border=box)
ws.merge_cells("B9:C9")
put(ws, "B9",
    '=IF(B8=0,"データを入力してください",'
    'IF(B8<10,"× n<10：極値解析には不足しています",'
    'IF(B8<30,"△ n<30：100年再現値の信頼性は低い（参考値扱い）",'
    'IF(B8<50,"○ n≧30：概ね実用可（n≧50を推奨）",'
    '"◎ n≧50：良好"))))', f_bold, border=box)
ws["C9"].border = box

stats = [
    ("平均 (℃)", '=IFERROR(AVERAGE(%s),"")' % RNG_T, "0.00"),
    ("標準偏差 (℃)", '=IFERROR(STDEV(%s),"")' % RNG_T, "0.000"),
    ("最大 (℃)", '=IFERROR(MAX(%s),"")' % RNG_T, "0.0"),
    ("最大値の年", '=IFERROR(INDEX(%s,MATCH(MAX(%s),%s,0)),"")'
                   % (RNG_Y, RNG_T, RNG_T), "0"),
    ("最小 (℃)", '=IFERROR(MIN(%s),"")' % RNG_T, "0.0"),
    ("トレンド (℃/年)", '=IFERROR(SLOPE(%s,%s),"")' % (RNG_T, RNG_Y), "0.0000"),
    ("100年あたり (℃)", "=IFERROR(E9*100,\"\")", "0.00"),
]
put(ws, "D3", "標本統計量（自動計算）", f_h2, fill_band, border=box)
put(ws, "E3", "", f_base, fill_band, border=box)
r = 4
for lab, fml, fmt in stats:
    put(ws, "D%d" % r, lab, f_base, border=box)
    put(ws, "E%d" % r, fml, f_bold, border=box, fmt=fmt)
    r += 1
ws["E10"].comment = Comment(
    "回帰直線の傾き×100年。値が大きい（おおむね +1℃/100年 以上）場合、"
    "母集団が定常でない可能性があります。期間を直近30〜50年に絞ることを検討してください。",
    "極値統計解析ブック", width=330, height=110)

ws.merge_cells("I11:P11")
put(ws, "I11", "※ グラフの縦軸が 0 から始まる場合は、縦軸を右クリック →"
               "［軸の書式設定］で最小値を調整してください。", f_small, align=wrap)
headers(ws, 11, ["年", "年最高気温 (℃)", "備考", "", "", "", "", ""])
for col in "DEFGH":
    ws["%s11" % col].fill = PatternFill("solid", fgColor="FFFFFF")
    ws["%s11" % col].border = Border()
put(ws, "A11", "年", f_h1, fill_head, align=ctr, border=box)
put(ws, "B11", "年最高気温 (℃)", f_h1, fill_head, align=ctr, border=box)
put(ws, "C11", "備考", f_h1, fill_head, align=ctr, border=box)

for i in range(NDATA):
    r = DROW0 + i
    if i < len(SAMPLE):
        put(ws, "A%d" % r, SAMPLE_START + i, f_input, fill_input,
            fmt="0", align=ctr, border=box)
        put(ws, "B%d" % r, SAMPLE[i], f_input, fill_input, fmt="0.0",
            align=ctr, border=box)
        put(ws, "C%d" % r, "ダミー" if i == 0 else "", f_small,
            fill_input, border=box)
    else:
        put(ws, "A%d" % r, None, f_input, fill_input, fmt="0",
            align=ctr, border=box)
        put(ws, "B%d" % r, None, f_input, fill_input, fmt="0.0",
            align=ctr, border=box)
        put(ws, "C%d" % r, None, f_small, fill_input, border=box)

ws.freeze_panes = "A12"
ws.sheet_view.showGridLines = False

ch = ScatterChart()
ch.title = "年最高気温の経年変化"
ch.style = 2
ch.height, ch.width = 7.0, 16
ch.x_axis.title = "年"
ch.y_axis.title = "年最高気温 (℃)"
ch.x_axis.majorGridlines = None
xref = Reference(ws, min_col=1, min_row=DROW0, max_row=DROW1)
yref = Reference(ws, min_col=2, min_row=DROW0, max_row=DROW1)
s = Series(yref, xref, title="年最高気温")
s.marker = Marker(symbol="circle", size=5)
s.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
ch.series.append(s)
ws.add_chart(ch, "G12")
printing(ws, "A1:P%d" % DROW1, title_rows="11:11")

# ================================================================ 3) 極値統計計算
ws = wb.create_sheet("極値統計計算")
widths(ws, {"A": 7, "B": 6, "C": 15, "D": 11, "E": 15, "F": 11, "G": 13,
            "H": 13, "I": 12, "J": 12, "K": 2,
            "L": 3, "M": 26, "N": 14, "O": 2, "P": 3, "Q": 20, "R": 12,
            "S": 2, "T": 11, "U": 12, "V": 12, "W": 12, "X": 11, "Y": 2,
            "Z": 11, "AA": 12, "AB": 12, "AC": 12, "AD": 11, "AE": 2,
            "AF": 11, "AG": 13, "AH": 13, "AI": 12})
banner(ws, "極値統計計算（順位統計・確率プロット・分布母数の推定）", "AI")
ws.merge_cells("A2:J2")
put(ws, "A2",
    "このシートは自動計算です。通常は編集しないでください。"
    "灰色の列（T〜AD）は上限母数 ω を探索するための作業領域です。",
    f_small, fill_note)

NREF = "データ入力!$B$8"           # 標本数 n
FLG = "$B$%d:$B$%d" % (RROW0, RROW1)
XCOL = "$C$%d:$C$%d" % (RROW0, RROW1)
YCOL = "$H$%d:$H$%d" % (RROW0, RROW1)
UCOL = "$G$%d:$G$%d" % (RROW0, RROW1)

# --- 順位統計表
put(ws, "A4", "【順位統計表】降順に並べ替え、ワイブルのプロッティング・ポジションを与える",
    f_h2)
hd = ["順位 m", "有効", "x(m) 降順 (℃)", "対応年(参考)", "F = 1−m/(n+1)",
      "T = 1/(1−F)", "u = −ln(−ln F)", "y = ln(−ln F)",
      "残差 グンベル", "残差 ワイブル"]
headers(ws, 6, hd)
for i in range(NDATA):
    r = RROW0 + i
    put(ws, "A%d" % r, i + 1, f_base, fmt="0", align=ctr, border=box)
    put(ws, "B%d" % r, "=IF($A%d<=%s,1,0)" % (r, NREF), f_small,
        fmt="0", align=ctr, border=box)
    put(ws, "C%d" % r, '=IF($B%d=1,LARGE(%s,$A%d),0)' % (r, RNG_T, r),
        f_base, fmt="0.0", align=ctr, border=box)
    put(ws, "D%d" % r,
        '=IF($B%d=1,IFERROR(INDEX(%s,MATCH($C%d,%s,0)),""),"")'
        % (r, RNG_Y, r, RNG_T), f_base, fmt="0", align=ctr, border=box)
    put(ws, "E%d" % r, "=IF($B%d=1,1-$A%d/(%s+1),0)" % (r, r, NREF),
        f_base, fmt="0.00000", align=ctr, border=box)
    put(ws, "F%d" % r, '=IF($B%d=1,(%s+1)/$A%d,"")' % (r, NREF, r),
        f_base, fmt="0.00", align=ctr, border=box)
    put(ws, "G%d" % r, "=IF($B%d=1,-LN(-LN($E%d)),0)" % (r, r),
        f_base, fmt="0.0000", align=ctr, border=box)
    put(ws, "H%d" % r, "=IF($B%d=1,LN(-LN($E%d)),0)" % (r, r),
        f_base, fmt="0.0000", align=ctr, border=box)
    put(ws, "I%d" % r,
        '=IF($B%d=1,IFERROR($C%d-($N$10+$N$9*$G%d),""),"")' % (r, r, r),
        f_base, fmt="0.00", align=ctr, border=box)
    put(ws, "J%d" % r,
        '=IF($B%d=1,IFERROR($C%d-($N$21-$N$26*EXP(-$G%d/$N$24)),""),"")'
        % (r, r, r), f_base, fmt="0.00", align=ctr, border=box)
ws["D6"].comment = Comment(
    "同じ気温の年が複数あるときは最初に見つかった年を返します（参考表示）。"
    "計算には使っていません。", "極値統計解析ブック", width=320, height=90)
ws.freeze_panes = "A7"

# --- 母数の推定ブロック (M:N)
def par(row, label, formula, fmt="0.0000", bold=False, note=None):
    put(ws, "M%d" % row, label, f_bold if bold else f_base,
        fill_band if bold else None, border=box)
    put(ws, "N%d" % row, formula, f_bold if bold else f_base,
        fill_result if bold else None, border=box, fmt=fmt, comment=note)

put(ws, "M4", "【1】グンベル分布（第I種極値分布）― 確率プロット最小二乗法", f_h2)
put(ws, "M5", "x = β + α·u　　u = −ln(−ln F)", f_small)
par(6, "Σx", "=SUMPRODUCT(%s,%s)" % (FLG, XCOL), "0.000")
par(7, "Σu", "=SUMPRODUCT(%s,%s)" % (FLG, UCOL), "0.000")
par(8, "Σu²", "=SUMPRODUCT(%s,%s,%s)" % (FLG, UCOL, UCOL), "0.000")
par(9, "尺度母数 α (℃)",
    '=IFERROR((%s*SUMPRODUCT(%s,%s,%s)-N6*N7)/(%s*N8-N7^2),"")'
    % (NREF, FLG, XCOL, UCOL, NREF), "0.0000", bold=True)
par(10, "位置母数 β (℃)", '=IFERROR((N6-N9*N7)/%s,"")' % NREF,
    "0.0000", bold=True)
par(11, "相関係数 r",
    '=IFERROR((%s*SUMPRODUCT(%s,%s,%s)-N6*N7)/SQRT((%s*N8-N7^2)*'
    '(%s*SUMPRODUCT(%s,%s,%s)-N6^2)),"")'
    % (NREF, FLG, XCOL, UCOL, NREF, NREF, FLG, XCOL, XCOL), "0.00000",
    note="確率プロット上の直線あてはまりの良さ。1に近いほど良い。"
         "目安として 0.98 以上あれば実用的。")

put(ws, "M13", "【2】グンベル分布 ― 積率法（参考・検算用）", f_h2)
par(14, "尺度母数 α (℃)",
    '=IFERROR(STDEV(%s)*SQRT(6)/PI(),"")' % RNG_T, "0.0000")
par(15, "位置母数 β (℃)",
    '=IFERROR(AVERAGE(%s)-0.5772*N14,"")' % RNG_T, "0.0000",
    note="オイラーの定数 γ=0.5772 を用いた積率推定。"
         "最小二乗法の結果と大きく食い違う場合は外れ値の混入を疑ってください。")

put(ws, "M17", "【3】ワイブル分布（第III種極値分布・上限型）★主結果", f_h2)
put(ws, "M18", "F(x)=exp[−((ω−x)/a)^k]　ln(−ln F)=k·ln(ω−x)−k·ln a", f_small)
par(19, "Σ ln(ω−x)",
    '=IFERROR(SUMPRODUCT(%s,LN($N$21-%s)),"")' % (FLG, XCOL), "0.000")
par(20, "Σ [ln(ω−x)]²",
    '=IFERROR(SUMPRODUCT(%s,LN($N$21-%s),LN($N$21-%s)),"")'
    % (FLG, XCOL, XCOL), "0.000")
par(21, "上限母数 ω (℃)",
    '=IFERROR(INDEX($Z$%d:$Z$%d,MATCH(MAX($AD$%d:$AD$%d),$AD$%d:$AD$%d,0)),"")'
    % (GROW0, GROW1, GROW0, GROW1, GROW0, GROW1), "0.000", bold=True,
    note="相関係数 r が最大となる ω を2段階のグリッド探索（T〜AD列）で決めています。"
         "気温の物理的な上限に相当する母数です。")
par(22, "Σ y", "=SUMPRODUCT(%s,%s)" % (FLG, YCOL), "0.000")
par(23, "Σ y²", "=SUMPRODUCT(%s,%s,%s)" % (FLG, YCOL, YCOL), "0.000")
par(24, "形状母数 k",
    '=IFERROR((%s*SUMPRODUCT(%s,LN($N$21-%s),%s)-N19*N22)/(%s*N20-N19^2),"")'
    % (NREF, FLG, XCOL, YCOL, NREF), "0.0000", bold=True)
par(25, "切片 −k·ln a", '=IFERROR((N22-N24*N19)/%s,"")' % NREF, "0.0000")
par(26, "尺度母数 a (℃)", '=IFERROR(EXP(-N25/N24),"")', "0.0000", bold=True)
par(27, "相関係数 r",
    '=IFERROR((%s*SUMPRODUCT(%s,LN($N$21-%s),%s)-N19*N22)/'
    'SQRT((%s*N20-N19^2)*(%s*N23-N22^2)),"")'
    % (NREF, FLG, XCOL, YCOL, NREF, NREF), "0.00000")
put(ws, "M28", "分布形の判定", f_bold, fill_band, border=box)
ws.merge_cells("N28:R28")
put(ws, "N28",
    '=IF(%s<10,"― 標本数が不足しています（n≧10が必要）",'
    'IF(N21="","―",IF(N21>=$R$7-$R$5,'
    '"⚠ ωが探索上限に達しています：分布はグンベル(k→∞)に漸近。グンベルの値を採用してください",'
    'IF(N24>50,"△ k>50：実質的にグンベル分布に近い形です。グンベルの値と比較してください",'
    '"○ 上限のあるワイブル形として推定（ω−標本最大値="&TEXT(N21-MAX(%s),"0.0")'
    '&"℃、k="&TEXT(N24,"0.0")&"）"))))' % (NREF, RNG_T),
    f_base, border=box)
for c in "OPQR":
    ws["%s28" % c].border = box

# --- ω 探索の設定 (Q:R)
put(ws, "Q4", "【ω グリッド探索の設定】", f_h2)
put(ws, "Q5", "第1段 刻み Δ1 (℃)", f_base, border=box)
put(ws, "R5", 0.5, f_input, fill_input, border=box, fmt="0.000")
put(ws, "Q6", "探索下限 ω_min (℃)", f_base, border=box)
put(ws, "R6", '=IFERROR(MAX(%s)+0.05,"")' % RNG_T, f_base, border=box,
    fmt="0.000")
put(ws, "Q7", "探索上限 (℃)", f_base, border=box)
put(ws, "R7", '=IFERROR($R$6+$R$5*(%d-1),"")' % NGRID, f_base, border=box,
    fmt="0.000")
put(ws, "Q8", "第1段の最良 ω (℃)", f_base, border=box)
put(ws, "R8",
    '=IFERROR(INDEX($T$%d:$T$%d,MATCH(MAX($X$%d:$X$%d),$X$%d:$X$%d,0)),"")'
    % (GROW0, GROW1, GROW0, GROW1, GROW0, GROW1), f_base, border=box,
    fmt="0.000")
put(ws, "Q9", "第2段 刻み Δ2 (℃)", f_base, border=box)
put(ws, "R9", '=IFERROR(2*$R$5/(%d-1),"")' % NGRID, f_base, border=box,
    fmt="0.00000")
put(ws, "Q10", "第2段 探索下限 (℃)", f_base, border=box)
put(ws, "R10", '=IFERROR(MAX($R$8-$R$5,$R$6),"")', f_base, border=box,
    fmt="0.000")
ws["R5"].comment = Comment(
    "探索範囲は ω_min から ω_min+Δ1×199 まで。上限に張り付く場合は Δ1 を大きくしてください。",
    "極値統計解析ブック", width=320, height=90)

# --- ω 探索表（第1段 T:X / 第2段 Z:AD）
put(ws, "T4", "【第1段】粗い探索", f_h2)
put(ws, "Z4", "【第2段】細かい探索", f_h2)
g_hd = ["ω 候補 (℃)", "Σ ln(ω−x)", "Σ [ln(ω−x)]²", "Σ y·ln(ω−x)", "相関係数 r"]
headers(ws, 6, g_hd, col0=20)   # T
headers(ws, 6, g_hd, col0=26)   # Z

for i in range(NGRID):
    r = GROW0 + i
    for base, omin, step in (("T", "$R$6", "$R$5"), ("Z", "$R$10", "$R$9")):
        cols = {"T": ("T", "U", "V", "W", "X"),
                "Z": ("Z", "AA", "AB", "AC", "AD")}[base]
        o, s1, s2, s3, s4 = cols
        put(ws, "%s%d" % (o, r), '=IFERROR(%s+($A%d-1)*%s,"")'
            % (omin, r, step), f_small, fmt="0.000", align=ctr, border=box)
        put(ws, "%s%d" % (s1, r),
            '=IFERROR(SUMPRODUCT(%s,LN($%s%d-%s)),"")' % (FLG, o, r, XCOL),
            f_small, fmt="0.000", align=ctr, border=box)
        put(ws, "%s%d" % (s2, r),
            '=IFERROR(SUMPRODUCT(%s,LN($%s%d-%s),LN($%s%d-%s)),"")'
            % (FLG, o, r, XCOL, o, r, XCOL),
            f_small, fmt="0.000", align=ctr, border=box)
        put(ws, "%s%d" % (s3, r),
            '=IFERROR(SUMPRODUCT(%s,LN($%s%d-%s),%s),"")'
            % (FLG, o, r, XCOL, YCOL),
            f_small, fmt="0.000", align=ctr, border=box)
        put(ws, "%s%d" % (s4, r),
            '=IFERROR((%s*%s%d-%s%d*$N$22)/SQRT((%s*%s%d-%s%d^2)*'
            '(%s*$N$23-$N$22^2)),-1)'
            % (NREF, s3, r, s1, r, NREF, s2, r, s1, r, NREF),
            f_small, fmt="0.000000", align=ctr, border=box)

# --- 適合曲線用 u グリッド (AF:AI)
put(ws, "AF4", "【適合曲線の描画用】", f_h2)
headers(ws, 6, ["u", "グンベル適合 (℃)", "ワイブル適合 (℃)", "T (年)"], col0=32)
for i in range(NUGRID):
    r = UROW0 + i
    put(ws, "AF%d" % r, -2 + 0.2 * i, f_small, fmt="0.00", align=ctr,
        border=box)
    put(ws, "AG%d" % r, '=IFERROR($N$10+$N$9*$AF%d,"")' % r, f_small,
        fmt="0.00", align=ctr, border=box)
    put(ws, "AH%d" % r, '=IFERROR($N$21-$N$26*EXP(-$AF%d/$N$24),"")' % r,
        f_small, fmt="0.00", align=ctr, border=box)
    put(ws, "AI%d" % r, '=IFERROR(1/(1-EXP(-EXP(-$AF%d))),"")' % r,
        f_small, fmt="0.0", align=ctr, border=box)
ws.column_dimensions.group("T", "AI", outline_level=1, hidden=False)
printing(ws, "A1:R%d" % RROW1, title_rows="6:6")
ws["AH6"].comment = Comment(
    "確率紙上でのワイブル(第III種)の式：x = ω − a·exp(−u/k)。"
    "u = −ln(−ln F) なので (−ln F)^(1/k) = exp(−u/k) となります。",
    "極値統計解析ブック", width=340, height=100)
ws.sheet_view.showGridLines = False

# ================================================================ 4) 結果
ws = wb.create_sheet("結果")
widths(ws, {"A": 14, "B": 14, "C": 20, "D": 20, "E": 20, "F": 15, "G": 15,
            "H": 13, "I": 13, "J": 3, "K": 12, "L": 12})
banner(ws, "結果　― 再現期間に対する年最高外気温度 ―", "I")

put(ws, "A3", "観測地点", f_bold, fill_band, border=box)
ws.merge_cells("B3:C3")
put(ws, "B3", "=データ入力!B4", f_base, border=box)
ws["C3"].border = box
put(ws, "D3", "統計期間", f_bold, fill_band, border=box)
ws.merge_cells("E3:G3")
put(ws, "E3", "=データ入力!B7", f_base, border=box)
for c in "FG":
    ws["%s3" % c].border = box
put(ws, "H3", "標本数 n", f_bold, fill_band, border=box)
put(ws, "I3", "=データ入力!B8", f_bold, border=box, fmt="0")

ws.merge_cells("A5:I5")
put(ws, "A5", "【主結果】再現期間 100 年に対する年最高外気温度", f_h1, fill_head,
    align=Alignment(horizontal="left", vertical="center"))
ws.row_dimensions[5].height = 22

ws.merge_cells("A6:B8")
put(ws, "A6", "ワイブル分布\n（第III種極値分布）", f_bold, fill_result,
    align=Alignment(horizontal="center", vertical="center", wrap_text=True),
    border=box)
ws.merge_cells("C6:D8")
put(ws, "C6", '=IF(データ入力!$B$8<10,"―",'
              'IFERROR(TEXT(ROUND($C$%d,1),"0.0")&" ℃","―"))' % ROW_T100,
    f_big, fill_result, align=ctr, border=box)
ws.merge_cells("E6:F8")
put(ws, "E6", "グンベル分布\n（第I種・比較用）", f_bold, fill_band,
    align=Alignment(horizontal="center", vertical="center", wrap_text=True),
    border=box)
ws.merge_cells("G6:I8")
put(ws, "G6", '=IF(データ入力!$B$8<10,"―",'
              'IFERROR(TEXT(ROUND($D$%d,1),"0.0")&" ℃","―"))' % ROW_T100,
    Font(name=FONT, size=24, bold=True, color=C_TITLE), fill_band,
    align=ctr, border=box)
for rr in (6, 7, 8):
    ws.row_dimensions[rr].height = 20

ws.merge_cells("A9:I9")
put(ws, "A9",
    '=データ入力!B9&"　／　あてはまり r(ワイブル)="'
    '&IFERROR(TEXT(極値統計計算!N27,"0.0000"),"―")'
    '&"　r(グンベル)="&IFERROR(TEXT(極値統計計算!N11,"0.0000"),"―")'
    '&"　／　"&極値統計計算!N28', f_bold, fill_note, align=Alignment(vertical="center"))
ws.row_dimensions[9].height = 18

ws.merge_cells("A10:I10")
put(ws, "A10",
    '=IF(データ入力!E10="","※ データを入力してください。",'
    'IF(ABS(データ入力!E10)>=1,'
    '"⚠ 昇温トレンドが "&TEXT(データ入力!E10,"0.0")&" ℃/100年 と大きく、定常性の前提が'
    '崩れている可能性があります（過小評価側）。期間を直近30〜50年に絞った再計算と比較してください。",'
    '"※ 昇温トレンドは "&TEXT(データ入力!E10,"0.0")&" ℃/100年 です。"))'
    '&"　出荷時のデータはダミーです。設計に用いる前に気象庁の実測値へ差し替えてください。"',
    f_warn, align=wrap)
ws.row_dimensions[10].height = 30

# --- 再現期間表
put(ws, "A12", "【再現期間表】", f_h2)
ws.row_dimensions[12].height = 16

headers(ws, 13, ["再現期間 T\n(年)", "非超過確率 F", "ワイブル分布\n(第III種) ℃",
                 "グンベル分布\n(最小二乗) ℃", "グンベル分布\n(積率法) ℃",
                 "グンベル\n95%下限 ℃", "グンベル\n95%上限 ℃",
                 "頻度係数 K(T)", "標準誤差 (℃)"])
ROW0 = TBL_ROW0
for i, T in enumerate(RETURN_PERIODS):
    r = ROW0 + i
    hl = (T == 100)
    fnt = Font(name=FONT, size=10, bold=True) if hl else f_base
    fl = fill_result if hl else None
    put(ws, "A%d" % r, T, fnt, fl, fmt="0", align=ctr, border=box)
    put(ws, "B%d" % r, "=1-1/A%d" % r, fnt, fl, fmt="0.00000", align=ctr,
        border=box)
    put(ws, "C%d" % r,
        '=IFERROR(極値統計計算!$N$21-極値統計計算!$N$26*(-LN($B%d))^(1/極値統計計算!$N$24),"")'
        % r, fnt, fl, fmt="0.00", align=ctr, border=box)
    put(ws, "D%d" % r,
        '=IFERROR(極値統計計算!$N$10+極値統計計算!$N$9*(-LN(-LN($B%d))),"")' % r,
        fnt, fl, fmt="0.00", align=ctr, border=box)
    put(ws, "E%d" % r,
        '=IFERROR(極値統計計算!$N$15+極値統計計算!$N$14*(-LN(-LN($B%d))),"")' % r,
        fnt, fl, fmt="0.00", align=ctr, border=box)
    put(ws, "F%d" % r, '=IFERROR($D%d-1.96*$I%d,"")' % (r, r), fnt, fl,
        fmt="0.00", align=ctr, border=box)
    put(ws, "G%d" % r, '=IFERROR($D%d+1.96*$I%d,"")' % (r, r), fnt, fl,
        fmt="0.00", align=ctr, border=box)
    put(ws, "H%d" % r, '=IFERROR(0.779697*(-LN(-LN($B%d))-0.5772),"")' % r,
        f_small, fl, fmt="0.000", align=ctr, border=box)
    put(ws, "I%d" % r,
        '=IFERROR(STDEV(%s)*SQRT((1+1.1396*$H%d+1.1*$H%d^2)/データ入力!$B$8),"")'
        % (RNG_T, r, r), f_small, fl, fmt="0.000", align=ctr, border=box)

ROWEND = ROW0 + len(RETURN_PERIODS) - 1
ws["F13"].comment = Comment(
    "グンベル分布の再現期待値に対する近似95%信頼区間（Kiteの式）。\n"
    "Var[x(T)] = (s²/n)(1 + 1.1396·K(T) + 1.1·K(T)²)、K(T)=(√6/π)(u(T)−0.5772)。\n"
    "標本誤差のみを表し、母集団の非定常性（温暖化）の不確かさは含みません。",
    "極値統計解析ブック", width=380, height=130)

r = ROWEND + 2
put(ws, "A%d" % r, "推定された分布母数", f_h2)
r += 1
pars = [("ワイブル 上限母数 ω (℃)", "=極値統計計算!N21", "0.000"),
        ("ワイブル 尺度母数 a (℃)", "=極値統計計算!N26", "0.000"),
        ("ワイブル 形状母数 k", "=極値統計計算!N24", "0.000"),
        ("ワイブル 相関係数 r", "=極値統計計算!N27", "0.00000"),
        ("グンベル 尺度母数 α (℃)", "=極値統計計算!N9", "0.000"),
        ("グンベル 位置母数 β (℃)", "=極値統計計算!N10", "0.000"),
        ("グンベル 相関係数 r", "=極値統計計算!N11", "0.00000")]
for i, (lab, fml, fmt) in enumerate(pars):
    rr = r + i
    ws.merge_cells("A%d:B%d" % (rr, rr))
    put(ws, "A%d" % rr, lab, f_base, fill_band, border=box)
    ws["B%d" % rr].border = box
    put(ws, "C%d" % rr, fml, f_bold, border=box, fmt=fmt, align=ctr)

rn = r + len(pars) + 1
ws.merge_cells("A%d:I%d" % (rn, rn + 3))
put(ws, "A%d" % rn,
    "計算式\n"
    "　ワイブル（第III種極値分布）： F(x)=exp[−((ω−x)/a)^k]　→　x(T) = ω − a·[−ln(1−1/T)]^(1/k)\n"
    "　グンベル（第I種極値分布）　： F(x)=exp[−exp(−(x−β)/α)]　→　x(T) = β + α·[−ln(−ln(1−1/T))]\n"
    "　非超過確率はワイブルのプロッティング・ポジション F = 1 − m/(n+1)（m は降順順位）で与えています。",
    f_small, fill_note, align=wrap, border=box)
ws.sheet_view.showGridLines = False

ws.merge_cells("K2:R2")
put(ws, "K2", "※ グラフの縦軸が 0 から始まってしまう場合は、縦軸を右クリック →"
              "［軸の書式設定］で最小値を 30 などに変更してください。", f_small)

# --- 確率プロット
calc = wb["極値統計計算"]
ch1 = ScatterChart()
ch1.title = "確率プロット（グンベル確率紙）"
ch1.style = 2
ch1.height, ch1.width = 10, 17
ch1.x_axis.title = "基準化変数 u = −ln(−ln F)"
ch1.y_axis.title = "年最高気温 (℃)"
ch1.x_axis.majorGridlines = None

s_obs = Series(Reference(calc, min_col=3, min_row=RROW0, max_row=RROW1),
               Reference(calc, min_col=7, min_row=RROW0, max_row=RROW1),
               title="実測値")
s_obs.marker = Marker(symbol="circle", size=6)
s_obs.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
ch1.series.append(s_obs)

s_w = Series(Reference(calc, min_col=34, min_row=UROW0, max_row=UROW1),
             Reference(calc, min_col=32, min_row=UROW0, max_row=UROW1),
             title="ワイブル分布(第III種)")
s_w.marker = Marker(symbol="none")
s_w.graphicalProperties = GraphicalProperties(
    ln=LineProperties(w=22000, solidFill="C00000"))
ch1.series.append(s_w)

s_g = Series(Reference(calc, min_col=33, min_row=UROW0, max_row=UROW1),
             Reference(calc, min_col=32, min_row=UROW0, max_row=UROW1),
             title="グンベル分布(第I種)")
s_g.marker = Marker(symbol="none")
s_g.graphicalProperties = GraphicalProperties(
    ln=LineProperties(w=22000, solidFill="1F4E79", prstDash="dash"))
ch1.series.append(s_g)
ws.add_chart(ch1, "K3")

# --- 再現期間曲線
ch2 = ScatterChart()
ch2.title = "再現期間と年最高外気温度"
ch2.style = 2
ch2.height, ch2.width = 10, 17
ch2.x_axis.title = "再現期間 T (年)  ※対数軸"
ch2.y_axis.title = "年最高外気温度 (℃)"
ch2.x_axis.scaling.logBase = 10
ch2.x_axis.majorGridlines = None

s_p = Series(Reference(calc, min_col=3, min_row=RROW0, max_row=RROW1),
             Reference(calc, min_col=6, min_row=RROW0, max_row=RROW1),
             title="実測値")
s_p.marker = Marker(symbol="circle", size=6)
s_p.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
ch2.series.append(s_p)

s_cw = Series(Reference(ws, min_col=3, min_row=ROW0, max_row=ROWEND),
              Reference(ws, min_col=1, min_row=ROW0, max_row=ROWEND),
              title="ワイブル分布(第III種)")
s_cw.marker = Marker(symbol="none")
s_cw.graphicalProperties = GraphicalProperties(
    ln=LineProperties(w=22000, solidFill="C00000"))
ch2.series.append(s_cw)

s_cg = Series(Reference(ws, min_col=4, min_row=ROW0, max_row=ROWEND),
              Reference(ws, min_col=1, min_row=ROW0, max_row=ROWEND),
              title="グンベル分布(第I種)")
s_cg.marker = Marker(symbol="none")
s_cg.graphicalProperties = GraphicalProperties(
    ln=LineProperties(w=22000, solidFill="1F4E79", prstDash="dash"))
ch2.series.append(s_cg)
ws.add_chart(ch2, "K24")
printing(ws, "A1:R45")

# ================================================================ 5) サンプル
ws = wb.create_sheet("サンプルデータ")
widths(ws, {"A": 10, "B": 18, "C": 60})
banner(ws, "サンプルデータ（動作確認用のダミー・気象庁の実測値ではありません）", "C")
ws.merge_cells("A2:C2")
put(ws, "A2",
    "グンベル分布に緩やかな昇温トレンド（+0.02℃/年）を与えて発生させた人工標本です。"
    "「データ入力」シートを実データに置き換えたあと、動作を確認したくなったときの控えです。",
    f_small, fill_note, align=Alignment(vertical="center"))
ws.row_dimensions[2].height = 30
put(ws, "A4", "年", f_h1, fill_head, align=ctr, border=box)
put(ws, "B4", "年最高気温 (℃)", f_h1, fill_head, align=ctr, border=box)
put(ws, "C4", "備考", f_h1, fill_head, align=ctr, border=box)
for i, v in enumerate(SAMPLE):
    r = 5 + i
    put(ws, "A%d" % r, SAMPLE_START + i, f_base, fmt="0", align=ctr, border=box)
    put(ws, "B%d" % r, v, f_base, fmt="0.0", align=ctr, border=box)
    put(ws, "C%d" % r, "ダミー（実測値ではない）" if i == 0 else "", f_small,
        border=box)
ws.freeze_panes = "A5"
printing(ws, "A1:C%d" % (4 + len(SAMPLE)), title_rows="4:4", landscape=False)
ws.sheet_view.showGridLines = False

# ---------------------------------------------------------------- 名前定義
wb.defined_names.add(DefinedName("年次", attr_text=RNG_Y))
wb.defined_names.add(DefinedName("年最高気温", attr_text=RNG_T))
wb.defined_names.add(DefinedName("標本数", attr_text="データ入力!$B$8"))

wb.calculation.fullCalcOnLoad = True
wb.active = 0

out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "年最高外気温_ワイブル極値解析.xlsx")
wb.save(out)
print("saved:", out)
