# -*- coding: utf-8 -*-
"""特殊スラブ 問題集（図つき）Excel 生成スクリプト。
出力: docs/special_slab/特殊スラブ問題集.xlsx
1 ボイドスラブ構造特性(自重・等価スラブ厚) / 2 ボイドスラブ断面性能(A,Z,I)
3 FEM解析モデル / 4 ハーフPCa構造特性(合成効果) / 5 ハーフPCa施工時・完成時
6 施工方法と構造的留意点(支保工・養生・品質管理)
RC造マンションの設計担当を想定。数値は build 時に検算済み。
"""
import os
import math
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib import font_manager as fm

FONT_PATH = "/usr/share/fonts/truetype/fonts-japanese-gothic.ttf"
jp = fm.FontProperties(fname=FONT_PATH)
fm.fontManager.addfont(FONT_PATH)
plt.rcParams["font.family"] = jp.get_name()
plt.rcParams["axes.unicode_minus"] = False

OUT = "docs/special_slab"
FIG = os.path.join(OUT, "figures")
os.makedirs(FIG, exist_ok=True)

C_BLUE = "#2a78d6"
C_PINK = "#d55181"


def save(fig, name):
    p = os.path.join(FIG, name)
    fig.savefig(p, dpi=130, bbox_inches="tight")
    plt.close(fig)
    return p


# ===========================================================================
# 図 1: ボイドスラブの断面と等価スラブ厚
# ===========================================================================
def fig_void():
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 5.0),
                              gridspec_kw={"width_ratios": [1.3, 1.0]})
    # (a) 断面
    ax = axes[0]
    D, dv, pitch = 500, 320, 400
    Wsec = pitch * 3
    ax.add_patch(mpatches.Rectangle((0, 0), Wsec, D, fc="#d9d9d9", ec="k",
                                     lw=1.2))
    for i in range(3):
        cx = pitch * i + pitch / 2
        ax.add_patch(plt.Circle((cx, D / 2), dv / 2, fc="white", ec="#1f4e79",
                                 lw=1.2, zorder=3))
        ax.text(cx, D / 2, "ボイド管\nφ320", fontproperties=jp, ha="center",
                va="center", fontsize=7.5, color="#1f4e79")
    # 上下スラブ（かぶり部）
    ax.annotate("", xy=(-0.3 * pitch, 0), xytext=(-0.3 * pitch, D),
                arrowprops=dict(arrowstyle="<|-|>", color="k", lw=1))
    ax.text(-0.42 * pitch, D / 2, "D=500", fontproperties=jp, rotation=90,
            va="center", fontsize=9)
    ax.annotate("", xy=(pitch / 2, D + 30), xytext=(pitch * 1.5, D + 30),
                arrowprops=dict(arrowstyle="<|-|>", color="#c00000", lw=1))
    ax.text(pitch, D + 55, "@400", fontproperties=jp, ha="center",
            fontsize=8.5, color="#c00000")
    ax.text(Wsec / 2, -70,
            "円筒ボイドを並べ、中間のコンクリートを抜いて軽量化\n"
            "上下の板が曲げに効く（I 形断面に近い）",
            fontproperties=jp, ha="center", fontsize=9, color="#444")
    ax.set_xlim(-0.6 * pitch, Wsec + 0.2 * pitch)
    ax.set_ylim(-160, D + 110)
    ax.set_aspect("equal")
    ax.axis("off")
    ax.set_title("(a) ボイドスラブ断面（厚500・φ320@400）",
                 fontproperties=jp, fontsize=11, fontweight="bold")

    # (b) 2つの等価スラブ厚
    ax = axes[1]
    cats = ["中実\nD=500", "等価厚\n(重量)\n299", "等価厚\n(曲げ剛性)\n478"]
    vals = [500, 299, 478]
    colors = ["#c9c9c9", C_PINK, C_BLUE]
    x = np.arange(3)
    for xi, v, c in zip(x, vals, colors):
        ax.add_patch(mpatches.Rectangle((xi - 0.35, 0), 0.7, v, fc=c, ec="k",
                                         lw=0.8))
        ax.text(xi, v + 18, f"{v}", fontproperties=jp, ha="center",
                fontsize=10, fontweight="bold")
    ax.set_xticks(x)
    ax.set_xticklabels(cats, fontproperties=jp, fontsize=9)
    ax.set_ylim(0, 580)
    ax.set_ylabel("等価スラブ厚 (mm)", fontproperties=jp)
    ax.text(1.0, -135,
            "★ 用途で使い分ける：\n"
            "自重・質量 → 重量等価厚 299（純断面積 / 幅）\n"
            "たわみ・剛性 → 剛性等価厚 478（I の 3 乗根換算）",
            fontproperties=jp, ha="center", fontsize=8.5, color="#c00000")
    ax.set_title("(b) 2 種類の等価スラブ厚",
                 fontproperties=jp, fontsize=11, fontweight="bold")
    ax.spines[["top", "right"]].set_visible(False)
    fig.suptitle("図 1  ボイドスラブの構造特性 ── 自重と等価スラブ厚",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    fig.subplots_adjust(bottom=0.22)
    return save(fig, "fig1_void.png")


# ===========================================================================
# 図 2: 断面性能の計算（中実 − ボイド円）
# ===========================================================================
def fig_section():
    fig, ax = plt.subplots(figsize=(11, 4.6))
    ax.text(0.5, 0.92, "ボイドスラブの断面性能（1 ピッチ b=400mm で評価）",
            transform=ax.transAxes, ha="center", fontproperties=jp,
            fontsize=12, fontweight="bold", color="#1f4e79")
    # 左：中実
    ax1 = fig.add_axes([0.02, 0.15, 0.24, 0.62])
    ax1.add_patch(mpatches.Rectangle((0, 0), 400, 500, fc="#c9c9c9", ec="k"))
    ax1.set_xlim(-60, 460); ax1.set_ylim(-60, 560)
    ax1.set_aspect("equal"); ax1.axis("off")
    ax1.set_title("中実", fontproperties=jp, fontsize=10)
    ax1.text(200, -40, "A1=b·D\nI1=b·D^3/12", fontproperties=jp, ha="center",
             fontsize=8)
    # 中
    fig.text(0.29, 0.45, "−", fontsize=30, ha="center")
    ax2 = fig.add_axes([0.32, 0.15, 0.24, 0.62])
    ax2.add_patch(plt.Circle((200, 250), 160, fc="white", ec="#1f4e79",
                             lw=1.5))
    ax2.set_xlim(-60, 460); ax2.set_ylim(-60, 560)
    ax2.set_aspect("equal"); ax2.axis("off")
    ax2.set_title("ボイド円", fontproperties=jp, fontsize=10)
    ax2.text(200, -40, "A2=π·dv^2/4\nI2=π·dv^4/64", fontproperties=jp,
             ha="center", fontsize=8)
    fig.text(0.59, 0.45, "=", fontsize=30, ha="center")
    # 右：結果表
    fig.text(0.63, 0.72,
             "純断面（図心は中央・対称）",
             fontproperties=jp, fontsize=10, fontweight="bold",
             color="#c00000")
    fig.text(0.63, 0.30,
             "A = A1 - A2 = 200,000 - 80,425 = 119,575 mm2\n"
             "I = I1 - I2 = 4.167e9 - 0.515e9 = 3.65e9 mm4\n"
             "Z = I / (D/2) = 3.65e9 / 250 = 1.46e7 mm3\n\n"
             "【1m 幅換算】 x(1000/400)\n"
             "I = 9.13e9 mm4/m、Z = 3.65e7 mm3/m\n"
             "曲げ剛性は中実の 88%（円を中心に抜くので I の低下は小さい）",
             fontproperties=jp, fontsize=9, color="#333")
    return save(fig, "fig2_section.png")


# ===========================================================================
# 図 3: FEM 解析モデル
# ===========================================================================
def fig_fem():
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 5.2))
    # (a) 要素分割と境界条件
    ax = axes[0]
    Lx, Ly = 8, 6
    n = 8
    for i in range(n + 1):
        ax.plot([0, Lx], [Ly * i / n, Ly * i / n], color="#bbb", lw=0.6)
        ax.plot([Lx * i / n, Lx * i / n], [0, Ly], color="#bbb", lw=0.6)
    ax.add_patch(mpatches.Rectangle((0, 0), Lx, Ly, fc="none", ec="k",
                                     lw=1.5))
    # 境界（梁＝ピン支持を三角で）
    for x in np.linspace(0.5, Lx - 0.5, 6):
        ax.plot(x, 0, "^", color="#c00000", ms=7)
        ax.plot(x, Ly, "^", color="#c00000", ms=7)
    for y in np.linspace(0.5, Ly - 0.5, 5):
        ax.plot(0, y, ">", color="#c00000", ms=7)
        ax.plot(Lx, y, "<", color="#c00000", ms=7)
    ax.text(Lx / 2, -0.7, "四周を大梁で支持（周辺支持スラブ）",
            fontproperties=jp, ha="center", fontsize=9, color="#c00000")
    ax.text(Lx / 2, Ly + 0.5, "板要素（シェル要素）でメッシュ分割",
            fontproperties=jp, ha="center", fontsize=9, color="#444")
    ax.set_xlim(-1, Lx + 1); ax.set_ylim(-1.3, Ly + 1.1)
    ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("(a) 境界条件と要素分割", fontproperties=jp,
                 fontsize=11, fontweight="bold")

    # (b) 直交異方性剛性
    ax = axes[1]
    ax.text(0.5, 0.9, "ボイドスラブ ＝ 直交異方性の板",
            transform=ax.transAxes, ha="center", fontproperties=jp,
            fontsize=11, fontweight="bold", color="#1f4e79")
    # ボイド方向と直交方向で剛性が違う
    ax.annotate("", xy=(0.75, 0.6), xytext=(0.25, 0.6),
                xycoords="axes fraction",
                arrowprops=dict(arrowstyle="<|-|>", color=C_BLUE, lw=3))
    ax.text(0.3, 0.68, "ボイド方向（管に沿う）\n剛性 大", transform=ax.transAxes,
            ha="center", fontproperties=jp, fontsize=9, color=C_BLUE)
    ax.annotate("", xy=(0.5, 0.75), xytext=(0.5, 0.35),
                xycoords="axes fraction",
                arrowprops=dict(arrowstyle="<|-|>", color=C_PINK, lw=3))
    ax.text(0.68, 0.5, "直交方向\n剛性 やや小", transform=ax.transAxes,
            ha="center", fontproperties=jp, fontsize=9, color=C_PINK)
    ax.text(0.5, 0.16,
            "入力：荷重（自重は重量等価厚で／積載・仕上）\n"
            "剛性：曲げ剛性等価厚 または 直交異方性の Dx・Dy\n"
            "境界：周辺大梁＝支持、連続端＝回転拘束を考慮",
            transform=ax.transAxes, ha="center", fontproperties=jp,
            fontsize=8.5, color="#333")
    ax.axis("off")
    ax.set_title("(b) 剛性の入力（等価厚 or 異方性）",
                 fontproperties=jp, fontsize=11, fontweight="bold")
    fig.suptitle("図 3  ボイドスラブの FEM 解析モデル",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig3_fem.png")


# ===========================================================================
# 図 4: ハーフPCaスラブの合成
# ===========================================================================
def fig_halfpca():
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 5.0))
    # (a) 断面構成
    ax = axes[0]
    b = 6.0
    ax.add_patch(mpatches.Rectangle((0, 0), b, 0.6, fc="#9ec6e8", ec="k",
                                     lw=1.2))
    ax.add_patch(mpatches.Rectangle((0, 0.6), b, 1.4, fc="#e8e8e8", ec="k",
                                     lw=1.2, hatch=".."))
    ax.annotate("PCa 板 60mm\n（工場製作・型枠兼用）", xy=(b*0.25, 0.3),
                xytext=(-0.3, -1.35), fontproperties=jp, fontsize=8.5,
                color="#1f4e79",
                arrowprops=dict(arrowstyle="->", color="#1f4e79"))
    ax.annotate("現場打ちコンクリート\n140mm（トッピング）", xy=(b*0.75, 1.5),
                xytext=(b*0.58, 2.35), fontproperties=jp, fontsize=8.5,
                color="#444",
                arrowprops=dict(arrowstyle="->", color="#444"))
    # トラス筋（ラチス）
    xs = np.linspace(0.6, b - 0.6, 9)
    for i in range(len(xs) - 1):
        y0, y1 = (0.5, 1.7) if i % 2 == 0 else (1.7, 0.5)
        ax.plot([xs[i], xs[i + 1]], [y0, y1], color="#c00000", lw=1.3)
    ax.plot([0.4, b - 0.4], [1.7, 1.7], color="#c00000", lw=1.5)
    ax.text(b + 0.1, 1.7, "トラス筋\n（ラチス）", fontproperties=jp,
            fontsize=8, color="#c00000", va="center")
    # 接合面
    ax.plot([0, b], [0.6, 0.6], color="#7a3b3b", lw=2, ls="--")
    ax.annotate("接合面（合成のカギ）\n目荒し＋トラス筋で\nずれ止め",
                xy=(1.5, 0.6), xytext=(2.6, -1.35), fontproperties=jp,
                fontsize=8, color="#7a3b3b",
                arrowprops=dict(arrowstyle="->", color="#7a3b3b"))
    ax.set_xlim(-0.5, b + 1.3); ax.set_ylim(-1.6, 2.4)
    ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("(a) ハーフPCaスラブの構成（合成200mm）",
                 fontproperties=jp, fontsize=11, fontweight="bold")

    # (b) 合成効果
    ax = axes[1]
    cats = ["施工時\nPCa単独\n60mm", "完成時\n合成\n200mm"]
    I = [1.8e7, 6.667e8]
    for xi, (c, v, col) in enumerate(zip(cats, I, [C_PINK, C_BLUE])):
        h = (v / 6.667e8) * 500
        ax.add_patch(mpatches.Rectangle((xi - 0.35, 0), 0.7, h, fc=col,
                                        ec="k", lw=0.8))
        ax.text(xi, h + 15, f"I={v:.2e}", fontproperties=jp, ha="center",
                fontsize=9, fontweight="bold")
    ax.set_xticks([0, 1])
    ax.set_xticklabels(cats, fontproperties=jp, fontsize=9)
    ax.set_ylim(0, 580)
    ax.set_ylabel("断面二次モーメント I (相対)", fontproperties=jp)
    ax.text(0.5, -140,
            "合成でトッピングが一体化 → I が 37 倍\n"
            "施工時（PCa単独）と完成時（合成）で別断面として検討",
            fontproperties=jp, ha="center", fontsize=8.5, color="#c00000")
    ax.set_title("(b) 合成効果（施工時 → 完成時）",
                 fontproperties=jp, fontsize=11, fontweight="bold")
    ax.spines[["top", "right"]].set_visible(False)
    fig.suptitle("図 4  ハーフPCaスラブの構造特性 ── 合成効果",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    fig.subplots_adjust(bottom=0.24)
    return save(fig, "fig4_halfpca.png")


# ===========================================================================
# 図 5: 施工時 vs 完成時
# ===========================================================================
def fig_stages():
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 4.8))
    # (a) 施工時
    ax = axes[0]
    L = 6
    ax.add_patch(mpatches.Rectangle((0, 1.0), L, 0.4, fc="#9ec6e8", ec="k",
                                     lw=1.2))
    ax.plot(0, 1.0, "^", color="k", ms=12)
    ax.plot(L, 1.0, "o", color="k", ms=10)
    # サポート（支保工）任意
    for x in [L / 3, 2 * L / 3]:
        ax.plot([x, x], [0, 1.0], color="#1f7a1f", lw=2)
        ax.plot(x, 0, "s", color="#1f7a1f", ms=8)
    ax.text(L / 2, 0.3, "支保工（サポート）", fontproperties=jp, ha="center",
            fontsize=8.5, color="#1f7a1f")
    for x in np.linspace(0.4, L - 0.4, 7):
        ax.annotate("", xy=(x, 1.4), xytext=(x, 1.9),
                    arrowprops=dict(arrowstyle="-|>", color="#c00000",
                                    lw=1.2))
    ax.text(L / 2, 2.1, "PCa自重 + 打設コンクリート + 作業荷重",
            fontproperties=jp, ha="center", fontsize=8.5, color="#c00000")
    ax.text(L / 2, -0.6,
            "施工時：PCa板 60mm 単独（未硬化）で支持\n"
            "→ 支保工の有無で PCa の検討応力が決まる",
            fontproperties=jp, ha="center", fontsize=9, color="#444")
    ax.set_xlim(-0.5, L + 0.5); ax.set_ylim(-1.2, 2.5)
    ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("(a) 施工時（PCa 単独）", fontproperties=jp,
                 fontsize=11, fontweight="bold")

    # (b) 完成時
    ax = axes[1]
    ax.add_patch(mpatches.Rectangle((0, 1.0), L, 0.35, fc="#9ec6e8", ec="k",
                                     lw=1.0))
    ax.add_patch(mpatches.Rectangle((0, 1.35), L, 0.85, fc="#e8e8e8", ec="k",
                                     lw=1.0, hatch=".."))
    ax.plot(0, 1.0, "^", color="k", ms=12)
    ax.plot(L, 1.0, "o", color="k", ms=10)
    for x in np.linspace(0.4, L - 0.4, 7):
        ax.annotate("", xy=(x, 2.2), xytext=(x, 2.7),
                    arrowprops=dict(arrowstyle="-|>", color="#c00000",
                                    lw=1.2))
    ax.text(L / 2, 2.9, "自重 + 仕上げ + 積載荷重", fontproperties=jp,
            ha="center", fontsize=8.5, color="#c00000")
    ax.text(L / 2, -0.6,
            "完成時：合成断面 200mm（支保工撤去後）\n"
            "→ 合成後の全荷重を合成断面で検討",
            fontproperties=jp, ha="center", fontsize=9, color="#444")
    ax.set_xlim(-0.5, L + 0.5); ax.set_ylim(-1.2, 3.2)
    ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("(b) 完成時（合成断面）", fontproperties=jp,
                 fontsize=11, fontweight="bold")
    fig.suptitle("図 5  ハーフPCaスラブ ── 施工時と完成時の検討",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig5_stages.png")


figs = {
    "void": fig_void(),
    "section": fig_section(),
    "fem": fig_fem(),
    "halfpca": fig_halfpca(),
    "stages": fig_stages(),
}
print("figures:", list(figs.keys()))

# ===========================================================================
# Excel 構築
# ===========================================================================
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

wb = Workbook()
C_TITLE = "1F4E79"; C_HEAD = "2E75B6"; C_ANS = "E2EFDA"
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
f_title = Font(name="MS PGothic", size=15, bold=True, color="FFFFFF")
f_head = Font(name="MS PGothic", size=11, bold=True, color="FFFFFF")
f_body = Font(name="MS PGothic", size=10)
f_ans = Font(name="MS PGothic", size=10, color="375623")
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)


def setup(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False


def title_row(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text)
    c.font = f_title; c.fill = PatternFill("solid", fgColor=C_TITLE)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[row].height = 30


def head(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text)
    c.font = f_head; c.fill = PatternFill("solid", fgColor=C_HEAD)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[row].height = 22


def body(ws, row, text, span=8, ans=False, h=None):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text)
    c.font = f_ans if ans else f_body
    c.alignment = wrap
    if ans:
        c.fill = PatternFill("solid", fgColor=C_ANS)
    if h:
        ws.row_dimensions[row].height = h


def table(ws, start_row, headers, rows, col1=1):
    r = start_row
    for j, htxt in enumerate(headers):
        c = ws.cell(r, col1 + j, htxt)
        c.font = Font(name="MS PGothic", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C_HEAD)
        c.alignment = center; c.border = border
    for data in rows:
        r += 1
        for j, v in enumerate(data):
            c = ws.cell(r, col1 + j, v)
            c.font = f_body
            c.alignment = center if j > 0 else wrap
            c.border = border
            if r % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F2F7FC")
    return r


def put_img(ws, path, anchor, w=None):
    img = XLImage(path)
    if w:
        ratio = w / img.width
        img.width = w; img.height = int(img.height * ratio)
    ws.add_image(img, anchor)


# ---- 目次 ----
ws = wb.active
ws.title = "目次"
setup(ws, [4, 24, 58, 14])
title_row(ws, 1, "特殊スラブ 問題集（RC マンション設計担当・新人向け）", span=4)
body(ws, 2,
     "目標：特殊スラブ（ボイドスラブ・ハーフPCaスラブ）の計算書が確認できること。"
     "大スパン・省人化のために多用される。自重・断面性能・FEM・合成効果・施工時検討・"
     "施工留意点を通す。",
     span=4, h=44)
r = table(ws, 4,
          ["No.", "シート", "到達目標", "図"],
          [["1", "1 ボイド構造特性",
            "自重設定・各等価スラブ厚を理解している", "断面・等価厚"],
           ["2", "2 ボイド断面性能",
            "断面性能 A・Z・I を理解している", "中実−ボイド"],
           ["3", "3 FEM 解析モデル",
            "境界条件・荷重・剛性の設定ができる", "メッシュ・異方性"],
           ["4", "4 ハーフPCa構造特性",
            "合成効果（PC・トラス等）を理解している", "合成断面"],
           ["5", "5 ハーフPCa施工/完成",
            "施工時・完成時の検討応力・断面性能を理解", "施工/完成"],
           ["6", "6 施工方法・留意点",
            "支保工・養生・品質管理を理解している", "施工留意"]])
body(ws, r + 2,
     "共通モデル：ボイドスラブ＝厚 500・円筒ボイド φ320@400。"
     "ハーフPCa＝PCa 板 60 ＋ トッピング 140 ＝ 合成 200。γc=24 kN/m³。"
     "数値は本教材作成時に検算済みだが、実務では製品メーカーの技術資料・"
     "各認定・RC規準で確認すること。",
     span=4, h=58)

# ---- 1 ボイド構造特性 ----
ws = wb.create_sheet("1 ボイド構造特性")
setup(ws, [8, 18, 16, 14, 14, 12, 12])
title_row(ws, 1, "1  ボイドスラブの構造特性（自重・等価スラブ厚）")
head(ws, 3, "■ 図 1  断面と 2 種類の等価スラブ厚")
put_img(ws, figs["void"], "A4", w=800)
head(ws, 28, "■ 問題 1  ボイドスラブとは")
body(ws, 29, "ボイドスラブの目的を 2 つ挙げよ（自重軽減による大スパン化／"
             "梁を減らしフラットな天井）。中間のコンクリートを抜いても曲げに効く理由を、"
             "断面の『上下の板』（I 形に近い）から説明せよ。", h=44)
head(ws, 31, "■ 問題 2  自重（重量等価厚）の計算")
body(ws, 32, "厚 500・φ320@400。1m 幅あたりの純断面積と自重を求めよ（γc=24）。",
     h=22)
r = table(ws, 34,
          ["項目", "式", "値（記入）"],
          [["中実断面積/m", "500×1000", ""],
           ["ボイド 1 本の面積", "π×320²/4", ""],
           ["1m 幅のボイド本数", "1000/400", ""],
           ["純断面積/m", "中実 − 本数×ボイド", ""],
           ["重量等価厚 t_w", "純断面積 / 1000", ""],
           ["自重 w", "純断面積×24×10⁻⁶", ""]])
body(ws, r + 2, "中実スラブ（500mm）の自重と比べ、何 % 軽くなるか。", h=20)
head(ws, r + 4, "■ 問題 3  2 つの等価スラブ厚の違い")
body(ws, r + 5, "『重量等価厚』と『曲げ剛性等価厚』はなぜ値が違うのか"
                "（重量は面積に比例、剛性は I ＝ 断面 2 次モーメントで中立軸から遠い"
                "コンクリートが効く）。どちらをどの検討に使うか整理せよ。", h=58)
head(ws, r + 7, "■ 問題 4  計算書チェック")
body(ws, r + 8, "ボイドスラブの計算書で、設計者が『自重』にうっかり中実厚（500mm）を"
                "使っていた。何が過大・過小になるか（自重過大→応力過大で不経済、"
                "ただし安全側）。逆に剛性に重量等価厚（299）を使うと"
                "たわみを過大評価する点も述べよ。", h=58)

# ---- 2 ボイド断面性能 ----
ws = wb.create_sheet("2 ボイド断面性能")
setup(ws, [8, 18, 16, 14, 14, 12, 12])
title_row(ws, 1, "2  ボイドスラブの断面性能（A・Z・I）")
head(ws, 3, "■ 図 2  断面性能の計算（中実 − ボイド円）")
put_img(ws, figs["section"], "A4", w=760)
head(ws, 26, "■ 問題 1  断面性能の計算（1 ピッチ b=400）")
body(ws, 27, "中実からボイド円を差し引いて A・I・Z を求めよ。"
             "図心はスラブ中央（上下対称）とする。", h=22)
r = table(ws, 29,
          ["項目", "式", "値（記入）"],
          [["A1（中実）", "400×500", ""],
           ["I1（中実）", "400×500³/12", ""],
           ["A2（ボイド）", "π×320²/4", ""],
           ["I2（ボイド）", "π×320⁴/64", ""],
           ["A = A1−A2", "", ""],
           ["I = I1−I2", "", ""],
           ["Z = I/(D/2)", "I/250", ""]])
head(ws, r + 2, "■ 問題 2  1m 幅への換算")
body(ws, r + 3, "上の値（400mm あたり）を 1m 幅に換算せよ（×1000/400）。"
                "I/m、Z/m を求め、中実スラブ（I=1000×500³/12）の何 % か。", h=32)
head(ws, r + 5, "■ 問題 3  応力度の照査")
body(ws, r + 6, "設計曲げモーメント M=80 kN·m/m のとき、縁応力度 σ=M/Z を求めよ。"
                "Z=3.65×10⁷ mm³/m を使用。コンクリートの許容曲げ圧縮応力度と比較する"
                "考え方も述べよ。", h=40)
head(ws, r + 8, "■ 問題 4  なぜ I の低下が小さいか")
body(ws, r + 9, "断面積は 40% も抜けるのに、曲げ剛性 I の低下は 12% にとどまる。"
                "この理由を『中立軸付近のコンクリートは曲げにあまり効かない』"
                "（I は中立軸からの距離の 2 乗で効く）という観点で説明せよ。"
                "ボイド管をスラブ中央に置く合理性がここにある。", h=58)

# ---- 3 FEM ----
ws = wb.create_sheet("3 FEM解析モデル")
setup(ws, [8, 18, 16, 14, 14, 12, 12])
title_row(ws, 1, "3  FEM 解析モデルの設定（境界条件・荷重・剛性）")
head(ws, 3, "■ 図 3  FEM モデル")
put_img(ws, figs["fem"], "A4", w=800)
head(ws, 30, "■ 問題 1  3 大入力（穴埋め）")
body(ws, 31, "FEM 板解析の 3 大入力は【 ① 】（メッシュ・要素種別）、"
             "【 ② 】（自重・仕上・積載）、【 ③ 】（板の曲げ剛性）である。"
             "境界条件は周辺の大梁を【 ④ 】支持、連続端は【 ⑤ 】を考慮する。",
     h=44)
head(ws, 33, "■ 問題 2  荷重の入力")
body(ws, 34, "ボイドスラブの FEM で自重を入力するとき、どの等価厚を使うか。"
             "また積載荷重（住宅 1.8 kN/m²）・仕上げ荷重を加えた設計荷重を組み立てよ。",
     h=32)
head(ws, 36, "■ 問題 3  剛性の入力（直交異方性）")
body(ws, 37, "ボイドスラブはボイド方向と直交方向で剛性が異なる（直交異方性）。"
             "(1) どちらの方向が剛性が大きいか。(2) 簡易には曲げ剛性等価厚"
             "（478mm）の等方板として入力することも多いが、その場合の注意点は。", h=44)
head(ws, 39, "■ 問題 4  境界条件の影響")
body(ws, 40, "同じスラブを (a) 四周ピン支持 と (b) 四周固定（連続）でモデル化すると、"
             "中央たわみ・スパン中央 M・端部 M はどう変わるか。"
             "実際のマンションのスラブはどちらに近いか（隣接スラブと連続）。", h=44)
head(ws, 42, "■ 問題 4b  計算書チェック")
body(ws, 43, "FEM 計算書で確認すべき点を 3 つ挙げよ"
             "（メッシュの粗密／荷重の合計値が手計算と合うか／境界条件が"
             "実状と整合／変形図が想定通りか）。", h=40)

# ---- 4 ハーフPCa構造特性 ----
ws = wb.create_sheet("4 ハーフPCa構造特性")
setup(ws, [8, 18, 16, 14, 14, 12, 12])
title_row(ws, 1, "4  ハーフPCaスラブの構造特性（合成効果）")
head(ws, 3, "■ 図 4  ハーフPCa の構成と合成効果")
put_img(ws, figs["halfpca"], "A4", w=800)
head(ws, 28, "■ 問題 1  ハーフPCa とは")
body(ws, 29, "ハーフPCaスラブの構成（工場製作の PCa 板＋現場打ちトッピング）と、"
             "メリットを 2 つ挙げよ（型枠・支保工の省略／工期短縮・品質安定）。", h=32)
head(ws, 31, "■ 問題 2  合成効果")
body(ws, 32, "(1) PCa 板とトッピングが一体化して 1 枚の断面として働くことを"
             "『合成』という。合成が成立するために接合面に必要なものを 2 つ"
             "（目荒し／トラス筋（ラチス）によるずれ止め）。", h=32)
body(ws, 33, "(2) 合成前（PCa 60mm 単独）と合成後（200mm）の I を比較せよ。"
             "I=b·t³/12 で計算し、何倍になるか（37 倍）。t が 3 乗で効くことを確認。",
     h=32)
r = table(ws, 36,
          ["状態", "断面厚", "I = 1000·t³/12（記入）"],
          [["施工時（PCa 単独）", "60mm", ""],
           ["完成時（合成）", "200mm", ""],
           ["比", "—", ""]])
head(ws, r + 2, "■ 問題 3  PC 方式とトラス筋方式")
body(ws, r + 3, "ハーフPCa には (a) PC 鋼材でプレストレスを与える方式と "
                "(b) トラス筋（ラチス）で補強する方式がある。"
                "それぞれの特徴（PC＝ひび割れ抑制・長スパン／トラス＝合成のずれ止め"
                "＋施工時補強）を述べよ。", h=44)
head(ws, r + 5, "■ 問題 4  計算書チェック")
body(ws, r + 6, "合成スラブの計算書で『完成時のみ』を検討し施工時を見落とすと"
                "何が危険か。逆に施工時の断面（60mm）で完成時荷重を検討すると"
                "どうなるか（過大）。両方の検討が必要な理由を述べよ。", h=44)

# ---- 5 ハーフPCa 施工/完成 ----
ws = wb.create_sheet("5 ハーフPCa施工完成")
setup(ws, [8, 18, 16, 14, 14, 12, 12])
title_row(ws, 1, "5  ハーフPCaスラブの施工時・完成時の検討")
head(ws, 3, "■ 図 5  施工時（PCa 単独）と完成時（合成）")
put_img(ws, figs["stages"], "A4", w=800)
head(ws, 28, "■ 問題 1  施工時の検討")
body(ws, 29, "施工時、PCa 板 60mm が単独で（未硬化のトッピングを載せて）支持する。"
             "支保工なし・スパン L=3.0m のとき、PCa に生じる M を求めよ。"
             "施工時荷重 w=PCa 自重＋打設コンクリート＋作業荷重。", h=44)
r = table(ws, 32,
          ["項目", "式・値", "記入"],
          [["PCa 自重", "0.06×24", ""],
           ["トッピング（生）", "0.14×24", ""],
           ["作業荷重", "1.5（目安）", ""],
           ["施工時 w 合計", "", ""],
           ["M = wL²/8（L=3m・単純梁）", "", ""],
           ["Z(PCa 60mm) = 1000×60²/6", "6.0×10⁵", ""],
           ["σ = M/Z", "", ""]])
head(ws, r + 2, "■ 問題 2  支保工の効果")
body(ws, r + 3, "施工時に支保工（サポート）を入れると、PCa のスパンが分割され M が"
                "激減する。L=3m を中央 1 点支持で 2 分割（L=1.5m）した場合、"
                "M は何分の 1 になるか（wL²/8 で L が半分→1/4）。", h=40)
head(ws, r + 5, "■ 問題 3  完成時の検討")
body(ws, r + 6, "完成時は合成断面 200mm で全荷重を支持する。"
                "設計荷重 w=自重 4.8＋仕上 2.0＋積載 1.8=8.6 kN/m²、L=6m 連続スラブ。"
                "M≒wL²/10（連続）で概算し、合成断面 Z=6.67×10⁶ mm³/m で σ を確認せよ。",
     h=44)
head(ws, r + 8, "■ 問題 4  応力の引き継ぎ")
body(ws, r + 9, "施工時に PCa に生じた応力は、合成後もそのまま残る（応力の履歴）。"
                "完成時の検討では『施工時応力＋合成後の追加応力』を重ねる必要がある"
                "ことを説明せよ。支保工の有無でこの初期応力が変わる。", h=44)

# ---- 6 施工方法・留意点 ----
ws = wb.create_sheet("6 施工方法・留意点")
setup(ws, [8, 18, 20, 16, 12, 12])
title_row(ws, 1, "6  特殊スラブの施工方法と構造的留意点")
head(ws, 3, "■ 問題 1  支保工")
body(ws, 4, "(1) ハーフPCa の支保工の目的を 2 つ（施工時応力の低減／たわみ・"
            "むくりの管理）。(2) 支保工の撤去時期は何で判断するか"
            "（トッピングコンクリートの強度発現＝合成断面が有効になる）。", h=44)
head(ws, 6, "■ 問題 2  ボイド管の施工留意点")
body(ws, 7, "(1) ボイド管の『浮き上がり』防止が重要な理由（打設時の浮力で管が浮くと"
            "かぶり・断面性能が狂う）。対策を 1 つ（浮き止め金物・打設速度管理）。"
            "(2) ボイド周りのコンクリートの充填不良（空洞）を防ぐ打設・締固めの注意。",
     h=44)
head(ws, 9, "■ 問題 3  養生")
body(ws, 10, "(1) トッピングコンクリートの養生期間が合成の成立に直結する理由。"
             "(2) 湿潤養生・強度確認（テストピース・現場養生供試体）の役割を述べよ。",
     h=40)
head(ws, 12, "■ 問題 4  品質管理")
r = table(ws, 13,
          ["管理項目", "確認内容（記入）"],
          [["PCa 板の製作精度", ""],
           ["接合面の目荒し", ""],
           ["トッピングの強度", ""],
           ["ボイド管の固定・かぶり", ""],
           ["トラス筋の定着・継手", ""]])
head(ws, r + 2, "■ 問題 5  計算書と施工の整合")
body(ws, r + 3, "設計（計算書）で仮定した『支保工あり／なし』『合成の成立』が、"
                "実際の施工で守られないと危険。設計者が施工計画で確認すべき点を"
                "3 つ挙げよ（支保工の位置・盤・撤去時期／養生期間／接合面処理）。",
     h=44)

# ---- 解答 ----
ws = wb.create_sheet("解答")
setup(ws, [4, 22, 60, 14])
title_row(ws, 1, "解答・解説", span=4)
row = 2


def ah(t):
    global row
    head(ws, row, t, span=4); row += 1


def an(t, h=None):
    global row
    body(ws, row, t, span=4, ans=True, h=h); row += 1


ah("1  ボイド構造特性")
an("問1：目的＝①自重軽減で大スパン・梁レス化 ②フラットスラブで天井すっきり・"
   "階高低減。中間コンクリートを抜いても、曲げは上下の板（圧縮側・引張側）が"
   "負担するので効率が落ちない（I 形断面に近い）。", h=44)
an("問2：中実 500×1000=500,000。ボイド 1 本 π×320²/4=80,425。本数 1000/400=2.5。"
   "純断面積=500,000−2.5×80,425=298,938 mm²/m。重量等価厚=298,938/1000≒299mm。"
   "自重=298,938×24×10⁻⁶=7.17 kN/m²。中実 12.0 に対し約 40% 軽い。", h=58)
an("問3：重量は断面積（体積）に比例するので純断面積そのまま（299mm 相当）。"
   "剛性 I は中立軸から遠いコンクリートが効き、抜いたのは中央付近なので低下が小さく、"
   "I の 3 乗根換算で 478mm 相当になる。"
   "→ 自重・質量は重量等価厚（299）、たわみ・剛性・周期は剛性等価厚（478）。", h=58)
an("問4：自重に中実 500 を使う→自重過大→応力・配筋過大で不経済（ただし安全側）。"
   "剛性に重量等価厚 299 を使う→I を過小評価→たわみを過大に見積もる"
   "（不経済または不要な増厚）。等価厚は用途で正しく使い分ける。", h=58)

ah("2  ボイド断面性能")
an("問1：A1=400×500=200,000。I1=400×500³/12=4.167×10⁹。"
   "A2=π×320²/4=80,425。I2=π×320⁴/64=0.515×10⁹。"
   "A=119,575 mm²、I=3.65×10⁹ mm⁴、Z=I/250=1.46×10⁷ mm³（1 ピッチ 400 あたり）。",
   h=58)
an("問2：×1000/400=×2.5。I=9.13×10⁹ mm⁴/m、Z=3.65×10⁷ mm³/m。"
   "中実 I=1000×500³/12=1.04×10¹⁰。比=9.13/10.4=0.88 → 88%。", h=44)
an("問3：σ=M/Z=80×10⁶/(3.65×10⁷)=2.19 N/mm²。"
   "これをコンクリートの許容曲げ圧縮応力度（長期 Fc/3 程度）と比較し、"
   "引張側は鉄筋が負担（ひび割れ許容）として配筋を決める。", h=44)
an("問4：I は中立軸からの距離 y の 2 乗（∫y²dA）で効く。中立軸付近（中央）の"
   "コンクリートは y が小さく I への寄与が小さいので、そこを抜いても I はあまり"
   "減らない。断面積は 40% 減っても I は 12% 減にとどまる。"
   "だからボイド管は断面の中央に置くのが合理的。", h=58)

ah("3  FEM 解析モデル")
an("問1：①要素（メッシュ・要素種別） ②荷重 ③剛性 ④（ピン）支持 ⑤回転拘束（連続）。",
   h=32)
an("問2：自重は重量等価厚（299mm→7.17 kN/m²）で入力。"
   "設計荷重=自重 7.17＋仕上 2.0＋積載 1.8=約 11 kN/m²（住宅）。"
   "※自重を剛性等価厚 478 で入れると重すぎる（誤り）。", h=44)
an("問3：(1)ボイド方向（管に沿う方向）が剛性大。(2)等方板（478mm）で近似する場合、"
   "本来は異方性なので、応力の方向性（弱軸方向のたわみ・配筋）を別途確認する。"
   "メーカー技術資料の異方性データがあればそれを用いるのが正確。", h=44)
an("問4：(a)ピン支持→中央たわみ・中央 M が大、端部 M=0。"
   "(b)固定→中央たわみ・中央 M が小、端部 M が発生。"
   "実際のマンションは隣接スラブと連続するので (b) 連続端に近く、"
   "端部上端筋・中央下端筋の両方を配筋する。", h=58)
an("問4b：①メッシュが応力集中部で十分細かいか ②荷重の総和が"
   "（面積×等分布荷重）と一致するか ③境界条件が連続・支持の実状と合うか "
   "④変形図・応力コンター図が想定通りか。", h=44)

ah("4  ハーフPCa構造特性")
an("問1：構成＝工場製作の PCa 板（型枠兼用）＋現場打ちトッピング。"
   "メリット＝①型枠・支保工の省略で省人化 ②工場製作で品質安定・工期短縮。", h=32)
an("問2：(1)接合面の目荒し（付着）＋トラス筋（ラチス）によるずれ止め（水平せん断伝達）。"
   "(2)施工時 I=1000×60³/12=1.8×10⁷。完成時 I=1000×200³/12=6.67×10⁸。"
   "比=37 倍。t が 3 乗で効くので厚さ 3.3 倍で剛性 37 倍。", h=58)
an("問3：PC 方式＝プレストレスでひび割れを抑え、長スパン・薄型に有利。"
   "トラス筋方式＝ラチスが施工時の PCa 補強＋合成のずれ止めを兼ね、"
   "一般的な中小スパンで多用。用途・スパンで使い分ける。", h=44)
an("問4：完成時のみ検討→施工時（PCa 60mm 単独）が荷重に耐えられず"
   "施工中のひび割れ・破損を見落とす。施工時断面で完成時を検討→過大で不経済。"
   "施工時（PCa 単独）と完成時（合成）は別断面・別荷重なので両方必要。", h=58)

ah("5  ハーフPCa 施工/完成")
an("問1：PCa 自重 0.06×24=1.44。トッピング 0.14×24=3.36。作業 1.5。"
   "w=1.44+3.36+1.5=6.3 kN/m²。M=wL²/8=6.3×3²/8=7.09 kN·m/m。"
   "Z=1000×60²/6=6.0×10⁵。σ=7.09×10⁶/6.0×10⁵=11.8 N/mm² → PCa の強度・"
   "PC/トラスで検討（施工時応力が大きいことがわかる）。", h=72)
an("問2：中央 1 点支持で L=1.5m に。M=wL²/8 は L² に比例するので (1.5/3)²=1/4。"
   "支保工で施工時応力を 1/4 に低減できる → 支保工の有無が PCa 設計を左右する。",
   h=44)
an("問3：M≒wL²/10=8.6×6²/10=31.0 kN·m/m。"
   "σ=31.0×10⁶/(6.67×10⁶)=4.6 N/mm²。合成断面（200mm）なら Z が大きく"
   "施工時（60mm）よりはるかに小さい応力で収まる。", h=44)
an("問4：施工時に PCa に生じた曲げ応力は硬化後も残留し、合成後の追加応力と重なる。"
   "よって完成時の PCa 下縁応力＝施工時応力＋合成後応力で照査する。"
   "支保工を入れて施工時応力を小さくすれば、この初期応力が減り完成時も有利。",
   h=58)

ah("6  施工方法・留意点")
an("問1：(1)支保工の目的＝①施工時の PCa 応力・たわみを低減 ②むくり・レベルの管理。"
   "(2)撤去時期＝トッピングコンクリートが所定強度に達し合成断面が有効になってから"
   "（強度確認供試体で判断）。", h=44)
an("問2：(1)打設時の生コンの浮力でボイド管が浮くと、下側かぶり増・上側減で断面性能・"
   "耐力が狂う→浮き止め金物で固定、打設速度・打上がり高さを管理。"
   "(2)ボイド管下・側面に空洞ができないよう、バイブレータで確実に充填・締固める。",
   h=58)
an("問3：(1)トッピングが所定強度に達して初めて PCa と一体（合成）になるため、"
   "養生不足は合成不成立＝設計耐力未達に直結。"
   "(2)湿潤養生でひび割れ・強度不足を防ぎ、現場養生供試体で実強度を確認して"
   "支保工撤去・後続工事の可否を判断。", h=58)
an("問4：PCa 製作精度＝寸法・そり・埋込金物位置。目荒し＝所定の粗度・レイタンス除去。"
   "トッピング強度＝供試体で Fc 確認。ボイド固定＝浮き上がり・かぶり。"
   "トラス筋＝定着・継手・かぶり。", h=58)
an("問5：①支保工の位置・受け盤（沈下しないか）・撤去時期が計算の仮定通りか "
   "②トッピングの養生期間・強度確認 ③接合面の目荒し・清掃が合成の前提を満たすか。"
   "設計の仮定（支保工あり/なし、合成成立）と施工計画を必ず突き合わせる。", h=58)

XLSX = os.path.join(OUT, "特殊スラブ問題集.xlsx")
wb.save(XLSX)
print("saved:", XLSX)
