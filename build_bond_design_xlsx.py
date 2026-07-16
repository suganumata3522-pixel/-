# -*- coding: utf-8 -*-
"""付着設計 問題集（図つき）Excel 生成スクリプト。
出力: docs/bond_design/付着設計問題集.xlsx
No.8-1 付着の応力伝達機構 / 8-2 許容付着応力度fa・付着割裂の基準となる強度fb
8-3 付着設計法 / 8-4 各規準の目的と選択 / 8-5 カットオフ筋・通し筋の検討
RC造マンションの設計担当を想定。
数値・式は web 検証済み（平12建告1450号、RC規準2010/2018、靭性指針1999）。
"""
import os
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib import font_manager as fm

FONT_PATH = "/usr/share/fonts/truetype/fonts-japanese-gothic.ttf"
jp = fm.FontProperties(fname=FONT_PATH)
fm.fontManager.addfont(FONT_PATH)
plt.rcParams["font.family"] = jp.get_name()
plt.rcParams["axes.unicode_minus"] = False

OUT = "docs/bond_design"
FIG = os.path.join(OUT, "figures")
os.makedirs(FIG, exist_ok=True)

# dataviz 検証済みチャート配色（2系列）
C_BLUE = "#2a78d6"
C_PINK = "#d55181"


def save(fig, name):
    p = os.path.join(FIG, name)
    fig.savefig(p, dpi=130, bbox_inches="tight")
    plt.close(fig)
    return p


# ===========================================================================
# 図 8-1: 付着の応力伝達機構（3段階）と破壊モード
# ===========================================================================
def fig_8_1():
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 5.2),
                              gridspec_kw={"width_ratios": [1.2, 1.0]})

    # --- (a) 3段階の抵抗機構 ---
    ax = axes[0]
    stages = [
        ("① 化学的付着", "セメントペーストの接着。\nごく小さい滑りで切れる", "#9ec6e8"),
        ("② 摩擦", "界面の摩擦抵抗。\n滑りが進むと低下", "#ffd180"),
        ("③ 節の支圧（機械的咬合）", "節がコンクリートを押す支圧。\n異形鉄筋の付着の主役", "#f4a261"),
    ]
    y = 4.4
    for name, desc, color in stages:
        ax.add_patch(mpatches.FancyBboxPatch(
            (0.3, y), 8.6, 1.5, boxstyle="round,pad=0.06",
            fc=color, ec="k", lw=1))
        ax.text(0.7, y + 1.05, name, fontproperties=jp, fontsize=11,
                fontweight="bold", va="center")
        ax.text(4.6, y + 0.55, desc, fontproperties=jp, fontsize=8.5,
                va="center", color="#333")
        y -= 2.0
    ax.annotate("", xy=(9.6, 0.6), xytext=(9.6, 5.7),
                arrowprops=dict(arrowstyle="-|>", color="#c00000", lw=2.5))
    ax.text(9.95, 3.2, "滑りの進行とともに移行", fontproperties=jp,
            rotation=90, fontsize=9, color="#c00000", va="center")
    ax.set_xlim(0, 10.8)
    ax.set_ylim(0, 6.4)
    ax.axis("off")
    ax.set_title("(a) 付着抵抗の 3 段階（異形鉄筋）",
                 fontproperties=jp, fontsize=11, fontweight="bold")

    # --- (b) 2つの破壊モード（断面図） ---
    ax = axes[1]
    # 引抜き破壊（左）：かぶり大
    cx1, cy = 2.2, 3.4
    ax.add_patch(mpatches.Rectangle((cx1 - 1.7, cy - 1.7), 3.4, 3.4,
                                     fc="#e8e8e8", ec="k", lw=1.0))
    ax.add_patch(plt.Circle((cx1, cy), 0.35, fc="#888", ec="k", zorder=4))
    # 円筒状のせん断面
    ax.add_patch(plt.Circle((cx1, cy), 0.62, fc="none", ec="#c00000",
                             lw=1.6, ls="--", zorder=5))
    ax.text(cx1, cy - 2.45, "引抜き破壊\n（かぶり・横補強が十分\n→ 鉄筋周囲で滑る）",
            fontproperties=jp, ha="center", fontsize=8.5, color="#444")
    # 割裂破壊（右）：かぶり小、放射ひび割れ
    cx2 = 7.2
    ax.add_patch(mpatches.Rectangle((cx2 - 1.7, cy - 1.7), 3.4, 3.4,
                                     fc="#e8e8e8", ec="k", lw=1.0))
    ax.add_patch(plt.Circle((cx2, cy), 0.35, fc="#888", ec="k", zorder=4))
    for ang in [90, 210, 330]:
        a = np.deg2rad(ang)
        x1 = cx2 + 0.38 * np.cos(a); y1 = cy + 0.38 * np.sin(a)
        x2 = cx2 + 1.65 * np.cos(a); y2 = cy + 1.65 * np.sin(a)
        ax.plot([x1, x2], [y1, y2], color="#c00000", lw=2, zorder=5)
    ax.text(cx2, cy - 2.45, "付着割裂破壊\n（かぶり・あき小\n→ かぶりが割れて剥落）",
            fontproperties=jp, ha="center", fontsize=8.5, color="#c00000")
    ax.text(4.7, 6.0, "節の支圧の反力（内圧）がコンクリートを押し広げる",
            fontproperties=jp, ha="center", fontsize=9, color="#1f4e79")
    ax.set_xlim(0, 9.4)
    ax.set_ylim(0.2, 6.6)
    ax.set_aspect("equal")
    ax.axis("off")
    ax.set_title("(b) 付着の 2 つの破壊モード（断面）",
                 fontproperties=jp, fontsize=11, fontweight="bold")
    fig.suptitle("図 8-1  付着の応力伝達機構と破壊モード",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig8-1_mechanism.png")


# ===========================================================================
# 図 8-2: 許容付着応力度 fa（告示1450号）のグラフ ＋ fb の位置づけ
# ===========================================================================
def fig_8_2():
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 5.0),
                              gridspec_kw={"width_ratios": [1.15, 1.0]})

    # --- (a) fa チャート（長期、異形鉄筋） ---
    ax = axes[0]
    Fc = np.linspace(15, 45, 200)
    fa_other = np.where(Fc <= 22.5, Fc / 10, 1.35 + Fc / 25)
    fa_top = np.where(Fc <= 22.5, Fc / 15, 0.9 + 2 * Fc / 75)
    ax.plot(Fc, fa_other, color=C_BLUE, lw=2)
    ax.plot(Fc, fa_top, color=C_PINK, lw=2)
    # 直接ラベル（線の右端）
    ax.text(45.4, fa_other[-1], "その他の鉄筋", fontproperties=jp,
            fontsize=9, color=C_BLUE, va="center")
    ax.text(45.4, fa_top[-1], "上端筋", fontproperties=jp,
            fontsize=9, color=C_PINK, va="center")
    ax.axvline(22.5, color="#bbb", lw=1, ls=":")
    ax.text(22.5, 3.05, "Fc=22.5 で式が切替", fontproperties=jp,
            fontsize=8, color="#888", ha="center")
    # 式の注記
    ax.text(31, 1.25, "上端筋：Fc/15（Fc>22.5 は 0.9+2Fc/75）",
            fontproperties=jp, fontsize=8.5, color=C_PINK)
    ax.text(17, 2.55, "その他：Fc/10\n（Fc>22.5 は 1.35+Fc/25）",
            fontproperties=jp, fontsize=8.5, color=C_BLUE)
    ax.set_xlabel("コンクリート設計基準強度 Fc (N/mm²)", fontproperties=jp)
    ax.set_ylabel("長期許容付着応力度 fa (N/mm²)", fontproperties=jp)
    ax.set_xlim(15, 52)
    ax.set_ylim(0.8, 3.2)
    ax.grid(alpha=0.25)
    ax.set_title("(a) 告示1450号の長期 fa（異形鉄筋）── 短期はこの 2 倍",
                 fontproperties=jp, fontsize=11, fontweight="bold")

    # --- (b) fa と fb の関係図 ---
    ax = axes[1]
    boxes = [
        (0.76, "#cfe0f0",
         "許容付着応力度 fa（法令）",
         "平12建告1450号（令91条に基づく）\n"
         "長期：上端筋 Fc/15、その他 Fc/10（Fc≦22.5）\n"
         "短期＝長期の 2 倍。かぶり<1.5db で低減あり"),
        (0.44, "#fde2c4",
         "付着割裂の基準となる強度 fb（RC規準）",
         "RC規準2010/2018 の付着検定に使う基準強度\n"
         "その他の鉄筋：fb = Fc/40 + 0.9\n"
         "上端筋：その 0.8 倍。多段配筋 2 段目以降は 0.6 倍"),
        (0.12, "#cfead4",
         "許容付着応力度（RC規準）＝ K × fb",
         "K：かぶり・あき(C)と横補強筋(W)による修正係数\n"
         "K = 0.3(C+W)/db + 0.4 ≦ 2.5\n"
         "かぶり・帯筋が多いほど付着を高く評価できる"),
    ]
    for y, color, title, desc in boxes:
        ax.add_patch(mpatches.FancyBboxPatch(
            (0.03, y), 0.94, 0.24, boxstyle="round,pad=0.01",
            fc=color, ec="k", lw=1, transform=ax.transAxes))
        ax.text(0.06, y + 0.185, title, transform=ax.transAxes,
                fontproperties=jp, fontsize=10, fontweight="bold")
        ax.text(0.06, y + 0.075, desc, transform=ax.transAxes,
                fontproperties=jp, fontsize=8, color="#333")
    ax.axis("off")
    ax.set_title("(b) fa（法令）と fb（RC規準）の位置づけ",
                 fontproperties=jp, fontsize=11, fontweight="bold")
    fig.suptitle("図 8-2  許容付着応力度 fa と 付着割裂の基準となる強度 fb",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig8-2_fa_fb.png")


# ===========================================================================
# 図 8-4: 規準マップ
# ===========================================================================
def fig_8_4():
    fig, ax = plt.subplots(figsize=(12.5, 5.6))

    def box(x, y, w, h, title, desc, color, tfs=10):
        ax.add_patch(mpatches.FancyBboxPatch(
            (x, y), w, h, boxstyle="round,pad=0.05", fc=color, ec="k", lw=1))
        ax.text(x + w / 2, y + h - 0.38, title, fontproperties=jp,
                ha="center", fontsize=tfs, fontweight="bold")
        ax.text(x + w / 2, y + (h - 0.5) / 2 - 0.05, desc, fontproperties=jp,
                ha="center", fontsize=8, color="#333", va="center")

    box(0.2, 3.6, 3.6, 1.9, "建築基準法・令91条\n＋平12建告1450号",
        "【最低基準（法令）】\n付着の許容応力度 fa を規定\n検討方法自体は規定しない", "#f8c0c0")
    box(4.4, 3.6, 3.6, 1.9, "RC規準 2010/2018\n（日本建築学会）",
        "【許容応力度設計の標準】\n長期=使用性／短期=損傷制御\n＋大地震=安全性の 3 段階検定", "#cfe0f0")
    box(8.6, 3.6, 3.6, 1.9, "靭性保証型耐震設計指針\n（1999）",
        "【大地震時の靭性保証】\nヒンジ域の付着劣化を制御\nτf ≦ τbu（付着信頼強度）", "#cfead4")
    box(4.4, 0.7, 3.6, 1.9, "技術基準解説書（黄色本）",
        "【確認審査の実務運用】\nRC規準・靭性指針の式を引用\nカットオフ筋の扱い等を Q&A で補完", "#fde2c4")
    ax.annotate("", xy=(4.4, 4.55), xytext=(3.8, 4.55),
                arrowprops=dict(arrowstyle="-|>", color="#444", lw=1.5))
    ax.annotate("", xy=(8.6, 4.55), xytext=(8.0, 4.55),
                arrowprops=dict(arrowstyle="-|>", color="#444", lw=1.5))
    ax.annotate("", xy=(6.2, 2.6), xytext=(6.2, 3.6),
                arrowprops=dict(arrowstyle="-|>", color="#444", lw=1.5))
    ax.text(4.1, 4.85, "詳細化", fontproperties=jp, fontsize=8, color="#444",
            ha="center")
    ax.text(8.3, 4.85, "靭性へ拡張", fontproperties=jp, fontsize=8,
            color="#444", ha="center")
    ax.text(6.55, 3.0, "実務運用", fontproperties=jp, fontsize=8, color="#444")
    ax.text(6.2, 0.15, "一般の RC マンション：告示 fa ＋ RC規準/黄色本が基本。"
                       "靭性依存の設計（高層・Ds 低減）では靭性指針を参照",
            fontproperties=jp, ha="center", fontsize=9.5, color="#1f4e79")
    ax.set_xlim(0, 12.5)
    ax.set_ylim(-0.2, 6.0)
    ax.axis("off")
    ax.set_title("図 8-4  付着に関する規準の全体マップと使い分け",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig8-4_codes.png")


# ===========================================================================
# 図 8-5: カットオフ筋と通し筋
# ===========================================================================
def fig_8_5():
    fig, axes = plt.subplots(1, 2, figsize=(14, 5.6),
                              gridspec_kw={"width_ratios": [1.5, 1.0]})

    # --- (a) カットオフ筋 ---
    ax = axes[0]
    L = 10.0
    yb = 3.0
    for xc in [0, L]:
        ax.add_patch(mpatches.Rectangle((xc - 0.45, yb - 2.0), 0.9, 4.2,
                                         fc="#d9d9d9", ec="k", lw=1.0))
    ax.add_patch(mpatches.Rectangle((0.45, yb - 0.45), L - 0.9, 0.9,
                                     fc="#bcd2ea", ec="k", lw=1.0))
    # M図（上端引張側のみ表示：端部負曲げ）
    xs = np.linspace(0.45, L - 0.45, 100)
    M = -1.0 + 6.0 * (xs / L) * (1 - xs / L)
    ax.plot(xs, yb - M * 1.0, color="#999", lw=1.4, ls="--")
    # 通し筋（上端 1段目）
    ax.plot([0.45, L - 0.45], [yb + 0.28, yb + 0.28], color="#1f4e79", lw=2.5)
    ax.text(L / 2, yb + 0.5, "通し筋（1 段目：全長通す）", fontproperties=jp,
            ha="center", fontsize=8.5, color="#1f4e79")
    # カットオフ筋（上端 2段目、左端から途中まで）
    xcut_th = 2.6   # 理論カットオフ点（計算上不要となる断面）
    xcut = 4.0      # 実際の切断位置
    ax.plot([0.45, xcut], [yb + 0.12, yb + 0.12], color="#c00000", lw=2.5)
    ax.plot(xcut, yb + 0.12, "o", color="#c00000", ms=5)
    ax.axvline(xcut_th, color="#1f7a1f", lw=1.2, ls=":")
    ax.annotate("理論カットオフ点\n（計算上不要となる断面）",
                xy=(xcut_th, yb - 1.35), xytext=(xcut_th + 0.4, yb - 2.4),
                fontproperties=jp, fontsize=8, color="#1f7a1f",
                arrowprops=dict(arrowstyle="->", color="#1f7a1f"))
    ax.annotate("", xy=(xcut_th, yb + 0.75), xytext=(xcut, yb + 0.75),
                arrowprops=dict(arrowstyle="<|-|>", color="#c00000", lw=1.2))
    ax.text((xcut_th + xcut) / 2, yb + 0.95,
            "規準：有効せい d 以上延長", fontproperties=jp, ha="center",
            fontsize=8.5, color="#c00000")
    ax.annotate("", xy=(0.45, yb + 1.45), xytext=(xcut, yb + 1.45),
                arrowprops=dict(arrowstyle="<|-|>", color="#7a3b3b", lw=1.2))
    ax.text((0.45 + xcut) / 2, yb + 1.65,
            "標準配筋の慣用：柱面から L0/4 + 15d", fontproperties=jp,
            ha="center", fontsize=8.5, color="#7a3b3b")
    ax.text(L / 2, yb - 2.9,
            "カットオフ筋は末端が自由端 → 付着検定が必要（存在応力を σy とし、"
            "有効付着長さ ld−d で平均付着応力度を検定）",
            fontproperties=jp, ha="center", fontsize=8.5, color="#444")
    ax.set_xlim(-0.8, L + 0.8)
    ax.set_ylim(yb - 3.3, yb + 2.4)
    ax.set_aspect("equal")
    ax.axis("off")
    ax.set_title("(a) カットオフ筋 ── 延長規定と付着検定",
                 fontproperties=jp, fontsize=11, fontweight="bold")

    # --- (b) 通し筋の仕口内付着 ---
    ax = axes[1]
    Dc = 2.6   # 柱せい
    ax.add_patch(mpatches.Rectangle((0, 0), Dc, 6.0, fc="#d9d9d9",
                                     ec="k", lw=1.2))
    ax.add_patch(mpatches.Rectangle((0, 2.4), Dc, 1.4, fc="#c9d9c0",
                                     ec="k", lw=0.8, hatch="..", zorder=1))
    for x0, x1 in [(-2.2, 0), (Dc, Dc + 2.2)]:
        ax.add_patch(mpatches.Rectangle((x0, 2.4), x1 - x0, 1.4,
                                         fc="#bcd2ea", ec="k", lw=0.8))
    # 通し筋
    ax.plot([-2.2, Dc + 2.2], [3.5, 3.5], color="#c00000", lw=3)
    # 左＝引張、右＝圧縮
    ax.annotate("", xy=(-2.2, 3.5), xytext=(-3.3, 3.5),
                arrowprops=dict(arrowstyle="<|-", color="#c00000", lw=2.5))
    ax.text(-3.4, 3.5, "引張\nσy", fontproperties=jp, ha="right",
            va="center", fontsize=9, color="#c00000")
    ax.annotate("", xy=(Dc + 2.2, 3.5), xytext=(Dc + 3.3, 3.5),
                arrowprops=dict(arrowstyle="<|-", color=C_BLUE, lw=2.5))
    ax.text(Dc + 3.4, 3.5, "圧縮", fontproperties=jp, ha="left",
            va="center", fontsize=9, color=C_BLUE)
    ax.annotate("", xy=(0, 4.35), xytext=(Dc, 4.35),
                arrowprops=dict(arrowstyle="<|-|>", color="#1f7a1f", lw=1.3))
    ax.text(Dc / 2, 4.6, "柱せい D の区間で\n応力が引張→圧縮に反転",
            fontproperties=jp, ha="center", fontsize=8.5, color="#1f7a1f")
    ax.text(Dc / 2, 0.8, "仕口内の平均付着応力度が\n非常に大きくなる",
            fontproperties=jp, ha="center", fontsize=8.5, color="#c00000")
    ax.text(Dc / 2, -1.1,
            "→ 柱せい D と梁主筋径 db の比を制限\n（D/db ≧ 20 程度が目安。"
            "σy が高いほど、Fc が低いほど厳しい）",
            fontproperties=jp, ha="center", fontsize=9, color="#1f4e79")
    ax.set_xlim(-4.6, Dc + 4.6)
    ax.set_ylim(-1.8, 6.4)
    ax.set_aspect("equal")
    ax.axis("off")
    ax.set_title("(b) 通し筋 ── 仕口内の付着", fontproperties=jp,
                 fontsize=11, fontweight="bold")
    fig.suptitle("図 8-5  カットオフ筋と通し筋の付着検討",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig8-5_cutoff.png")


figs = {
    "f1": fig_8_1(),
    "f2": fig_8_2(),
    "f4": fig_8_4(),
    "f5": fig_8_5(),
}
print("figures:", list(figs.keys()))

# ===========================================================================
# Excel 構築
# ===========================================================================
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

wb = Workbook()
C_TITLE = "1F4E79"; C_HEAD = "2E75B6"; C_ANS = "E2EFDA"
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
f_title = Font(name="MS PGothic", size=15, bold=True, color="FFFFFF")
f_head = Font(name="MS PGothic", size=11, bold=True, color="FFFFFF")
f_body = Font(name="MS PGothic", size=10)
f_ans = Font(name="MS PGothic", size=10, color="375623")
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)


def setup(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False


def title_row(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text)
    c.font = f_title; c.fill = PatternFill("solid", fgColor=C_TITLE)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[row].height = 30


def head(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text)
    c.font = f_head; c.fill = PatternFill("solid", fgColor=C_HEAD)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[row].height = 22


def body(ws, row, text, span=8, ans=False, h=None):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text)
    c.font = f_ans if ans else f_body
    c.alignment = wrap
    if ans:
        c.fill = PatternFill("solid", fgColor=C_ANS)
    if h:
        ws.row_dimensions[row].height = h


def table(ws, start_row, headers, rows, col1=1):
    r = start_row
    for j, htxt in enumerate(headers):
        c = ws.cell(r, col1 + j, htxt)
        c.font = Font(name="MS PGothic", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C_HEAD)
        c.alignment = center; c.border = border
    for data in rows:
        r += 1
        for j, v in enumerate(data):
            c = ws.cell(r, col1 + j, v)
            c.font = f_body
            c.alignment = center if j > 0 else wrap
            c.border = border
            if r % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F2F7FC")
    return r


def put_img(ws, path, anchor, w=None):
    img = XLImage(path)
    if w:
        ratio = w / img.width
        img.width = w; img.height = int(img.height * ratio)
    ws.add_image(img, anchor)


# ---- 目次 ----
ws = wb.active
ws.title = "目次"
setup(ws, [4, 22, 60, 14])
title_row(ws, 1, "付着設計 問題集（RC マンション設計担当・新人向け）", span=4)
body(ws, 2,
     "テーマ No.8：鉄筋とコンクリートの付着。伝達機構、法令の許容付着応力度 fa と"
     "RC規準の付着割裂の基準となる強度 fb、付着設計法、規準の使い分け、"
     "カットオフ筋・通し筋の検討までを通す。定着（No.6）・継手（No.7）とセットで学ぶこと。",
     span=4, h=44)
r = table(ws, 4,
          ["No.", "シート", "到達目標", "図"],
          [["8-1", "8-1 応力伝達機構",
            "付着の応力伝達機構を説明できる", "3 段階・破壊モード"],
           ["8-2", "8-2 fa と fb",
            "許容付着応力度 fa と付着割裂の基準となる強度 fb を理解している",
            "fa チャート"],
           ["8-3", "8-3 付着設計法",
            "付着設計法（許容付着応力度・付着割裂強度）を説明できる", "—"],
           ["8-4", "8-4 規準の選択",
            "各規準の目的を理解し準拠基準の選択ができる", "規準マップ"],
           ["8-5", "8-5 カットオフ・通し筋",
            "カットオフ筋・通し筋の付着検討ができる", "M図・仕口"]])
body(ws, r + 2,
     "凡例：fa＝許容付着応力度（平12建告1450号）、fb＝付着割裂の基準となる強度"
     "（RC規準2010/2018）、K＝付着割裂の修正係数、db＝鉄筋径、d＝有効せい。"
     "数値・式は告示・学会資料等の公開情報で照合済みだが、設計では必ず最新の"
     "規準原文・自社基準で確認すること。",
     span=4, h=58)

# ---- 8-1 ----
ws = wb.create_sheet("8-1 応力伝達機構")
setup(ws, [8, 16, 18, 16, 14, 12, 12, 12])
title_row(ws, 1, "8-1  付着の応力伝達機構")
head(ws, 3, "■ 図 8-1  付着抵抗の 3 段階と破壊モード")
put_img(ws, figs["f1"], "A4", w=820)
head(ws, 30, "■ 問題 1  3 段階の穴埋め")
body(ws, 31, "異形鉄筋の付着抵抗は、滑りの進行とともに【 ① 】的付着 → 【 ② 】→ "
             "節の【 ③ 】（機械的咬合）へと移行する。異形鉄筋の付着の主役は③である。"
             "丸鋼は③がないため付着に頼れず、原則【 ④ 】による定着が必要となる。",
     h=44)
head(ws, 33, "■ 問題 2  破壊モードの判別")
body(ws, 34, "次の状況では「引抜き破壊」「付着割裂破壊」のどちらが起こりやすいか。",
     h=20)
r = table(ws, 36,
          ["No.", "状況", "破壊モード（記入）"],
          [["(1)", "かぶりが厚く、帯筋も十分な柱の主筋", ""],
           ["(2)", "かぶりが薄く、鉄筋のあきが小さい梁の隅角部の主筋", ""],
           ["(3)", "太径鉄筋を密に並べた梁の 1 段目主筋", ""]])
head(ws, r + 2, "■ 問題 3  節の支圧と割裂の関係")
body(ws, r + 3, "節の支圧の反力は、鉄筋を中心とした放射方向の圧力（内圧）として"
                "コンクリートに作用する。この内圧がかぶりコンクリートに"
                "【 ① 】応力を生じさせ、限界を超えると鉄筋に沿った割裂ひび割れが発生する。"
                "割裂を抑える要素を 3 つ挙げよ（かぶり厚さ／鉄筋のあき／横補強筋）。",
     h=58)
head(ws, r + 5, "■ 問題 4  定着・継手・付着の関係整理")
body(ws, r + 6, "定着（No.6）・継手（No.7）・付着（No.8）は同じ現象の別の顔である。"
                "『付着』を共通の土台として、定着長 Lb・重ね継手長 L1 がどう決まるかを"
                "1〜2 行で整理せよ（力の釣合い T = 付着抵抗）。", h=44)

# ---- 8-2 ----
ws = wb.create_sheet("8-2 fa と fb")
setup(ws, [8, 14, 14, 14, 14, 14, 12, 12])
title_row(ws, 1, "8-2  許容付着応力度 fa と 付着割裂の基準となる強度 fb")
head(ws, 3, "■ 図 8-2  fa のグラフと fb の位置づけ")
put_img(ws, figs["f2"], "A4", w=820)
head(ws, 30, "■ 問題 1  fa の計算（平12建告1450号・異形鉄筋）")
body(ws, 31, "長期許容付着応力度：上端筋 fa = Fc/15（Fc>22.5 では 0.9+2Fc/75）、"
             "その他 fa = Fc/10（Fc>22.5 では 1.35+Fc/25）。短期は長期の 2 倍。"
             "下表を計算せよ（N/mm²、小数 2 桁）。※ Fc=24 以上は式の切替に注意！",
     h=44)
r = table(ws, 34,
          ["Fc", "長期・上端筋（記入）", "長期・その他（記入）",
           "短期・上端筋（記入）", "短期・その他（記入）"],
          [["21", "", "", "", ""],
           ["24", "", "", "", ""],
           ["27", "", "", "", ""],
           ["30", "", "", "", ""],
           ["36", "", "", "", ""]])
head(ws, r + 2, "■ 問題 2  上端筋の定義")
body(ws, r + 3, "告示・規準でいう『上端筋』の定義を答えよ"
                "（曲げ材で、その鉄筋の下に【 何mm 】以上のコンクリートが"
                "打ち込まれる水平鉄筋）。また fa が低い理由（ブリーディング）も述べよ。",
     h=44)
head(ws, r + 5, "■ 問題 3  fb の計算（RC規準2010/2018）")
body(ws, r + 6, "付着割裂の基準となる強度：その他の鉄筋 fb = Fc/40 + 0.9、"
                "上端筋はその 0.8 倍。Fc=24・27・30 について fb を計算せよ。", h=32)
r = table(ws, r + 8,
          ["Fc", "fb・その他（記入）", "fb・上端筋（記入）"],
          [["24", "", ""],
           ["27", "", ""],
           ["30", "", ""]])
head(ws, r + 2, "■ 問題 4  修正係数 K")
body(ws, r + 3, "RC規準の許容付着応力度は K·fb で表され、K = 0.3(C+W)/db + 0.4（上限 2.5）。"
                "C はかぶり・あきから決まる寸法（5db 以下）、W は横補強筋の効果（2.5db 以下）。"
                "db=25mm、C=60mm、W=40mm のとき K を計算せよ。"
                "また、かぶり・帯筋が増えると付着をどう評価できるか 1 行で。", h=58)
head(ws, r + 5, "■ 問題 5  法令と学会規準の違い（注意点）")
body(ws, r + 6, "短期の扱いは、法令（告示1450号）では長期の【 ① 】倍、"
                "RC規準では長期の【 ② 】倍と食い違いがある。"
                "実務でどちらを使うかは何で決めるか（審査上の位置づけ・設計方針書）。",
     h=44)

# ---- 8-3 ----
ws = wb.create_sheet("8-3 付着設計法")
setup(ws, [8, 18, 20, 16, 14, 12, 12])
title_row(ws, 1, "8-3  付着設計法（許容付着応力度・付着割裂強度）")
head(ws, 3, "■ 設計の枠組み（RC規準2010 の 3 段階）")
r = table(ws, 5,
          ["段階", "対象", "検定内容"],
          [["使用性（長期）", "常時荷重", "存在付着応力度 τ ≦ 長期の許容付着応力度"],
           ["損傷制御（短期）", "中地震", "存在付着応力度 τ ≦ 短期の許容付着応力度"],
           ["安全性（大地震）", "大地震", "鉄筋応力 1.1σy に対し付着割裂強度 K·fb で検討"]])
head(ws, r + 2, "■ 問題 1  検定式の理解（穴埋め）")
body(ws, r + 3, "平均付着応力度の検定式：τ = σt·db / (4·ld) ≦ fa（または K·fb）。"
                "この式は、鉄筋の引張力 T = σt·(π·db²/4) を、付着面積 =【 ① 】× 付着長さ ld "
                "で除したものである。同値変形すると、必要付着長さ ld ≧ σt·db / (4·【 ② 】) となる。",
     h=44)
head(ws, r + 5, "■ 問題 2  必要付着長さの計算")
body(ws, r + 6, "D25（db=25）、σt = 345 N/mm²（短期・降伏想定）、"
                "許容付着応力度 K·fb = 3.0 N/mm² のとき、必要付着長さ ld を求めよ（mm と db 倍数）。",
     h=32)
head(ws, r + 8, "■ 問題 3  テンションシフト")
body(ws, r + 9, "大地震時の検討では、端部のせん断ひび割れにより鉄筋の引張域が"
                "広がる（テンションシフト）を考慮し、付着の有効長さを ld から"
                "【 ① 】を差し引いた（ld − d）とする。なぜ安全側になるか 1 行で。", h=44)
head(ws, r + 11, "■ 問題 4  付着検討の省略")
body(ws, r + 12, "RC規準2018 では、どのような場合に付着の検討を省略できる扱いがあるか"
                 "（全数通し配筋でせん断の安全性検討を行う場合等）。"
                 "省略できる理由も 1 行で述べよ。", h=44)
head(ws, r + 14, "■ 問題 5  設計フロー")
body(ws, r + 15, "梁主筋の付着設計の手順を並べ替えよ："
                 "(A) K を算定（かぶり・あき・横補強筋） (B) 存在応力度 σt を求める "
                 "(C) fb を算定（Fc・上端筋か否か） (D) τ ≦ K·fb を検定 "
                 "(E) NG なら径を細く・本数増・カットオフ見直し。", h=44)

# ---- 8-4 ----
ws = wb.create_sheet("8-4 規準の選択")
setup(ws, [8, 18, 20, 16, 14, 12, 12])
title_row(ws, 1, "8-4  各規準の目的と準拠基準の選択")
head(ws, 3, "■ 図 8-4  規準マップ")
put_img(ws, figs["f4"], "A4", w=780)
head(ws, 28, "■ 問題 1  規準と目的の対応")
r = table(ws, 29,
          ["規準", "目的・位置づけ（記入）"],
          [["建築基準法・令91条＋告示1450号", ""],
           ["RC規準 2010/2018（日本建築学会）", ""],
           ["靭性保証型耐震設計指針（1999）", ""],
           ["技術基準解説書（黄色本）", ""]])
head(ws, r + 2, "■ 問題 2  靭性指針の付着設計")
body(ws, r + 3, "(1) 靭性指針が付着劣化を重視する理由を、塑性ヒンジの変形能力・"
                "エネルギー吸収の観点から述べよ。", h=32)
body(ws, r + 4, "(2) 検定の形は τf ≦ τbu。τbu（付着信頼強度）はどんな破壊に対する"
                "下限強度か（付着割裂）。", h=32)
head(ws, r + 6, "■ 問題 3  準拠基準の選択")
body(ws, r + 7, "次のプロジェクトで付着検討の準拠基準を選び、理由を述べよ。", h=20)
r = table(ws, r + 9,
          ["No.", "プロジェクト", "準拠基準（記入）", "理由（記入）"],
          [["(1)", "10 階建 RC マンション（ルート3・保有水平耐力計算）", "", ""],
           ["(2)", "高層 RC（靭性に依存、Ds を小さく設定）", "", ""],
           ["(3)", "5 階建 RC マンション（ルート1）", "", ""]])
head(ws, r + 2, "■ 問題 4  規準間の食い違いへの向き合い方")
body(ws, r + 3, "法令と学会規準で数値・扱いが異なる場合（例：短期 fa の倍率）、"
                "新人としてどう対処すべきか（設計方針の確認・上司/審査機関との整合・"
                "勝手にいいとこ取りしない）を 2 行で。", h=44)

# ---- 8-5 ----
ws = wb.create_sheet("8-5 カットオフ・通し筋")
setup(ws, [8, 16, 18, 16, 14, 12, 12, 12])
title_row(ws, 1, "8-5  カットオフ筋・通し筋の付着検討")
head(ws, 3, "■ 図 8-5  カットオフ筋の延長規定と通し筋の仕口内付着")
put_img(ws, figs["f5"], "A4", w=820)
head(ws, 30, "■ 問題 1  用語の説明")
body(ws, 31, "(1) カットオフ筋とは何か。なぜ全長通さず途中で切るのか（経済性）。", h=32)
body(ws, 32, "(2) カットオフ筋の末端はなぜ付着上不利か（自由端・応力集中）。", h=32)
head(ws, 34, "■ 問題 2  延長規定の区別（重要）")
body(ws, 35, "カットオフ筋の切断位置について、次の 2 つを区別して説明せよ：\n"
             "(A) RC規準の規定：計算上不要となる断面（理論カットオフ点）を超えて"
             "【 ① 】以上延長する。\n"
             "(B) 標準配筋の慣用：梁の 2 段目上端筋は柱面から【 ② 】+ 15d。\n"
             "(A) は付着の規準、(B) は配筋標準図の慣用寸法 ── 混同しないこと。", h=72)
head(ws, 37, "■ 問題 3  カットオフ筋の付着検定（計算）")
body(ws, 38, "梁：d=640mm、上端 2 段目 D25（db=25）をカットオフ。"
             "付着長さ ld = 1,800mm（危険断面から鉄筋末端まで）、"
             "鉄筋応力 σt = 345 N/mm²（降伏想定）、K·fb = 3.0 N/mm²。", h=32)
r = table(ws, 41,
          ["項目", "式", "値（記入）"],
          [["有効付着長さ", "ld − d", ""],
           ["平均付着応力度 τ", "σt·db / (4·(ld−d))", ""],
           ["判定", "τ ≦ K·fb = 3.0", ""]])
head(ws, r + 2, "■ 問題 4  通し筋の仕口内付着（計算）")
body(ws, r + 3, "内柱の仕口を貫通する梁主筋は、片側で引張降伏・反対側で圧縮となり、"
                "柱せい D の区間で応力が反転する → 平均付着応力度が非常に大きい。"
                "目安として D/db ≧ 20 を確認する。次の組合せを判定せよ。", h=44)
r = table(ws, r + 5,
          ["No.", "柱せい D", "梁主筋", "D/db（記入）", "判定（記入）"],
          [["(1)", "700", "D25", "", ""],
           ["(2)", "700", "D29", "", ""],
           ["(3)", "600", "D32", "", ""]])
body(ws, r + 2, "(3) が NG の場合の対策を 2 つ挙げよ（主筋径を細く・柱せいを大きく・"
                "鉄筋強度を下げる 等）。", h=32)
head(ws, r + 4, "■ 問題 5  マンション実務まとめ")
body(ws, r + 5, "配筋図チェックで付着に関して確認する項目を 4 つ挙げよ："
                "①カットオフ位置（理論カットオフ点＋d 以上、慣用 L0/4+15d との整合） "
                "②カットオフ筋の付着検定の要否 ③通し筋の D/db "
                "④上端筋の fa 低減の反映。", h=44)

# ---- 解答 ----
ws = wb.create_sheet("解答")
setup(ws, [4, 22, 60, 14])
title_row(ws, 1, "解答・解説", span=4)
row = 2


def ah(t):
    global row
    head(ws, row, t, span=4); row += 1


def an(t, h=None):
    global row
    body(ws, row, t, span=4, ans=True, h=h); row += 1


ah("8-1  応力伝達機構")
an("問1：①化学 ②摩擦 ③支圧 ④フック。丸鋼（付着が弱い）は令91条の低い付着値"
   "（長期 0.7 N/mm²）しか使えず、フック定着が原則。", h=44)
an("問2：(1)引抜き破壊（かぶり・拘束が十分 → 鉄筋周囲の円筒面で滑る）。"
   "(2)付着割裂破壊（かぶり小 → かぶりが割れて剥落）。"
   "(3)付着割裂破壊（太径・密配筋は割裂が先行しやすい）。"
   "設計上問題になるのは主に割裂側 → だから『付着割裂の基準となる強度 fb』が規準化された。",
   h=72)
an("問3：①引張（リング状の引張応力）。割裂を抑える要素：かぶり厚さを増す／"
   "鉄筋のあきを確保する／横補強筋（帯筋・あばら筋）を増やす。"
   "→ RC規準の修正係数 K = 0.3(C+W)/db+0.4 はまさにこの 3 要素を式にしたもの。", h=58)
an("問4：付着は「鉄筋の力をコンクリートへ伝える界面の応力」。定着長 Lb も重ね継手長 L1 も、"
   "鉄筋の引張力 T = σt·(π·db²/4) を付着抵抗 fa·(π·db)·L で受け止める釣合いから"
   "L = σt·db/(4·fa) の形で決まる。付着が全ての土台。", h=58)

ah("8-2  fa と fb")
an("問1：Fc21（≦22.5 なので F/15・F/10）：上端 1.40／その他 2.10。"
   "Fc24（>22.5！）：上端 0.9+48/75=1.54／その他 1.35+0.96=2.31。"
   "Fc27：上端 0.9+0.72=1.62／その他 1.35+1.08=2.43。"
   "Fc30：上端 0.9+0.80=1.70／その他 1.35+1.20=2.55。"
   "Fc36：上端 0.9+0.96=1.86／その他 1.35+1.44=2.79。"
   "短期は各 2 倍（例 Fc24 その他 4.62）。"
   "※ Fc24 から式が切り替わる点が最大の引っかけ。F/10 のまま計算すると 2.40 と"
   "危険側に過大評価する。", h=130)
an("問2：その鉄筋の下に 300mm 以上のコンクリートが打ち込まれる曲げ材の水平鉄筋。"
   "打設時のブリーディング（水の上昇）で鉄筋下面に水膜・空隙ができ付着が低下するため "
   "fa が低く設定される。", h=44)
an("問3：fb（その他）= Fc/40+0.9：Fc24→1.50、Fc27→1.58、Fc30→1.65。"
   "上端筋は 0.8 倍：1.20、1.26、1.32。"
   "（多段配筋の 2 段目以降はさらに 0.6 倍）", h=44)
an("問4：K = 0.3×(60+40)/25 + 0.4 = 0.3×4.0 + 0.4 = 1.6。"
   "かぶり・あき（C）や帯筋（W）が増えるほど K が大きくなり、付着を高く評価できる"
   "（上限 2.5）。付着は「配筋ディテールで稼げる」ことを式が示している。", h=58)
an("問5：①2 ②1.5。法令（告示1450号・令91条2項）は短期＝長期の 2 倍、"
   "RC規準は 1.5 倍と食い違う。どちらに拠るかはプロジェクトの設計方針"
   "（確認申請での位置づけ・審査機関との整合）で決め、設計方針書に明記する。", h=58)

ah("8-3  付着設計法")
an("問1：①周長 ψ = π·db ②fa（または K·fb）。"
   "T = σt·(π·db²/4) を 付着面積 π·db·ld で除すと τ = σt·db/(4·ld)。", h=44)
an("問2：ld ≧ σt·db/(4·K·fb) = 345×25/(4×3.0) = 8,625/12 = 719mm ≒ 720mm（≒29db）。",
   h=32)
an("問3：①有効せい d。端部のせん断ひび割れで引張域が柱面より広がるため、"
   "付着に使える長さを d だけ短く見る＝検定が厳しくなる＝安全側。", h=44)
an("問4：全数通し配筋とし（カットオフなし）、せん断の安全性検討を行う場合等は"
   "付着検討を省略できる扱いがある（2018年版で整理）。通し筋は末端の応力集中がなく、"
   "せん断設計で部材の脆性破壊が防がれていれば付着割裂が先行しにくいため。", h=58)
an("問5：B → C → A → D → E。存在応力を出し、fb・K を揃えて検定、NG なら配筋・断面の見直し。",
   h=32)

ah("8-4  規準の選択")
an("問1：法令＝最低基準。付着の許容応力度 fa を定める（検討方法は定めない）。"
   "RC規準＝許容応力度設計の標準。使用性/損傷制御/安全性の 3 段階の付着検定。"
   "靭性指針＝大地震時の靭性保証。ヒンジ域の付着劣化を τf≦τbu で制御。"
   "黄色本＝確認審査の実務運用。規準の式を引用し、カットオフ筋の扱い等を Q&A で補完。",
   h=86)
an("問2：(1)塑性ヒンジに変形能力を期待する設計では、主筋の付着が劣化すると"
   "復元力特性が痩せて（スリップ形状）エネルギー吸収が激減し、想定した靭性が"
   "発揮できないため。(2)付着割裂破壊に対する信頼できる下限強度。", h=58)
an("問3：(1)告示 fa＋RC規準/黄色本（一般的な保有水平耐力計算のルート）。"
   "(2)靭性指針（Ds を小さく設定＝靭性に依存する分、ヒンジ域の付着劣化を厳密に管理）。"
   "(3)告示 fa＋RC規準（ルート1 は許容応力度計算のみ。仕様規定＋標準配筋の遵守が中心）。",
   h=72)
an("問4：規準間の食い違いは「どの規準体系で設計全体を組み立てるか」の問題。"
   "個々の数値のいいとこ取りはせず、設計方針書で準拠規準を明示し、"
   "上司・審査機関と整合を取ってから計算に反映する。", h=44)

ah("8-5  カットオフ・通し筋")
an("問1：(1)曲げモーメントが小さくなる区間では上端筋の一部が計算上不要になるため、"
   "途中で切って鉄筋量を節約する筋。(2)末端が自由端で応力がゼロから立ち上がる"
   "＝末端付近の付着に応力伝達が集中し、通し筋より付着上不利。", h=58)
an("問2：(A)＝有効せい d 以上（RC規準。ld ≧ 理論カットオフ点までの長さ + d）。"
   "(B)＝L0/4（柱面から L0/4 + 15d が標準配筋の慣用）。"
   "『15d』は規準の付着規定ではなく標準図の慣用寸法。付着検定の結果、"
   "L0/4+15d より長く必要になる場合もある（検定が優先）。", h=72)
an("問3：有効付着長さ = 1,800−640 = 1,160mm。"
   "τ = 345×25/(4×1,160) = 8,625/4,640 = 1.86 N/mm²。"
   "1.86 ≦ 3.0 → OK。", h=44)
an("問4：(1)700/25 = 28 ≧ 20 OK。(2)700/29 = 24.1 ≧ 20 OK。"
   "(3)600/32 = 18.8 < 20 NG。対策：主筋を D29 以下に細径化（本数を増やす）／"
   "柱せいを 650 以上に拡大／（高強度化はかえって不利：σy が上がると必要 D/db も増える）。",
   h=58)
an("問5：①カットオフ位置（理論カットオフ点＋d 以上と L0/4+15d の両方を満たすか） "
   "②カットオフ筋がある梁の付着検定（σt=σy で実施したか） "
   "③通し筋の D/db ≧ 20 目安（外柱は投影定着長さ側の検討） "
   "④上端筋の fa・fb 低減（0.8 倍・多段 0.6 倍）が計算に反映されているか。", h=72)

XLSX = os.path.join(OUT, "付着設計問題集.xlsx")
wb.save(XLSX)
print("saved:", XLSX)
