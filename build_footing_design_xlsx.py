# -*- coding: utf-8 -*-
"""基礎フーチング(直接・杭)の設計 問題集（図つき）Excel。
出力: docs/footing_design/基礎フーチング設計問題集.xlsx
1 地反力・杭反力 / 2 独立フーチングの設計（曲げ・パンチング）
3 連続・べた基礎・基礎梁と接地圧(e/l) / 4 杭基礎・偏心基礎の設計
※ 許容応力度・パンチング許容せん断は規準（RC規準/告示/基礎指針）で確認要
"""
import os
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib import font_manager as fm
FONT = "/usr/share/fonts/truetype/fonts-japanese-gothic.ttf"
jp = fm.FontProperties(fname=FONT); fm.fontManager.addfont(FONT)
plt.rcParams["font.family"] = jp.get_name(); plt.rcParams["axes.unicode_minus"] = False
OUT = "docs/footing_design"; FIG = os.path.join(OUT, "figures"); os.makedirs(FIG, exist_ok=True)


def save(fig, n):
    p = os.path.join(FIG, n); fig.savefig(p, dpi=130, bbox_inches="tight"); plt.close(fig); return p


def fig_reaction():
    fig, axes = plt.subplots(1, 2, figsize=(13.0, 4.8))
    # (a) 地反力（接地圧）
    ax = axes[0]
    ax.plot([0, 0], [0, 3], color="#2e5b8a", lw=4)
    ax.add_patch(mpatches.Rectangle((-2.5, -0.6), 5.0, 0.6, fc="#d9d9d9", ec="k", lw=1.2))
    for x in np.linspace(-2.5, 2.5, 12):
        ax.annotate("", xy=(x, -0.6), xytext=(x, -1.4), arrowprops=dict(arrowstyle="-|>", color="#7a3b00", lw=1.3))
    ax.plot([-2.5, 2.5], [-1.4, -1.4], color="#7a3b00", lw=1.2)
    ax.annotate("", xy=(0, 3), xytext=(0, 3.9), arrowprops=dict(arrowstyle="-|>", color="#c00000", lw=2.5))
    ax.text(0.2, 3.5, "N", fontproperties=jp, fontsize=10, color="#c00000")
    ax.text(0, -2.0, "地反力（接地圧）q = N/A ± M/Z\n地盤が面で支える", ha="center",
            fontproperties=jp, fontsize=9, color="#7a3b00")
    ax.set_xlim(-3, 3); ax.set_ylim(-2.6, 4.2); ax.axis("off")
    ax.set_title("(a) 直接基礎：地反力（接地圧）", fontproperties=jp, fontsize=9.5, fontweight="bold")
    # (b) 杭反力
    ax = axes[1]
    ax.plot([0, 0], [0.5, 3], color="#2e5b8a", lw=4)
    ax.add_patch(mpatches.Rectangle((-2.2, 0), 4.4, 0.6, fc="#d9d9d9", ec="k", lw=1.2))
    for px in [-1.5, -0.5, 0.5, 1.5]:
        ax.add_patch(mpatches.Rectangle((px - 0.15, -1.6), 0.3, 1.6, fc="#b0b0b0", ec="k"))
        r = 900 + 260 * px  # 反力の大小イメージ
        ax.annotate("", xy=(px, -1.7), xytext=(px, -1.7 - r / 900), arrowprops=dict(arrowstyle="-|>", color="#548235", lw=2))
    ax.annotate("", xy=(0, 3), xytext=(0, 3.9), arrowprops=dict(arrowstyle="-|>", color="#c00000", lw=2.5))
    ax.text(0.2, 3.5, "N", fontproperties=jp, fontsize=10, color="#c00000")
    ax.text(0, -3.4, "杭反力 Ri = N/n ± M・xi/Σxi²\n杭が点で支える", ha="center",
            fontproperties=jp, fontsize=9, color="#548235")
    ax.set_xlim(-3, 3); ax.set_ylim(-4.0, 4.2); ax.axis("off")
    ax.set_title("(b) 杭基礎：杭反力", fontproperties=jp, fontsize=9.5, fontweight="bold")
    fig.suptitle("図 1  地反力（接地圧）と杭反力", fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig1_reaction.png")


def fig_isolated():
    fig, axes = plt.subplots(1, 2, figsize=(13.0, 5.0))
    # (a) 断面（曲げ）
    ax = axes[0]
    ax.add_patch(mpatches.Rectangle((0, 1.0), 6, 0.8, fc="#d9d9d9", ec="k", lw=1.2))  # フーチング
    ax.add_patch(mpatches.Rectangle((2.5, 1.8), 1.0, 1.6, fc="#bcd2ea", ec="k", lw=1.2))  # 柱
    ax.text(3.0, 2.6, "柱", ha="center", fontproperties=jp, fontsize=9)
    for x in np.linspace(0.2, 5.8, 12):
        ax.annotate("", xy=(x, 1.0), xytext=(x, 0.4), arrowprops=dict(arrowstyle="-|>", color="#7a3b00", lw=1.2))
    ax.annotate("", xy=(0, 0.1), xytext=(2.5, 0.1), arrowprops=dict(arrowstyle="<->", color="#c00000"))
    ax.text(1.25, -0.15, "片持ち長", ha="center", fontproperties=jp, fontsize=8, color="#c00000")
    ax.annotate("柱面で曲げ最大\nM=q・B・(片持ち長)²/2", xy=(2.5, 1.4), xytext=(3.6, 0.2),
                fontproperties=jp, fontsize=8.5, color="#c00000",
                arrowprops=dict(arrowstyle="->", color="#c00000"))
    ax.set_xlim(-0.3, 6.3); ax.set_ylim(-0.5, 3.6); ax.axis("off")
    ax.set_title("(a) 曲げ：片持ち版として", fontproperties=jp, fontsize=9.5, fontweight="bold")
    # (b) パンチング（平面）
    ax = axes[1]
    ax.add_patch(mpatches.Rectangle((0, 0), 6, 6, fc="#eef2f7", ec="k", lw=1.2))  # フーチング平面
    ax.add_patch(mpatches.Rectangle((2.4, 2.4), 1.2, 1.2, fc="#bcd2ea", ec="k", lw=1.2))  # 柱
    ax.add_patch(mpatches.Rectangle((1.9, 1.9), 2.2, 2.2, fill=False, ec="#c00000", lw=1.8, ls="--"))  # 危険断面
    ax.text(3.0, 3.0, "柱", ha="center", va="center", fontproperties=jp, fontsize=9)
    ax.annotate("パンチング危険断面\n（柱面から d/2）\nbo=4(c+d)", xy=(4.1, 4.1), xytext=(4.3, 5.3),
                fontproperties=jp, fontsize=8.5, color="#c00000",
                arrowprops=dict(arrowstyle="->", color="#c00000"))
    ax.text(3.0, -0.6, "2方向せん断（パンチング）\nVp=N-q(c+d)², τ=Vp/(bo・d)", ha="center",
            fontproperties=jp, fontsize=8.5, color="#1f4e79")
    ax.set_xlim(-0.3, 6.3); ax.set_ylim(-1.1, 6.3); ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("(b) せん断：パンチング（2方向）", fontproperties=jp, fontsize=9.5, fontweight="bold")
    fig.suptitle("図 2  独立フーチングの設計（曲げ・パンチングせん断）",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig2_isolated.png")


def fig_mat():
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 5.0))
    # (a) べた基礎・基礎梁
    ax = axes[0]
    ax.add_patch(mpatches.Rectangle((0, 1.0), 8, 0.5, fc="#d9d9d9", ec="k", lw=1.2))  # スラブ
    for cx in [1.5, 4.0, 6.5]:
        ax.plot([cx, cx], [1.5, 3.0], color="#2e5b8a", lw=3)
        ax.add_patch(mpatches.Rectangle((cx - 0.5, 1.5), 1.0, 0.5, fc="#e8e0d0", ec="k"))  # 基礎梁
    for x in np.linspace(0.3, 7.7, 16):
        ax.annotate("", xy=(x, 1.0), xytext=(x, 0.5), arrowprops=dict(arrowstyle="-|>", color="#7a3b00", lw=1.1))
    ax.text(4.0, 0.0, "接地圧をスラブが受け、基礎梁で柱へ伝える", ha="center",
            fontproperties=jp, fontsize=8.5, color="#7a3b00")
    ax.text(4.0, 2.5, "基礎梁", ha="center", fontproperties=jp, fontsize=8, color="#7a5a2a")
    ax.set_xlim(-0.3, 8.3); ax.set_ylim(-0.5, 3.3); ax.axis("off")
    ax.set_title("(a) べた基礎・基礎梁", fontproperties=jp, fontsize=9.5, fontweight="bold")
    # (b) 接地圧 e/l
    ax = axes[1]
    # 台形 e<l/6
    ax.add_patch(mpatches.Rectangle((0, 2.3), 3.2, 0.3, fc="#d9d9d9", ec="k"))
    ax.plot([0, 0, 3.2, 3.2], [2.3 - 0.25, 2.3, 2.3, 2.3 - 0.42], color="#7a3b00", lw=1.8)
    ax.fill_between([0, 3.2], [2.3, 2.3], [2.3 - 0.25, 2.3 - 0.42], color="#f0e0c8", alpha=0.7)
    ax.text(1.6, 2.75, "e <= l/6：全面圧縮（台形）", ha="center", fontproperties=jp, fontsize=8.5)
    ax.text(3.4, 2.0, "qmin>0", fontproperties=jp, fontsize=7.5, color="#7a3b00")
    # 三角 e>l/6
    ax.add_patch(mpatches.Rectangle((0, 0.6), 3.2, 0.3, fc="#d9d9d9", ec="k"))
    ax.plot([0, 0, 2.4, 3.2], [0.6 - 0.6, 0.6, 0.6, 0.6], color="#c00000", lw=1.8)
    ax.plot([0, 3.2], [0.0, 0.6], color="#c00000", lw=1.8)
    ax.fill([0, 0, 3.2], [0.0, 0.6, 0.6], color="#f8d0d0", alpha=0.6)
    ax.text(1.6, 1.05, "e > l/6：一部浮上り（三角）", ha="center", fontproperties=jp, fontsize=8.5, color="#c00000")
    ax.text(0.05, 1.5, "偏心 e=M/N。q=N/A(1±6e/l)。例 e=0.5<l/6=2.0→全面圧縮\n"
            "qmax=41.7, qmin=25.0 kPa（べた12×15m, N=6000, M=3000）",
            fontproperties=jp, fontsize=8, color="#1f4e79")
    ax.set_xlim(-0.3, 4.4); ax.set_ylim(-0.3, 3.2); ax.axis("off")
    ax.set_title("(b) べた基礎の接地圧（e/l 判定）", fontproperties=jp, fontsize=9.5, fontweight="bold")
    fig.suptitle("図 3  連続・べた基礎・基礎梁と接地圧（e/l）",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig3_mat.png")


def fig_pile_ecc():
    fig, axes = plt.subplots(1, 2, figsize=(13.0, 5.0))
    # (a) 杭基礎（パイルキャップ）
    ax = axes[0]
    ax.add_patch(mpatches.Rectangle((0, 2.0), 6, 1.0, fc="#e8e0d0", ec="k", lw=1.2))  # キャップ
    ax.plot([3, 3], [3.0, 4.4], color="#2e5b8a", lw=4)
    ax.annotate("", xy=(3, 4.4), xytext=(3, 5.2), arrowprops=dict(arrowstyle="-|>", color="#c00000", lw=2.5))
    ax.text(3.2, 4.8, "N・M", fontproperties=jp, fontsize=9, color="#c00000")
    for px in [1.0, 2.33, 3.66, 5.0]:
        ax.add_patch(mpatches.Rectangle((px - 0.15, 0.4), 0.3, 1.6, fc="#b0b0b0", ec="k"))
        ax.annotate("", xy=(px, 0.3), xytext=(px, -0.4), arrowprops=dict(arrowstyle="-|>", color="#548235", lw=2))
    ax.text(3.0, -0.9, "杭反力でパイルキャップを曲げ・パンチング照査", ha="center",
            fontproperties=jp, fontsize=8.5, color="#1f4e79")
    ax.set_xlim(-0.3, 6.3); ax.set_ylim(-1.3, 5.4); ax.axis("off")
    ax.set_title("(a) 杭基礎（パイルキャップ）", fontproperties=jp, fontsize=9.5, fontweight="bold")
    # (b) 偏心基礎
    ax = axes[1]
    ax.add_patch(mpatches.Rectangle((0, 1.0), 5, 0.7, fc="#d9d9d9", ec="k", lw=1.2))  # フーチング
    ax.plot([1.2, 1.2], [1.7, 3.4], color="#2e5b8a", lw=4)  # 柱（偏心）
    ax.plot([2.5, 2.5], [1.0, 0.4], color="#333", lw=1, ls=":")  # 基礎中心
    ax.annotate("", xy=(1.2, 0.6), xytext=(2.5, 0.6), arrowprops=dict(arrowstyle="<->", color="#c00000"))
    ax.text(1.85, 0.35, "偏心 e0", ha="center", fontproperties=jp, fontsize=8, color="#c00000")
    ax.annotate("", xy=(1.2, 3.4), xytext=(1.2, 4.2), arrowprops=dict(arrowstyle="-|>", color="#c00000", lw=2.5))
    ax.text(1.4, 3.8, "N", fontproperties=jp, fontsize=9, color="#c00000")
    # 偏った接地圧
    for i, x in enumerate(np.linspace(0.2, 4.8, 10)):
        h = 0.7 - 0.09 * i
        ax.annotate("", xy=(x, 1.0), xytext=(x, 1.0 - max(h, 0.1)), arrowprops=dict(arrowstyle="-|>", color="#7a3b00", lw=1.1))
    ax.text(2.5, -0.5, "柱が基礎中心からずれる→偏心モーメント Me=N・e0\n基礎梁で負担、または接地圧が偏る",
            ha="center", fontproperties=jp, fontsize=8.5, color="#1f4e79")
    ax.set_xlim(-0.3, 5.3); ax.set_ylim(-1.0, 4.4); ax.axis("off")
    ax.set_title("(b) 偏心基礎", fontproperties=jp, fontsize=9.5, fontweight="bold")
    fig.suptitle("図 4  杭基礎（パイルキャップ）・偏心基礎の応力",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig4_pile_ecc.png")


figs = {"reaction": fig_reaction(), "isolated": fig_isolated(),
        "mat": fig_mat(), "pile_ecc": fig_pile_ecc()}
print("figs:", list(figs.keys()))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
wb = Workbook()
C_TITLE = "1F4E79"; C_HEAD = "2E75B6"; C_ANS = "E2EFDA"; C_WARN = "FCE4D6"
thin = Side(style="thin", color="BFBFBF"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
f_title = Font(name="MS PGothic", size=15, bold=True, color="FFFFFF")
f_head = Font(name="MS PGothic", size=11, bold=True, color="FFFFFF")
f_body = Font(name="MS PGothic", size=10); f_ans = Font(name="MS PGothic", size=10, color="375623")
f_warn = Font(name="MS PGothic", size=9, color="833C00", italic=True)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)


def setup(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False


def title_row(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_title; c.fill = PatternFill("solid", fgColor=C_TITLE)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1); ws.row_dimensions[row].height = 30


def head(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_head; c.fill = PatternFill("solid", fgColor=C_HEAD)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1); ws.row_dimensions[row].height = 22


def body(ws, row, text, span=8, ans=False, h=None):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_ans if ans else f_body; c.alignment = wrap
    if ans:
        c.fill = PatternFill("solid", fgColor=C_ANS)
    if h:
        ws.row_dimensions[row].height = h


def warn(ws, row, text, span=8, h=None):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_warn; c.alignment = wrap
    c.fill = PatternFill("solid", fgColor=C_WARN)
    if h:
        ws.row_dimensions[row].height = h


def tbl(ws, start_row, headers, rows, col1=1):
    r = start_row
    for j, htxt in enumerate(headers):
        c = ws.cell(r, col1 + j, htxt); c.font = Font(name="MS PGothic", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C_HEAD); c.alignment = center; c.border = border
    for data in rows:
        r += 1
        for j, v in enumerate(data):
            c = ws.cell(r, col1 + j, v); c.font = f_body
            c.alignment = wrap if j == 0 else center; c.border = border
            if r % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F2F7FC")
    return r


def put_img(ws, path, anchor, w=None):
    img = XLImage(path)
    if w:
        ratio = w / img.width; img.width = w; img.height = int(img.height * ratio)
    ws.add_image(img, anchor)


# ===== 目次 =====
ws = wb.active; ws.title = "目次"; setup(ws, [4, 24, 54, 16])
title_row(ws, 1, "基礎フーチング(直接・杭)の設計 問題集（RC マンション設計担当・新人向け）", span=4)
body(ws, 2, "目標：地反力・杭反力を理解し、直接基礎（独立・連続・べた）と杭基礎、偏心基礎の"
            "応力計算・断面算定（曲げ・パンチング・接地圧e/l）ができること。", span=4, h=32)
r = tbl(ws, 4, ["No.", "シート", "到達目標", "図"],
        [["1", "1 地反力・杭反力", "地反力・杭反力の分布を理解", "反力"],
         ["2", "2 独立フーチング", "曲げ・パンチングの断面算定", "曲げ/せん断"],
         ["3", "3 連続・べた・基礎梁", "べた基礎の接地圧(e/l)・基礎梁", "べた/e-l"],
         ["4", "4 杭基礎・偏心基礎", "杭基礎・偏心基礎の応力計算", "杭/偏心"]])
warn(ws, r + 2, "※ 許容応力度・パンチング許容せん断・安全率は規準（RC規準/告示/基礎指針）で確認。"
                "自重・土被りは簡略化した例示。", span=4, h=36)

# ===== 1 地反力・杭反力 =====
ws = wb.create_sheet("1 地反力杭反力"); setup(ws, [8, 16, 18, 16, 12, 12])
title_row(ws, 1, "1  地反力（接地圧）と杭反力")
head(ws, 3, "■ 図 1  地反力・杭反力"); put_img(ws, figs["reaction"], "A4", w=840)
head(ws, 27, "■ 問題 1  地反力と杭反力")
body(ws, 28, "直接基礎の地反力（接地圧 q=N/A±M/Z、面で支持）と杭基礎の杭反力"
             "（Ri=N/n±M·xi/Σxi²、点で支持）の違いを説明せよ。", h=36)
head(ws, 30, "■ 問題 2  反力が設計外力になる")
body(ws, 31, "基礎（フーチング・パイルキャップ・基礎梁）の設計では、地反力・杭反力が"
             "『下から上向きに作用する外力』になることを説明せよ。上部反力とのつり合いに触れよ。", h=40)
head(ws, 33, "■ 問題 3  偏心の影響")
body(ws, 34, "軸力Nに加えて曲げMが作用すると、地反力・杭反力が偏る（±M項）ことを説明せよ。"
             "最大反力が許容値を超えないか確認する必要性を述べよ。", h=36)

# ===== 2 独立フーチング =====
ws = wb.create_sheet("2 独立フーチング"); setup(ws, [8, 16, 18, 16, 12, 12])
title_row(ws, 1, "2  独立フーチングの設計（曲げ・パンチング）")
head(ws, 3, "■ 図 2  曲げ・パンチング"); put_img(ws, figs["isolated"], "A4", w=840)
head(ws, 28, "■ 問題 1  曲げの計算")
body(ws, 29, "フーチング B=L=3.0m、柱0.7角、N=1500kN（接地圧 q=N/A=166.7kPa、自重略）のとき、"
             "片持ち長と柱面の曲げモーメント M=q·B·(片持ち長)²/2 を求めよ。", h=40)
head(ws, 31, "■ 問題 2  パンチングせん断")
body(ws, 32, "同じフーチングで有効せい d=0.5m のとき、パンチング（2方向せん断）の危険断面周長"
             "bo=4(c+d)、せん断力 Vp=N-q(c+d)²、せん断応力 τ=Vp/(bo·d) を求めよ。", h=44)
head(ws, 34, "■ 問題 3  断面算定")
body(ws, 35, "曲げに対する必要鉄筋（下端筋）の考え方、パンチングせん断が許容を超える場合の対応"
             "（せいを増す・せん断補強）を述べよ。", h=40)
head(ws, 37, "■ 問題 4  1方向せん断")
body(ws, 38, "パンチング（2方向）のほかに、柱面から距離dの断面での1方向せん断も検討することを述べよ。"
             "フーチング設計で確認する応力（曲げ・1方向せん断・パンチング）を整理せよ。", h=40)
warn(ws, 40, "※ 許容せん断応力度・危険断面の取り方は規準で確認要。", h=22)

# ===== 3 連続・べた・基礎梁 =====
ws = wb.create_sheet("3 べた基礎e-l"); setup(ws, [8, 16, 18, 16, 12, 12])
title_row(ws, 1, "3  連続・べた基礎・基礎梁と接地圧（e/l）")
head(ws, 3, "■ 図 3  べた基礎・接地圧e/l"); put_img(ws, figs["mat"], "A4", w=860)
head(ws, 28, "■ 問題 1  べた基礎と基礎梁")
body(ws, 29, "べた基礎で接地圧をスラブが受け、基礎梁で柱に伝える仕組みを説明せよ。"
             "連続基礎・べた基礎・独立基礎の使い分け（地盤・荷重・沈下）にも触れよ。", h=40)
head(ws, 31, "■ 問題 2  接地圧の e/l 判定")
body(ws, 32, "べた基礎 12×15m、鉛直合力 N=6000kN、偏心モーメント M=3000kN·m のとき、"
             "偏心 e=M/N と l/6 を比較し、接地圧分布（全面圧縮/一部浮上り）を判定せよ。", h=40)
head(ws, 34, "■ 問題 3  最大・最小接地圧")
body(ws, 35, "問2で q=N/A(1±6e/l) により qmax・qmin を求めよ（A=12×15）。"
             "qmin>0 で全面圧縮となることを確認せよ。", h=36)
head(ws, 37, "■ 問題 4  基礎梁の役割")
body(ws, 38, "基礎梁が接地圧・杭反力を集めて柱・杭に伝え、不同沈下を抑える役割を説明せよ。"
             "基礎梁の応力計算・断面算定（曲げ・せん断）の考え方に触れよ。", h=40)

# ===== 4 杭基礎・偏心基礎 =====
ws = wb.create_sheet("4 杭基礎偏心基礎"); setup(ws, [8, 16, 18, 16, 12, 12])
title_row(ws, 1, "4  杭基礎・偏心基礎の設計")
head(ws, 3, "■ 図 4  杭基礎・偏心基礎"); put_img(ws, figs["pile_ecc"], "A4", w=840)
head(ws, 28, "■ 問題 1  杭基礎の応力計算")
body(ws, 29, "4本杭（ピッチ2.5m）、N=4000kN・M=1500kN·m のとき、杭反力 Ri=N/n±M·xi/Σxi² を求め、"
             "その杭反力でパイルキャップの曲げ・パンチングを照査する流れを述べよ。", h=44)
head(ws, 31, "■ 問題 2  パイルキャップの断面算定")
body(ws, 32, "杭反力を外力として、パイルキャップ（杭を結ぶフーチング）を曲げ・パンチングで断面算定する"
             "考え方を述べよ。杭頭補強筋の定着にも触れよ。", h=40)
head(ws, 34, "■ 問題 3  偏心基礎の応力")
body(ws, 35, "柱が基礎中心から e0=0.4m 偏心する偏心基礎で、偏心モーメント Me=N·e0（N=1200kN）を求めよ。"
             "この偏心モーメントを基礎梁で負担する／接地圧が偏ることを説明せよ。", h=44)
head(ws, 37, "■ 問題 4  設計の流れ")
body(ws, 38, "基礎フーチング設計の流れをまとめよ（①地反力・杭反力の算定→②基礎形式（独立/連続/べた/杭）→"
             "③曲げ・せん断（1方向・パンチング）・接地圧e/l の照査→④断面算定・配筋→⑤偏心・干渉の確認）。", h=44)
warn(ws, 40, "※ 許容応力度・断面算定・杭頭定着は規準（RC規準/告示/基礎指針）で確認要。", h=24)

# ===== 解答 =====
ws = wb.create_sheet("解答"); setup(ws, [4, 22, 60, 14])
title_row(ws, 1, "解答・解説", span=4); row = 2


def ah(t):
    global row; head(ws, row, t, span=4); row += 1


def an(t, h=None):
    global row; body(ws, row, t, span=4, ans=True, h=h); row += 1


def aw(t, h=None):
    global row; warn(ws, row, t, span=4, h=h); row += 1


ah("1  地反力・杭反力")
an("問1：直接基礎は地盤が面で支え、接地圧 q=N/A±M/Z が分布する。杭基礎は杭が点で支え、"
   "各杭に杭反力 Ri=N/n±M·xi/Σxi² が生じる。前者は連続分布、後者は離散点の反力。", h=44)
an("問2：基礎部材（フーチング・キャップ・基礎梁）にとって、地反力・杭反力は下から上向きに作用する"
   "外力。上部からの軸力・曲げとつり合い、その差で部材に曲げ・せん断が生じる。", h=40)
an("問3：曲げMが加わると反力は±M項で偏り、片側が大きくなる。最大接地圧qmax・最大杭反力Rmaxが"
   "地盤の許容支持力・杭の許容支持力を超えないか確認する必要がある。", h=40)

ah("2  独立フーチング")
an("問1：片持ち長=(3.0-0.7)/2=1.15m。M=q·B·(片持ち長)²/2=166.7×3.0×1.15²/2≒331kN·m（全幅）、"
   "単位幅あたり約110kN·m/m。柱面で曲げ最大、下端引張となる。", h=44)
an("問2：bo=4(c+d)=4×(0.7+0.5)=4.8m。Vp=N-q(c+d)²=1500-166.7×1.2²=1500-240=1260kN。"
   "τ=Vp/(bo·d)=1260/(4.8×0.5)=525kPa≒0.53N/mm²。許容せん断応力度と比較して照査。", h=44)
an("問3：曲げに対しては柱面のMから下端筋を算定（M<=許容曲げ）。パンチングが許容を超えるなら"
   "フーチングのせいdを増す、または（やむを得ない場合）せん断補強を行う。", h=40)
an("問4：柱面から距離dの断面で1方向（梁状）せん断も検討する。フーチングは①曲げ（柱面）"
   "②1方向せん断 ③2方向せん断（パンチング）の3つを確認する。", h=40)

ah("3  連続・べた・基礎梁")
an("問1：べた基礎はスラブ全面で接地圧を受け、基礎梁で柱位置に集約して伝える。"
   "軟弱地盤・重荷重・不同沈下抑制にはべた基礎、良質地盤・軽荷重なら独立基礎が経済的。", h=40)
an("問2：e=M/N=3000/6000=0.5m。l/6=12/6=2.0m。e=0.5<l/6=2.0 なので全面圧縮（台形分布、浮上りなし）。", h=36)
an("問3：A=12×15=180m²、N/A=33.3kPa。qmax=33.3×(1+6×0.5/12)=33.3×1.25=41.7kPa。"
   "qmin=33.3×0.75=25.0kPa。qmin>0で全面圧縮を確認。", h=40)
an("問4：基礎梁は接地圧・杭反力を集めて柱・杭に伝達し、基礎の一体性を高めて不同沈下を抑える。"
   "接地圧/杭反力を分布荷重・集中反力として曲げ・せん断を計算し、断面算定・配筋する。", h=40)

ah("4  杭基礎・偏心基礎")
an("問1：Σxi²=4×1.25²=6.25。Ri=4000/4±1500×1.25/6.25=1000±300 → Rmax=1300, Rmin=700kN。"
   "この杭反力を外力としてパイルキャップの曲げ・パンチング（杭列間・柱周）を照査する。", h=44)
an("問2：杭反力を上向き外力とし、柱面での曲げ、柱周・杭周のパンチングでキャップを断面算定する。"
   "杭頭補強筋をキャップに定着し、杭頭曲げ・引抜きを伝達する（前教材の干渉確認と連携）。", h=44)
an("問3：Me=N·e0=1200×0.4=480kN·m。柱が基礎中心からずれると偏心モーメントが生じ、"
   "基礎梁がこれを負担する（または接地圧が偏りqmaxが増える）。境界杭・境界基礎で生じやすい。", h=44)
an("問4：①地反力・杭反力の算定→②基礎形式の選定（独立/連続/べた/杭）→③曲げ・1方向せん断・"
   "パンチング・接地圧e/l の照査→④断面算定・配筋→⑤偏心モーメント・杭頭補強筋の干渉確認。", h=44)
aw("※ 許容応力度・断面算定・杭頭定着は規準（RC規準/告示/基礎指針）で確認要。", h=24)

XLSX = os.path.join(OUT, "基礎フーチング設計問題集.xlsx"); wb.save(XLSX); print("saved:", XLSX)
