# -*- coding: utf-8 -*-
"""コンクリート配合 年間区分表（S値補正 × 混和剤の種類）を生成する。

前提条件（S値の補正期間・混和剤の適用期間）はユーザー提示条件をそのまま
「前提条件」シートの入力欄に持たせ、他シートは数式で自動計算する。
"""
import calendar
import datetime as dt

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = "/home/user/-/concrete-mix/コンクリート配合_年間区分表.xlsx"

FONT = "Meiryo"
BASE_YEAR = 2026  # 平年。スケジュールは毎年同じ（うるう年は 2/29 が 1 日加算）

# ---- 色 -------------------------------------------------------------------
C_TITLE = "1F4E79"
C_HDR = "D9E1F2"
C_HDR2 = "F2F2F2"
C_INPUT = "FFF2CC"      # 入力欄（黄）
C_CALC = "F2F2F2"       # 計算用（グレー）
F_S6 = PatternFill("solid", fgColor="FFD966")   # S値+6
F_S3 = PatternFill("solid", fgColor="BDD7EE")   # S値+3
# 色相＝混和剤の適用期、濃淡＝S値補正（濃＝+6／淡＝+3）
DARK = {"標": "A9D08E", "夏": "F4B183", "冬": "9DC3E6"}
LIGHT = {"標": "E2EFDA", "夏": "FBE5D6", "冬": "DEEBF7"}
TERM_CH = {"標準期": "標", "夏期": "夏", "冬期": "冬"}
F_STD = PatternFill("solid", fgColor=DARK["標"])  # 標準期
F_SUM = PatternFill("solid", fgColor=DARK["夏"])  # 夏期
F_WIN = PatternFill("solid", fgColor=DARK["冬"])  # 冬期

THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")


def style(cell, *, bold=False, size=11, color="000000", fill=None,
          border=True, align=CENTER, fmt=None, wrap=False):
    cell.font = Font(name=FONT, bold=bold, size=size, color=color)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if border:
        cell.border = BOX
    cell.alignment = Alignment(horizontal=align.horizontal,
                               vertical="center", wrap_text=wrap)
    if fmt:
        cell.number_format = fmt
    return cell


def put(ws, ref, value, **kw):
    ws[ref] = value
    return style(ws[ref], **kw)


def title(ws, text, sub=None):
    put(ws, "A1", text, bold=True, size=14, color=C_TITLE, border=False, align=LEFT)
    if sub:
        put(ws, "A2", sub, size=9, color="595959", border=False, align=LEFT)


wb = Workbook()

# ===========================================================================
# シート「前提条件」（入力シート）
# ===========================================================================
ps = wb.create_sheet("前提条件")
title(ps, "前提条件（入力シート）",
      "黄色のセルが入力欄です。月・日を書き換えると「日別データ」「年間カレンダー」「年間区分表」が自動で再計算されます。")

put(ps, "A4", "■ スランプ（S値）の補正期間", bold=True, border=False, align=LEFT)
s_hdr = ["区分", "S値補正", "開始月", "開始日", "終了月", "終了日", "開始キー", "終了キー"]
for i, h in enumerate(s_hdr):
    put(ps, f"{get_column_letter(1+i)}5", h, bold=True,
        fill=C_CALC if i >= 6 else C_HDR)

s_rows = [("S値+6　期間①", 6, 6, 30, 9, 14),
          ("S値+6　期間②（年をまたぐ）", 6, 11, 27, 2, 20)]
for r, (name, sv, sm, sd, em, ed) in enumerate(s_rows, start=6):
    put(ps, f"A{r}", name, align=LEFT)
    put(ps, f"B{r}", sv, fmt='"+"0')
    for col, v in zip("CDEF", (sm, sd, em, ed)):
        put(ps, f"{col}{r}", v, fill=C_INPUT)
    put(ps, f"G{r}", f"=C{r}*100+D{r}", fill=C_CALC)
    put(ps, f"H{r}", f"=E{r}*100+F{r}", fill=C_CALC)

put(ps, "A8", "S値+3（上記以外の期間）", align=LEFT)
put(ps, "B8", 3, fmt='"+"0')
for col in "CDEFGH":
    put(ps, f"{col}8", "－", color="808080")

put(ps, "A10", "■ 混和剤の適用期間（AE減水剤／高性能AE減水剤 共通）", bold=True,
    border=False, align=LEFT)
a_hdr = ["区分", "混和剤の種類（形式）", "適用期", "開始月", "開始日", "終了月",
         "終了日", "開始キー", "終了キー"]
for i, h in enumerate(a_hdr):
    put(ps, f"{get_column_letter(1+i)}11", h, bold=True,
        fill=C_CALC if i >= 7 else C_HDR)

a_rows = [("①", "標準形", "標準期", 3, 8, 6, 15),
          ("②", "遅延形", "夏期", 6, 16, 9, 23),
          ("③", "標準形", "標準期", 9, 24, 12, 7),
          ("④", "標準形", "冬期", 12, 8, 3, 7)]
for r, (no, kind, term, sm, sd, em, ed) in enumerate(a_rows, start=12):
    put(ps, f"A{r}", no)
    put(ps, f"B{r}", kind)
    put(ps, f"C{r}", term)
    for col, v in zip("DEFG", (sm, sd, em, ed)):
        put(ps, f"{col}{r}", v, fill=C_INPUT)
    put(ps, f"H{r}", f"=D{r}*100+E{r}", fill=C_CALC)
    put(ps, f"I{r}", f"=F{r}*100+G{r}", fill=C_CALC)

notes = [
    "【前提・注記】",
    "・出典：ご提示条件（2026-08-24 受領）。S値+6＝6/30～9/14, 11/27～2/20／それ以外はS値+3。",
    "・混和剤はAE減水剤・高性能AE減水剤とも同一日程（標準期 3/8～6/15・9/24～12/7、夏期 6/16～9/23、冬期 12/8～3/7）。",
    "・「開始キー／終了キー」は月×100＋日で表した判定用の数値（例：6/30→630）。年をまたぐ期間も判定できます。",
    "・日数は平年（365日）で集計しています。うるう年は 2/29 が区分②（S値+3・冬期／標準形）に入り、同区分が1日増えます。",
]
for i, t in enumerate(notes):
    put(ps, f"A{18+i}", t, size=9, bold=(i == 0), border=False, align=LEFT,
        color="000000" if i == 0 else "404040")

ps["C6"].comment = Comment("入力欄：ご提示条件どおり 6/30～9/14 を設定。", "Claude")
ps["D12"].comment = Comment("入力欄：ご提示条件どおり 標準期① 3/8～6/15 を設定。", "Claude")
for col, w in zip("ABCDEFGHI", (28, 20, 9, 9, 9, 9, 11, 11, 11)):
    ps.column_dimensions[col].width = w

# ===========================================================================
# シート「日別データ」（365日 × 判定結果）
# ===========================================================================
ds = wb.create_sheet("日別データ")
title(ds, "日別データ（平年365日の自動判定）",
      "「前提条件」シートの期間から、1日ごとのS値補正・混和剤の種類・適用期を数式で判定しています。")

d_hdr = ["月", "日", "日付", "判定キー", "S値補正", "混和剤の種類（形式）", "適用期", "組合せ"]
for i, h in enumerate(d_hdr):
    put(ds, f"{get_column_letter(1+i)}3", h, bold=True, fill=C_HDR)
h_hdr = ["S①", "S②", "混①", "混②", "混③", "混④"]
for i, h in enumerate(h_hdr):
    put(ds, f"{get_column_letter(10+i)}3", h, bold=True, fill=C_CALC, size=9)
put(ds, "J2", "▼ 判定用（1=該当）", size=9, color="808080", border=False, align=LEFT)


def in_range(key_cell, s_ref, e_ref):
    """年をまたぐ期間にも対応する範囲判定式（1/0）。"""
    return (f"=IF(IF({s_ref}<={e_ref},"
            f"AND({key_cell}>={s_ref},{key_cell}<={e_ref}),"
            f"OR({key_cell}>={s_ref},{key_cell}<={e_ref})),1,0)")


row = 4
for m in range(1, 13):
    for d in range(1, calendar.monthrange(BASE_YEAR, m)[1] + 1):
        put(ds, f"A{row}", m)
        put(ds, f"B{row}", d)
        put(ds, f"C{row}", f'=A{row}&"/"&B{row}')
        put(ds, f"D{row}", f"=A{row}*100+B{row}")
        put(ds, f"E{row}", f"=IF(SUM(J{row}:K{row})>0,前提条件!$B$6,前提条件!$B$8)",
            fmt='"+"0')
        put(ds, f"F{row}",
            f"=IFERROR(INDEX(前提条件!$B$12:$B$15,MATCH(1,L{row}:O{row},0)),\"未設定\")")
        put(ds, f"G{row}",
            f"=IFERROR(INDEX(前提条件!$C$12:$C$15,MATCH(1,L{row}:O{row},0)),\"未設定\")")
        put(ds, f"H{row}", f'="S値+"&E{row}&"／"&F{row}&"（"&G{row}&"）"', align=LEFT)
        put(ds, f"J{row}", in_range(f"$D{row}", "前提条件!$G$6", "前提条件!$H$6"),
            fill=C_CALC, size=9)
        put(ds, f"K{row}", in_range(f"$D{row}", "前提条件!$G$7", "前提条件!$H$7"),
            fill=C_CALC, size=9)
        for i, ar in enumerate(range(12, 16)):
            col = get_column_letter(12 + i)
            put(ds, f"{col}{row}",
                in_range(f"$D{row}", f"前提条件!$H${ar}", f"前提条件!$I${ar}"),
                fill=C_CALC, size=9)
        row += 1
LAST_D = row - 1

for col, w in zip("ABCDEFGH", (6, 6, 9, 10, 10, 20, 10, 30)):
    ds.column_dimensions[col].width = w
for col in "JKLMNO":
    ds.column_dimensions[col].width = 6
ds.freeze_panes = "A4"
ds.auto_filter.ref = f"A3:H{LAST_D}"

# ===========================================================================
# シート「年間区分表」（メイン）
# ===========================================================================
ws = wb.create_sheet("年間区分表", 0)
title(ws, "コンクリート配合　年間期間区分表（S値補正 × 混和剤の種類）",
      "S値の補正期間と混和剤の切替期間が一致しないため、両方を重ね合わせて年間9区分に整理したものです。"
      "（1/1～2/20 と 12/8～12/31 は年をまたいで連続＝12/8～2/20）")

hdr = ["区分", "開始", "終了", "日数\n（平年）", "S値補正", "混和剤の\n適用期",
       "AE減水剤", "高性能AE減水剤", "備考", "検証"]
for i, h in enumerate(hdr):
    put(ws, f"{get_column_letter(1+i)}4", h, bold=True, fill=C_HDR, wrap=True)
for i, h in enumerate(["開始キー", "終了キー", "組合せ"]):
    put(ws, f"{get_column_letter(12+i)}4", h, bold=True, fill=C_CALC, size=9)
put(ws, "L3", "▼ 計算用", size=9, color="808080", border=False, align=LEFT)

segments = [
    ("①", "1/1", "2/20", "冬期。12/8からの継続期間"),
    ("②", "2/21", "3/7", "S値のみ+3へ切替（混和剤は冬期のまま）"),
    ("③", "3/8", "6/15", "冬期→標準期へ切替"),
    ("④", "6/16", "6/29", "混和剤のみ遅延形へ先行切替（S値は+3のまま）"),
    ("⑤", "6/30", "9/14", "夏期のピーク期間"),
    ("⑥", "9/15", "9/23", "S値のみ+3へ戻る（混和剤は夏期のまま）"),
    ("⑦", "9/24", "11/26", "夏期→標準期へ切替"),
    ("⑧", "11/27", "12/7", "S値のみ+6へ先行切替（混和剤は標準期のまま）"),
    ("⑨", "12/8", "12/31", "標準期→冬期へ切替。翌年2/20まで継続"),
]

FIRST = 5
for i, (no, start, end, note) in enumerate(segments):
    r = FIRST + i
    put(ws, f"A{r}", no, bold=True)
    put(ws, f"B{r}", start, color="0000FF")
    put(ws, f"C{r}", end, color="0000FF")
    put(ws, f"L{r}",
        f'=VALUE(LEFT(B{r},FIND("/",B{r})-1))*100+VALUE(MID(B{r},FIND("/",B{r})+1,5))',
        fill=C_CALC, size=9)
    put(ws, f"M{r}",
        f'=VALUE(LEFT(C{r},FIND("/",C{r})-1))*100+VALUE(MID(C{r},FIND("/",C{r})+1,5))',
        fill=C_CALC, size=9)
    put(ws, f"N{r}", f"=INDEX(日別データ!$H:$H,MATCH($L{r},日別データ!$D:$D,0))",
        fill=C_CALC, size=9, align=LEFT)
    put(ws, f"D{r}",
        f"=IF($L{r}<=$M{r},"
        f'COUNTIFS(日別データ!$D:$D,">="&$L{r},日別データ!$D:$D,"<="&$M{r}),'
        f'COUNTIFS(日別データ!$D:$D,">="&$L{r})+COUNTIFS(日別データ!$D:$D,"<="&$M{r}))',
        fmt='0"日"')
    put(ws, f"E{r}", f"=INDEX(日別データ!$E:$E,MATCH($L{r},日別データ!$D:$D,0))",
        bold=True, fmt='"+"0')
    put(ws, f"F{r}", f"=INDEX(日別データ!$G:$G,MATCH($L{r},日別データ!$D:$D,0))")
    put(ws, f"G{r}", f"=INDEX(日別データ!$F:$F,MATCH($L{r},日別データ!$D:$D,0))")
    put(ws, f"H{r}", f"=$G{r}")
    put(ws, f"I{r}", note, align=LEFT, size=10)
    put(ws, f"J{r}",
        f"=IF(IF($L{r}<=$M{r},"
        f'COUNTIFS(日別データ!$D:$D,">="&$L{r},日別データ!$D:$D,"<="&$M{r},日別データ!$H:$H,$N{r}),'
        f'COUNTIFS(日別データ!$D:$D,">="&$L{r},日別データ!$H:$H,$N{r})'
        f'+COUNTIFS(日別データ!$D:$D,"<="&$M{r},日別データ!$H:$H,$N{r}))'
        f'=$D{r},"OK","要確認")', size=10)

LAST = FIRST + len(segments) - 1
TOT = LAST + 1
put(ws, f"A{TOT}", "合計", bold=True, fill=C_HDR2)
put(ws, f"B{TOT}", "", fill=C_HDR2)
put(ws, f"C{TOT}", "", fill=C_HDR2)
put(ws, f"D{TOT}", f"=SUM(D{FIRST}:D{LAST})", bold=True, fill=C_HDR2, fmt='0"日"')
for col in "EFGHI":
    put(ws, f"{col}{TOT}", "", fill=C_HDR2)
put(ws, f"I{TOT}", "平年の年間日数と一致すればOK", fill=C_HDR2, align=LEFT, size=9)
put(ws, f"J{TOT}", f'=IF(D{TOT}=365,"OK","要確認")', bold=True, fill=C_HDR2, size=10)

cmp_top = TOT + 2
put(ws, f"A{cmp_top}", "【S値と混和剤の切替日の比較】", bold=True, size=10,
    border=False, align=LEFT)
put(ws, f"A{cmp_top+1}",
    f'="S値の切替日　　：　"&B6&"（+6→+3）　"&B9&"（+3→+6）　"&B10&"（+6→+3）　"&B12&"（+3→+6）"',
    size=10, border=False, align=LEFT)
put(ws, f"A{cmp_top+2}",
    f'="混和剤の切替日　：　"&B7&"（冬期→標準期）　"&B8&"（標準期→夏期）　"&B11&"（夏期→標準期）　"&B13&"（標準期→冬期）"',
    size=10, border=False, align=LEFT)
put(ws, f"A{cmp_top+3}",
    f'="→ 両者がずれる期間（備考が黄色の区分）：②"&B6&"～"&C6&"（"&D6&"日）　④"&B8&"～"&C8&"（"&D8&"日）　'
    f'⑥"&B10&"～"&C10&"（"&D10&"日）　⑧"&B12&"～"&C12&"（"&D12&"日）"',
    size=10, bold=True, border=False, align=LEFT, color="C00000")

ws.conditional_formatting.add(
    f"I{FIRST}:I{LAST}",
    FormulaRule(formula=[f'ISNUMBER(SEARCH("のみ",$I{FIRST}))'],
                fill=PatternFill("solid", fgColor="FFF2CC")))

memo = [
    "【この表の見方・前提】",
    "・S値補正：+6＝6/30～9/14, 11/27～2/20／+3＝それ以外（ご提示条件）。",
    "・混和剤：AE減水剤・高性能AE減水剤とも、標準期 3/8～6/15・9/24～12/7＝標準形、夏期 6/16～9/23＝遅延形、冬期 12/8～3/7＝標準形。両剤の日程は同一のため H列は G列を参照しています。",
    "・青字の開始／終了日は、上記2つの条件を重ね合わせて切り替わる日（区切り日）です。「前提条件」シートの期間を変更した場合は、この青字の日付も合わせて見直してください。",
    "・J列「検証」は、その区分の全日が同じ組合せかどうかを日別データから数式で照合した結果です（すべてOKで整合）。",
    "・日数は平年（365日）。うるう年は 2/29 が区分②に入り、区分②が16日・合計366日になります。",
    "・備考欄が黄色の区分（②④⑥⑧）は、S値と混和剤のどちらか一方だけが切り替わる「ずれ期間」です。",
    "・出典：ご提示条件（2026-08-24 受領）。",
]
for i, t in enumerate(memo):
    put(ws, f"A{cmp_top+5+i}", t, size=9, bold=(i == 0), border=False, align=LEFT,
        color="000000" if i == 0 else "404040")

for col, w in zip("ABCDEFGHIJ", (6, 10, 10, 10, 10, 12, 14, 16, 42, 8)):
    ws.column_dimensions[col].width = w
for col in "LMN":
    ws.column_dimensions[col].width = 10
ws.row_dimensions[4].height = 32
ws.freeze_panes = "A5"

rng = f"E{FIRST}:E{LAST}"
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=["6"], fill=F_S6))
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=["3"], fill=F_S3))
for col in ("F", "G", "H"):
    r = f"{col}{FIRST}:{col}{LAST}"
    ws.conditional_formatting.add(r, CellIsRule(operator="equal", formula=['"夏期"'], fill=F_SUM))
    ws.conditional_formatting.add(r, CellIsRule(operator="equal", formula=['"冬期"'], fill=F_WIN))
    ws.conditional_formatting.add(r, CellIsRule(operator="equal", formula=['"標準期"'], fill=F_STD))
    ws.conditional_formatting.add(r, CellIsRule(operator="equal", formula=['"遅延形"'], fill=F_SUM))

# ===========================================================================
# シート「年間カレンダー」（S値と混和剤を1つのカレンダーに統合）
# ===========================================================================
cs = wb.create_sheet("年間カレンダー", 1)
title(cs, "年間カレンダー（S値補正 × 混和剤の適用期）",
      "1マスに「数字＝S値補正／文字＝混和剤の適用期」をまとめて表示しています。"
      "色は 緑＝標準期・橙＝夏期・青＝冬期、濃い色＝S値+6、淡い色＝S値+3。値は「日別データ」から自動参照しています。")

GTOP = 4              # 見出し行
GHDR = GTOP + 1       # 日付ヘッダー行
GFIRST = GHDR + 1     # 1月の行
GLAST = GFIRST + 11   # 12月の行

put(cs, f"A{GTOP}", "■ 数字＝S値補正（+6／+3）　　文字＝混和剤の適用期（標＝標準期／夏＝夏期／冬＝冬期）",
    bold=True, fill=C_HDR2, align=LEFT)
put(cs, f"A{GHDR}", "月＼日", bold=True, fill=C_HDR, size=9)
for d in range(1, 32):
    put(cs, f"{get_column_letter(1+d)}{GHDR}", d, bold=True, fill=C_HDR, size=9)

for i, m in enumerate(range(1, 13)):
    r = GFIRST + i
    put(cs, f"A{r}", m, bold=True, fill=C_HDR, fmt='0"月"', size=9)
    for d in range(1, 32):
        col = get_column_letter(1 + d)
        k = f"$A{r}*100+{col}${GHDR}"
        put(cs, f"{col}{r}",
            f"=IFERROR(INDEX(日別データ!$E$4:$E${LAST_D},MATCH({k},日別データ!$D$4:$D${LAST_D},0))"
            f"&LEFT(INDEX(日別データ!$G$4:$G${LAST_D},MATCH({k},日別データ!$D$4:$D${LAST_D},0)),1),\"\")",
            size=9)

# 6通りの組合せをセルの文字で判定して着色（濃＝S値+6／淡＝S値+3）
for term_ch in ("標", "夏", "冬"):
    for sv, palette in ((6, DARK), (3, LIGHT)):
        cs.conditional_formatting.add(
            f"B{GFIRST}:AF{GLAST}",
            CellIsRule(operator="equal", formula=[f'"{sv}{term_ch}"'],
                       fill=PatternFill("solid", fgColor=palette[term_ch])))

# --- 凡例（S値 × 適用期 のマトリクス） ------------------------------------
LG = GLAST + 2
put(cs, f"A{LG}", "【凡例】", bold=True, size=9, border=False, align=LEFT)

legend_cols = [("B", "H", "標"), ("I", "O", "夏"), ("P", "V", "冬")]
legend_head = {"標": "標準期（標準形）", "夏": "夏期（遅延形）", "冬": "冬期（標準形）"}


def legend_cell(c1, c2, row, text, fill, bold=False):
    cs.merge_cells(f"{c1}{row}:{c2}{row}")
    for c in range(cs[f"{c1}{row}"].column, cs[f"{c2}{row}"].column + 1):
        style(cs.cell(row=row, column=c), fill=fill, size=9, bold=bold)
    cs[f"{c1}{row}"] = text
    style(cs[f"{c1}{row}"], fill=fill, size=9, bold=bold)


put(cs, f"A{LG+1}", "", fill=C_HDR2, size=9)
for c1, c2, ch in legend_cols:
    legend_cell(c1, c2, LG + 1, legend_head[ch], C_HDR, bold=True)
for j, (sv, palette) in enumerate(((6, DARK), (3, LIGHT))):
    r = LG + 2 + j
    put(cs, f"A{r}", f"S値+{sv}", bold=True, fill=C_HDR, size=9)
    for c1, c2, ch in legend_cols:
        legend_cell(c1, c2, r, f"{sv}{ch}", palette[ch])

cal_notes = [
    "※ 例：「6夏」＝S値+6・夏期（遅延形）、「3標」＝S値+3・標準期（標準形）。",
    "※ 混和剤の種類はAE減水剤・高性能AE減水剤とも共通で、標準期・冬期＝標準形、夏期＝遅延形です。",
    "※ 空欄はその月に存在しない日（例：2/30）です。日数は平年基準のため 2/29 は表示していません。",
    "※ 色や文字が切り替わる日が配合の切替日です。区分ごとの一覧は「年間区分表」シートを参照してください。",
]
for i, t in enumerate(cal_notes):
    put(cs, f"A{LG+5+i}", t, size=9, color="404040", border=False, align=LEFT)

cs.column_dimensions["A"].width = 7.5
for d in range(1, 32):
    cs.column_dimensions[get_column_letter(1 + d)].width = 4.3
cs.freeze_panes = f"B{GFIRST}"

# ===========================================================================
del wb["Sheet"]
for sh in wb.worksheets:
    sh.sheet_view.showGridLines = False
    sh.page_setup.orientation = "landscape"
    sh.page_setup.fitToWidth = 1
    sh.page_setup.fitToHeight = 0
    sh.sheet_properties.pageSetUpPr.fitToPage = True
wb.active = 0
wb.save(OUT)
print("saved:", OUT, "| 日別データ最終行:", LAST_D)
