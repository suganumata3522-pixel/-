# -*- coding: utf-8 -*-
"""場所打ち杭の支持力 問題集（図つき）Excel。出力: docs/castpile/場所打ち杭の支持力問題集.xlsx
1 支持力式（告示/指針・先端+周面）/ 2 先端支持力（N値範囲・上限）と周面摩擦（Ns・qu上限）
3 許容圧縮力（地盤・杭体 fc=Fc/4）/ 4 引抜き耐力と算定
※ 先端支持力係数α・周面摩擦係数・N上限・qu上限・安全率・fc=Fc/4は告示1113号/基礎指針で確認要
"""
import os
import math
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib import font_manager as fm
FONT = "/usr/share/fonts/truetype/fonts-japanese-gothic.ttf"
jp = fm.FontProperties(fname=FONT); fm.fontManager.addfont(FONT)
plt.rcParams["font.family"] = jp.get_name(); plt.rcParams["axes.unicode_minus"] = False
OUT = "docs/castpile"; FIG = os.path.join(OUT, "figures"); os.makedirs(FIG, exist_ok=True)


def save(fig, n):
    p = os.path.join(FIG, n); fig.savefig(p, dpi=130, bbox_inches="tight"); plt.close(fig); return p


def fig_mech():
    fig, ax = plt.subplots(figsize=(11.5, 6.2))
    # 地層
    ax.add_patch(mpatches.Rectangle((0, 0), 10, 8, fc="#f0e6d2", ec="none"))
    ax.add_patch(mpatches.Rectangle((0, 5.5), 10, 2.5, fc="#e7d9bf", ec="none"))  # 砂
    ax.add_patch(mpatches.Rectangle((0, 2.5), 10, 3.0, fc="#d9c9a8", ec="none"))  # 粘土
    ax.add_patch(mpatches.Rectangle((0, 0), 10, 2.5, fc="#c8b48c", ec="none"))  # 支持層
    ax.text(0.3, 6.6, "砂質土 Ns", fontproperties=jp, fontsize=9, color="#7a5a2a")
    ax.text(0.3, 3.8, "粘性土 qu", fontproperties=jp, fontsize=9, color="#7a5a2a")
    ax.text(0.3, 1.0, "支持層（先端N値）", fontproperties=jp, fontsize=9, color="#5a4020")
    # 杭
    ax.add_patch(mpatches.Rectangle((4.3, 1.5), 1.4, 6.7, fc="#c0c0c0", ec="k", lw=1.5))
    ax.text(5.0, 7.9, "場所打ちコンクリート杭", ha="center", fontproperties=jp, fontsize=9, color="#333")
    # 荷重
    ax.annotate("", xy=(5.0, 8.2), xytext=(5.0, 9.2), arrowprops=dict(arrowstyle="-|>", color="#c00000", lw=3))
    ax.text(5.2, 8.8, "軸力 N（圧縮）", fontproperties=jp, fontsize=10, color="#c00000")
    # 先端支持力
    ax.annotate("", xy=(5.0, 1.3), xytext=(5.0, 0.4), arrowprops=dict(arrowstyle="-|>", color="#548235", lw=3))
    ax.text(5.7, 0.7, "先端支持力 Rp=qp・Ap", fontproperties=jp, fontsize=9.5, color="#548235")
    # 周面摩擦
    for yy in [2.5, 3.5, 4.5, 5.5, 6.5]:
        ax.annotate("", xy=(4.2, yy), xytext=(3.8, yy - 0.35), arrowprops=dict(arrowstyle="->", color="#c55a11"))
        ax.annotate("", xy=(5.8, yy), xytext=(6.2, yy - 0.35), arrowprops=dict(arrowstyle="->", color="#c55a11"))
    ax.text(1.7, 5.0, "周面摩擦\nRf", ha="center", fontproperties=jp, fontsize=9.5, color="#c55a11")
    # 式
    ax.text(7.2, 5.5, "極限支持力\nRu = Rp + Rf\n\n長期許容\nRa = Ru/3", fontproperties=jp,
            fontsize=10, color="#1f4e79", va="center",
            bbox=dict(boxstyle="round", fc="#eaf1fb", ec="#2e75b6"))
    ax.set_xlim(0, 10.5); ax.set_ylim(0, 9.4); ax.axis("off")
    ax.set_title("図 1  場所打ち杭の支持機構（先端支持＋周面摩擦）",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig1_mech.png")


def fig_terms():
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 5.0))
    # (a) 先端N値の取り方
    ax = axes[0]
    ax.text(0.5, 0.95, "先端支持力 qp = α・N", transform=ax.transAxes, ha="center",
            fontproperties=jp, fontsize=10.5, fontweight="bold", color="#1f4e79")
    # N値プロファイル模式
    d = np.array([9, 10, 11, 12, 13, 14, 15, 16])
    Nv = np.array([12, 18, 30, 48, 52, 55, 58, 60])
    ax.plot(Nv, d, "o-", color="#c00000", transform=ax.transData)
    ax.axhline(13.0, color="#548235", ls="--", lw=1.2)
    ax.text(5, 12.6, "杭先端", fontproperties=jp, fontsize=8, color="#548235")
    ax.axvspan(0, 60, ymin=0, ymax=0, alpha=0)
    ax.fill_between([0, 65], 12, 14, color="#cfe0c0", alpha=0.5)
    ax.text(30, 14.6, "先端付近の平均N値をとる\n（上限あり、例 N<=60）",
            fontproperties=jp, fontsize=8.5, color="#333", ha="center")
    ax.set_xlim(0, 65); ax.set_ylim(17, 8)
    ax.set_xlabel("N値", fontproperties=jp, fontsize=9)
    ax.set_ylabel("深度 (m)", fontproperties=jp, fontsize=9)
    ax.set_title("(a) 先端N値の取り方と上限", fontproperties=jp, fontsize=9.5, fontweight="bold")
    ax.grid(alpha=0.3)
    # (b) 周面摩擦
    ax = axes[1]
    ax.text(0.5, 0.95, "周面摩擦力 Rf = ψ・Σ(fs・L)", transform=ax.transAxes, ha="center",
            fontproperties=jp, fontsize=10.5, fontweight="bold", color="#1f4e79")
    ax.text(0.04, 0.8,
            "周長 ψ = π・D\n\n"
            "【砂質土層】 fs = (10/3)・Ns  (kN/m2)\n"
            "   Ns：平均N値（上限あり）\n\n"
            "【粘性土層】 fs = (1/2)・qu = c  (kN/m2)\n"
            "   qu：一軸圧縮強さ（上限あり）\n\n"
            "各層の fs に層厚 L と周長 ψ を掛けて合計\n"
            "Rf = ψ・( fs_s・Ls + fs_c・Lc + ... )",
            transform=ax.transAxes, fontproperties=jp, fontsize=9, va="top")
    ax.text(0.04, 0.08, "※ 係数・N上限・qu上限は告示/指針で確認要。",
            transform=ax.transAxes, fontproperties=jp, fontsize=7.8, color="#833c00")
    ax.axis("off")
    ax.set_title("(b) 周面摩擦（砂・粘土の層別）", fontproperties=jp, fontsize=9.5, fontweight="bold")
    fig.suptitle("図 2  先端支持力（N値・上限）と周面摩擦（Ns・qu）",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig2_terms.png")


def fig_comp():
    fig, ax = plt.subplots(figsize=(11.5, 5.4))
    ax.text(0.5, 0.95, "許容圧縮力＝地盤と杭体の小さい方", transform=ax.transAxes, ha="center",
            fontproperties=jp, fontsize=11, fontweight="bold", color="#1f4e79")
    ax.add_patch(mpatches.FancyBboxPatch((0.03, 0.5), 0.44, 0.34, boxstyle="round,pad=0.01",
                 transform=ax.transAxes, fc="#eaf7ea", ec="#548235", lw=1.2))
    ax.text(0.25, 0.78, "(1) 地盤で決まる許容支持力", transform=ax.transAxes, ha="center",
            fontproperties=jp, fontsize=9.5, fontweight="bold", color="#548235")
    ax.text(0.25, 0.62, "Ra = Ru/3 = (Rp+Rf)/3\n例 Ru=7514kN → Ra≒2505kN",
            transform=ax.transAxes, ha="center", va="center", fontproperties=jp, fontsize=9)
    ax.add_patch(mpatches.FancyBboxPatch((0.53, 0.5), 0.44, 0.34, boxstyle="round,pad=0.01",
                 transform=ax.transAxes, fc="#e8f0fa", ec="#2e75b6", lw=1.2))
    ax.text(0.75, 0.78, "(2) 杭体で決まる許容圧縮力", transform=ax.transAxes, ha="center",
            fontproperties=jp, fontsize=9.5, fontweight="bold", color="#1f4e79")
    ax.text(0.75, 0.62, "Rc = fc・Ap,  fc = Fc/4（長期）\n例 Fc=24→fc=6N/mm2, Rc≒4712kN",
            transform=ax.transAxes, ha="center", va="center", fontproperties=jp, fontsize=9)
    ax.annotate("", xy=(0.5, 0.42), xytext=(0.25, 0.5), xycoords="axes fraction",
                arrowprops=dict(arrowstyle="-|>", color="#c00000", lw=1.6))
    ax.annotate("", xy=(0.5, 0.42), xytext=(0.75, 0.5), xycoords="axes fraction",
                arrowprops=dict(arrowstyle="-|>", color="#c00000", lw=1.6))
    ax.add_patch(mpatches.FancyBboxPatch((0.22, 0.24), 0.56, 0.14, boxstyle="round,pad=0.01",
                 transform=ax.transAxes, fc="#f8d0d0", ec="#c00000", lw=1.2))
    ax.text(0.5, 0.31, "許容圧縮力 = min( 地盤Ra, 杭体Rc ) = min(2505, 4712) = 2505kN（地盤支配）",
            transform=ax.transAxes, ha="center", va="center", fontproperties=jp, fontsize=9.5,
            fontweight="bold", color="#c00000")
    ax.text(0.5, 0.1, "※ 場所打ち杭の長期許容圧縮応力度は fc=Fc/4（施工品質のばらつきを考慮）。\n"
            "既製杭・鋼管杭とは異なる。安全率・係数は告示/指針で確認要。",
            transform=ax.transAxes, ha="center", va="center", fontproperties=jp, fontsize=8.3, color="#833c00")
    ax.axis("off")
    ax.set_title("図 3  許容圧縮力（地盤 と 杭体 fc=Fc/4）",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig3_comp.png")


def fig_pullout():
    fig, axes = plt.subplots(1, 2, figsize=(13.0, 5.2), gridspec_kw={"width_ratios": [1.0, 1.1]})
    # (a) 引抜き機構
    ax = axes[0]
    ax.add_patch(mpatches.Rectangle((0, 0), 6, 7.5, fc="#f0e6d2", ec="none"))
    ax.add_patch(mpatches.Rectangle((2.3, 0.5), 1.4, 6.5, fc="#c0c0c0", ec="k", lw=1.5))
    ax.annotate("", xy=(3.0, 7.3), xytext=(3.0, 8.3), arrowprops=dict(arrowstyle="-|>", color="#c00000", lw=3))
    ax.text(3.2, 7.9, "引抜き力 T", fontproperties=jp, fontsize=10, color="#c00000")
    for yy in [1.5, 2.5, 3.5, 4.5, 5.5]:
        ax.annotate("", xy=(2.2, yy - 0.35), xytext=(1.8, yy), arrowprops=dict(arrowstyle="->", color="#c55a11"))
        ax.annotate("", xy=(3.8, yy - 0.35), xytext=(4.2, yy), arrowprops=dict(arrowstyle="->", color="#c55a11"))
    ax.text(0.8, 4.0, "周面\n摩擦\nRf", ha="center", fontproperties=jp, fontsize=9, color="#c55a11")
    ax.annotate("", xy=(3.0, 2.0), xytext=(3.0, 3.0), arrowprops=dict(arrowstyle="-|>", color="#333", lw=2))
    ax.text(3.2, 2.4, "杭自重 Wp", fontproperties=jp, fontsize=9, color="#333")
    ax.text(3.0, 0.0, "先端支持は引抜きに効かない", ha="center", fontproperties=jp, fontsize=8, color="#548235")
    ax.set_xlim(-0.3, 6.3); ax.set_ylim(-0.5, 8.6); ax.axis("off")
    ax.set_title("(a) 引抜き機構（周面摩擦＋自重）", fontproperties=jp, fontsize=9.5, fontweight="bold")
    # (b) 圧縮 vs 引抜き
    ax = axes[1]
    ax.text(0.5, 0.95, "圧縮と引抜きの対比", transform=ax.transAxes, ha="center",
            fontproperties=jp, fontsize=10.5, fontweight="bold", color="#1f4e79")
    ax.text(0.04, 0.8,
            "【圧縮】\n"
            "  Ru = Rp（先端）+ Rf（周面）\n"
            "  Ra = Ru/3 = 2505kN（例）\n\n"
            "【引抜き】※先端支持は効かない\n"
            "  極限 Tu = Rf + Wp（杭自重）\n"
            "  長期 Ta = Rf/3 + Wp\n"
            "  例 Rf=1623, Wp=283kN\n"
            "    Ta = 1623/3 + 283 ≒ 824kN\n\n"
            "  引抜きは周面摩擦と杭自重で抵抗\n"
            "  （地震時の転倒・浮上りで重要）",
            transform=ax.transAxes, fontproperties=jp, fontsize=9, va="top")
    ax.text(0.04, 0.06, "※ 引抜きの安全率・周面摩擦の扱いは告示/指針で確認要。",
            transform=ax.transAxes, fontproperties=jp, fontsize=7.8, color="#833c00")
    ax.axis("off")
    ax.set_title("(b) 圧縮・引抜きの算定", fontproperties=jp, fontsize=9.5, fontweight="bold")
    fig.suptitle("図 4  引抜き耐力（周面摩擦＋杭自重）と圧縮の対比",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig4_pullout.png")


figs = {"mech": fig_mech(), "terms": fig_terms(), "comp": fig_comp(), "pullout": fig_pullout()}
print("figs:", list(figs.keys()))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
wb = Workbook()
C_TITLE = "1F4E79"; C_HEAD = "2E75B6"; C_ANS = "E2EFDA"; C_WARN = "FCE4D6"
thin = Side(style="thin", color="BFBFBF"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
f_title = Font(name="MS PGothic", size=15, bold=True, color="FFFFFF")
f_head = Font(name="MS PGothic", size=11, bold=True, color="FFFFFF")
f_body = Font(name="MS PGothic", size=10); f_ans = Font(name="MS PGothic", size=10, color="375623")
f_warn = Font(name="MS PGothic", size=9, color="833C00", italic=True)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)


def setup(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False


def title_row(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_title; c.fill = PatternFill("solid", fgColor=C_TITLE)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1); ws.row_dimensions[row].height = 30


def head(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_head; c.fill = PatternFill("solid", fgColor=C_HEAD)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1); ws.row_dimensions[row].height = 22


def body(ws, row, text, span=8, ans=False, h=None):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_ans if ans else f_body; c.alignment = wrap
    if ans:
        c.fill = PatternFill("solid", fgColor=C_ANS)
    if h:
        ws.row_dimensions[row].height = h


def warn(ws, row, text, span=8, h=None):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_warn; c.alignment = wrap
    c.fill = PatternFill("solid", fgColor=C_WARN)
    if h:
        ws.row_dimensions[row].height = h


def tbl(ws, start_row, headers, rows, col1=1):
    r = start_row
    for j, htxt in enumerate(headers):
        c = ws.cell(r, col1 + j, htxt); c.font = Font(name="MS PGothic", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C_HEAD); c.alignment = center; c.border = border
    for data in rows:
        r += 1
        for j, v in enumerate(data):
            c = ws.cell(r, col1 + j, v); c.font = f_body
            c.alignment = wrap if j == 0 else center; c.border = border
            if r % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F2F7FC")
    return r


def put_img(ws, path, anchor, w=None):
    img = XLImage(path)
    if w:
        ratio = w / img.width; img.width = w; img.height = int(img.height * ratio)
    ws.add_image(img, anchor)


# ===== 目次 =====
ws = wb.active; ws.title = "目次"; setup(ws, [4, 24, 54, 16])
title_row(ws, 1, "場所打ち杭の支持力 問題集（RC マンション設計担当・新人向け）", span=4)
body(ws, 2, "目標：場所打ちコンクリート杭の支持力式（告示/基礎指針）を理解し、先端支持力・周面摩擦・"
            "許容圧縮力（地盤/杭体）・引抜き耐力を手計算で求められること。", span=4, h=32)
r = tbl(ws, 4, ["No.", "シート", "到達目標", "図"],
        [["1", "1 支持力式", "先端＋周面の式構造を理解", "支持機構"],
         ["2", "2 先端・周面", "N値範囲・上限、Ns・qu上限を理解", "N値/摩擦"],
         ["3", "3 許容圧縮力", "地盤・杭体(fc=Fc/4)で決定", "min判定"],
         ["4", "4 引抜き耐力", "周面摩擦＋自重で算定", "引抜き"]])
warn(ws, r + 2, "※ 先端支持力係数α、周面摩擦係数、N上限・qu上限、安全率、fc=Fc/4 は"
                "告示1113号・建築基礎構造設計指針の最新版で必ず確認すること。数値は例示。", span=4, h=40)

# ===== 1 支持力式 =====
ws = wb.create_sheet("1 支持力式"); setup(ws, [8, 16, 18, 16, 12, 12])
title_row(ws, 1, "1  場所打ち杭の支持力式（先端＋周面）")
head(ws, 3, "■ 図 1  支持機構"); put_img(ws, figs["mech"], "A4", w=760)
head(ws, 30, "■ 問題 1  支持力式の構造")
body(ws, 31, "場所打ち杭の極限支持力 Ru=Rp+Rf（先端支持力＋周面摩擦力）の意味を説明せよ。"
             "長期許容支持力 Ra=Ru/3 の考え方（安全率）にも触れよ。", h=40)
head(ws, 33, "■ 問題 2  先端支持と周面摩擦")
body(ws, 34, "先端支持力 Rp（先端地盤の抵抗）と周面摩擦力 Rf（杭側面と地盤の摩擦）が"
             "どのように荷重を支えるか、図1を使って説明せよ。", h=36)
head(ws, 36, "■ 問題 3  告示/指針の位置づけ")
body(ws, 37, "杭の支持力算定に告示1113号・建築基礎構造設計指針を用いること、"
             "係数や上限値は必ず最新の規準で確認することの重要性を述べよ。", h=36)

# ===== 2 先端・周面 =====
ws = wb.create_sheet("2 先端と周面"); setup(ws, [8, 16, 18, 16, 12, 12])
title_row(ws, 1, "2  先端支持力（N値・上限）と周面摩擦（Ns・qu）")
head(ws, 3, "■ 図 2  先端N値・周面摩擦"); put_img(ws, figs["terms"], "A4", w=860)
head(ws, 27, "■ 問題 1  先端支持力")
body(ws, 28, "先端支持力 qp=α・N（α：先端支持力係数、N：先端付近の平均N値）の考え方を述べよ。"
             "先端N値の取り方（先端付近の平均）と上限（例 N<=60）に触れよ。", h=40)
head(ws, 30, "■ 問題 2  周面摩擦")
r = tbl(ws, 31, ["土質", "周面摩擦度 fs", "備考"],
        [["砂質土", "(10/3)・Ns", "Ns：平均N値（上限あり）"],
         ["粘性土", "(1/2)・qu ＝ c", "qu：一軸圧縮強さ（上限あり）"]])
body(ws, r + 1, "Rf = ψ・Σ(fs・L)（ψ=π・D：杭周長、L：各層厚）で合計する。", h=24)
head(ws, r + 3, "■ 問題 3  上限値の意味")
body(ws, r + 4, "先端N値・周面のNs・qu に上限を設ける理由を述べよ（過大評価の防止、"
                "極端に硬い層でも支持力を頭打ちにして安全側に評価する）。", h=40)
head(ws, r + 6, "■ 問題 4  周面摩擦の計算")
body(ws, r + 7, "φ1000（ψ=3.14m）で、砂質土層 Ns=10・厚さ8m、粘性土層 qu=100・厚さ5m のとき、"
                "各層の fs と周面摩擦力 Rf を求めよ。", h=36)

# ===== 3 許容圧縮力 =====
ws = wb.create_sheet("3 許容圧縮力"); setup(ws, [8, 16, 18, 16, 12, 12])
title_row(ws, 1, "3  許容圧縮力（地盤 と 杭体 fc=Fc/4）")
head(ws, 3, "■ 図 3  許容圧縮力の決定"); put_img(ws, figs["comp"], "A4", w=780)
head(ws, 29, "■ 問題 1  地盤で決まる許容支持力")
body(ws, 30, "φ1000（Ap=0.785m2）、先端N=50・α=150（qp=α・N）、Rf=1623kN のとき、"
             "先端支持力 Rp・極限支持力 Ru・長期許容支持力 Ra=Ru/3 を求めよ。", h=40)
head(ws, 32, "■ 問題 2  杭体で決まる許容圧縮力")
body(ws, 33, "場所打ち杭の長期許容圧縮応力度 fc=Fc/4（Fc=24N/mm2）のとき、"
             "杭体で決まる許容圧縮力 Rc=fc・Ap を求めよ。なぜ場所打ち杭は Fc/4 と厳しいか述べよ。", h=40)
head(ws, 35, "■ 問題 3  許容圧縮力の決定")
body(ws, 36, "許容圧縮力＝min（地盤Ra, 杭体Rc）である理由を述べよ。問1・問2の値で"
             "どちらが支配するか判定せよ（Ra=2505 vs Rc=4712）。", h=36)
head(ws, 38, "■ 問題 4  杭種による違い")
body(ws, 39, "場所打ち杭（fc=Fc/4）と既製コンクリート杭・鋼管杭で許容応力度の考え方が異なることに"
             "触れ、杭体の許容圧縮力が支配するのはどんな場合か述べよ。", h=36)
warn(ws, 41, "※ α・安全率・fc=Fc/4 は告示1113号/基礎指針で確認要。", h=24)

# ===== 4 引抜き耐力 =====
ws = wb.create_sheet("4 引抜き耐力"); setup(ws, [8, 16, 18, 16, 12, 12])
title_row(ws, 1, "4  引抜き耐力（周面摩擦＋杭自重）")
head(ws, 3, "■ 図 4  引抜き機構と対比"); put_img(ws, figs["pullout"], "A4", w=840)
head(ws, 28, "■ 問題 1  引抜き抵抗の要素")
body(ws, 29, "杭の引抜き抵抗が『周面摩擦力 Rf ＋ 杭自重 Wp』で構成され、先端支持力は効かないことを"
             "説明せよ。地震時の転倒・浮上りで引抜き検討が重要になることに触れよ。", h=40)
head(ws, 31, "■ 問題 2  引抜き耐力の計算")
body(ws, 32, "Rf=1623kN、杭自重 Wp=Ap・L・γc（φ1000・L=15m・γc=24kN/m3）のとき、"
             "極限引抜き Tu=Rf+Wp と長期許容引抜き Ta=Rf/3+Wp を求めよ。", h=40)
head(ws, 34, "■ 問題 3  圧縮と引抜きの対比")
body(ws, 35, "同じ杭で圧縮（Ra≒2505kN）と引抜き（Ta≒824kN）の許容値を比べ、"
             "なぜ引抜きの方が小さいか説明せよ（先端支持が使えない・自重のみ加算）。", h=40)
head(ws, 37, "■ 問題 4  設計の流れ")
body(ws, 38, "場所打ち杭の設計の流れをまとめよ（①地盤・杭配置→②先端支持力・周面摩擦→"
             "③許容圧縮力＝min(地盤,杭体)→④引抜き・水平力の検討→⑤杭体断面・配筋）。"
             "液状化層は周面摩擦を見込まない等の配慮にも触れよ。", h=44)
warn(ws, 40, "※ 引抜きの安全率・液状化層の扱いは告示1113号/基礎指針で確認要。", h=24)

# ===== 解答 =====
ws = wb.create_sheet("解答"); setup(ws, [4, 22, 60, 14])
title_row(ws, 1, "解答・解説", span=4); row = 2


def ah(t):
    global row; head(ws, row, t, span=4); row += 1


def an(t, h=None):
    global row; body(ws, row, t, span=4, ans=True, h=h); row += 1


def aw(t, h=None):
    global row; warn(ws, row, t, span=4, h=h); row += 1


ah("1  支持力式")
an("問1：Ru=Rp+Rf。先端支持力Rp（杭先端の地盤反力）と周面摩擦力Rf（杭側面と地盤の摩擦）の和が"
   "極限支持力。長期許容支持力Ra=Ru/3は、極限に安全率3を見込んだ常時の許容値。", h=44)
an("問2：軸力（圧縮）は、杭側面で周面摩擦Rfとして地盤に伝わりつつ、残りが先端Rpで支持層に伝わる。"
   "深い杭ほど周面摩擦の寄与が大きくなる。", h=40)
an("問3：杭の支持力は告示1113号・建築基礎構造設計指針の式で算定する。先端支持力係数・周面摩擦係数・"
   "上限値・安全率は改定され、杭種でも異なるため、必ず最新の規準を確認する。", h=40)

ah("2  先端と周面")
an("問1：qp=α・N。αは先端支持力係数（杭種で異なる。場所打ちは打込み杭より小さい）、"
   "Nは先端付近の平均N値。極端に大きいN値は上限（例 N<=60）で頭打ちにして安全側に評価する。", h=44)
an("問2：砂質土 fs=(10/3)Ns、粘性土 fs=(1/2)qu=c。Rf=ψ・Σ(fs・L)、ψ=π・D。"
   "各層の摩擦度に層厚と周長を掛けて合計する。", h=36)
an("問3：N値やquは局所的に過大な値が出ることがあり、そのまま使うと支持力を過大評価する。"
   "上限を設けて頭打ちにすることで、安全側かつばらつきに強い評価とする。", h=40)
an("問4：ψ=π×1.0=3.14m。砂 fs=(10/3)×10=33.3kN/m2、粘土 fs=100/2=50kN/m2。"
   "Rf=3.14×(33.3×8+50×5)=3.14×(266.4+250)=3.14×516.4≒1623kN。", h=40)

ah("3  許容圧縮力")
an("問1：Rp=qp・Ap=(150×50)×0.785=7500×0.785≒5890kN。Ru=Rp+Rf=5890+1623=7514kN。"
   "Ra=Ru/3=7514/3≒2505kN。", h=40)
an("問2：fc=Fc/4=24/4=6N/mm2=6000kN/m2。Rc=fc・Ap=6000×0.785≒4712kN。"
   "場所打ち杭は現場施工で品質のばらつきが大きいため、既製杭より厳しい Fc/4 とする。", h=40)
an("問3：杭の支持力は地盤と杭体の弱い方で決まるため min（地盤Ra, 杭体Rc）。"
   "Ra=2505 < Rc=4712 なので地盤支配で許容圧縮力=2505kN。杭径が大きく地盤が良いと杭体支配になり得る。", h=44)
an("問4：既製コンクリート杭（PHC等）や鋼管杭は工場製作で品質が安定し、許容応力度の考え方が異なる"
   "（fcの割り方や鋼材の許容応力度）。杭径が大きく短い・良質地盤では杭体（材料強度）が支配しやすい。", h=44)
aw("※ α・安全率・fc=Fc/4 は告示1113号/基礎指針で確認要。", h=22)

ah("4  引抜き耐力")
an("問1：引抜きでは先端支持は働かず、周面摩擦Rfと杭自重Wpだけが抵抗する。"
   "地震時の転倒モーメントで引抜き軸力が生じる杭（外周・隅）で引抜き検討が重要。", h=40)
an("問2：Wp=Ap・L・γc=0.785×15×24≒283kN。Tu=Rf+Wp=1623+283=1906kN。"
   "Ta=Rf/3+Wp=1623/3+283=541+283≒824kN。", h=36)
an("問3：圧縮は先端Rp＋周面Rfで支持できるが、引抜きは先端が効かず周面Rfと自重Wpのみ。"
   "支配的な先端支持が使えない分、引抜き許容値（824kN）は圧縮（2505kN）よりかなり小さい。", h=44)
an("問4：①地盤・杭配置→②先端支持力・周面摩擦の算定→③許容圧縮力=min(地盤,杭体fc=Fc/4)→"
   "④引抜き（周面＋自重）・水平力の検討→⑤杭体断面・配筋。液状化層は周面摩擦を見込まない"
   "（または低減）等の配慮が必要。", h=44)
aw("※ 引抜きの安全率・液状化層の扱いは告示1113号/基礎指針で確認要。", h=24)

XLSX = os.path.join(OUT, "場所打ち杭の支持力問題集.xlsx"); wb.save(XLSX); print("saved:", XLSX)
