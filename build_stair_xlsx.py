# -*- coding: utf-8 -*-
"""階段の設計 問題集（図つき）Excel 生成スクリプト。
出力: docs/stair/階段設計問題集.xlsx
1 構造形式と特性 / 2 自重の算出 / 3 片持ち階段の応力・断面算定
4 片持ち階段支持壁の耐震設計(両方向) / 5 一方向階段の応力(水平投影・実長)・断面算定
6 一方向階段支持梁の耐震設計
RC造マンションの設計担当を想定。数値は build 時に検算済み。
"""
import os
import math
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib import font_manager as fm

FONT_PATH = "/usr/share/fonts/truetype/fonts-japanese-gothic.ttf"
jp = fm.FontProperties(fname=FONT_PATH)
fm.fontManager.addfont(FONT_PATH)
plt.rcParams["font.family"] = jp.get_name()
plt.rcParams["axes.unicode_minus"] = False

OUT = "docs/stair"
FIG = os.path.join(OUT, "figures")
os.makedirs(FIG, exist_ok=True)

C_BLUE = "#2a78d6"
C_PINK = "#d55181"


def save(fig, name):
    p = os.path.join(FIG, name)
    fig.savefig(p, dpi=130, bbox_inches="tight")
    plt.close(fig)
    return p


def draw_steps(ax, x0, y0, n, tread, riser, color="#bcd2ea", ec="k"):
    """階段のギザギザ断面を描く（右上がり）。"""
    pts = [(x0, y0)]
    x, y = x0, y0
    for i in range(n):
        x += tread
        pts.append((x, y))
        y += riser
        pts.append((x, y))
    # 下面（斜めスラブ）
    x_end, y_end = x, y
    return pts, (x_end, y_end)


# ===========================================================================
# 図 1: 階段の構造形式
# ===========================================================================
def fig_types():
    fig, axes = plt.subplots(1, 3, figsize=(14.5, 4.6))
    # (a) 一方向階段（両端の梁で支持）
    ax = axes[0]
    tread, riser = 0.28, 0.18
    x, y = 0, 0
    top_pts = [(x, y)]
    for i in range(7):
        x += tread; top_pts.append((x, y))
        y += riser; top_pts.append((x, y))
    xs = [p[0] for p in top_pts]; ys = [p[1] for p in top_pts]
    # 斜めスラブ下面
    ax.plot(xs, ys, color="#1f4e79", lw=1.5)
    ax.plot([0, x], [0 - 0.25, y - 0.25], color="#1f4e79", lw=1.5)
    ax.fill_between([0, x], [-0.25, y - 0.25], [0, y], color="#cfe0c0",
                   alpha=0.3)
    # 両端支持梁
    ax.add_patch(mpatches.Rectangle((-0.35, -0.5), 0.35, 0.6, fc="#d9d9d9",
                                     ec="k"))
    ax.add_patch(mpatches.Rectangle((x, y - 0.5), 0.35, 0.6, fc="#d9d9d9",
                                     ec="k"))
    ax.text(x / 2, -1.0, "一方向階段\n（上下端の梁で支持）", fontproperties=jp,
            ha="center", fontsize=9, color="#444")
    ax.set_xlim(-0.6, x + 0.6); ax.set_ylim(-1.4, y + 0.4)
    ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("(a) 一方向階段", fontproperties=jp, fontsize=10,
                 fontweight="bold")

    # (b) 片持ち階段（壁から踏板が片持ち）
    ax = axes[1]
    ax.add_patch(mpatches.Rectangle((0, -0.3), 0.4, 3.4, fc="#d9d9d9", ec="k",
                                     lw=1.2))  # 壁
    ax.text(0.2, 3.3, "支持壁", fontproperties=jp, ha="center", fontsize=8)
    for i in range(6):
        yy = i * 0.5
        ax.add_patch(mpatches.Rectangle((0.4, yy), 1.5, 0.12, fc="#bcd2ea",
                                         ec="k", lw=0.8))  # 踏板
    ax.annotate("踏板が壁から\n片持ち", xy=(1.4, 2.5), xytext=(2.2, 2.8),
                fontproperties=jp, fontsize=8, color="#c00000",
                arrowprops=dict(arrowstyle="->", color="#c00000"))
    ax.text(1.2, -1.0, "片持ち階段\n（各踏板が壁から片持ち）",
            fontproperties=jp, ha="center", fontsize=9, color="#444")
    ax.set_xlim(-0.3, 3.2); ax.set_ylim(-1.4, 3.4)
    ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("(b) 片持ち階段", fontproperties=jp, fontsize=10,
                 fontweight="bold")

    # (c) 折り返し階段（中間に踊り場）
    ax = axes[2]
    # 下段
    x, y = 0, 0
    for i in range(4):
        ax.add_patch(mpatches.Rectangle((x, y), 0.28, 0.12, fc="#bcd2ea",
                                        ec="k", lw=0.6))
        x += 0.28; y += 0.18
    # 踊り場
    ax.add_patch(mpatches.Rectangle((x, y), 0.9, 0.12, fc="#9ec6e8", ec="k",
                                     lw=0.8))
    ax.text(x + 0.45, y + 0.25, "踊り場", fontproperties=jp, ha="center",
            fontsize=7)
    # 上段
    x2 = x + 0.9; y2 = y
    for i in range(4):
        ax.add_patch(mpatches.Rectangle((x2, y2), 0.28, 0.12, fc="#bcd2ea",
                                        ec="k", lw=0.6))
        x2 += 0.28; y2 += 0.18
    ax.text((x2) / 2, -1.0, "折り返し階段\n（踊り場・中間梁付き）",
            fontproperties=jp, ha="center", fontsize=9, color="#444")
    ax.set_xlim(-0.3, x2 + 0.3); ax.set_ylim(-1.4, y2 + 0.4)
    ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("(c) 折り返し階段", fontproperties=jp, fontsize=10,
                 fontweight="bold")
    fig.suptitle("図 1  階段の構造形式",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig1_types.png")


# ===========================================================================
# 図 2: 自重の算定（斜めスラブ＋段形）
# ===========================================================================
def fig_weight():
    fig, ax = plt.subplots(figsize=(10.5, 5.2))
    tread, riser = 2.8, 1.8   # 拡大（mよりデシ）
    n = 4
    x, y = 0, 0
    top = [(x, y)]
    for i in range(n):
        x += tread; top.append((x, y))
        y += riser; top.append((x, y))
    xs = [p[0] for p in top]; ys = [p[1] for p in top]
    # 斜めスラブ（厚 t、下面）
    t = 1.5
    dx = t * riser / math.hypot(tread, riser)
    dy = -t * tread / math.hypot(tread, riser)
    ax.plot(xs, ys, color="k", lw=1.2)
    ax.plot([0 + dx, x + dx], [0 + dy, y + dy], color="k", lw=1.2)
    ax.fill(
        [0 + dx, x + dx] + xs[::-1],
        [0 + dy, y + dy] + ys[::-1],
        color="#bcd2ea", alpha=0.5)
    # 段形三角部（ハッチ）
    ax.fill_between([p[0] for p in top], [p[1] for p in top],
                    [top[0][1]] * len(top), color="#f4a261", alpha=0.3)
    # 斜め厚 t 注記
    mid = (xs[3], ys[3])
    ax.annotate("斜めスラブ厚 t\n（勾配に直角）", xy=(mid[0] + dx / 2, mid[1] + dy / 2),
                xytext=(mid[0] - 3.0, mid[1] + 1.5), fontproperties=jp,
                fontsize=9, color="#1f4e79",
                arrowprops=dict(arrowstyle="->", color="#1f4e79"))
    # 勾配角
    ax.plot([0, tread], [0, 0], color="gray", lw=0.8, ls=":")
    ax.annotate("θ", xy=(tread * 0.7, 0.15), fontproperties=jp, fontsize=11)
    ax.text(tread * 0.5, riser * 0.5, "蹴上180\n踏面280", fontproperties=jp,
            fontsize=8, color="#7a3b3b")
    # 説明
    ax.text(x / 2, -2.2,
            "水平投影 1m² あたりの自重：\n"
            "① 斜めスラブ = γ · t / cosθ（水平換算で厚くなる）\n"
            "② 段形三角 = γ · (蹴上/2)   ③ 仕上げ・積載を加算",
            fontproperties=jp, ha="center", fontsize=9.5, color="#444")
    ax.set_xlim(-3.5, x + 1); ax.set_ylim(-3.4, y + 1)
    ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("図 2  階段の自重（斜めスラブ換算＋段形）",
                 fontproperties=jp, fontsize=12, fontweight="bold")
    return save(fig, "fig2_weight.png")


# ===========================================================================
# 図 3: 片持ち階段の応力
# ===========================================================================
def fig_cantilever():
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 5.0))
    # (a) 片持ち踏板の力
    ax = axes[0]
    ax.add_patch(mpatches.Rectangle((0, 0), 0.5, 2.5, fc="#d9d9d9", ec="k",
                                     lw=1.2))  # 壁
    ax.add_patch(mpatches.Rectangle((0.5, 1.0), 2.8, 0.25, fc="#bcd2ea",
                                     ec="k", lw=1.2))  # 踏板
    for xx in np.linspace(0.8, 3.1, 6):
        ax.annotate("", xy=(xx, 1.25), xytext=(xx, 1.85),
                    arrowprops=dict(arrowstyle="-|>", color="#c00000", lw=1.2))
    ax.text(1.9, 2.1, "自重＋積載 w", fontproperties=jp, ha="center",
            fontsize=9, color="#c00000")
    ax.annotate("", xy=(0.5, 1.0), xytext=(3.3, 1.0),
                arrowprops=dict(arrowstyle="<|-|>", color="#1f4e79", lw=1))
    ax.text(1.9, 0.7, "片持ち出 Lc=1.3m", fontproperties=jp, ha="center",
            fontsize=9, color="#1f4e79")
    ax.annotate("根元（固定端）\nM=w·Lc²/2", xy=(0.55, 1.1), xytext=(1.0, -0.3),
                fontproperties=jp, fontsize=9, color="#c00000",
                arrowprops=dict(arrowstyle="->", color="#c00000"))
    ax.text(1.9, -1.0, "各踏板は壁から片持ち → 根元で最大曲げ・せん断\n"
                       "引張は上側 → 主筋は踏板の上側に配置",
            fontproperties=jp, ha="center", fontsize=9, color="#444")
    ax.set_xlim(-0.3, 3.6); ax.set_ylim(-1.5, 2.6)
    ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("(a) 片持ち踏板の応力", fontproperties=jp, fontsize=10,
                 fontweight="bold")

    # (b) モーメント図
    ax = axes[1]
    L = 1.3
    x = np.linspace(0, L, 50)
    w = 7.6
    M = w * (L - x)**2 / 2
    ax.plot(x, M, color="#c00000", lw=2)
    ax.fill_between(x, 0, M, color="#c00000", alpha=0.15)
    ax.plot([0, L], [0, 0], color="k", lw=1)
    ax.text(0.05, M.max() * 0.8, f"根元 M={M.max():.1f}\nkN·m/m",
            fontproperties=jp, fontsize=9, color="#c00000")
    ax.set_xlabel("壁からの距離 (m)", fontproperties=jp)
    ax.set_ylabel("曲げモーメント (kN·m/m)", fontproperties=jp)
    ax.set_title("(b) 片持ち踏板の M 図", fontproperties=jp, fontsize=10,
                 fontweight="bold")
    ax.grid(alpha=0.25)
    ax.invert_xaxis()
    fig.suptitle("図 3  片持ち階段の応力計算",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig3_cantilever.png")


# ===========================================================================
# 図 4: 片持ち階段支持壁の両方向
# ===========================================================================
def fig_wall():
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 5.2))
    # (a) 面外（踏板反力によるねじり・面外曲げ）
    ax = axes[0]
    ax.add_patch(mpatches.Rectangle((0, 0), 0.5, 4.0, fc="#d9d9d9", ec="k",
                                     lw=1.2))
    for i in range(7):
        yy = 0.3 + i * 0.5
        ax.add_patch(mpatches.Rectangle((0.5, yy), 1.3, 0.1, fc="#bcd2ea",
                                        ec="k", lw=0.6))
        ax.annotate("", xy=(1.1, yy), xytext=(1.1, yy + 0.4),
                    arrowprops=dict(arrowstyle="-|>", color="#c00000", lw=0.9))
    # ねじり矢印
    ax.annotate("", xy=(0.25, 1.0), xytext=(0.25, 0.5),
                arrowprops=dict(arrowstyle="-|>", color="#7a3b3b", lw=2,
                                connectionstyle="arc3,rad=0.5"))
    ax.text(-0.9, 2.0, "各踏板の根元 M が\n壁に『ねじり』＋\n『面外曲げ』として\n"
                       "累積作用",
            fontproperties=jp, fontsize=8.5, color="#7a3b3b", va="center")
    ax.text(1.1, -0.7, "面外方向（踏板と直角）", fontproperties=jp,
            ha="center", fontsize=9, color="#c00000")
    ax.set_xlim(-1.8, 2.2); ax.set_ylim(-1.2, 4.3)
    ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("(a) 面外：踏板反力によるねじり・曲げ",
                 fontproperties=jp, fontsize=10, fontweight="bold")

    # (b) 面内（地震時 壁のせん断・曲げ）
    ax = axes[1]
    ax.add_patch(mpatches.Rectangle((0, 0), 2.5, 4.0, fc="#e8c9ce", ec="k",
                                     lw=1.2, hatch="//", alpha=0.6))
    ax.annotate("", xy=(2.9, 3.8), xytext=(0.6, 3.8),
                arrowprops=dict(arrowstyle="-|>", color="#c00000", lw=2.5))
    ax.text(1.75, 4.2, "地震力（面内）", fontproperties=jp, ha="center",
            fontsize=9, color="#c00000")
    ax.text(1.25, -0.7, "面内方向（壁の長さ方向）\n耐震壁として面内せん断・曲げ",
            fontproperties=jp, ha="center", fontsize=9, color="#7a3b3b")
    ax.set_xlim(-0.5, 3.4); ax.set_ylim(-1.2, 4.7)
    ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("(b) 面内：耐震壁としての抵抗",
                 fontproperties=jp, fontsize=10, fontweight="bold")
    fig.suptitle("図 4  片持ち階段 支持壁の両方向の耐震設計",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig4_wall.png")


# ===========================================================================
# 図 5: 一方向階段（水平投影長 vs 実長）
# ===========================================================================
def fig_oneway():
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 5.0))
    # (a) 水平投影長と実長
    ax = axes[0]
    Lh, H = 4.0, 2.4
    ang = math.atan2(H, Lh)
    ax.plot([0, Lh], [0, H], color="#1f4e79", lw=6, alpha=0.5,
            solid_capstyle="round")
    ax.plot([0, Lh], [0, 0], color="#c00000", lw=1.5, ls="--")
    ax.plot([Lh, Lh], [0, H], color="gray", lw=1, ls=":")
    ax.annotate("", xy=(0, -0.4), xytext=(Lh, -0.4),
                arrowprops=dict(arrowstyle="<|-|>", color="#c00000", lw=1.2))
    ax.text(Lh / 2, -0.75, "水平投影長 Lh=4.0m（← 応力計算に使う）",
            fontproperties=jp, ha="center", fontsize=9, color="#c00000")
    ax.text(Lh / 2 - 0.3, H / 2 + 0.35, "実長 L=Lh/cosθ=4.76m",
            fontproperties=jp, fontsize=9, color="#1f4e79", rotation=31)
    ax.text(0.5, 0.12, "θ", fontproperties=jp, fontsize=12)
    # 鉛直荷重
    for xx in np.linspace(0.4, Lh - 0.4, 6):
        yy = xx * H / Lh
        ax.annotate("", xy=(xx, yy), xytext=(xx, yy + 0.6),
                    arrowprops=dict(arrowstyle="-|>", color="#7a3b3b", lw=1))
    ax.text(Lh / 2, H + 0.5, "鉛直荷重 w（水平投影あたり）", fontproperties=jp,
            ha="center", fontsize=8.5, color="#7a3b3b")
    ax.text(Lh / 2, -1.6,
            "★ 鉛直荷重による曲げは『水平投影長 Lh』で M=w·Lh²/8\n"
            "実長 L で解くと過大（誤り）。ただし自重は斜め版を水平換算",
            fontproperties=jp, ha="center", fontsize=9, color="#c00000")
    ax.set_xlim(-0.5, Lh + 0.6); ax.set_ylim(-2.1, H + 1.0)
    ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("(a) 水平投影長 と 実長", fontproperties=jp, fontsize=10,
                 fontweight="bold")

    # (b) M図
    ax = axes[1]
    x = np.linspace(0, Lh, 50)
    w = 10.44
    M = w * x * (Lh - x) / 2
    ax.plot(x, M, color="#c00000", lw=2)
    ax.fill_between(x, 0, M, color="#c00000", alpha=0.15)
    ax.plot([0, Lh], [0, 0], color="k", lw=1)
    ax.plot(0, 0, "^", color="k", ms=10)
    ax.plot(Lh, 0, "^", color="k", ms=10)
    ax.text(Lh / 2, M.max() + 1, f"中央 M=w·Lh²/8={M.max():.1f} kN·m/m",
            fontproperties=jp, ha="center", fontsize=9, color="#c00000")
    ax.set_xlabel("水平投影長 (m)", fontproperties=jp)
    ax.set_ylabel("曲げモーメント (kN·m/m)", fontproperties=jp)
    ax.set_title("(b) 一方向階段の M 図（単純支持）", fontproperties=jp,
                 fontsize=10, fontweight="bold")
    ax.grid(alpha=0.25)
    ax.set_ylim(0, M.max() + 4)
    fig.suptitle("図 5  一方向階段の応力計算（水平投影長・実長）",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig5_oneway.png")


# ===========================================================================
# 図 6: 一方向階段支持梁の耐震設計
# ===========================================================================
def fig_beam():
    fig, ax = plt.subplots(figsize=(10, 5.0))
    # 階段と踊り場、支持梁
    ax.add_patch(mpatches.Rectangle((-0.4, 1.5), 0.5, 0.6, fc="#d9d9d9",
                                     ec="k", lw=1.2))  # 上梁
    ax.add_patch(mpatches.Rectangle((4.0, 0), 0.5, 0.6, fc="#d9d9d9", ec="k",
                                     lw=1.2))  # 下梁
    ax.plot([0.1, 4.0], [2.0, 0.3], color="#1f4e79", lw=5, alpha=0.5)
    ax.text(2.0, 1.5, "階段スラブ", fontproperties=jp, fontsize=9,
            color="#1f4e79", rotation=-24)
    # 反力
    ax.annotate("", xy=(0.1, 2.1), xytext=(0.1, 2.7),
                arrowprops=dict(arrowstyle="-|>", color="#c00000", lw=2))
    ax.annotate("", xy=(4.0, 0.65), xytext=(4.0, 1.25),
                arrowprops=dict(arrowstyle="-|>", color="#c00000", lw=2))
    ax.text(-0.1, 2.9, "階段反力 R", fontproperties=jp, fontsize=9,
            color="#c00000")
    ax.text(3.5, 1.4, "階段反力 R", fontproperties=jp, fontsize=9,
            color="#c00000")
    # 短スパン梁＝短柱化に注意
    ax.annotate("階段の中間梁・受け梁は\nスパンが短く剛性大\n→ 地震時に力が集中\n"
                "（短梁・短柱化に注意）",
                xy=(4.25, 0.3), xytext=(5.0, 1.8), fontproperties=jp,
                fontsize=8.5, color="#7a3b3b",
                arrowprops=dict(arrowstyle="->", color="#7a3b3b"))
    ax.text(2.5, -0.9,
            "支持梁は 階段反力（鉛直）＋ 地震時応力 を受ける。\n"
            "踊り場レベルの中間梁は階高の途中に付き、短スパンで剛性が高い\n"
            "→ 地震力が集中しやすく、せん断設計に注意",
            fontproperties=jp, ha="center", fontsize=9, color="#444")
    ax.set_xlim(-1.0, 8.5); ax.set_ylim(-1.6, 3.2)
    ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("図 6  一方向階段 支持梁の耐震設計",
                 fontproperties=jp, fontsize=12, fontweight="bold")
    return save(fig, "fig6_beam.png")


figs = {
    "types": fig_types(),
    "weight": fig_weight(),
    "cant": fig_cantilever(),
    "wall": fig_wall(),
    "oneway": fig_oneway(),
    "beam": fig_beam(),
}
print("figures:", list(figs.keys()))

# ===========================================================================
# Excel 構築
# ===========================================================================
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

wb = Workbook()
C_TITLE = "1F4E79"; C_HEAD = "2E75B6"; C_ANS = "E2EFDA"
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
f_title = Font(name="MS PGothic", size=15, bold=True, color="FFFFFF")
f_head = Font(name="MS PGothic", size=11, bold=True, color="FFFFFF")
f_body = Font(name="MS PGothic", size=10)
f_ans = Font(name="MS PGothic", size=10, color="375623")
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)


def setup(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False


def title_row(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text)
    c.font = f_title; c.fill = PatternFill("solid", fgColor=C_TITLE)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[row].height = 30


def head(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text)
    c.font = f_head; c.fill = PatternFill("solid", fgColor=C_HEAD)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[row].height = 22


def body(ws, row, text, span=8, ans=False, h=None):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text)
    c.font = f_ans if ans else f_body
    c.alignment = wrap
    if ans:
        c.fill = PatternFill("solid", fgColor=C_ANS)
    if h:
        ws.row_dimensions[row].height = h


def table(ws, start_row, headers, rows, col1=1):
    r = start_row
    for j, htxt in enumerate(headers):
        c = ws.cell(r, col1 + j, htxt)
        c.font = Font(name="MS PGothic", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C_HEAD)
        c.alignment = center; c.border = border
    for data in rows:
        r += 1
        for j, v in enumerate(data):
            c = ws.cell(r, col1 + j, v)
            c.font = f_body
            c.alignment = center if j > 0 else wrap
            c.border = border
            if r % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F2F7FC")
    return r


def put_img(ws, path, anchor, w=None):
    img = XLImage(path)
    if w:
        ratio = w / img.width
        img.width = w; img.height = int(img.height * ratio)
    ws.add_image(img, anchor)


# ---- 目次 ----
ws = wb.active
ws.title = "目次"
setup(ws, [4, 22, 60, 14])
title_row(ws, 1, "階段の設計 問題集（RC マンション設計担当・新人向け）", span=4)
body(ws, 2,
     "目標：RC 階段の設計ができること。構造形式と特性、自重の算出、"
     "片持ち階段と一方向階段の応力・断面算定、支持壁・支持梁の耐震設計までを通す。"
     "『水平投影長さと実長の使い分け』が最大の勘所。",
     span=4, h=44)
r = table(ws, 4,
          ["No.", "シート", "到達目標", "図"],
          [["1", "1 構造形式と特性", "構造形式（片持ち・一方向等）を理解する",
            "形式3種"],
           ["2", "2 自重の算出", "階段の自重を算出できる", "斜め換算＋段形"],
           ["3", "3 片持ち応力・断面", "片持ち階段の応力・断面算定ができる",
            "片持ちM図"],
           ["4", "4 支持壁の耐震", "片持ち支持壁の両方向耐震を理解する",
            "面内・面外"],
           ["5", "5 一方向応力・断面", "水平投影長・実長を用い応力・断面算定",
            "投影長vs実長"],
           ["6", "6 支持梁の耐震", "一方向階段支持梁の耐震設計ができる",
            "反力・短梁"]])
body(ws, r + 2,
     "共通条件：蹴上 180・踏面 280（勾配 32.7°、cosθ=0.841）、γc=24、SD295。"
     "一方向：水平投影スパン Lh=4.0m、段板厚 150。片持ち：出 Lc=1.3m、踏板厚 150。"
     "積載 3.0 kN/m²（共用階段）。数値は本教材作成時に検算済み。実務は"
     "建築基準法（階段寸法・積載）と RC 規準で確認すること。",
     span=4, h=58)

# ---- 1 構造形式 ----
ws = wb.create_sheet("1 構造形式と特性")
setup(ws, [8, 18, 18, 16, 12, 12])
title_row(ws, 1, "1  階段の構造形式と特性")
head(ws, 3, "■ 図 1  階段の構造形式")
put_img(ws, figs["types"], "A4", w=820)
head(ws, 26, "■ 問題 1  形式と特性")
r = table(ws, 27,
          ["形式", "支持のしかた（記入）", "特性・適用（記入）"],
          [["一方向階段", "", ""],
           ["片持ち階段", "", ""],
           ["折り返し階段（踊り場付き）", "", ""],
           ["中間梁で支持する階段", "", ""]])
body(ws, r + 2, "選択肢（支持）：上下端の梁で両端支持／片側の壁から片持ち／"
                "踊り場で分割し中間梁で支持。", h=32)
head(ws, r + 4, "■ 問題 2  力の流れ")
body(ws, r + 5, "各形式で、階段の荷重が最終的にどこへ伝わるか答えよ"
                "（一方向＝上下の梁→柱、片持ち＝支持壁、折り返し＝踊り場梁・中間梁）。",
     h=40)
head(ws, r + 7, "■ 問題 3  マンションでの採用")
body(ws, r + 8, "RC マンションの共用階段でよく使う形式と、その理由"
                "（折り返し階段が階高を稼ぎつつ省スペース／片持ちは意匠性）を述べよ。"
                "階段室が耐震壁で囲まれる場合の利点にも触れよ。", h=44)

# ---- 2 自重 ----
ws = wb.create_sheet("2 自重の算出")
setup(ws, [8, 18, 16, 14, 14, 12, 12])
title_row(ws, 1, "2  階段の自重の算出")
head(ws, 3, "■ 図 2  斜めスラブ換算＋段形")
put_img(ws, figs["weight"], "A4", w=720)
head(ws, 28, "■ 問題 1  勾配と cosθ")
body(ws, 29, "蹴上 180・踏面 280 の階段の勾配角 θ と cosθ を求めよ"
             "（θ=atan(180/280)）。実長は水平投影長の何倍か（1/cosθ）。", h=32)
head(ws, 31, "■ 問題 2  自重の算定（水平投影 1m² あたり）")
body(ws, 32, "段板（斜めスラブ）厚 t=150。次の各項を計算し合計せよ（γc=24）。",
     h=22)
r = table(ws, 34,
          ["項目", "式", "値（記入）"],
          [["① 斜めスラブ", "γ·t/cosθ = 24×0.15/0.841", ""],
           ["② 段形三角", "γ·(蹴上/2) = 24×0.09", ""],
           ["③ 仕上げ", "（例）1.0", ""],
           ["固定荷重 小計 (①+②+③)", "", ""],
           ["④ 積載（共用階段）", "3.0", ""],
           ["全荷重 w", "固定＋積載", ""]])
head(ws, r + 2, "■ 問題 3  なぜ斜めスラブを割増すか")
body(ws, r + 3, "斜めスラブの自重を『水平投影 1m² あたり』で見ると、なぜ γ·t ではなく"
                "γ·t/cosθ になるのか（水平 1m の区間に、実長 1/cosθ 分の斜め版がある）"
                "を図で説明せよ。", h=40)
head(ws, r + 5, "■ 問題 4  段形の重さ")
body(ws, r + 6, "段形（ステップの三角形）の平均厚が『蹴上/2』になる理由を述べよ"
                "（三角形の平均高さ）。仕上げが石張り等で重い場合の影響にも触れよ。",
     h=40)

# ---- 3 片持ち応力・断面 ----
ws = wb.create_sheet("3 片持ち応力・断面")
setup(ws, [8, 18, 16, 14, 14, 12, 12])
title_row(ws, 1, "3  片持ち階段の応力計算・断面算定")
head(ws, 3, "■ 図 3  片持ち踏板の応力")
put_img(ws, figs["cant"], "A4", w=800)
head(ws, 28, "■ 問題 1  片持ち踏板の応力")
body(ws, 29, "踏板厚 150、出 Lc=1.3m。w=自重(0.15×24)＋仕上 1.0＋積載 3.0。"
             "根元の曲げ M=w·Lc²/2 と せん断 Q=w·Lc を求めよ。", h=32)
r = table(ws, 32,
          ["項目", "式", "値（記入）"],
          [["w", "0.15×24＋1.0＋3.0", ""],
           ["根元 M", "w·Lc²/2 = w×1.3²/2", ""],
           ["根元 Q", "w·Lc = w×1.3", ""]])
head(ws, r + 2, "■ 問題 2  断面算定（必要鉄筋）")
body(ws, r + 3, "根元 M に対する必要鉄筋 As=M/(ft·j)。d=150−30=120、ft=195、"
                "j=(7/8)d。As を求め配筋を選べ（引張は上側！）。"
                "参考：D13@150=847、D16@150=1327 mm²/m。", h=40)
head(ws, r + 5, "■ 問題 3  片持ちの注意点")
body(ws, r + 6, "(1) 片持ち踏板の主筋を『上側』に配する理由（引張側）。"
                "(2) 主筋の定着（壁の中へののみ込み）が特に重要な理由"
                "（片持ちは定着切れ＝即崩壊）。(3) たわみ・振動の検討"
                "（片持ちは変形が大きく、人の昇降で揺れやすい）。", h=58)
head(ws, r + 8, "■ 問題 4  段板一体か 1 枚ごとか")
body(ws, r + 9, "片持ち階段には (a) 踏板 1 枚ずつ独立して壁から片持ちする形式と "
                "(b) 段板スラブが一体で片持ちする形式がある。"
                "力の伝わり方の違いを述べよ。", h=40)

# ---- 4 支持壁の耐震 ----
ws = wb.create_sheet("4 支持壁の耐震")
setup(ws, [8, 18, 18, 16, 12, 12])
title_row(ws, 1, "4  片持ち階段 支持壁の耐震設計（両方向）")
head(ws, 3, "■ 図 4  支持壁の面外・面内")
put_img(ws, figs["wall"], "A4", w=820)
head(ws, 30, "■ 問題 1  両方向の意味")
body(ws, 31, "片持ち階段の支持壁は 2 方向で検討する。それぞれ何を受けるか答えよ。",
     h=20)
r = table(ws, 33,
          ["方向", "作用（記入）", "検討内容（記入）"],
          [["面外（踏板と直角）", "", ""],
           ["面内（壁の長さ方向）", "", ""]])
head(ws, r + 2, "■ 問題 2  面外の検討")
body(ws, r + 3, "各踏板の根元 M（問題 3）が、壁に『ねじり』＋『面外曲げ』として"
                "高さ方向に累積する。壁厚が薄いと面外で不利になる理由と、"
                "対策（壁厚確保・壁端部の柱型）を述べよ。", h=44)
head(ws, r + 5, "■ 問題 3  面内の検討")
body(ws, r + 4 + 2, "支持壁は地震時に耐震壁として面内せん断・曲げも負担する。"
                "面内・面外の両方を同時に満たす必要があることを、"
                "片持ち階段特有の設計の難しさとして述べよ。", h=44)
head(ws, r + 8, "■ 問題 4  実務上の配慮")
body(ws, r + 9, "片持ち階段の支持壁で実務上配慮すべき点を 3 つ挙げよ"
                "（壁厚と配筋／踏板と壁の接合部（定着・ハンチ）／"
                "壁の開口（扉等）による欠損の影響）。", h=44)

# ---- 5 一方向応力・断面 ----
ws = wb.create_sheet("5 一方向応力・断面")
setup(ws, [8, 18, 16, 14, 14, 12, 12])
title_row(ws, 1, "5  一方向階段の応力計算（水平投影長・実長）・断面算定")
head(ws, 3, "■ 図 5  水平投影長と実長・M 図")
put_img(ws, figs["oneway"], "A4", w=820)
head(ws, 30, "■ 問題 1  ★水平投影長で解く（最重要）")
body(ws, 31, "一方向階段（両端の梁で単純支持）。水平投影スパン Lh=4.0m、"
             "全荷重 w=10.44 kN/m²（問題 2 の値）。"
             "鉛直荷重による曲げは『水平投影長』で計算する。"
             "M=w·Lh²/8、Q=w·Lh/2 を求めよ。", h=44)
r = table(ws, 34,
          ["項目", "式", "値（記入）"],
          [["中央 M", "w·Lh²/8 = 10.44×4²/8", ""],
           ["端部 Q", "w·Lh/2 = 10.44×4/2", ""]])
head(ws, r + 2, "■ 問題 2  なぜ実長で解くと誤りか")
body(ws, r + 3, "実長 L=Lh/cosθ=4.76m で M=w·L²/8 と解くと過大になる。"
                "その理由を、鉛直荷重は水平投影に対して分布しているという観点で"
                "説明せよ（水平投影長が正しい。ただし自重の算定は斜め版を"
                "水平換算済み）。", h=44)
head(ws, r + 5, "■ 問題 3  断面算定")
body(ws, r + 6, "中央 M に対する必要鉄筋 As=M/(ft·j)。段板厚 150、d=120、"
                "ft=195、j=(7/8)d。As を求め配筋を選べ"
                "（参考：D16@200=995、D16@150=1327、D13@100=1270）。", h=40)
head(ws, r + 8, "■ 問題 4  実長を使う場面")
body(ws, r + 9, "『水平投影長で解く』のが原則だが、実長（斜め長）を使うのは"
                "どんな場面か（配筋の実寸法・鉄筋の定尺・仕上げ数量など）。"
                "応力と数量で使い分けることを整理せよ。", h=44)

# ---- 6 支持梁の耐震 ----
ws = wb.create_sheet("6 支持梁の耐震")
setup(ws, [8, 18, 18, 16, 12, 12])
title_row(ws, 1, "6  一方向階段 支持梁の耐震設計")
head(ws, 3, "■ 図 6  支持梁の反力と短梁の注意")
put_img(ws, figs["beam"], "A4", w=720)
head(ws, 28, "■ 問題 1  支持梁の反力")
body(ws, 29, "一方向階段の両端の支持梁が受ける階段反力 R を求めよ"
             "（R=w·Lh/2 ×階段幅）。階段幅 1.2m のとき、梁 1m あたりの反力に"
             "換算する考え方も述べよ。", h=40)
head(ws, 31, "■ 問題 2  踊り場レベルの中間梁")
body(ws, 32, "折り返し階段では、踊り場レベルに中間梁が付く。この梁は階高の途中"
             "（床レベルでない位置）に取り付くため、地震時に問題になる。"
             "何が起こるか説明せよ（短スパン・剛性大 → 力が集中、短梁のせん断破壊）。",
     h=44)
head(ws, 34, "■ 問題 3  短柱・短梁化の回避")
body(ws, 35, "階段室まわりの柱・梁が短柱・短梁になりやすい理由と、"
             "その対策を述べよ（耐震スリットで縁を切る／せん断補強を密にする／"
             "階段を独立構造にする）。No.1-4（耐震スリット）教材とも関連。", h=44)
head(ws, 37, "■ 問題 4  階段室の耐震的役割")
body(ws, 38, "階段室・EV シャフトは壁で囲まれ剛性が高い。"
             "(1) 建物全体の耐震要素として有利な点。"
             "(2) 一方で偏心（平面の隅に寄ると剛心がずれる）の注意。"
             "No.10（バランス配置）教材とも関連づけて述べよ。", h=44)

# ---- 解答 ----
ws = wb.create_sheet("解答")
setup(ws, [4, 22, 60, 14])
title_row(ws, 1, "解答・解説", span=4)
row = 2


def ah(t):
    global row
    head(ws, row, t, span=4); row += 1


def an(t, h=None):
    global row
    body(ws, row, t, span=4, ans=True, h=h); row += 1


ah("1  構造形式と特性")
an("問1：一方向＝上下端の梁で両端支持／大スパンは段板厚増。"
   "片持ち＝片側の壁から片持ち／意匠性高いが変形・定着に注意。"
   "折り返し＝踊り場で分割・中間梁で支持／省スペースで階高を稼ぐ。"
   "中間梁支持＝踊り場梁で受ける／スパンを分割して段板を薄く。", h=58)
an("問2：一方向＝段板→上下の梁→柱。片持ち＝踏板→支持壁→基礎。"
   "折り返し＝段板→踊り場梁・中間梁→柱。荷重の『行き先』を必ず追う。", h=44)
an("問3：マンション共用階段は折り返し階段が主流（省スペースで階高を消化）。"
   "階段室を耐震壁で囲むと、階段が建物の耐震要素（コア壁）として働き有利。"
   "片持ちは意匠性が高いが変形・支持壁の負担が大きく採用は限定的。", h=44)

ah("2  自重の算出")
an("問1：θ=atan(180/280)=32.7°。cosθ=280/√(280²+180²)=0.841。"
   "実長/水平長=1/cosθ=1.189（約 1.19 倍）。", h=32)
an("問2：①斜めスラブ=24×0.15/0.841=4.28。②段形三角=24×0.09=2.16。"
   "③仕上=1.0。固定小計=7.44。④積載=3.0。全荷重 w=10.44 kN/m²。", h=44)
an("問3：水平 1m の区間には、勾配のため実長 1/cosθ 分（=1.19m）の斜め版が"
   "存在する。だから水平投影 1m² あたりの重さは γ·t を 1/cosθ 倍した"
   "γ·t/cosθ になる。", h=44)
an("問4：段形はステップごとの直角三角形（底辺=踏面、高さ=蹴上）で、"
   "水平面に均すと平均高さ＝蹴上/2 になる。石張り等の重い仕上げは"
   "③の仕上げ荷重を大きくし、全荷重・断面に効く。", h=44)

ah("3  片持ち応力・断面")
an("問1：w=0.15×24＋1.0＋3.0=3.6＋4.0=7.6 kN/m²。"
   "根元 M=7.6×1.3²/2=6.42 kN·m/m。根元 Q=7.6×1.3=9.88 kN/m。", h=32)
an("問2：d=120、j=(7/8)×120=105。As=6.42×10⁶/(195×105)=314 mm²/m。"
   "→ D13@150（847）で十分（最小鉄筋・かぶりに注意）。引張は上側なので"
   "上端に主筋を配置。", h=44)
an("問3：(1)片持ちは固定端上側が引張→主筋は上側。"
   "(2)片持ちは定着が切れると即座に落下＝崩壊するため、壁内への定着"
   "（のみ込み長さ）を確実に確保する。(3)片持ちは変形が大きく、"
   "人の昇降で振動しやすいのでたわみ・固有振動数も確認。", h=58)
an("問4：(a)踏板 1 枚ずつ独立片持ち＝各踏板の根元 M が個別に壁へ。"
   "(b)段板一体片持ち＝スラブ全体が片持ち版として働き、幅方向にも応力が回る。"
   "(b)の方が一体性が高く変形も小さいが、壁との接合が長く必要。", h=44)

ah("4  支持壁の耐震")
an("問1：面外（踏板と直角）＝各踏板の反力による壁の面外曲げ＋ねじり。"
   "面内（壁の長さ方向）＝地震時の耐震壁としてのせん断・曲げ。", h=32)
an("問2：踏板の根元 M が高さ方向に累積し、壁にねじり＋面外曲げを与える。"
   "壁が薄いと面外曲げ・ねじりに対する剛性・耐力が不足する。"
   "対策＝壁厚の確保、壁端部に柱型（そで壁付き柱）を設けて面外剛性を上げる。",
   h=44)
an("問3：支持壁は面内（耐震壁としての地震力）と面外（踏板反力）を"
   "同時に満たす必要がある。面内配筋（縦横壁筋）に加え、面外曲げ・ねじり用の"
   "補強が要り、両立が片持ち階段特有の難しさ。", h=44)
an("問4：①壁厚と縦横壁筋＋面外補強 ②踏板と壁の接合部の定着・ハンチ"
   "（応力集中部）③扉等の開口による壁の欠損・応力集中の影響を確認。", h=44)

ah("5  一方向応力・断面")
an("問1：中央 M=w·Lh²/8=10.44×4²/8=10.44×2=20.88 kN·m/m。"
   "端部 Q=w·Lh/2=10.44×4/2=20.88 kN/m。", h=32)
an("問2：鉛直荷重 w は水平投影面に対して分布しているので、曲げは水平投影長 Lh で"
   "評価する（M=w·Lh²/8）。実長 L=4.76m で解くと (4.76/4)²=1.42 倍過大になる。"
   "『鉛直荷重×水平投影長』が傾斜部材の定石。ただし自重の算定時は"
   "斜め版を水平換算（t/cosθ）済みなので二重補正しないこと。", h=58)
an("問3：As=20.88×10⁶/(195×105)=1,020 mm²/m。"
   "→ D16@150（1,327）または D13@100（1,270）。D16@200（995）は僅かに不足。", h=44)
an("問4：応力計算は水平投影長。実長（斜め長）は、鉄筋の実寸法・定尺・"
   "コンクリート/仕上げの数量（面積）算出に使う。『応力＝水平投影、"
   "数量＝実長』で使い分ける。", h=44)

ah("6  支持梁の耐震")
an("問1：R=w·Lh/2=10.44×4/2=20.88 kN/m（階段幅方向 1m あたり）。"
   "階段幅 1.2m なら 1 端の全反力=20.88×1.2=25.1 kN、これを支持梁の"
   "スパンで分布荷重（または集中）として梁設計に加える。", h=44)
an("問2：踊り場中間梁は階高の途中（床レベルでない）に付き、上下の柱を"
   "短く分割する。短スパンで剛性が高いため地震力が集中し、"
   "短梁・短柱のせん断破壊（脆性）が起こりやすい。", h=44)
an("問3：階段室まわりは腰壁・踊り場梁で柱・梁が短く拘束され短柱・短梁化する。"
   "対策＝耐震スリットで雑壁の縁を切る／せん断補強筋を密にする／"
   "階段を本体と分離した独立構造（エキスパンション）にする。"
   "（No.1-4 耐震スリット教材参照）", h=58)
an("問4：(1)階段室・EV シャフトは壁で囲まれ剛性が高く、コア耐震壁として"
   "建物の水平力に有効。(2)ただし平面の隅に寄ると剛心が偏り、"
   "ねじれ振動を起こす（偏心率悪化）。配置バランスに注意"
   "（No.10 バランス配置教材参照）。", h=58)

XLSX = os.path.join(OUT, "階段設計問題集.xlsx")
wb.save(XLSX)
print("saved:", XLSX)
