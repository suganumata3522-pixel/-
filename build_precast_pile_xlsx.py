# -*- coding: utf-8 -*-
"""既成杭の計算書内容確認 問題集（図つき）Excel。出力: docs/precast_pile/既成杭の計算書確認問題集.xlsx
1 既成杭の種類 / 2 各杭工法の特徴 / 3 杭計算結果の確認と基礎設計への反映
4 杭頭補強筋と基礎梁配筋の干渉確認
※ 各社の杭仕様・認定支持力・根固め諸元はメーカー認定・カタログで確認要
"""
import os
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib import font_manager as fm
FONT = "/usr/share/fonts/truetype/fonts-japanese-gothic.ttf"
jp = fm.FontProperties(fname=FONT); fm.fontManager.addfont(FONT)
plt.rcParams["font.family"] = jp.get_name(); plt.rcParams["axes.unicode_minus"] = False
OUT = "docs/precast_pile"; FIG = os.path.join(OUT, "figures"); os.makedirs(FIG, exist_ok=True)


def save(fig, n):
    p = os.path.join(FIG, n); fig.savefig(p, dpi=130, bbox_inches="tight"); plt.close(fig); return p


def fig_types():
    fig, axes = plt.subplots(1, 4, figsize=(14.0, 4.2))
    # PHC杭
    ax = axes[0]
    ax.add_patch(plt.Circle((0.5, 0.55), 0.38, fc="#d9d9d9", ec="k", lw=1.5))
    ax.add_patch(plt.Circle((0.5, 0.55), 0.20, fc="white", ec="k", lw=1))
    for a in np.linspace(0, 2 * np.pi, 10, endpoint=False):
        ax.add_patch(plt.Circle((0.5 + 0.29 * np.cos(a), 0.55 + 0.29 * np.sin(a)), 0.025, fc="k"))
    ax.text(0.5, 0.08, "遠心プレストレスト\nコンクリート・中空", ha="center", va="top",
            fontproperties=jp, fontsize=8)
    ax.set_xlim(0, 1); ax.set_ylim(-0.4, 1.0); ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("(a) PHC杭", fontproperties=jp, fontsize=9.5, fontweight="bold")
    # SC杭
    ax = axes[1]
    ax.add_patch(plt.Circle((0.5, 0.55), 0.38, fc="#8a8a8a", ec="k", lw=1.5))
    ax.add_patch(plt.Circle((0.5, 0.55), 0.30, fc="#d9d9d9", ec="k", lw=1))
    ax.add_patch(plt.Circle((0.5, 0.55), 0.15, fc="white", ec="k", lw=1))
    ax.text(0.5, 0.08, "外周鋼管+コンクリート\n杭頭部(曲げ・せん断大)", ha="center", va="top",
            fontproperties=jp, fontsize=8)
    ax.set_xlim(0, 1); ax.set_ylim(-0.4, 1.0); ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("(b) SC杭", fontproperties=jp, fontsize=9.5, fontweight="bold")
    # 鋼管杭
    ax = axes[2]
    ax.add_patch(plt.Circle((0.5, 0.55), 0.38, fc="#8a8a8a", ec="k", lw=1.5))
    ax.add_patch(plt.Circle((0.5, 0.55), 0.32, fc="white", ec="k", lw=1))
    ax.text(0.5, 0.08, "鋼管・じん性大\n継手溶接・回転圧入", ha="center", va="top",
            fontproperties=jp, fontsize=8)
    ax.set_xlim(0, 1); ax.set_ylim(-0.4, 1.0); ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("(c) 鋼管杭", fontproperties=jp, fontsize=9.5, fontweight="bold")
    # 節杭
    ax = axes[3]
    ax.add_patch(mpatches.Rectangle((0.35, 0.15), 0.3, 0.7, fc="#d9d9d9", ec="k", lw=1.2))
    for yy in [0.3, 0.5, 0.7]:
        ax.add_patch(mpatches.Ellipse((0.5, yy), 0.5, 0.08, fc="#c8c8c8", ec="k", lw=1))
    ax.text(0.5, 0.08, "周面に節→周面摩擦大\n埋込み工法向き", ha="center", va="top",
            fontproperties=jp, fontsize=8)
    ax.set_xlim(0, 1); ax.set_ylim(-0.4, 1.0); ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("(d) 節杭", fontproperties=jp, fontsize=9.5, fontweight="bold")
    fig.suptitle("図 1  既成杭の種類（PHC・SC・鋼管・節杭）", fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig1_types.png")


def fig_methods():
    fig, axes = plt.subplots(1, 3, figsize=(14.0, 4.6))
    # 打込み杭
    ax = axes[0]
    ax.add_patch(mpatches.Rectangle((0, 0), 6, 4, fc="#f0e6d2", ec="none"))
    ax.add_patch(mpatches.Rectangle((2.5, 0.5), 1.0, 4.0, fc="#c8c8c8", ec="k", lw=1.2))
    ax.add_patch(mpatches.Rectangle((2.3, 4.5), 1.4, 0.8, fc="#8a8a8a", ec="k"))  # ハンマー
    ax.annotate("", xy=(3.0, 4.5), xytext=(3.0, 5.6), arrowprops=dict(arrowstyle="-|>", color="#c00000", lw=2.5))
    ax.text(3.0, -0.6, "打込み杭（打撃・圧入）\n支持力大／騒音・振動大", ha="center",
            fontproperties=jp, fontsize=8.5, color="#333")
    ax.set_xlim(-0.3, 6.3); ax.set_ylim(-1.2, 6.0); ax.axis("off")
    ax.set_title("(a) 打込み工法", fontproperties=jp, fontsize=9.5, fontweight="bold")
    # プレボーリング根固め
    ax = axes[1]
    ax.add_patch(mpatches.Rectangle((0, 0), 6, 4, fc="#f0e6d2", ec="none"))
    ax.add_patch(mpatches.Rectangle((2.5, 0.9), 1.0, 3.6, fc="#c8c8c8", ec="k", lw=1.2))
    ax.add_patch(plt.Circle((3.0, 0.7), 0.8, fc="#b0a080", ec="k", lw=1))  # 根固め球根
    ax.text(3.0, 0.7, "根固め\n球根", ha="center", va="center", fontproperties=jp, fontsize=7)
    ax.text(3.0, -0.6, "プレボーリング拡大根固め\nセメントミルクで根固め／低振動", ha="center",
            fontproperties=jp, fontsize=8.5, color="#333")
    ax.set_xlim(-0.3, 6.3); ax.set_ylim(-1.2, 6.0); ax.axis("off")
    ax.set_title("(b) 埋込み（プレボーリング）", fontproperties=jp, fontsize=9.5, fontweight="bold")
    # 中掘り・回転圧入
    ax = axes[2]
    ax.add_patch(mpatches.Rectangle((0, 0), 6, 4, fc="#f0e6d2", ec="none"))
    ax.add_patch(mpatches.Rectangle((2.5, 0.5), 1.0, 4.0, fc="#8a8a8a", ec="k", lw=1.2))
    ax.plot([2.5, 2.0], [0.5, 0.1], color="k", lw=2)
    ax.plot([3.5, 4.0], [0.5, 0.1], color="k", lw=2)  # 先端翼
    ax.annotate("回転", xy=(4.2, 2.5), xytext=(4.6, 3.2), fontproperties=jp, fontsize=8,
                arrowprops=dict(arrowstyle="->", connectionstyle="arc3,rad=0.4"))
    ax.text(3.0, -0.6, "中掘り／回転圧入\n低騒音・低振動・排土少", ha="center",
            fontproperties=jp, fontsize=8.5, color="#333")
    ax.set_xlim(-0.3, 6.3); ax.set_ylim(-1.2, 6.0); ax.axis("off")
    ax.set_title("(c) 中掘り・回転圧入", fontproperties=jp, fontsize=9.5, fontweight="bold")
    fig.suptitle("図 2  各杭工法の特徴（打込み・埋込み・回転圧入）",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig2_methods.png")


def fig_reflect():
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 5.0), gridspec_kw={"width_ratios": [1.0, 1.1]})
    # (a) 群杭の反力分配
    ax = axes[0]
    pitch = 2.5
    for (xx, yy) in [(-pitch / 2, -pitch / 2), (pitch / 2, -pitch / 2),
                     (-pitch / 2, pitch / 2), (pitch / 2, pitch / 2)]:
        ax.add_patch(plt.Circle((xx, yy), 0.35, fc="#c8c8c8", ec="k", lw=1.2))
    ax.add_patch(mpatches.Rectangle((-pitch / 2 - 0.7, -pitch / 2 - 0.7), pitch + 1.4, pitch + 1.4,
                 fill=False, ec="#1f4e79", lw=1.2, ls="--"))
    ax.annotate("", xy=(-0.9, 1.9), xytext=(0.9, 1.9), arrowprops=dict(arrowstyle="<->", color="#c00000"))
    ax.text(0, 2.1, "曲げ M", ha="center", fontproperties=jp, fontsize=9, color="#c00000")
    ax.text(pitch / 2, -pitch / 2 - 0.65, "Rmax=1300", ha="center", fontproperties=jp, fontsize=8, color="#c00000")
    ax.text(-pitch / 2, -pitch / 2 - 0.65, "Rmin=700", ha="center", fontproperties=jp, fontsize=8, color="#548235")
    ax.text(0, -2.9, "杭反力 Ri = N/n ± M・xi/Σxi²\n例 N=4000, M=1500 → 1000±300",
            ha="center", va="top", fontproperties=jp, fontsize=8.5, color="#1f4e79")
    ax.set_xlim(-2.6, 2.6); ax.set_ylim(-3.6, 2.6); ax.set_aspect("equal"); ax.axis("off")
    ax.set_title("(a) 群杭の反力分配", fontproperties=jp, fontsize=9.5, fontweight="bold")
    # (b) 反映チェックリスト
    ax = axes[1]
    ax.text(0.5, 0.96, "杭計算結果→基礎設計への反映", transform=ax.transAxes, ha="center",
            fontproperties=jp, fontsize=10.5, fontweight="bold", color="#1f4e79")
    ax.text(0.03, 0.85,
            "□ 杭反力（鉛直・水平・曲げ）→基礎梁・フーチング\n"
            "□ 杭頭曲げ・せん断→杭頭接合部・基礎梁の設計\n"
            "□ 杭頭変位→上部構造の変形との整合\n"
            "□ 群杭の偏心・最大反力（Rmax）→許容支持力照査\n"
            "□ 引抜き力（Rmin<0）→引抜き耐力・杭頭定着\n"
            "□ 液状化時の低減→水平・鉛直の再検討\n"
            "□ 認定支持力・杭仕様（各社）→杭種・本数の妥当性\n"
            "□ 施工性（近接・低空頭・排土）→工法選定\n\n"
            "杭業者と協議：支持層深さ・傾斜、杭長、支持力、\n"
            "施工制約、試験杭、杭頭処理、コスト",
            transform=ax.transAxes, fontproperties=jp, fontsize=8.6, va="top")
    ax.axis("off")
    ax.set_title("(b) 反映チェックと業者協議", fontproperties=jp, fontsize=9.5, fontweight="bold")
    fig.suptitle("図 3  杭計算結果の確認と基礎設計への反映",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig3_reflect.png")


def fig_interfere():
    fig, ax = plt.subplots(figsize=(11.5, 5.6))
    # 杭
    ax.add_patch(mpatches.Rectangle((3.5, 0), 2.0, 3.0, fc="#c8c8c8", ec="k", lw=1.5))
    ax.text(4.5, 1.4, "既成杭", ha="center", fontproperties=jp, fontsize=9)
    # フーチング/基礎梁
    ax.add_patch(mpatches.Rectangle((1.0, 3.0), 7.0, 1.8, fc="#e8e0d0", ec="k", lw=1.2))
    ax.text(1.4, 4.5, "基礎梁／フーチング", fontproperties=jp, fontsize=9, color="#7a5a2a")
    # 杭頭補強筋（差し筋）
    for xx in [3.9, 4.3, 4.7, 5.1]:
        ax.plot([xx, xx], [2.7, 4.6], color="#c00000", lw=2)
        ax.plot([xx, xx + 0.3], [4.6, 4.6], color="#c00000", lw=2)  # フック
    ax.annotate("杭頭補強筋（差し筋）\n杭を基礎に定着", xy=(4.5, 4.0), xytext=(6.2, 5.3),
                fontproperties=jp, fontsize=9, color="#c00000",
                arrowprops=dict(arrowstyle="->", color="#c00000"))
    # 基礎梁主筋（水平）
    for yy in [3.4, 4.4]:
        ax.plot([1.2, 7.8], [yy, yy], color="#1f7a1f", lw=2.5)
    ax.annotate("基礎梁主筋（上端・下端）", xy=(2.0, 4.4), xytext=(0.3, 5.3),
                fontproperties=jp, fontsize=9, color="#1f7a1f",
                arrowprops=dict(arrowstyle="->", color="#1f7a1f"))
    # 干渉ゾーン
    ax.add_patch(mpatches.Ellipse((4.5, 4.4), 1.8, 0.7, fill=False, ec="#c00000", lw=1.8, ls="--"))
    ax.text(4.5, -0.5, "杭頭補強筋（縦）と基礎梁主筋（横）が同じ領域で干渉\n"
            "→ 定着長・かぶり・あきを確保し、通り・段を調整して納める",
            ha="center", fontproperties=jp, fontsize=9, color="#c00000")
    ax.set_xlim(0, 9); ax.set_ylim(-1.1, 5.8); ax.axis("off")
    ax.set_title("図 4  杭頭補強筋と基礎梁配筋の干渉確認",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig4_interfere.png")


figs = {"types": fig_types(), "methods": fig_methods(),
        "reflect": fig_reflect(), "interfere": fig_interfere()}
print("figs:", list(figs.keys()))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
wb = Workbook()
C_TITLE = "1F4E79"; C_HEAD = "2E75B6"; C_ANS = "E2EFDA"; C_WARN = "FCE4D6"
thin = Side(style="thin", color="BFBFBF"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
f_title = Font(name="MS PGothic", size=15, bold=True, color="FFFFFF")
f_head = Font(name="MS PGothic", size=11, bold=True, color="FFFFFF")
f_body = Font(name="MS PGothic", size=10); f_ans = Font(name="MS PGothic", size=10, color="375623")
f_warn = Font(name="MS PGothic", size=9, color="833C00", italic=True)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)


def setup(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False


def title_row(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_title; c.fill = PatternFill("solid", fgColor=C_TITLE)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1); ws.row_dimensions[row].height = 30


def head(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_head; c.fill = PatternFill("solid", fgColor=C_HEAD)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1); ws.row_dimensions[row].height = 22


def body(ws, row, text, span=8, ans=False, h=None):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_ans if ans else f_body; c.alignment = wrap
    if ans:
        c.fill = PatternFill("solid", fgColor=C_ANS)
    if h:
        ws.row_dimensions[row].height = h


def warn(ws, row, text, span=8, h=None):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_warn; c.alignment = wrap
    c.fill = PatternFill("solid", fgColor=C_WARN)
    if h:
        ws.row_dimensions[row].height = h


def tbl(ws, start_row, headers, rows, col1=1):
    r = start_row
    for j, htxt in enumerate(headers):
        c = ws.cell(r, col1 + j, htxt); c.font = Font(name="MS PGothic", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C_HEAD); c.alignment = center; c.border = border
    for data in rows:
        r += 1
        for j, v in enumerate(data):
            c = ws.cell(r, col1 + j, v); c.font = f_body
            c.alignment = wrap if j == 0 else center; c.border = border
            if r % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F2F7FC")
    return r


def put_img(ws, path, anchor, w=None):
    img = XLImage(path)
    if w:
        ratio = w / img.width; img.width = w; img.height = int(img.height * ratio)
    ws.add_image(img, anchor)


# ===== 目次 =====
ws = wb.active; ws.title = "目次"; setup(ws, [4, 24, 54, 16])
title_row(ws, 1, "既成杭の計算書内容確認 問題集（RC マンション設計担当・新人向け）", span=4)
body(ws, 2, "目標：既成杭の種類・工法の特徴を理解し、杭計算書の内容を確認して基礎設計に反映でき、"
            "杭業者と協議し、杭頭補強筋と基礎梁配筋の干渉を確認できること。", span=4, h=32)
r = tbl(ws, 4, ["No.", "シート", "到達目標", "図"],
        [["1", "1 既成杭の種類", "PHC/SC/鋼管/節杭を理解", "杭断面"],
         ["2", "2 各杭工法の特徴", "打込み/埋込み/回転圧入の特徴", "工法"],
         ["3", "3 計算結果の反映", "杭反力→基礎、抽出・協議", "反力/CL"],
         ["4", "4 杭頭補強筋の干渉", "杭頭筋と基礎梁筋の干渉確認", "干渉"]])
warn(ws, r + 2, "※ 各社の杭仕様・認定支持力・根固め諸元・継手はメーカー認定・カタログで確認。"
                "杭は大臣認定工法が多く、認定内容の範囲で使用すること。", span=4, h=40)

# ===== 1 既成杭の種類 =====
ws = wb.create_sheet("1 既成杭の種類"); setup(ws, [8, 16, 18, 16, 12, 12])
title_row(ws, 1, "1  既成杭の種類")
head(ws, 3, "■ 図 1  既成杭の種類"); put_img(ws, figs["types"], "A4", w=880)
head(ws, 26, "■ 問題 1  杭種と特徴")
r = tbl(ws, 27, ["杭種", "特徴（記入）"],
        [["PHC杭", "遠心成形プレストレストコンクリート・中空（記入）"],
         ["PRC杭", "PC鋼材＋鉄筋で曲げ・じん性を確保（記入）"],
         ["SC杭", "外周鋼管＋コンクリート・杭頭部（曲げ大）（記入）"],
         ["鋼管杭", "鋼管・じん性大・継手溶接（記入）"],
         ["節杭", "周面に節・周面摩擦大（記入）"]])
head(ws, r + 2, "■ 問題 2  杭種の使い分け")
body(ws, r + 3, "曲げ・せん断が大きい杭頭部に SC杭・PRC杭を用い、中間部に PHC杭を継ぐ"
                "『組合せ杭』の考え方を説明せよ。杭種を曲げ・軸力の大きさで使い分ける意義を述べよ。", h=40)
head(ws, r + 5, "■ 問題 3  既成杭と場所打ち杭")
body(ws, r + 6, "既成杭（工場製作・品質安定）と場所打ち杭（現場施工・大径可）の長所短所を対比せよ"
                "（品質・径・施工・残土・工期・コスト）。", h=40)

# ===== 2 各杭工法の特徴 =====
ws = wb.create_sheet("2 各杭工法の特徴"); setup(ws, [8, 16, 18, 16, 12, 12])
title_row(ws, 1, "2  各杭工法の特徴")
head(ws, 3, "■ 図 2  杭工法"); put_img(ws, figs["methods"], "A4", w=880)
head(ws, 27, "■ 問題 1  工法の分類")
r = tbl(ws, 28, ["工法", "特徴"],
        [["打込み杭", "打撃・圧入。支持力大だが騒音・振動大"],
         ["埋込み（プレボーリング）", "掘削後に杭挿入・根固め。低振動・低騒音（記入）"],
         ["中掘り杭", "杭中空部を掘削しながら沈設。（記入）"],
         ["回転圧入（鋼管）", "先端翼で回転貫入。低騒音・低振動・無排土（記入）"]])
head(ws, r + 2, "■ 問題 2  工法選定の観点")
body(ws, r + 3, "杭工法を選ぶ観点を挙げよ（支持力、騒音・振動、近接・低空頭、排土、施工深度、"
                "地盤条件、コスト・工期）。市街地マンションで配慮すべき点にも触れよ。", h=40)
head(ws, r + 4, "■ 問題 3  根固め・支持力")
body(ws, r + 5, "埋込み杭の根固め球根（セメントミルク）が先端支持力を確保する仕組みを説明せよ。"
                "認定工法では支持力が大臣認定で定まることに触れよ。", h=40)

# ===== 3 計算結果の反映 =====
ws = wb.create_sheet("3 計算結果の反映"); setup(ws, [8, 16, 18, 16, 12, 12])
title_row(ws, 1, "3  杭計算結果の確認と基礎設計への反映")
head(ws, 3, "■ 図 3  反力分配・反映チェック"); put_img(ws, figs["reflect"], "A4", w=860)
head(ws, 28, "■ 問題 1  群杭の反力分配")
body(ws, 29, "4本杭（2×2、ピッチ2.5m）で柱軸力 N=4000kN・基礎に作用する曲げ M=1500kN·m のとき、"
             "各杭反力 Ri=N/n±M·xi/Σxi² を求め、最大・最小杭反力を計算せよ（xi=±1.25m）。", h=44)
head(ws, 31, "■ 問題 2  最大反力の照査")
body(ws, 32, "問1の最大杭反力 Rmax を杭の許容支持力（各社の認定支持力）と比較する意味を述べよ。"
             "最小反力 Rmin が負なら引抜き検討が必要なことにも触れよ。", h=40)
head(ws, 34, "■ 問題 3  基礎設計への反映項目")
body(ws, 35, "杭計算書から基礎の設計に反映が必要な項目を抽出せよ（杭反力→基礎梁・フーチング、"
             "杭頭曲げ・せん断、杭頭変位、引抜き、液状化時低減、認定支持力・杭仕様）。", h=44)
head(ws, 37, "■ 問題 4  杭業者との協議")
body(ws, 38, "杭業者担当者と協議すべき内容を挙げよ（支持層深さ・傾斜、杭長、支持力、施工制約"
             "（近接・低空頭）、試験杭、杭頭処理、コスト）。協議で設計条件を確定する意義を述べよ。", h=44)

# ===== 4 杭頭補強筋の干渉 =====
ws = wb.create_sheet("4 杭頭補強筋の干渉"); setup(ws, [8, 16, 18, 16, 12, 12])
title_row(ws, 1, "4  杭頭補強筋と基礎梁配筋の干渉確認")
head(ws, 3, "■ 図 4  干渉確認"); put_img(ws, figs["interfere"], "A4", w=760)
head(ws, 30, "■ 問題 1  杭頭補強筋の役割")
body(ws, 31, "杭頭補強筋（差し筋）が杭を基礎（フーチング・基礎梁）に定着し、杭頭の曲げ・引抜きを"
             "伝達する役割を説明せよ。定着長・本数の考え方に触れよ。", h=40)
head(ws, 33, "■ 問題 2  干渉の内容")
body(ws, 34, "杭頭補強筋（縦）と基礎梁主筋（横）が同じ領域で交差・干渉することを説明せよ。"
             "干渉すると定着長・かぶり・あきが確保できず施工不能になる問題を述べよ。", h=40)
head(ws, 36, "■ 問題 3  干渉の回避")
body(ws, 37, "干渉を回避する方法を述べよ（補強筋の本数・配置・段の調整、基礎梁主筋の通り・段の調整、"
             "杭芯と柱芯・梁芯のずれの確認、納まり図での事前チェック）。", h=40)
head(ws, 39, "■ 問題 4  事前確認の重要性")
body(ws, 40, "杭頭補強筋と基礎梁配筋の干渉を『施工前に図面で確認する』ことの重要性を述べよ"
             "（現場での手戻り防止、かぶり・定着不足による耐力低下の回避）。", h=40)

# ===== 解答 =====
ws = wb.create_sheet("解答"); setup(ws, [4, 22, 60, 14])
title_row(ws, 1, "解答・解説", span=4); row = 2


def ah(t):
    global row; head(ws, row, t, span=4); row += 1


def an(t, h=None):
    global row; body(ws, row, t, span=4, ans=True, h=h); row += 1


def aw(t, h=None):
    global row; warn(ws, row, t, span=4, h=h); row += 1


ah("1  既成杭の種類")
an("問1：PHC＝遠心成形プレストレストコンクリート杭・中空（一般部）。PRC＝PC鋼材＋鉄筋で曲げ・"
   "じん性を確保。SC＝外周鋼管＋コンクリートで曲げ・せん断大（杭頭部）。鋼管杭＝じん性大・継手溶接。"
   "節杭＝周面の節で周面摩擦を増す（埋込み向き）。", h=48)
an("問2：地震時に曲げ・せん断が大きい杭頭部にSC杭・PRC杭、中間・下部に安価なPHC杭を継ぐ"
   "『組合せ杭』が一般的。応力の大きさに応じて杭種を使い分け、経済性と耐力を両立する。", h=44)
an("問3：既成杭＝工場製作で品質安定・施工速い・残土少、ただし径・長さに制限、継手・運搬制約。"
   "場所打ち杭＝大径・長尺・現場条件対応、ただし品質は施工管理依存・残土多い。条件で選ぶ。", h=44)

ah("2  各杭工法の特徴")
an("問1：打込み＝支持力大だが騒音・振動大（市街地は不利）。埋込み（プレボーリング）＝掘削後に"
   "杭挿入し根固め、低騒音・低振動。中掘り＝杭中空部を掘削しながら沈設。回転圧入＝先端翼で"
   "回転貫入、低騒音・低振動・無排土。", h=48)
an("問2：支持力、騒音・振動（近隣）、近接・低空頭、排土（残土処分）、施工深度、地盤条件、"
   "コスト・工期。市街地マンションでは低騒音・低振動・低排土の埋込み/回転圧入が選ばれやすい。", h=44)
an("問3：埋込み杭は先端でセメントミルクを注入・撹拌して根固め球根を造り、杭先端と支持層を"
   "一体化して先端支持力を発揮する。支持力は工法ごとに大臣認定で定まり、認定範囲で使う。", h=44)

ah("3  計算結果の反映")
an("問1：Σxi²=4×1.25²=6.25m²。Ri=N/n±M·xi/Σxi²=4000/4±1500×1.25/6.25=1000±300。"
   "最大杭反力 Rmax=1300kN、最小 Rmin=700kN（ともに圧縮）。", h=40)
an("問2：Rmax=1300kNを各社の認定（許容）支持力と比べ、Rmax<=許容支持力を確認する。"
   "Rminが負（引抜き）なら、周面摩擦＋杭自重による引抜き耐力と杭頭定着を別途検討する。", h=40)
an("問3：杭反力（鉛直・水平・曲げ）→基礎梁・フーチング、杭頭曲げ・せん断→接合部、杭頭変位→"
   "上部変形との整合、群杭の偏心・Rmax→支持力照査、引抜き、液状化時低減、認定支持力・杭仕様。", h=44)
an("問4：支持層深さ・傾斜、杭長、支持力（認定）、施工制約（近接・低空頭・排土）、試験杭、"
   "杭頭処理、コスト・工期。協議で設計条件（杭長・支持力・施工性）を確定し、手戻りを防ぐ。", h=44)

ah("4  杭頭補強筋の干渉")
an("問1：杭頭補強筋（差し筋）は杭頭を基礎（フーチング・基礎梁）に定着し、杭頭の曲げ・せん断・"
   "引抜き力を基礎に伝達する。所定の定着長・本数を杭頭曲げ・引抜きに応じて確保する。", h=44)
an("問2：杭頭補強筋（縦筋）と基礎梁主筋（横筋）が杭頭直上の同じ領域で交差・干渉する。"
   "干渉すると補強筋の定着長・かぶり・あきが取れず、コンクリートも充填しにくく施工不能になる。", h=44)
an("問3：補強筋の本数・配置・段を調整、基礎梁主筋の通り・段を調整、杭芯と柱芯・梁芯のずれ（偏心）を"
   "確認し、納まり図（配筋詳細図）で事前に干渉をチェックして納める。", h=44)
an("問4：現場で干渉が判明すると鉄筋の切断・移動など手戻りが生じ、かぶり・定着不足で耐力低下を招く。"
   "施工前に配筋図・納まり図で干渉を確認し、必要なら設計変更しておくことが重要。", h=44)
aw("※ 各社の杭仕様・認定支持力・杭頭補強筋の仕様はメーカー認定・規準で確認要。", h=24)

XLSX = os.path.join(OUT, "既成杭の計算書確認問題集.xlsx"); wb.save(XLSX); print("saved:", XLSX)
