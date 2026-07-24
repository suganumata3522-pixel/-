# -*- coding: utf-8 -*-
"""上下分離モデル 問題集（図つき）Excel。出力: docs/updown_model/上下分離モデル問題集.xlsx
1 分離モデルと一体モデル / 2 仮定条件と基礎形状別の支点条件
3 上部・下部の検討内容と応力伝達 / 4 支点条件の影響と接地圧計算例
"""
import os
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib import font_manager as fm
FONT = "/usr/share/fonts/truetype/fonts-japanese-gothic.ttf"
jp = fm.FontProperties(fname=FONT); fm.fontManager.addfont(FONT)
plt.rcParams["font.family"] = jp.get_name(); plt.rcParams["axes.unicode_minus"] = False
OUT = "docs/updown_model"; FIG = os.path.join(OUT, "figures"); os.makedirs(FIG, exist_ok=True)


def save(fig, n):
    p = os.path.join(FIG, n); fig.savefig(p, dpi=130, bbox_inches="tight"); plt.close(fig); return p


def frame(ax, x0, y0, w=3.0, storyh=1.2, stories=3, bays=2, color="#2e5b8a"):
    bw = w / bays
    for s in range(stories + 1):
        ax.plot([x0, x0 + w], [y0 + s * storyh, y0 + s * storyh], color=color, lw=2)
    for b in range(bays + 1):
        ax.plot([x0 + b * bw, x0 + b * bw], [y0, y0 + stories * storyh], color=color, lw=2)


def springs(ax, x0, y0, n, w, depth=0.5):
    xs = np.linspace(x0, x0 + w, n)
    for x in xs:
        zz = np.linspace(0, depth, 20)
        xx = x + 0.06 * np.sin(zz / depth * 4 * np.pi)
        ax.plot(xx, y0 - zz, color="#7a3b00", lw=0.9)


def fig_models():
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 5.6))
    # (a) 一体モデル
    ax = axes[0]
    frame(ax, 0.5, 1.2, w=3.0, stories=3)
    ax.add_patch(mpatches.Rectangle((0.2, 0.7), 3.6, 0.5, fc="#d9d9d9", ec="k"))  # 基礎
    springs(ax, 0.4, 0.7, 8, 3.2, depth=0.5)
    ax.axhline(0.2, color="#8B5A2B", lw=1, ls=":")
    ax.text(2.0, 0.95, "基礎", ha="center", fontproperties=jp, fontsize=8)
    ax.text(2.0, -0.05, "地盤ばね", ha="center", fontproperties=jp, fontsize=8, color="#7a3b00")
    ax.text(2.0, 5.3, "上部・基礎・地盤ばねを\n一つのモデルで解析\n（相互作用を考慮）",
            ha="center", fontproperties=jp, fontsize=9, color="#1f4e79")
    ax.set_xlim(-0.3, 4.3); ax.set_ylim(-0.6, 6.0); ax.axis("off")
    ax.set_title("(a) 一体モデル", fontproperties=jp, fontsize=11, fontweight="bold")
    # (b) 分離モデル
    ax = axes[1]
    frame(ax, 0.5, 2.6, w=3.0, stories=3)
    # 支点（ピン/固定）
    for bx in [0.5, 2.0, 3.5]:
        ax.add_patch(mpatches.Polygon([[bx, 2.6], [bx - 0.18, 2.3], [bx + 0.18, 2.3]],
                     closed=True, fc="none", ec="k", lw=1.2))
    ax.text(2.0, 6.7, "上部構造（支点で切離し）", ha="center", fontproperties=jp, fontsize=9, color="#1f4e79")
    # 反力矢印
    for bx in [0.5, 2.0, 3.5]:
        ax.annotate("", xy=(bx, 1.5), xytext=(bx, 2.2),
                    arrowprops=dict(arrowstyle="-|>", color="#c00000", lw=1.6))
    ax.text(2.0, 1.75, "柱脚反力 N・Q・M を下部へ", ha="center", fontproperties=jp,
            fontsize=8.5, color="#c00000")
    # 下部モデル
    ax.add_patch(mpatches.Rectangle((0.2, 0.9), 3.6, 0.5, fc="#d9d9d9", ec="k"))
    springs(ax, 0.4, 0.9, 8, 3.2, depth=0.5)
    ax.axhline(0.4, color="#8B5A2B", lw=1, ls=":")
    ax.text(2.0, 0.2, "下部（基礎・地盤ばね）", ha="center", fontproperties=jp, fontsize=8.5, color="#7a3b00")
    ax.set_xlim(-0.3, 4.3); ax.set_ylim(-0.3, 7.2); ax.axis("off")
    ax.set_title("(b) 上下分離モデル", fontproperties=jp, fontsize=11, fontweight="bold")
    fig.suptitle("図 1  一体モデルと上下分離モデル", fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig1_models.png")


def fig_support():
    fig, axes = plt.subplots(1, 3, figsize=(14.0, 4.4))
    # (a) 独立フーチング
    ax = axes[0]
    ax.plot([1.5, 1.5], [1.2, 3.2], color="#2e5b8a", lw=3)
    ax.add_patch(mpatches.Rectangle((0.5, 0.6), 2.0, 0.6, fc="#d9d9d9", ec="k"))
    ax.add_patch(mpatches.Polygon([[1.5, 0.6], [1.3, 0.3], [1.7, 0.3]], fc="none", ec="k", lw=1.2))
    ax.axhline(0.3, color="#8B5A2B", lw=1, ls=":")
    ax.text(1.5, -0.15, "支点：ピン〜固定\n（フーチング剛性・根入れによる）", ha="center",
            fontproperties=jp, fontsize=8, color="#333")
    ax.set_xlim(0, 3); ax.set_ylim(-0.7, 3.4); ax.axis("off")
    ax.set_title("(a) 独立フーチング", fontproperties=jp, fontsize=10, fontweight="bold")
    # (b) べた基礎
    ax = axes[1]
    for cx in [0.9, 1.5, 2.1]:
        ax.plot([cx, cx], [1.2, 3.2], color="#2e5b8a", lw=3)
    ax.add_patch(mpatches.Rectangle((0.3, 0.6), 2.4, 0.6, fc="#d9d9d9", ec="k"))
    springs(ax, 0.4, 0.6, 7, 2.2, depth=0.4)
    ax.axhline(0.2, color="#8B5A2B", lw=1, ls=":")
    ax.text(1.5, -0.25, "支点：面で支持\n（連続ばね・固定に近い）", ha="center",
            fontproperties=jp, fontsize=8, color="#333")
    ax.set_xlim(0, 3); ax.set_ylim(-0.7, 3.4); ax.axis("off")
    ax.set_title("(b) べた基礎", fontproperties=jp, fontsize=10, fontweight="bold")
    # (c) 杭基礎
    ax = axes[2]
    ax.plot([1.5, 1.5], [1.4, 3.2], color="#2e5b8a", lw=3)
    ax.add_patch(mpatches.Rectangle((0.7, 0.9), 1.6, 0.5, fc="#d9d9d9", ec="k"))  # フーチング
    for px in [1.0, 2.0]:
        ax.add_patch(mpatches.Rectangle((px - 0.1, -0.8), 0.2, 1.7, fc="#b0b0b0", ec="k"))
    ax.axhline(0.9, color="#8B5A2B", lw=0.8, ls=":")
    ax.text(1.5, -1.25, "支点：杭頭の固定度\n（剛結=固定/ピン、群杭）", ha="center",
            fontproperties=jp, fontsize=8, color="#333")
    ax.set_xlim(0, 3); ax.set_ylim(-1.7, 3.4); ax.axis("off")
    ax.set_title("(c) 杭基礎", fontproperties=jp, fontsize=10, fontweight="bold")
    fig.suptitle("図 2  基礎形状に応じた支点条件", fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig2_support.png")


def fig_transfer():
    fig, ax = plt.subplots(figsize=(11.5, 5.6))
    # 柱
    ax.plot([5, 5], [3.0, 6.0], color="#2e5b8a", lw=4)
    # N,Q,M
    ax.annotate("", xy=(5, 3.2), xytext=(5, 4.6), arrowprops=dict(arrowstyle="-|>", color="#c00000", lw=2.5))
    ax.text(5.25, 4.0, "N（軸力）", fontproperties=jp, fontsize=10, color="#c00000")
    ax.annotate("", xy=(6.2, 5.5), xytext=(5, 5.5), arrowprops=dict(arrowstyle="-|>", color="#2a78d6", lw=2))
    ax.text(6.3, 5.4, "Q（せん断）", fontproperties=jp, fontsize=10, color="#2a78d6")
    ax.annotate("", xy=(4.3, 5.0), xytext=(4.3, 5.8),
                arrowprops=dict(arrowstyle="-|>", color="#548235", lw=2, connectionstyle="arc3,rad=0.5"))
    ax.text(3.2, 5.4, "M（曲げ）", fontproperties=jp, fontsize=10, color="#548235")
    # フーチング
    ax.add_patch(mpatches.Rectangle((2.5, 2.0), 5.0, 1.0, fc="#d9d9d9", ec="k", lw=1.2))
    ax.text(5, 2.5, "フーチング（基礎）", ha="center", fontproperties=jp, fontsize=9)
    # 接地圧（台形）
    xs = np.linspace(2.5, 7.5, 10)
    for i, x in enumerate(xs):
        h = 0.3 + 1.1 * (i / (len(xs) - 1))  # 台形（右ほど大）
        ax.annotate("", xy=(x, 2.0), xytext=(x, 2.0 - h),
                    arrowprops=dict(arrowstyle="-|>", color="#7a3b00", lw=1.3))
    ax.plot([2.5, 7.5], [0.6, -0.4], color="#7a3b00", lw=1.2)
    ax.text(5, -0.9, "地盤反力（接地圧 q）：偏心で台形/三角分布", ha="center",
            fontproperties=jp, fontsize=9, color="#7a3b00")
    ax.axhline(2.0, color="#8B5A2B", lw=0.5, ls=":")
    # 流れの注記
    ax.text(8.2, 4.5, "力の流れ\n上部柱脚\n（N・Q・M）\n↓\n基礎\n↓\n地盤（反力）",
            ha="center", va="center", fontproperties=jp, fontsize=9, color="#1f4e79",
            bbox=dict(boxstyle="round", fc="#eaf1fb", ec="#2e75b6"))
    ax.text(5, 6.3, "つり合い：上部反力 = 基礎底の地盤反力の合力", ha="center",
            fontproperties=jp, fontsize=9.5, color="#c00000")
    ax.set_xlim(0.5, 10.0); ax.set_ylim(-1.3, 6.8); ax.axis("off")
    ax.set_title("図 3  応力伝達（上部柱脚 → 基礎 → 地盤）",
                 fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig3_transfer.png")


def fig_effect():
    fig, axes = plt.subplots(1, 2, figsize=(13.5, 5.0))
    # (a) 支点条件による柱脚M
    ax = axes[0]
    # ピン支点フレーム
    frame(ax, 0.3, 1.0, w=2.0, storyh=2.2, stories=1, bays=1)
    ax.add_patch(mpatches.Polygon([[0.3, 1.0], [0.15, 0.7], [0.45, 0.7]], fc="none", ec="k", lw=1.2))
    ax.add_patch(mpatches.Polygon([[2.3, 1.0], [2.15, 0.7], [2.45, 0.7]], fc="none", ec="k", lw=1.2))
    ax.text(1.3, 0.35, "柱脚ピン → 柱脚 M=0", ha="center", fontproperties=jp, fontsize=8.5, color="#c00000")
    ax.text(1.3, 3.7, "ピン支点", ha="center", fontproperties=jp, fontsize=9, fontweight="bold")
    # 固定支点フレーム
    frame(ax, 3.3, 1.0, w=2.0, storyh=2.2, stories=1, bays=1)
    for bx in [3.3, 5.3]:
        ax.plot([bx - 0.25, bx + 0.25], [1.0, 1.0], color="k", lw=2)
        for hx in np.linspace(bx - 0.22, bx + 0.18, 4):
            ax.plot([hx, hx - 0.1], [1.0, 0.8], color="k", lw=0.8)
    ax.text(4.3, 0.35, "柱脚固定 → 柱脚 M 大", ha="center", fontproperties=jp, fontsize=8.5, color="#c00000")
    ax.text(4.3, 3.7, "固定支点", ha="center", fontproperties=jp, fontsize=9, fontweight="bold")
    ax.set_xlim(-0.2, 5.8); ax.set_ylim(0, 4.1); ax.axis("off")
    ax.set_title("(a) 支点条件で柱脚モーメントが変わる", fontproperties=jp, fontsize=9.5, fontweight="bold")
    # (b) 接地圧分布（台形）
    ax = axes[1]
    B = 3.0
    qmax, qmin = 352.2, 41.1
    ax.add_patch(mpatches.Rectangle((0, 0), B, 0.4, fc="#d9d9d9", ec="k"))
    ax.plot([0, 0, B, B], [-qmin / 120, 0, 0, -qmax / 120], color="#7a3b00", lw=2)
    ax.fill_between([0, B], [0, 0], [-qmin / 120, -qmax / 120], color="#f0e0c8", alpha=0.7)
    ax.text(0.1, -qmin / 120 - 0.25, f"qmin={qmin:.0f}", fontproperties=jp, fontsize=9, color="#7a3b00")
    ax.text(B - 0.9, -qmax / 120 - 0.25, f"qmax={qmax:.0f}", fontproperties=jp, fontsize=9, color="#c00000")
    ax.text(B / 2, 0.6, "偏心 e<B/6 → 全面圧縮（台形）", ha="center", fontproperties=jp, fontsize=9)
    ax.text(B / 2, -3.5, "q = Np/A ± Mp/Z（kPa）\nNp=1770kN, Mp=700kN·m", ha="center",
            fontproperties=jp, fontsize=8.5, color="#1f4e79")
    ax.set_xlim(-0.5, 3.5); ax.set_ylim(-4.2, 1.1); ax.axis("off")
    ax.set_title("(b) 接地圧分布（応力伝達の結果）", fontproperties=jp, fontsize=9.5, fontweight="bold")
    fig.suptitle("図 4  支点条件の影響と接地圧の計算", fontproperties=jp, fontsize=13, fontweight="bold")
    return save(fig, "fig4_effect.png")


figs = {"models": fig_models(), "support": fig_support(),
        "transfer": fig_transfer(), "effect": fig_effect()}
print("figs:", list(figs.keys()))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
wb = Workbook()
C_TITLE = "1F4E79"; C_HEAD = "2E75B6"; C_ANS = "E2EFDA"
thin = Side(style="thin", color="BFBFBF"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
f_title = Font(name="MS PGothic", size=15, bold=True, color="FFFFFF")
f_head = Font(name="MS PGothic", size=11, bold=True, color="FFFFFF")
f_body = Font(name="MS PGothic", size=10); f_ans = Font(name="MS PGothic", size=10, color="375623")
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)


def setup(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False


def title_row(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_title; c.fill = PatternFill("solid", fgColor=C_TITLE)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1); ws.row_dimensions[row].height = 30


def head(ws, row, text, span=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_head; c.fill = PatternFill("solid", fgColor=C_HEAD)
    c.alignment = Alignment(vertical="center", horizontal="left", indent=1); ws.row_dimensions[row].height = 22


def body(ws, row, text, span=8, ans=False, h=None):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row, 1, text); c.font = f_ans if ans else f_body; c.alignment = wrap
    if ans:
        c.fill = PatternFill("solid", fgColor=C_ANS)
    if h:
        ws.row_dimensions[row].height = h


def tbl(ws, start_row, headers, rows, col1=1):
    r = start_row
    for j, htxt in enumerate(headers):
        c = ws.cell(r, col1 + j, htxt); c.font = Font(name="MS PGothic", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C_HEAD); c.alignment = center; c.border = border
    for data in rows:
        r += 1
        for j, v in enumerate(data):
            c = ws.cell(r, col1 + j, v); c.font = f_body
            c.alignment = wrap if j == 0 else center; c.border = border
            if r % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F2F7FC")
    return r


def put_img(ws, path, anchor, w=None):
    img = XLImage(path)
    if w:
        ratio = w / img.width; img.width = w; img.height = int(img.height * ratio)
    ws.add_image(img, anchor)


# ===== 目次 =====
ws = wb.active; ws.title = "目次"; setup(ws, [4, 24, 56, 16])
title_row(ws, 1, "上下分離モデル 問題集（RC マンション設計担当・新人向け）", span=4)
body(ws, 2, "目標：上部構造と下部（基礎）を分離して解くモデルの考え方・仮定条件・支点条件を理解し、"
            "上部柱脚力が基礎・地盤へどう伝達されるかを説明できること。", span=4, h=32)
r = tbl(ws, 4, ["No.", "シート", "到達目標", "図"],
        [["1", "1 分離と一体モデル", "分離/一体の違いと使い分けを理解", "モデル対比"],
         ["2", "2 仮定条件と支点条件", "基礎形状別の支点条件を理解", "基礎3形式"],
         ["3", "3 検討内容と応力伝達", "上部/下部の検討と力の流れを理解", "応力伝達"],
         ["4", "4 支点条件の影響と接地圧", "支点で応力が変わる・接地圧計算", "M・接地圧"]])
body(ws, r + 2, "実務では上部構造を支点で切離して解き、その柱脚反力を下部（基礎・杭・地盤）の"
                "設計外力として用いる『上下分離』が一般的。仮定の妥当性が設計精度を左右する。", span=4, h=32)

# ===== 1 分離と一体 =====
ws = wb.create_sheet("1 分離と一体"); setup(ws, [8, 16, 18, 16, 12, 12])
title_row(ws, 1, "1  上下分離モデルと一体モデル")
head(ws, 3, "■ 図 1  一体モデルと分離モデル"); put_img(ws, figs["models"], "A4", w=840)
head(ws, 30, "■ 問題 1  分離と一体の違い")
r = tbl(ws, 31, ["項目", "一体モデル", "上下分離モデル"],
        [["解析対象", "上部+基礎+地盤ばねを一体", "上部と下部を分けて解く"],
         ["相互作用", "考慮する", "支点条件で近似（記入）"],
         ["規模・手間", "大きい・複雑", "小さい・実務的（記入）"],
         ["適用", "重要建物・杭-地盤相互作用が大", "一般的な建物（記入）"]])
head(ws, r + 2, "■ 問題 2  分離モデルの手順")
body(ws, r + 3, "上下分離モデルで設計する手順を述べよ（①上部を支点で切離して解析→②柱脚反力N・Q・Mを"
                "求める→③その反力を下部（基礎・杭）の外力として基礎を設計）。", h=40)
head(ws, r + 5, "■ 問題 3  分離の利点と注意")
body(ws, r + 6, "上下分離が実務で広く使われる利点（モデルが単純・分業しやすい）と、注意点"
                "（支点条件の仮定が実状と合わないと上部・下部の応力がずれる）を述べよ。", h=40)

# ===== 2 仮定条件と支点条件 =====
ws = wb.create_sheet("2 仮定と支点条件"); setup(ws, [8, 16, 18, 16, 12, 12])
title_row(ws, 1, "2  仮定条件と基礎形状に応じた支点条件")
head(ws, 3, "■ 図 2  基礎形状別の支点条件"); put_img(ws, figs["support"], "A4", w=880)
head(ws, 26, "■ 問題 1  支点条件（ピン/固定/ばね）")
body(ws, 27, "上部構造の柱脚を『ピン』『固定』『回転ばね』とする仮定の違いを説明せよ。"
             "支点条件によって柱脚の曲げモーメントがどう変わるか述べよ。", h=40)
head(ws, 29, "■ 問題 2  基礎形状と支点条件")
r = tbl(ws, 30, ["基礎形状", "支点条件の目安", "考え方（記入）"],
        [["独立フーチング", "ピン〜固定", ""],
         ["べた基礎", "固定に近い/連続ばね", ""],
         ["杭基礎（剛結）", "固定に近い", ""],
         ["杭基礎（ピン接合）", "ピン", ""]])
head(ws, r + 2, "■ 問題 3  仮定条件の整理")
body(ws, r + 3, "上下分離で用いる主な仮定を挙げよ（柱脚の支点条件、剛床仮定、基礎の剛性、"
                "地盤ばねの評価、上部反力の受け渡し）。仮定が結果に与える影響に触れよ。", h=40)

# ===== 3 検討内容と応力伝達 =====
ws = wb.create_sheet("3 検討と応力伝達"); setup(ws, [8, 16, 18, 16, 12, 12])
title_row(ws, 1, "3  上部・下部の検討内容と応力伝達")
head(ws, 3, "■ 図 3  応力伝達（柱脚→基礎→地盤）"); put_img(ws, figs["transfer"], "A4", w=760)
head(ws, 29, "■ 問題 1  上部・下部の検討内容")
r = tbl(ws, 30, ["区分", "検討内容"],
        [["上部構造", "骨組の応力・断面算定・層間変形・保有水平耐力（柱脚を支点で支持）"],
         ["下部構造", "基礎スラブ・基礎梁・杭・地盤（上部反力＋接地圧/杭反力）（記入）"]])
head(ws, r + 2, "■ 問題 2  応力伝達の流れ")
body(ws, r + 3, "上部構造の柱脚力（軸力N・せん断Q・曲げM）が、基礎（フーチング/杭）を通じて"
                "地盤へ伝達される流れを、図3を使って説明せよ。つり合い（上部反力＝地盤反力の合力）に触れよ。", h=44)
head(ws, r + 5, "■ 問題 3  力のつり合い")
body(ws, r + 6, "基礎底面で『上部からの外力（N・Q・M＋基礎自重）』と『地盤反力』がつり合うことを説明せよ。"
                "この関係が接地圧・杭反力の算定の基礎になることを述べよ。", h=40)

# ===== 4 支点条件の影響と接地圧 =====
ws = wb.create_sheet("4 影響と接地圧"); setup(ws, [8, 16, 18, 16, 12, 12])
title_row(ws, 1, "4  支点条件の影響と接地圧の計算例")
head(ws, 3, "■ 図 4  支点条件の影響・接地圧"); put_img(ws, figs["effect"], "A4", w=840)
head(ws, 28, "■ 問題 1  支点条件が応力に与える影響")
body(ws, 29, "柱脚をピンとした場合と固定とした場合で、上部構造の応力（特に柱脚曲げ・柱頭曲げ）が"
             "どう変わるか述べよ。設計では安全側・実状に応じた仮定を選ぶ理由を説明せよ。", h=40)
head(ws, 31, "■ 問題 2  接地圧の計算")
body(ws, 32, "独立フーチング B=L=3.0m、根入れDf=2.0m。上部柱脚 N=1500kN・Q=200kN・M=300kN·m、"
             "フーチング＋土の自重Wf=270kN。底面の平均圧・偏心e・qmax・qminを求めよ"
             "（Mp=M+Q·Df、A=B·L、Z=B·L^2/6、e=Mp/Np）。", h=48)
head(ws, 34, "■ 問題 3  偏心と接地圧分布")
body(ws, 35, "偏心 e と B/6 の関係で接地圧分布（全面台形／一部三角＝浮上り）が決まることを説明せよ。"
             "e<B/6 のとき全面圧縮となることを問2の値で確認せよ。", h=40)
head(ws, 37, "■ 問題 4  設計への反映")
body(ws, 38, "求めた qmax を地盤の許容支持力度 qa と比較する意味を述べよ。"
             "qmax>qa の場合の対応（フーチング拡大・杭基礎への変更）に触れよ。", h=40)

# ===== 解答 =====
ws = wb.create_sheet("解答"); setup(ws, [4, 22, 60, 14])
title_row(ws, 1, "解答・解説", span=4); row = 2


def ah(t):
    global row; head(ws, row, t, span=4); row += 1


def an(t, h=None):
    global row; body(ws, row, t, span=4, ans=True, h=h); row += 1


ah("1  分離と一体")
an("問1：一体は上部・基礎・地盤ばねを1モデルで解き相互作用を厳密に扱うが大規模・複雑。"
   "分離は上部と下部を分け、支点条件で相互作用を近似。手間が小さく分業しやすく実務的。"
   "一般建物は分離、杭-地盤相互作用が重要な建物等は一体を用いる。", h=48)
an("問2：①上部を柱脚の支点条件（ピン/固定等）で切離して骨組解析→②柱脚反力N・Q・Mを算出→"
   "③その反力＋基礎自重を外力として下部（基礎・杭・地盤）を設計。上部と下部を順に解く。", h=44)
an("問3：利点＝モデルが単純で計算・照査が容易、構造/基礎の分業がしやすい。"
   "注意点＝支点条件の仮定が実状（基礎の回転剛性・地盤変形）と合わないと、"
   "柱脚モーメントや基礎応力が過大/過小になる。仮定の妥当性確認が重要。", h=48)

ah("2  仮定と支点条件")
an("問1：ピン＝回転自由で柱脚M=0（基礎に曲げを伝えない）。固定＝回転拘束で柱脚Mが生じる。"
   "回転ばね＝中間（基礎の回転剛性を評価）。固定に近いほど柱脚Mは大きく柱頭Mは小さくなる。", h=44)
an("問2：独立フーチング＝ピン〜固定（フーチング剛性・根入れ・基礎梁拘束による）。"
   "べた基礎＝面で支持し固定に近い/連続ばね。杭基礎は杭頭が剛結なら固定、ピン接合ならピン。"
   "群杭の回転剛性も考慮。", h=44)
an("問3：主な仮定＝①柱脚の支点条件 ②各階の剛床仮定 ③基礎（梁・スラブ）の剛性 "
   "④地盤ばねの評価 ⑤上部反力の下部への受け渡し。特に支点条件は上部・下部双方の"
   "応力に直結するため、基礎形式に整合した仮定とする。", h=48)

ah("3  検討と応力伝達")
an("問1：上部＝骨組応力・断面算定・層間変形・保有水平耐力（柱脚を支点で支持して解析）。"
   "下部＝基礎スラブ・基礎梁・杭・地盤の設計（上部の柱脚反力＋基礎自重を外力に、"
   "接地圧や杭反力を算定して照査）。", h=44)
an("問2：柱脚の軸力N・せん断Q・曲げMがフーチング（または杭）に伝わり、"
   "フーチングでは接地圧、杭では杭頭反力（軸力・水平力・曲げ）として地盤に伝達される。"
   "各段でつり合い（上部反力＝地盤反力の合力）が成立する。", h=44)
an("問3：基礎底面で、上部からのN・Q・Mと基礎自重の合力が、地盤反力（接地圧の合力）と"
   "力・モーメントの両方でつり合う。このつり合いから接地圧分布・杭反力が定まり、"
   "地盤・基礎部材の照査に用いる。", h=44)

ah("4  影響と接地圧")
an("問1：柱脚ピンでは柱脚M=0で柱頭・梁に曲げが集中、固定では柱脚Mが大きくなる。"
   "支点条件で応力分布が変わるため、基礎形式に応じた実状に近い仮定（または安全側の仮定）を選ぶ。", h=40)
an("問2：Np=N+Wf=1500+270=1770kN。Mp=M+Q·Df=300+200×2.0=700kN·m。A=9m2、Z=3×3^2/6=4.5m3。"
   "平均 q=1770/9=196.7kPa。e=Mp/Np=700/1770=0.395m。"
   "qmax=Np/A(1+6e/B)=196.7×(1+6×0.395/3)=352kPa。qmin=196.7×(1-0.79)=41kPa。", h=52)
an("問3：e=0.395m < B/6=0.500m なので全面圧縮（台形分布、qmin>0で浮上りなし）。"
   "e>B/6 なら底面の一部が浮き上がり三角形分布（qmin=0の側が生じる）となる。", h=40)
an("問4：qmax（=352kPa）を地盤の許容支持力度qaと比べ、qmax<=qaを確認する。"
   "qmax>qaならフーチングを拡大して接地圧を下げる、または直接基礎では持たず杭基礎に変更する。", h=40)

XLSX = os.path.join(OUT, "上下分離モデル問題集.xlsx"); wb.save(XLSX); print("saved:", XLSX)
